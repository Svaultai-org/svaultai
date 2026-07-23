import asyncio
import os
import io
import json
import sys
import uuid
import logging
import base64
import re
import secrets
import string
import pathlib
import time
import traceback
from contextlib import asynccontextmanager
from typing import Optional
from auth_local import verify_session_token
from duplicate_detection import (
    find_existing_file_by_hash,
    find_name_conflict,
    next_available_versioned_name,
    normalize_content_sha256,
    normalize_duplicate_action,
    short_hash_for_log,
)
from fastapi import (
    FastAPI, HTTPException, Depends, Request, UploadFile, File, Form,
    BackgroundTasks,
)
from fastapi.responses import StreamingResponse, JSONResponse
from pydantic import BaseModel, Field
from dotenv import load_dotenv
from fastapi.middleware.cors import CORSMiddleware
from openai import AsyncOpenAI
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.errors import RateLimitExceeded
from slowapi.util import get_remote_address
from routes.auth_routes import router as auth_router
from routes.auth_zk_routes import router as auth_zk_router
from routes.login_routes import router as login_router
from routes.vault_manage_routes import router as vault_manage_router
from routes.vault_metadata_migration_routes import (
    router as vault_metadata_migration_router,
)
from routes.vault_ciphertext_write_routes import (
    router as vault_ciphertext_write_router,
)
from vault_chat_memory import (
    get_memory,
    remember_service,
    set_pending_edit,
    clear_pending,
    TTLDict,
)
import psycopg2
from psycopg2.extras import RealDictCursor
import uvicorn
import openpyxl
from vault_core import (
    get_db,
    derive_key,
    decrypt_message,
    encrypt_message,
    decrypt_bytes,
    encrypt_bytes,
    MAX_VAULT_BYTES,
    normalize_service,
    close_pool,
    verify_vault_pin,
    rotate_vault_kdf_if_needed,
    generate_pin_salt as _vault_generate_pin_salt,
    KDF_TARGET_ITERATIONS,
    KDF_LEGACY_ITERATIONS,
    MIN_PIN_LENGTH_NEW_VAULT,
)
from chunked_aead import decrypt_chunk, encrypt_chunk
from cryptography.exceptions import InvalidTag
from device_gate import verify_trusted_device
from tools import SYSTEM_PROMPT, VAULT_FUNCTIONS
from extractor import redact_message, extract_credentials, extract_multiple_credentials
from username_policy import (
    PolicyCache as UsernamePolicyCache,
    UsernamePolicy,
    conservative_default as conservative_username_policy,
    format_policy_note,
    generate_username as generate_username_for_policy,
    resolve_username_policy,
    tighten_policy as tighten_username_policy,
)
try:
    import PyPDF2
except Exception:
    PyPDF2 = None

try:
    import docx
except Exception:
    docx = None


logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

load_dotenv()
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
DATABASE_URL = os.getenv("DATABASE_URL")

if not OPENAI_API_KEY:
    raise RuntimeError("Missing OPENAI_API_KEY")

if not DATABASE_URL:
    raise RuntimeError("Missing DATABASE_URL")

OPENAI_TIMEOUT_SECONDS = float(os.getenv("OPENAI_TIMEOUT_SECONDS", "30"))
OPENAI_MAX_RETRIES = int(os.getenv("OPENAI_MAX_RETRIES", "2"))

                                                                         
try:
    from vault_config import upload as _upload_cfg
    _u = _upload_cfg()
    MAX_UPLOAD_BYTES = _u.max_upload_bytes
    MAX_JSON_BODY_BYTES = _u.max_json_body_bytes
except Exception:
    MAX_UPLOAD_BYTES = int(
        os.getenv("MAX_UPLOAD_BYTES", str(100 * 1024 * 1024)),
    )                                                                          
    MAX_JSON_BODY_BYTES = int(
        os.getenv("MAX_JSON_BODY_BYTES", str(1 * 1024 * 1024)),
    )                                                                        

                                                                       
CHUNK_UPLOAD_MAX_FRAME_BYTES = int(os.getenv(
    "CHUNK_UPLOAD_MAX_FRAME_BYTES", str(20 * 1024 * 1024)
))


def _env_bool(name: str) -> bool:
    return os.getenv(name, "").strip().lower() in ("1", "true", "yes", "on")


_UPLOAD_DEBUG = _env_bool("VAULTAI_UPLOAD_DEBUG")
_CHAT_DEBUG = _env_bool("VAULTAI_CHAT_DEBUG")


def _ulog(label: str, **fields: object) -> None:
    if not _UPLOAD_DEBUG:
        return
    payload = " ".join(f"{k}={v}" for k, v in fields.items())
    print(f"[UPLOAD-TIMING] {label} {payload}", flush=True)


def _clog(label: str, **fields: object) -> None:
    if not _CHAT_DEBUG:
        return
    payload = " ".join(f"{k}={v}" for k, v in fields.items())
    print(f"[CHAT-TIMING] {label} {payload}", flush=True)

client = AsyncOpenAI(
    api_key=OPENAI_API_KEY,
    timeout=OPENAI_TIMEOUT_SECONDS,
    max_retries=OPENAI_MAX_RETRIES,
)
limiter = Limiter(key_func=get_remote_address)

LAST_SERVICE_CACHE = TTLDict(
    max_size=int(os.getenv("LAST_SERVICE_CACHE_MAX_ENTRIES", "10000")),
    ttl_seconds=int(os.getenv("LAST_SERVICE_CACHE_TTL_SECONDS", str(60 * 60))),
)

GENERIC_UPLOAD_INSTRUCTIONS = {
    "analyze and save the uploaded file",
    "analyze and save the uploaded files",
    "analyze and save the uploaded file(s)",
    "save the uploaded file",
    "save the uploaded files",
    "save uploaded file",
    "save uploaded files",
    "analyze uploaded file",
    "analyze uploaded files",
    "upload file",
    "upload files",
    "uploaded file",
    "uploaded files",
    "analyze and save uploaded file",
    "analyze and save uploaded files",
}

                                                                        
@asynccontextmanager
async def lifespan(app: FastAPI):
                                                                     
                                                                    
    try:
        from vault_config import get_config as _vc_get
        _cfg = _vc_get()
        logger.info("[vault_config] active config: %s", _cfg.describe())
    except Exception:
                                                                  
                                                     
        logger.exception(
            "[vault_config] FATAL: config probe failed. Refusing to "
            "boot. Set the required env vars (see vault_config.py).",
        )
        raise

    logger.info("Initializing database...")
    init_db()
    logger.info("Database ready")

    # 2026-07-22 chat deep-fix — eagerly resolve the chat-state
    # backend at startup so the [CHAT-STATE] backend_resolved log
    # line fires when the container boots (not lazily on first
    # chat request). Operators grep for this line to confirm the
    # Uvicorn worker is actually using Redis and not the
    # in-memory fallback.
    try:
        from vault_chat_state_store import get_chat_state_backend
        get_chat_state_backend()
    except Exception:
        logger.exception(
            "[CHAT-STATE] eager backend resolution failed at "
            "startup — chat pending-save may degrade to per-"
            "process state",
        )

                                                                        
    try:
        from device_gate import gate_status_for_boot
        status = gate_status_for_boot()
        if status["enabled"]:
            logger.info(
                "[DEVICE-GATE] enabled=True reason=%s",
                status["reason"],
            )
        else:
            logger.warning(
                "[DEVICE-GATE] ***DISABLED*** reason=%s. This bypasses "
                "trusted-device enforcement on every vault route. NEVER "
                "ship to production with the gate disabled. Unset "
                "VAULTAI_DEVICE_GATE_DEV_DISABLE to re-enable.",
                status["reason"],
            )
    except Exception:
        logger.exception("[DEVICE-GATE] gate_status_for_boot failed; "
                         "continuing with default behavior")

                                                                    
    try:
        from device_gate import (
            _is_dev_environment,
            is_dev_auto_trust_enabled,
        )
        print(
            "[VAULT-DEBUG] DEV FLAGS",
            {
                "VAULTAI_ENV": os.getenv("VAULTAI_ENV"),
                "ENVIRONMENT": os.getenv("ENVIRONMENT"),
                "VAULTAI_DEVICE_GATE_DEV_AUTO_TRUST":
                    os.getenv("VAULTAI_DEVICE_GATE_DEV_AUTO_TRUST"),
                "dev_environment": _is_dev_environment(),
                "dev_auto_trust": is_dev_auto_trust_enabled(),
            },
            flush=True,
        )
    except Exception as exc:
        print(
            f"[VAULT-DEBUG] DEV FLAGS read failed: "
            f"{type(exc).__name__}: {exc}",
            flush=True,
        )

                                                                      
    try:
        from stripe_service import (
            get_stripe_api_key,
            get_stripe_block_price_id,
            get_stripe_webhook_secret,
        )
        _api_ok    = bool(get_stripe_api_key())
        _price_ok  = bool(get_stripe_block_price_id())
        _whsec_ok  = bool(get_stripe_webhook_secret())
        print(
            "[VAULT-DEBUG] STRIPE CONFIG",
            {
                "STRIPE_API_KEY":                "present" if _api_ok else "missing",
                "STRIPE_STORAGE_BLOCK_PRICE_ID": "present" if _price_ok else "missing",
                "STRIPE_WEBHOOK_SECRET":         "present" if _whsec_ok else "missing",
                "checkout_ready":                _api_ok and _price_ok,
                "webhook_ready":                 _api_ok and _whsec_ok,
            },
            flush=True,
        )
                                                                         
                                                                     
        if _api_ok and _price_ok:
            try:
                from stripe_service import probe_stripe_price_currency
                _price_probe = probe_stripe_price_currency()
                print(
                    "[VAULT-DEBUG] STRIPE PRICE PROBE",
                    _price_probe,
                    flush=True,
                )
                if _price_probe.get("ok") and not _price_probe.get(
                    "matches_expected_usd_2500"
                ):
                    print(
                        "[VAULT-DEBUG] STRIPE PRICE WARNING",
                        "Configured Price is NOT USD $25/month. App copy "
                        "shows '$25/month' but Stripe will charge in "
                        f"{_price_probe.get('currency')!r} at "
                        f"{_price_probe.get('unit_amount_cents')} "
                        f"cents/{_price_probe.get('interval')}. "
                        "Recreate the Stripe Price in USD or update the "
                        "app copy.",
                        flush=True,
                    )
            except Exception as exc:
                print(
                    f"[VAULT-DEBUG] STRIPE PRICE PROBE failed: "
                    f"{type(exc).__name__}: {exc}",
                    flush=True,
                )
    except Exception as exc:
        print(
            f"[VAULT-DEBUG] STRIPE CONFIG read failed: "
            f"{type(exc).__name__}: {exc}",
            flush=True,
        )

                                                               
    if os.getenv("VAULTAI_CHUNKED_UPLOADS", "false").lower() == "true":
        try:
            from routes.chunked_upload_routes import sweep_all_stale_uploads
            summary = sweep_all_stale_uploads()
            if summary["files_swept"] > 0:
                logger.warning(
                    "[CHUNKED-SWEEP] startup released %d stale upload(s) "
                    "across %d vault(s); %d byte(s) freed.",
                    summary["files_swept"],
                    summary["vaults_touched"],
                    summary["bytes_freed"],
                )
            else:
                logger.info("[CHUNKED-SWEEP] startup found no stale uploads.")
        except Exception:
                                                                        
            logger.exception("[CHUNKED-SWEEP] startup sweep failed; continuing")

                                                                  
    try:
        from vault_brain_worker import set_embedder
        from brain_embedder import make_query_embed_fn
        set_embedder(make_query_embed_fn(client))
        logger.info("[BRAIN] chunk embedder wired")
    except Exception:
        logger.exception(
            "[BRAIN] set_embedder wiring failed; chunks will land "
            "without vectors and brain retrieval will lexical-fallback",
        )

                                                        
    try:
        from vault_startup import schedule_background_startup_tasks
        _startup_tasks = await schedule_background_startup_tasks()
    except Exception:
        logger.exception(
            "[startup] scheduling background tasks failed; continuing",
        )
        _startup_tasks = None

                                                           
    try:
        yield
    finally:
                                                                    
                                                                  
        if _startup_tasks is not None:
            try:
                from vault_startup import cancel_background_startup_tasks
                await cancel_background_startup_tasks(_startup_tasks)
            except Exception:
                logger.exception(
                    "[startup] cancel_background_startup_tasks raised",
                )
        try:
            from vault_analysis_daemon import shutdown_daemon
            await shutdown_daemon()
        except Exception:
            logger.exception("[DAEMON] shutdown_daemon raised")
        try:
            from vault_intelligence_updater import (
                shutdown_intelligence_updater,
            )
            await shutdown_intelligence_updater()
        except Exception:
            logger.exception(
                "[INTEL-UPDATER] shutdown raised",
            )
        close_pool()


app = FastAPI(title="VaultAI Backend", lifespan=lifespan)
                                                                        
                                                                         
app.include_router(auth_router)
app.include_router(auth_zk_router)
app.include_router(vault_metadata_migration_router)
app.include_router(vault_ciphertext_write_router)
app.include_router(login_router)
app.include_router(vault_manage_router, prefix="/manage")

                                                                        
from routes.device_routes import router as device_router
app.include_router(device_router)

                                                                         
from routes.semantic_search_routes import router as semantic_search_router
app.include_router(semantic_search_router)

                                                                    
from routes.security_center_routes import router as security_center_router
app.include_router(security_center_router)

                                                          
from routes.expiry_routes import router as expiry_router
app.include_router(expiry_router)

                                                                     
from routes.memory_routes import router as memory_router
from routes.relationship_routes import router as relationship_router
app.include_router(memory_router)
app.include_router(relationship_router)


# Inheritance credential escrow (Phase 1) + release flow (Phase 2).
from routes.inheritance_credential_routes import (
    router as inheritance_credential_router,
)
from routes.inheritance_release_routes import (
    router as inheritance_release_router,
)
from inheritance_error_codes import install_inheritance_redaction
install_inheritance_redaction()
app.include_router(inheritance_credential_router)
app.include_router(inheritance_release_router)

                                                                      
from routes.billing_routes import router as billing_router
app.include_router(billing_router)

                                                                      
from routes.stripe_routes import router as stripe_router
app.include_router(stripe_router)

                                                                       
from routes.chunked_download_routes import (              
    router as chunked_download_router,
)
app.include_router(chunked_download_router)
logger.info("Download manifest endpoint registered (always-on)")

                                                               
from routes.file_text_routes import (              
    router as file_text_router,
)
app.include_router(file_text_router)
logger.info("File-text read route registered")

                                                                    
if os.getenv("VAULTAI_CHUNKED_UPLOADS", "false").lower() == "true":
    from routes.chunked_upload_routes import router as chunked_upload_router
    app.include_router(chunked_upload_router)
    logger.info("Chunked uploads enabled (Phase 9.8.B)")

                                                                    
from routes.import_batch_routes import router as import_batch_router              
app.include_router(import_batch_router)

                                                                           
from routes.folder_browser_routes import router as folder_browser_router              
app.include_router(folder_browser_router)

                                                               
from routes.vault_analysis_admin_routes import (              
    router as vault_analysis_admin_router,
)
app.include_router(vault_analysis_admin_router)

                                                          
from routes.deep_answer_routes import (              
    router as deep_answer_router,
)
app.include_router(deep_answer_router)


from routes.crypto_wallet_routes import (
    router as crypto_wallet_router,
)
app.include_router(crypto_wallet_router)


from routes.vault_delete_routes import (
    router as vault_delete_router,
)
app.include_router(vault_delete_router)

app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)


@app.exception_handler(Exception)
async def _global_exception_handler(request: Request, exc: Exception):
    import re as _re
    origin = request.headers.get("origin") or ""
                                                                  
                                                   
    headers = {}
    try:
        if origin and _re.fullmatch(CORS_ALLOWED_ORIGIN_REGEX, origin):
            headers["Access-Control-Allow-Origin"] = origin
            headers["Access-Control-Allow-Credentials"] = "true"
            headers["Vary"] = "Origin"
    except Exception:
        pass
    from fastapi.responses import JSONResponse
    logger.exception(
        "[CHAT-DEBUG] unhandled 500 path=%s origin=%s err=%s",
        request.url.path, origin, repr(exc),
    )
    return JSONResponse(
        status_code=500,
        headers=headers,
        content={
            "detail": "internal_server_error",
            "error_class": type(exc).__name__,
                                                           
                                                             
            "hint": (
                "Server-side error. Check the backend logs for "
                "the stack trace."
            ),
        },
    )


_CORS_DEV_LOCALHOST_ORIGIN_REGEX: str = (
    r"http://("
    r"localhost|"
    r"127\.0\.0\.1|"
    r"0\.0\.0\.0|"
    r"\[::1\]|"
    r"\[::\]|"
    r"10\.0\.2\.2"
    r"):\d+"
)


def _resolve_cors_origin_regex() -> str:
    """Return the CORS ``allow_origin_regex`` for this environment.

    Rules:

    * In **production** (``VAULTAI_ENV=production``/``prod``/``live``),
      an explicit ``CORS_ALLOWED_ORIGIN_REGEX`` is required. Without
      one the app refuses to boot. Localhost is **never** appended in
      production — the returned regex is exactly the operator-supplied
      value. This preserves the "production origins are restricted"
      guarantee and keeps the wildcard-with-credentials footgun off
      the table.

    * In **development / local / test / ci**, browsers running the
      Vite dev server (``http://localhost:5173`` etc.) must always be
      able to preflight against a local backend, even if the operator
      has also configured a production-shaped
      ``CORS_ALLOWED_ORIGIN_REGEX`` (a common local misconfiguration
      that produced the "OPTIONS /auth/login → 400 Bad Request" bug).
      We therefore return a union: the explicit regex (if any)
      ``|`` the loopback pattern. If no explicit regex is set, we
      return the loopback pattern alone.
    """
    explicit = os.getenv("CORS_ALLOWED_ORIGIN_REGEX", "").strip()
    from device_gate import _is_dev_environment
    dev = _is_dev_environment()

    if dev:
        if explicit:


            return f"({explicit})|({_CORS_DEV_LOCALHOST_ORIGIN_REGEX})"
        return _CORS_DEV_LOCALHOST_ORIGIN_REGEX

    if explicit:
        return explicit

    raise RuntimeError(
        "CORS_ALLOWED_ORIGIN_REGEX is required in production. "
        "Set it to a regex matching every legitimate frontend "
        "origin (e.g. r'https://(app|www)\\.example\\.com'). "
        "Refusing to start with a loopback-only default."
    )


CORS_ALLOWED_ORIGIN_REGEX = _resolve_cors_origin_regex()

                                                                           
CORS_ALLOWED_HEADERS = [
    "Authorization",
    "Content-Type",
    "Accept",
    "Origin",
    "X-Requested-With",
    "X-App-Locale",
    "X-Device-Id",
    # 2026-07-22 (3) CORS regression fix. The 7b34d92 client build
    # added X-App-Release to _defaultHeaders() as a diagnostic-only
    # log field, but this allowlist was not updated in the same
    # commit. Starlette's CORSMiddleware rejected every preflight
    # whose Access-Control-Request-Headers included x-app-release
    # with 400 Bad Request. Every authenticated route from
    # https://app.svaultai.com stopped working — /vault-meta,
    # /devices/register, /list-my-vaults, /notifications,
    # /vault-stats, /billing/me, /list-files, /folders,
    # /list-secure-items — because the real GET/POST never fired.
    # See test_cors_preflight_login_2026_07_09.py for the
    # regression tests.
    "X-App-Release",
    # 2026-07-22 (4) Safari-specific CORS regression follow-up.
    # After the (3) fix landed, the Safari production frontend still
    # failed OPTIONS /vault-meta with 400 because its preflight also
    # carries Cache-Control and Pragma in Access-Control-Request-
    # Headers (Safari includes them for GETs when the client sets
    # cache-invalidating headers on the real request). Chrome sent
    # only authorization,content-type,x-app-release,x-device-id and
    # worked; Safari sent
    # authorization,cache-control,pragma,x-app-release,x-device-id
    # and was 400'd. Adding both keeps parity with any browser that
    # forwards standard caching hints in the preflight header set.
    "Cache-Control",
    "Pragma",
    # 2026-07-22 (5) CORS regression fix. Commits 27c14e6 / 4c371ff
    # added X-Client-Request-Id to /beneficiary/list-mine as a
    # correlation id for nginx access-log alignment, but this
    # allowlist was not extended in the same commit. Starlette's
    # CORSMiddleware rejected every preflight whose Access-Control-
    # Request-Headers included x-client-request-id with 400 Bad
    # Request. On the owner-side inheritance panel that meant the
    # real POST /beneficiary/list-mine never fired, package:http
    # surfaced the blocked fetch as ClientException, and the new
    # _looksLikeNetworkAbort() branch silently swallowed it — the
    # panel rendered "No beneficiaries yet." with no error toast.
    # Same shape as (3)/(4) above; see
    # test_cors_preflight_login_2026_07_09.py for the regression
    # test family this belongs to.
    "X-Client-Request-Id",
]

                                                                          
CORS_ALLOWED_METHODS = ["GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"]

app.add_middleware(
    CORSMiddleware,
    allow_origin_regex=CORS_ALLOWED_ORIGIN_REGEX,
    allow_credentials=True,
    allow_methods=CORS_ALLOWED_METHODS,
    allow_headers=CORS_ALLOWED_HEADERS,
    max_age=600,
)


from security_headers import SecurityHeadersMiddleware
app.add_middleware(SecurityHeadersMiddleware)


@app.middleware("http")
async def limit_request_body_size(request: Request, call_next):


    content_length = request.headers.get("content-length")
    if content_length:
        try:
            length = int(content_length)
        except ValueError:
            length = 0

        path = request.url.path or ""
                                    
                                                                              
        is_upload = path == "/upload-file"
        if is_upload:
            limit = MAX_UPLOAD_BYTES
        elif path == "/upload-file/chunk":
            limit = CHUNK_UPLOAD_MAX_FRAME_BYTES
        else:
            limit = MAX_JSON_BODY_BYTES

        if length > limit:
            if is_upload:
                size_mb = length / (1024 * 1024)
                limit_mb = limit / (1024 * 1024)
                return JSONResponse(
                    status_code=413,
                    content={
                        "detail": {
                            "code": "upload_safety_cap_exceeded",
                            "message": (
                                f"This file is {size_mb:.0f} MB. The current "
                                f"single-upload safety cap is {limit_mb:.0f} MB "
                                f"— a temporary backend limit, not your vault "
                                f"storage quota."
                            ),
                            "max_upload_bytes": limit,
                        },
                    },
                )
            return JSONResponse(
                status_code=413,
                content={
                    "detail": (
                        f"Request body too large for {path or 'this endpoint'} "
                        f"(received {length} bytes, limit {limit})."
                    ),
                },
            )

    return await call_next(request)


_VAULT_DEBUG_PATHS_OF_INTEREST = frozenset({
    "/chat",
    "/expiry/active",
    "/memory/timeline",
    "/relationships/list",
    "/debug/routes",
    "/debug/ping",
    "/verify-pin",
    "/my-vault",
})


@app.middleware("http")
async def debug_request_logger(request: Request, call_next):
    method = request.method
    path = request.url.path
    origin = request.headers.get("origin") or "-"
    has_auth = "authorization" in request.headers
    has_did = "x-device-id" in request.headers
    of_interest = (
        path in _VAULT_DEBUG_PATHS_OF_INTEREST
        or method == "OPTIONS"
    )

    if of_interest:
        if method == "OPTIONS":
            acrm = request.headers.get("access-control-request-method") or "-"
            acrh = request.headers.get("access-control-request-headers") or "-"
            logger.info(
                "[VAULT-DEBUG] PREFLIGHT %s %s origin=%s req-method=%s req-headers=%s",
                method, path, origin, acrm, acrh,
            )
                                                                           
            print(
                f"[VAULT-DEBUG] PREFLIGHT {method} {path} origin={origin} "
                f"req-method={acrm} req-headers={acrh}",
                flush=True,
            )
        else:
            logger.info(
                "[VAULT-DEBUG] REQ %s %s origin=%s auth_present=%s x_device_id_present=%s",
                method, path, origin, has_auth, has_did,
            )
                                                                              
                                                                 
            print(
                f"[VAULT-DEBUG] REQ {method} {path} origin={origin} "
                f"auth_present={has_auth} x_device_id_present={has_did}",
                flush=True,
            )

    response = await call_next(request)

    if of_interest:
        if method == "OPTIONS":
            acao = response.headers.get("access-control-allow-origin") or "-"
            acah = response.headers.get("access-control-allow-headers") or "-"
            acam = response.headers.get("access-control-allow-methods") or "-"
            acac = response.headers.get("access-control-allow-credentials") or "-"
            logger.info(
                "[VAULT-DEBUG] PREFLIGHT-RESP %s %s status=%d allow-origin=%s "
                "allow-headers=%s allow-methods=%s allow-credentials=%s",
                method, path, response.status_code, acao, acah, acam, acac,
            )
        else:
            logger.info(
                "[VAULT-DEBUG] RESP %s %s status=%d",
                method, path, response.status_code,
            )

    return response


def _debug_endpoints_enabled() -> bool:


    try:
        from device_gate import _is_dev_environment
        return _is_dev_environment()
    except Exception:
        return False


@app.get("/debug/routes")
async def debug_routes_endpoint():


    if not _debug_endpoints_enabled():
        raise HTTPException(status_code=404, detail="Not Found")
    routes = []
    for r in app.routes:
        path = getattr(r, "path", None)
        if not isinstance(path, str):
            continue
        methods = sorted(list(getattr(r, "methods", set()) or set()))
        routes.append({"path": path, "methods": methods})
    routes.sort(key=lambda x: x["path"])
    return {"count": len(routes), "routes": routes}


@app.get("/debug/ping")
async def debug_ping_endpoint():


    if not _debug_endpoints_enabled():
        raise HTTPException(status_code=404, detail="Not Found")
    logger.info(
        "[VAULT-DEBUG] /debug/ping pid=%s cwd=%r file=%r",
        os.getpid(), os.getcwd(), __file__,
    )
    return {
        "ok": True,
        "pid": os.getpid(),
        "cwd": os.getcwd(),
        "file": __file__,
    }


def generate_pin_salt() -> str:
                                                                            
                                   
    return _vault_generate_pin_salt()


def get_or_create_pin_salt(vault_id: str) -> str:


    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)

        cursor.execute(
            """
            SELECT pin_salt
            FROM vaults
            WHERE vault_id = %s
            LIMIT 1
            """,
            (vault_id,),
        )
        row = cursor.fetchone()

        if not row:
            raise HTTPException(status_code=404, detail="Vault not found")

        if row.get("pin_salt"):
            return row["pin_salt"]

        new_salt = generate_pin_salt()
        cursor.execute(
            """
            UPDATE vaults
            SET pin_salt = %s
            WHERE vault_id = %s
            RETURNING pin_salt
            """,
            (new_salt, vault_id),
        )
        saved = cursor.fetchone()
        conn.commit()
        return saved["pin_salt"]
    finally:
        conn.close()


BACKEND_DIR = pathlib.Path(__file__).resolve().parent


def init_db():


    if os.getenv("RUN_MIGRATIONS_ON_STARTUP", "true").strip().lower() in ("false", "0", "no", "off"):
        logger.info("RUN_MIGRATIONS_ON_STARTUP=false → skipping startup migrations")
        return

    from alembic.config import Config
    from alembic import command

    alembic_cfg = Config(str(BACKEND_DIR / "alembic.ini"))
    alembic_cfg.set_main_option("script_location", str(BACKEND_DIR / "migrations"))

    logger.info("Running Alembic migrations to head...")
    try:
        command.upgrade(alembic_cfg, "head")
    except Exception as exc:
                                                                     
                                                                  
        banner = (
            "\n" + "=" * 72 + "\n"
            "[MIGRATION FAILURE] alembic upgrade head raised "
            f"{type(exc).__name__}: {exc}\n"
            "Backend startup is ABORTED. The schema may be partially\n"
            "migrated — inspect the traceback below before retrying.\n"
            "Run `pwsh scripts/check-migrations.ps1` (or the equivalent\n"
            "`python -m compileall migrations/versions` + `python -m\n"
            "alembic upgrade head`) to reproduce the failure locally\n"
            "without bringing up the full app shell.\n"
            + "=" * 72 + "\n"
        )
        print(banner, file=sys.stderr, flush=True)
        logger.exception("[MIGRATION FAILURE] alembic upgrade head raised")
        raise
    logger.info("Schema is at Alembic head.")

@app.get("/my-vault")
def get_my_vault(
    principal=Depends(verify_trusted_device),
):
                                                                         
                                                                        
    vault_id = principal["vault_id"]

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)

                                                                       
        cursor.execute(
            """
            SELECT vault_id, vault_name, pin_salt, pin_verifier,
                   must_reset, locked_until, inherited_from_label, created_at
            FROM vaults
            WHERE vault_id = %s
            LIMIT 1
            """,
            (vault_id,)
        )

        row = cursor.fetchone()

                                                                        
        logger.info(
            "[PIN-DEBUG] /my-vault vault=...%s has_vault=%s has_pin=%s "
            "vault_name=%r",
            (vault_id or "")[-8:],
            row is not None,
            bool(row["pin_verifier"]) if row else False,
            row["vault_name"] if row else None,
        )

        if row:
            return {
                "has_vault": True,
                "vault_name": row["vault_name"],
                "has_pin": bool(row["pin_verifier"]),
                "must_reset": bool(row.get("must_reset")),
                "locked_until": row["locked_until"].isoformat() if row.get("locked_until") else None,
                "is_orphaned": False,
                "orphan_data": None,
            }
    finally:
        conn.close()

    return {
        "has_vault": False,
        "vault_name": None,
        "has_pin": False,
        "is_orphaned": False,
        "orphan_data": None,
    }


class ChatRequest(BaseModel):
    encrypted_message: str
    vault_name: str
    pin: str
    uploaded_file_ids: list[str] = Field(default_factory=list)



    app_locale: Optional[str] = None

    # Selection hint from a chat card tap. When the user taps a
    # specific row in a login/file list card the frontend sends a
    # natural-language prompt like "Show my Gmail login" AND this
    # structured hint identifying the exact row the user picked.
    # Backend uses the hint to bypass ambiguity when two rows share
    # the same title (two logins named "Gmail", two files named
    # "videos", etc.). The id is never rendered in visible prose;
    # it stays as a structured field only.
    # Shape: {"kind": "login" | "file", "id": "<uuid>"}
    selection_hint: Optional[dict] = None

    # 2026-07-21 concurrency-safety fields for the "first chat forces
    # PIN" production incident. These carry the (pin_salt,
    # kdf_iterations) the client's cached vault key was ACTUALLY
    # derived from. The /chat handler compares them to the CURRENT
    # DB row BEFORE decryption: if they differ, the DB rotated under
    # the client (another tab's /rotate-vault-kdf, a maintenance
    # path, or a stale metadata cache the client derived against) and
    # the handler responds with 409 kdf_generation_stale carrying the
    # authoritative salt/iter for the client to re-derive against.
    #
    # Left Optional so the field is safely absent on requests from
    # older clients — server falls back to the legacy behavior of
    # deriving with the current DB state. Presence + mismatch is the
    # only path that returns 409; presence + match is the fast path
    # (matches what the server would have derived anyway); absence
    # is legacy-compat (same behavior as before this field existed).
    kdf_salt_used: Optional[str] = None
    kdf_iterations_used: Optional[int] = None

    # 2026-07-22 (2) protocol version — the AUTHORITATIVE boundary
    # for whether kdf_salt_used and kdf_iterations_used are
    # mandatory on this request. A client-controlled header is NOT
    # a valid enforcement gate (spoofable, stripable by proxies);
    # a request-body field the client MUST include is. See
    # vault_kdf_generation.client_requires_kdf_fields.
    #
    #   >= 2  => KDF fields mandatory; missing => 400
    #            missing_kdf_generation_fields (typed).
    #   1 or None => legacy pass-through (older builds).
    crypto_protocol_version: Optional[int] = None


class VaultNameCheck(BaseModel):
    vault_name: str


class VaultMetaRequest(BaseModel):
    vault_name: str
    pin: Optional[str] = None
    acknowledge_orphan_wipe: bool = False


class VerifyPinRequest(BaseModel):
    vault_name: str
    pin: str


class WipeOrphanRequest(BaseModel):
    confirm: bool = False


class RotateVaultKdfRequest(BaseModel):
    vault_name: str
    pin: str


class CreateBeneficiaryRequest(BaseModel):
    vault_name: str
    pin: str
    label: str


class LinkBeneficiaryRequest(BaseModel):
    vault_name: str                               
    pin: str                                    
    pairing_code: str


class BeneficiaryListRequest(BaseModel):
    vault_name: str
    pin: str


class DeleteBeneficiaryRequest(BaseModel):
    link_id: int
    vault_name: str                      
    pin: str


class TransferActionRequest(BaseModel):
    link_id: int
    vault_name: str
    pin: str


class ClaimTransferRequest(BaseModel):
    link_id: int
    vault_name: str                                       
    pin: str
    inherited_vault_name: Optional[str] = None                          


class ListFilesRequest(BaseModel):
    vault_name: str
    pin: str


class VaultStatsRequest(BaseModel):
    vault_name: str
    pin: str


class DownloadFileRequest(BaseModel):
    vault_name: str
    file_id: str
    pin: str


class DeleteFileRequest(BaseModel):
    vault_name: str
    file_id: str
    pin: str


def create_pin_verifier(pin: str, pin_salt: str, iterations: int) -> str:
    key = derive_key(pin, pin_salt, iterations=iterations)
    return encrypt_message("vaultai_pin_ok", key)


def get_verified_vault_key(vault_id: str, pin: str) -> bytes:
                                                                           
                                                                           
    return verify_vault_pin(vault_id, pin)


def populate_vault_key_cache_after_pin(
    *, vault_id: str, token_id: str, key: bytes,
) -> None:


    try:
        from vault_key_cache import get_cache
        get_cache().store(vault_id=vault_id, token_id=token_id, key=key)
    except Exception:
                                                                   
                                                             
        logger.exception(
            "[KEY-CACHE] store failed for vault=%s token=%s",
            (vault_id or "")[:8], (token_id or "")[:8],
        )

                                                      
    try:
        from vault_intelligence_updater import on_vault_unlock
        on_vault_unlock(vault_id)
    except Exception:
        pass


def generate_strong_password(length: int = 20) -> str:
    alphabet = string.ascii_letters + string.digits + "!@#$%^&*()-_=+"
    while True:
        password = "".join(secrets.choice(alphabet) for _ in range(length))
        if (
            any(c.islower() for c in password)
            and any(c.isupper() for c in password)
            and any(c.isdigit() for c in password)
            and any(c in "!@#$%^&*()-_=+" for c in password)
        ):
            return password


def _generate_username() -> str:


    return "".join(
        secrets.choice(string.ascii_lowercase + string.digits) for _ in range(10)
    )


_username_policy_cache: Optional[UsernamePolicyCache] = None


def _get_username_policy_cache() -> UsernamePolicyCache:
    global _username_policy_cache
    if _username_policy_cache is None:
        _username_policy_cache = UsernamePolicyCache(get_db)
    return _username_policy_cache


def password_strength_warning(password: Optional[str]) -> Optional[str]:
    if not password:
        return None

    weak_words = ["password", "123456", "qwerty", "admin", "letmein", "welcome"]
    lower = password.lower()

    if len(password) < 10:
        return "Your password looks short and may not be secure."

    if any(w in lower for w in weak_words):
        return "Your password looks easy to guess."

    if password.isdigit() or password.isalpha():
        return "Your password needs a mix of letters, numbers, and symbols."

    return None

async def detect_vault_intent(message: str, memory: dict, vault_name: str) -> dict:
    prompt = f"""
You are the intent detection brain for VaultAI.

Return ONLY valid JSON. No explanation.

Possible intents:
- save_login
- retrieve_login
- list_logins
- delete_login
- edit_login
- generate_password
- generate_login
- generated_login_repair
- confirm
- reject
- name_file
- retrieve_file
- search_memory
- list_by_tag
- understand_document
- expiry_alerts
- expiry_status
- remember_fact
- recall_memory
- related_items
- document_details
- general_chat
- identity
- summarize_vault
- list_files
- list_folders
- list_recent_uploads
- search_files_for_credentials
- extract_logins_from_file
- travel_readiness
- analyze_file
- read_file_text
- search_files_about
- related_files
- vault_clusters

Use generate_login when the user asks you to GENERATE or CREATE new
credentials. The user is asking VaultAI to PRODUCE the fields, not to
save fields they typed. The service name is usually inside the same
message — extract it; do NOT ask the user to repeat it. Set parts_wanted
based on the wording:
  - "create me a login for X"               -> generate_login, service=X, parts_wanted=full
  - "create me a login for my X bank"       -> generate_login, service=X bank, parts_wanted=full
  - "make a username and password for X"    -> generate_login, service=X, parts_wanted=username_password
  - "generate a strong password for X"      -> generate_password, service=X
  - "i want you to create me username,
     email, and password for X"             -> generate_login, service=X, parts_wanted=full
  - "create me a login for my BNB bank"     -> generate_login, service=BNB bank, parts_wanted=full
The service token can be ANY string (BNB, Chase, Gmail, a bank name in
any language, a website name, etc.). Treat every service token the same;
there is NO hardcoded list. Never set service=null when the user named
one in the same message — extract it verbatim.

Use generated_login_repair when the user is rejecting a field that
VaultAI just generated. Triggered by rejection phrases targeting one of
the generated fields. Examples — these are illustrative, not a closed
list; any synonymous phrase in any language counts:
  - "the username is not acceptable"
  - "username not accepted"
  - "username rejected"
  - "this username does not work"
  - "service says username is invalid"
  - "email not acceptable"
  - "password not acceptable"
  - "regenerate the username"
  - "change the username"
  - "make another username"
  - "that username doesn't work"
  - "give me a new password"
  - "BNB bank username is not acceptable"   (service explicitly named)
  - "the username and password are both bad" (multiple fields)
Set fields_to_regenerate to a list containing one or more of:
  ["username", "email", "password"]
based on which field(s) the user is rejecting. If the user explicitly
names the service ("BNB bank username..."), put it in service. Otherwise
leave service=null — the handler will fall back to the most-recent
generated login from conversation memory.
Set reason to a short tag like "not_accepted", "rejected_by_service",
"invalid_format", or "user_preference" when inferrable.

DO NOT use generated_login_repair when:
  - The user is rejecting a NON-generated stored value (use edit_login).
  - The user is starting a fresh generation for a different service
    (use generate_login).
  - There is no recent generated login in conversation memory AND the
    user did not name a service. In that case classify as general_chat;
    the handler will ask one clarification question.

The conversation memory dump below may include
"last_generated_login": {{service, generated_fields, ts}}. When present,
treat it as the default service for a follow-up repair. When absent,
DO NOT invent a service — leave service=null.

Use search_memory when the user is doing a FUZZY find across the vault
without naming an exact service or asset name. Examples:
  - "find my Qatar documents"
  - "show me my passport file"
  - "any voice note about investors"
  - "where's the Chase login I saved"
  - "do I have a document about taxes"
If the user names an exact saved service ("show my Gmail login"), prefer
retrieve_login. If they clearly name an exact saved asset, retrieve_file
is fine. When in doubt between retrieve_file and search_memory, prefer
search_memory: it is broader and falls back gracefully if nothing matches.

Use list_by_tag when the user asks for one of the 13 fixed categories:
travel, finance, legal, medical, business, education, identity,
government, security, personal, media, receipt, tax.
Also fire on natural phrasings such as:
  - "show my travel files"
  - "my travel stuff"
  - "finance documents"
  - "tax things"
  - "identity documents"
  - "list my receipts"

SEMANTIC LABEL MAPPING (this is critical — colloquial labels must map
to the right closed-set tag):
  - "show my IDs" / "all my IDs" / "my IDs"        -> tag=identity
  - "show identity documents"                       -> tag=identity
  - "show my id documents" / "my identity docs"     -> tag=identity
  - "show government documents" / "gov docs"        -> tag=government
  - "show my financial documents" / "money stuff"   -> tag=finance
  - "all my receipts"                               -> tag=receipt
  - "show my medical files" / "health docs"         -> tag=medical
  - "show legal documents"                          -> tag=legal
A user who says "show my IDs" is asking for the IDENTITY family
(passport, driver's license, national ID card, residence card,
etc.) — return list_by_tag with tag=identity, NOT retrieve_file with
asset_name="IDs". Retrieval by exact saved_name is the wrong call for
plural/family labels.

Return JSON with "tag" set to one of the 13 strings above (lowercase).
For fuzzy queries that don't fit a fixed category, prefer search_memory.

Use understand_document when the user asks about a SPECIFIC saved
document by type or by a field on a document. Examples:
  - "show passport info"
  - "when does my visa expire"
  - "what documents do I have"
  - "receipt amount"
  - "insurance expiry"
  - "passport expiry date"
  - "show my contracts"
Allowed doc_type values: passport, visa, boarding_pass, ticket,
hotel_itinerary, receipt, invoice, contract, agreement, degree,
certificate, insurance, driver_license, id_card, tax_document,
medical_record.
Allowed doc_field values: country, expiry_date, issue_date,
purchase_date, due_date, effective_date, renewal_date, departure_date,
event_date, record_date, check_in, check_out, amount, currency,
merchant, vendor, parties, airline, flight_number, origin, destination,
hotel, location, provider, institution, year, field, issuer, tax_year,
form_type, passport_number_last4, id_number_last4,
license_number_last4, policy_number_last4, license_class.
If the user asks by CATEGORY ("travel files", "finance docs"), prefer
list_by_tag. If they're asking about a specific document or a field
on one, use understand_document. If neither fits, fall back to
search_memory.

Use expiry_alerts when the user is asking what's EXPIRING, EXPIRED,
RENEWING, COMING DUE, or about a date that is approaching or past,
WITHOUT pinning to one specific doc_type. Examples:
  - "what expires soon"
  - "expired documents"
  - "contracts renewing"
  - "insurance expiring"
  - "documents expiring this month"
  - "anything expiring"
  - "what's coming due"
If the user names a specific doc_type AND asks about its expiry
("when does my passport expire", "my visa expiry", "passport expiry")
prefer understand_document with doc_type set and doc_field=expiry_date,
because that returns the specific document. Use expiry_alerts when the
question is broader (all docs, no specific type).

Use expiry_status (14.C) when the user is asking for a structured
review of OUTSTANDING actions / RENEWALS / DEADLINES grouped by
severity. Examples:
  - "do I need to renew anything"
  - "what should I act on"
  - "what's my expiry status"
  - "show urgent renewals"
  - "anything critical expiring"
Optionally set `expiry_filter` to one of: passport, visa, id_card,
driver_license, insurance, tax, contract, subscription, inheritance,
custom - when the user narrows the question to one type.
expiry_alerts (broad date-field sweep) and expiry_status (structured
alerts grouped by severity) cover overlapping intents; pick
expiry_status whenever the user is asking "is something urgent" or
"what to act on" rather than just "list dated things".

Use remember_fact ONLY when the user uses an EXPLICIT remember-style
trigger phrase. Examples that SHOULD trigger remember_fact:
  - "Remember that I work on Goufer"
  - "Remember my company is VaultAI"
  - "Save this about me: I travel to Qatar"
  - "Note that my goal is becoming an AI inventor"

Examples that MUST NOT trigger remember_fact (these are casual chat
statements, NOT save requests - classify as general_chat instead):
  - "I flew to Qatar last week"
  - "I like Emirates"
  - "I visited London"
  - "I use Chase"
  - "I want to go to Tokyo"
  - "My day was good"
If the message does not contain an explicit save trigger ("remember",
"save this", "note that", "don't forget"), DO NOT use remember_fact.

NEVER set memory_value to a password, PIN, card number, CVV, SSN,
seed phrase, mnemonic phrase, wallet address, private key, API key,
JWT, bearer token, OTP, or any credential. NEVER set memory_value to
a passport number, ID number, license number, or any government
identifier - those belong in the encrypted vault, not in memory. If
the message contains a secret, classify as general_chat instead.

Allowed memory_type values: identity, travel, preference, project,
company, goal, location, relationship, note, family, date, life_event.

Mapping hints (informal):
  people / coworkers / friends -> relationship
  spouse / parent / sibling    -> family
  business / employer          -> company
  destinations / cities        -> location  (or travel for a trip)
  trips / journeys             -> travel
  ambitions / targets          -> goal
  ongoing work / endeavours    -> project
  anniversaries / birthdays    -> date
  weddings, moves, milestones  -> life_event
  habits / likes / dislikes    -> preference

memory_key is a short slug (e.g. "name", "airline", "current_project",
or a slug derived from the value). It may contain non-Latin letters
(Arabic / CJK / Cyrillic / Hangul) - the slug normalizer keeps them.

memory_value is the human-readable fact in display form.

memory_event_date is OPTIONAL: an ISO date "YYYY-MM-DD" when the
memory is ABOUT (e.g. a meeting in 2024, a trip in 2023). Leave null
for facts without a specific date.

Use document_details (14.D) when the user asks for STRUCTURED
ENTITY summaries pulled from saved documents - country, expiry
date, merchant name, amount + currency, contract dates, hospital,
policy provider, etc. Examples:
  - "show passport details"
  - "what country is this visa"
  - "invoice amount"
  - "who issued this insurance"
  - "show contract dates"
  - "passport country and expiry"
  - "who is the merchant on this receipt"
Optional payload (any of):
  - `entity_doc_type` : one of passport, visa, id_card,
                        driver_license, boarding_pass, hotel_itinerary,
                        ticket, invoice, receipt, contract, agreement,
                        degree, certificate, insurance, tax_document,
                        medical_record. Filter to one doc family.
  - `entity_type`     : one of identity, travel, government, finance,
                        tax, business, medical, education, insurance,
                        legal, personal, other. Filter to one entity
                        category.
  - `entity_key`      : a single metadata key like "country",
                        "expiry_date", "amount", "merchant",
                        "provider", "parties". Returns just that
                        field across matching docs.
Prefer document_details over understand_document when the user is
asking ACROSS docs ("all my passport countries") or for a focused
entity field. Use understand_document for one-doc summaries.

Use related_items when the user asks what is CONNECTED / LINKED /
RELATED to a specific saved asset, doc type, country, company, or
service. Examples:
  - "show travel things linked to my passport"
  - "what belongs to this company"
  - "things related to Qatar"
  - "what connects to this invoice"
  - "related items for Acme"
  - "what's tied to my Gmail account"
Return JSON with "anchor_text" set to the noun the user is asking
about ("passport", "Qatar", "Acme", "invoice", "Gmail", etc.) - a
short free-form string. If the user is asking generally about
relationships without naming an anchor, prefer general_chat.

Use recall_memory when the user asks what you remember about them.
Examples:
  - "what do you remember about me"
  - "what company am I building"
  - "what airline do I prefer"
  - "what projects am I working on"
  - "who was the investor from last year"
  - "what business goal did I save"
  - "where did I travel before Morocco"
  - "what do you remember about my sister"
Optionally fill memory_type if the user asks about a specific category.

For 14.B ranked recall, also fill (any of) these optional fields when
the user is asking a focused question:
  - memory_query   : the noun(s) they are asking ABOUT, in any
                     language. Example: "investor" for the investor
                     question, "Goufer" for a project question.
                     Leave null for broad "what do you remember"
                     questions.
  - memory_anchor  : an entity name to position the query relative to.
                     For "where did I travel before Morocco" the
                     anchor is "Morocco".
  - memory_direction : "before" | "after" | "around" - the relation
                       to the anchor. Set null if no direction word
                       is present.

For generate_login, also set "parts_wanted":
  - "password_only"     — user asked only for a password
  - "username_password" — user asked for both a username and password
  - "full"              — user asked to generate the whole login, or said
                          "create credentials" without specifying parts

User vault name: {vault_name}

PENDING-NAMING CONTEXT:
If conversation memory contains a "pending_named_file" entry, the user
has JUST uploaded a file/audio/video/image and the assistant asked
"what should I save this as?" — the user is now answering that
question. In this case, even a short bare-noun reply ("vibing",
"holiday memories", "trip notes") or a "save the ... as X" / "call it
X" / "name it X" phrasing is naming THAT pending file. Set:
  intent="name_file"
  asset_name=<the name the user gave, with command words like "save
             the audio recording to" / "call it" / "name it" stripped>

BUT — and this is critical — the presence of a pending_named_file
does NOT lock the next message into name_file. The user is free to
change subject at any time, and a clearly different request must win.
When the new message is an EXPLICIT different intent, classify it as
that intent and IGNORE the pending file. Examples that must NOT be
classified as name_file even when a pending_named_file is present:
  - "show my saved logins"      -> list_logins
  - "all the logins"            -> list_logins
  - "list my files"             -> search_memory  (or list_by_tag)
  - "retrieve my documents"     -> search_memory
  - "show recordings"           -> search_memory
  - "find my Gmail password"    -> retrieve_login
  - "delete my Chase login"     -> delete_login
  - "what do you remember"      -> recall_memory
  - "open billing"              -> general_chat   (handled by frontend route)
  - "show settings"             -> general_chat
The rule of thumb: if the message contains a verb (show, list, find,
search, retrieve, delete, edit, open) paired with an item-type noun
(logins, files, recordings, documents, cards, notes, secrets), it is
a separate request. Classify it that way; do NOT name the pending
file. A backend guard also enforces this, but the LLM must make the
right call first so the user gets the correct response on the first
hop.

Conversation memory:
{json.dumps(memory)}

User message:
{message}

Use summarize_vault / list_files when the user asks for an overview
of EVERYTHING in their vault. Examples:
  - "what files do I have"
  - "what is in my vault"
  - "show me my files"
  - "list my files"
  - "show everything I uploaded"
Either intent is fine — the handler treats them identically and
returns the inventory summary (counts, top folders, recent files).
This is NOT a request for a specific saved login or a category —
those go to retrieve_login / list_by_tag.

Use list_folders when the user asks specifically about FOLDERS,
not files. Examples:
  - "what folders do I have"
  - "show my folders"
  - "list folders"
  - "what's in my folders"
  - "show the folders in my vault"

Use list_recent_uploads when the user asks what they uploaded
RECENTLY. Examples:
  - "show recent uploads"
  - "what did I upload today"
  - "what files did I add this week"
  - "show me my recent files"
  - "what changed recently"

Use read_file_text when the user wants VaultAI to READ THE TEXT of
a specific file and return what's inside. The user is asking
"what does this file say" — they want the file's actual content,
not a summary, not a search. Examples:
  - "what text is in agreement.pdf"
  - "what does invoice-69.pdf say"
  - "show me the text from canada.pdf"
  - "read this file" (when the user has just opened or named one)
  - "what's inside Marital_Separation_Agreement.pdf"
  - "open the text of NGL_Proposal.docx"
Set file_name to the filename token from the message verbatim
(e.g. "agreement.pdf"). Set file_name=null when the user says
"this file" / "the one I just opened" without naming one — the
handler will fall back to the most recent file in chat memory.

read_file_text is DIFFERENT from analyze_file. analyze_file produces
a SUMMARY / classification of a document. read_file_text returns
the RAW text of the file as it was extracted (PDF text / OCR /
DOCX content). Prefer read_file_text when the user is asking
"what does it say" — they want the content.

read_file_text is DIFFERENT from extract_logins_from_file.
extract_logins_from_file pulls structured credential records out
of a credentials-list file. read_file_text returns the file's
text and never tries to detect credentials inside it.

Use search_files_for_credentials when the user wants to find
UPLOADED FILES (PDFs, .env files, password manager exports, etc.)
that may contain login credentials. Examples:
  - "find files with usernames and passwords"
  - "which files contain login details"
  - "scan my files for passwords"
  - "do I have any files with credentials"
  - "show me files that have credentials saved in it"
  - "find env files"
This is NOT the same as retrieving a saved login. Saved logins live
in vault_items and are reached via retrieve_login / list_logins.
search_files_for_credentials returns file CARDS, not credential
fields. Never ask the user for a service name for this intent —
the answer is a list of files, not a single credential.

This intent is FIND-ONLY. It surfaces files and ranks them; it
does NOT extract or save anything. If the user wants to actually
pull credentials OUT of a named file and review them as vault
logins, route to extract_logins_from_file instead.

Use extract_logins_from_file when the user wants VaultAI to PULL
the login records OUT of a specific file and let them review the
result before saving anything. Triggered by phrases like:
  - "extract logins from this file"
  - "extract the credentials from passedwordtex.pdf"
  - "turn this password list into vault logins"
  - "save these credentials" (when referring to credentials INSIDE
    an uploaded file, not a credential the user just typed)
  - "pull the logins out of my Bitwarden export"
Set asset_name to the literal filename token from the message when
present ("passedwordtex.pdf"). Set asset_name=null when the user
refers to "this file" / "that file" without naming one — the
handler falls back to the last opened file in chat memory.

The handler builds a credential_extraction_review card that lists
the detected login records (service, username/email if safe,
password_present: true) for the user to review. It NEVER auto-
saves. The user must affirmatively confirm before any record
becomes a vault login.

extract_logins_from_file is DIFFERENT from
search_files_for_credentials: the latter ranks files, the former
operates on ONE named file. When the user names a file AND asks
to extract / pull / save the credentials in it, prefer
extract_logins_from_file. When the user asks the generic "which
files have credentials" / "show me files with credentials",
prefer search_files_for_credentials.

Use search_files_about when the user asks for files matching a free-
form noun / topic / entity / person / company / category.
Examples:
  - "find files about Wells Fargo"
  - "show files about Maureen"
  - "what files do I have about taxes"
  - "documents about Qatar"
  - "find documents mentioning health insurance"
  - "find saved-login files"
  - "show legal documents"
Set asset_name to the noun the user is asking about ("Wells Fargo",
"Maureen", "taxes", "saved-login", etc.) — a short free-form
string. The handler runs a ranked search across the Vault
Understanding Index (purpose / category / entity / topic / term /
summary / preview / content) and returns a file list.

search_files_about is DIFFERENT from list_by_tag. list_by_tag is
for the 13 fixed categories ("show my travel files", "tax things")
where the answer is a closed-set bucket lookup. search_files_about
is for the open-set "about X" queries where X is a person, company,
brand, topic, or free-form noun the user typed. When the message
fits a closed-set tag, prefer list_by_tag; otherwise route to
search_files_about.

search_files_about is ALSO different from search_files_for_credentials.
search_files_for_credentials is the dedicated credential-files
search that returns the credential-files card. search_files_about
is broader; it can find credential files too (purpose=saved_login_list),
but for the explicit "find files with credentials" phrasing
search_files_for_credentials is preferred.

Use analyze_file when the user asks WHAT IS IN a specific file, or
asks for an analysis / summary / classification of one file. The
file is named in the message OR referred to via "this file" / "that
file" (the handler will resolve the reference against the most
recently mentioned file). Examples:
  - "analyze passedwordtex.pdf"
  - "what is in passedwordtex.pdf"
  - "what's inside this file"
  - "tell me about this file"
  - "what kind of file is this"
  - "explain this document"
  - "summarize this file"
  - "what does this pdf contain"
Set asset_name to the literal filename token from the message when
present ("passedwordtex.pdf"). Set asset_name=null when the user
refers to "this file" / "that file" without naming one — the handler
falls back to the last opened file in chat memory.

This intent is DIFFERENT from retrieve_file. retrieve_file returns
the file (the View/Download card). analyze_file returns a content
analysis (purpose, structure, summary) without ever exposing
credential values inside the file. When the user asks BOTH (rare),
prefer analyze_file — the file card is still reachable from the
analysis reply via a tap.

NEVER use analyze_file when the user is asking VaultAI to extract
or surface specific credential values ("show me the AOL password
inside this file"). That stays general_chat (the handler refuses)
or retrieve_login (for saved logins). analyze_file is the
purpose+structure overview, never a credential-extraction trigger.

Use related_files when the user asks VaultAI to show files
RELATED / CONNECTED / GROUPED with a specific anchor file. The
answer is the Relationship Graph view, anchored on ONE file,
listing other files linked by entity / topic / category /
folder / import batch / duplicate hash / front/back filename
pattern / semantic similarity / archive content. Examples:
  - "what is related to this file"
  - "show files related to my Maureen ID"
  - "what documents belong with this application"
  - "show everything related to my Qatar trip"
  - "find files connected to this passport"
  - "show supporting documents for this form"
  - "what goes with passport_2024.pdf"
  - "find related items for backup.zip"
Set asset_name to the literal anchor filename when the user
named one ("passport_2024.pdf"). Set asset_name=null when the
user referred to the anchor via "this file" / "that one" — the
handler falls back to the last opened / last searched file in
chat memory. Multiple plausible anchors return a
file_disambiguation card.

related_files is DIFFERENT from search_files_about. search_files_about
matches files against a free-form noun the user typed
("find files about Wells Fargo"). related_files is anchored on
a SPECIFIC FILE and walks the persisted relationship graph
rather than re-ranking against a query. When the user names a
NOUN, route to search_files_about; when the user references a
SPECIFIC FILE (named or via "this file"), route to related_files.

Use vault_clusters when the user asks about WHOLE-VAULT connected
groups — they want a map of how documents cluster together, NOT
anchored on a single file. Examples:
  - "show me everything connected in my vault"
  - "show vault clusters"
  - "what groups of related files do I have"
  - "organize my vault by connected documents"
  - "show me my document groups"
  - "find all the related document groups"
  - "show me how my documents cluster"
The handler walks every persisted relationship edge in the vault
and groups files into connected components (identity / travel /
finance / tax / application / company / duplicates / archive /
same_person). Weak same-folder-only clusters are filtered out.

vault_clusters is DIFFERENT from related_files. related_files is
anchored on ONE file ("what is related to this passport").
vault_clusters takes the ENTIRE vault and surfaces every
meaningful group at once. When the user names a SPECIFIC FILE,
route to related_files; when the user asks about THE VAULT or
"my documents" / "everything" / "all", route to vault_clusters.

Use travel_readiness when the user wants a checklist of whether
they have what they need for travel. Examples:
  - "am I travel-ready"
  - "do I have everything for my trip"
  - "what travel documents do I have"
  - "check my travel docs"
  - "ready for travel"
  - "travel readiness"
The handler inspects passports, visas, tickets, boarding passes,
hotel itineraries, and their expiry dates. Use this intent when
the user is asking for the CHECK, not just a list of travel files
(that's list_by_tag with tag=travel).

Return JSON with this shape:
{{
  "intent": "...",
  "service": null,
  "field": null,
  "value": null,
  "asset_name": null,
  "parts_wanted": null,
  "fields_to_regenerate": null,
  "reason": null,
  "tag": null,
  "doc_type": null,
  "doc_field": null,
  "memory_type": null,
  "memory_key": null,
  "memory_value": null,
  "memory_event_date": null,
  "memory_query": null,
  "memory_anchor": null,
  "memory_direction": null,
  "anchor_text": null,
  "expiry_filter": null,
  "entity_doc_type": null,
  "entity_type": null,
  "entity_key": null,
  "confirmed": false
}}
"""

    try:
        from vault_ai_provider import chat_complete_with_fallback
        result = await chat_complete_with_fallback(
            messages=[
                {"role": "system", "content": "Return only valid JSON for VaultAI intent detection."},
                {"role": "user", "content": prompt},
            ],
            model_kind="intent",
            temperature=0,
        )

        raw = (result.content or "").strip()
        return json.loads(raw)

    except Exception:
        logger.error("Intent detection failed")
        return {
            "intent": "general_chat",
            "service": None,
            "field": None,
            "value": None,
            "asset_name": None,
            "parts_wanted": None,
            "fields_to_regenerate": None,
            "reason": None,
            "tag": None,
            "doc_type": None,
            "doc_field": None,
            "memory_type": None,
            "memory_key": None,
            "memory_value": None,
            "memory_event_date": None,
            "memory_query": None,
            "memory_anchor": None,
            "memory_direction": None,
            "anchor_text": None,
            "expiry_filter": None,
            "entity_doc_type": None,
            "entity_type": None,
            "entity_key": None,
            "confirmed": False,
        }

def _stream_single_message(message: str, key: bytes):
    async def generator():
        encrypted_chunk = encrypt_message(message, key)
        yield f"data: {encrypted_chunk}\n\n".encode("utf-8")
    return generator()


def _conversation_key(vault_id: str) -> str:
    return vault_id


def _set_last_service(vault_id: str, service: Optional[str]):
    normalized = _normalize_service_name(service)
    if normalized != "general":
        LAST_SERVICE_CACHE.set(_conversation_key(vault_id), normalized)


def _get_last_service(vault_id: str) -> Optional[str]:
    return LAST_SERVICE_CACHE.get(_conversation_key(vault_id))


def _normalize_service_name(service: Optional[str]) -> str:
    return normalize_service(service)


_NAMING_COMMAND_PATTERNS = [
                                                                     
                                                                     
    re.compile(
        r"^\s*save\s+(?:the|this|that)?\s*"
        r"(?:audio\s+recording|recording|audio|video|image|photo|file)\s*"
        r"(?:as|to|under|with\s+(?:the\s+)?name|named|called)\s+(.+?)\s*$",
        re.IGNORECASE,
    ),
    re.compile(r"^\s*save\s+(?:it|this|that)\s+as\s+(.+?)\s*$", re.IGNORECASE),
    re.compile(r"^\s*(?:call|name)\s+(?:it|this|that)\s+(.+?)\s*$", re.IGNORECASE),
    re.compile(r"^\s*name(?:d)?\s+(?:it|this|that)?\s*(.+?)\s*$", re.IGNORECASE),
    re.compile(r"^\s*as\s+(.+?)\s*$", re.IGNORECASE),
]


def _strip_naming_command(message: Optional[str]) -> str:


    if not message:
        return ""
    s = message.strip().rstrip(".")
    for pat in _NAMING_COMMAND_PATTERNS:
        m = pat.match(s)
        if m:
            return m.group(1).strip().rstrip(".") or s
    return s


_EXPLICIT_DIFFERENT_INTENT_PATTERNS = [
                                                                    
                                                                       
    re.compile(
        r"\b(?:show|list|see|view|display|find|search|retrieve|get|fetch|"
        r"open|read|delete|remove|edit|change|update|rename)\b"
        r".{0,40}?"
        r"\b(?:login|logins|password|passwords|credential|credentials|"
        r"file|files|recording|recordings|document|documents|"
        r"card|cards|note|notes|memory|memories|fact|facts|"
        r"upload|uploads|secret|secrets)\b",
        re.IGNORECASE,
    ),
                                                               
    re.compile(
        r"\ball\b.{0,15}?\b(?:login|logins|password|passwords|credential|"
        r"credentials|file|files|recording|recordings|document|documents|"
        r"card|cards|note|notes|secret|secrets)\b",
        re.IGNORECASE,
    ),
                                                                       
                                                         
    re.compile(
        r"\bmy\s+(?:saved\s+)?(?:login|logins|password|passwords|"
        r"credential|credentials|file|files|recording|recordings|"
        r"document|documents|card|cards|note|notes|secret|secrets)\b",
        re.IGNORECASE,
    ),
                                                                  
                                                                   
    re.compile(
        r"\b(?:billing|subscription|storage|account|settings|profile|"
        r"upgrade|plan|invoice|invoices)\b",
        re.IGNORECASE,
    ),
                                                      
    re.compile(
        r"\b(?:what(?:'s| is| do| are)|how many|do i have|"
        r"have i (?:got|saved))\b",
        re.IGNORECASE,
    ),
]


def message_signals_explicit_different_intent(message: Optional[str]) -> bool:


    if not message:
        return False
    for pat in _EXPLICIT_DIFFERENT_INTENT_PATTERNS:
        if pat.search(message):
            return True
    return False


LOGIN_INTENT_KEYWORDS = (
    "login", "logins", "password", "passwords",
    "credential", "credentials",
)


def decide_pending_file_intent_override(
    llm_intent: Optional[str],
    message: str,
    pending_file_present: bool,
) -> tuple[str, bool, Optional[str]]:


    final = llm_intent
    if not pending_file_present:
        return final or "", False, None

    stripped = (message or "").strip()
    looks_like_name = (
        0 < len(stripped) <= 80
        and "?" not in stripped
        and "\n" not in stripped
    )
    explicit_other_intent = message_signals_explicit_different_intent(message)
    chat_bucket = {"general_chat", "identity", None, ""}

    if llm_intent == "name_file":
        return "name_file", False, "llm_already_name_file"









    if looks_like_name and not explicit_other_intent:
        return "name_file", True, None



    if llm_intent not in chat_bucket:
        return llm_intent or "", False, "llm_explicit_other_intent"


    if explicit_other_intent:
        lower_msg = (message or "").lower()
        if any(k in lower_msg for k in LOGIN_INTENT_KEYWORDS):
            return "list_logins", False, "explicit_logins_intent_downgrade"
        return (
            llm_intent or "",
            False,
            "explicit_different_intent",
        )


    return llm_intent or "", False, "message_not_bare_name"


_ACCOMPANYING_TEXT_NAMING_PATTERNS = [
                                                                 
    re.compile(
        r"^\s*(?:save|store|keep|file|put)\s+"
        r"(?:this|that|it|my|the)?\s*"
        r"(?:audio\s+recording|recording|audio|video|image|photo|"
        r"file|document|doc|pic|picture|note|scan)?\s*"
        r"(?:as|under|named|called|with\s+(?:the\s+)?name)\s+(.+?)\s*$",
        re.IGNORECASE,
    ),
                                                       
    re.compile(
        r"^\s*(?:call|name|label)\s+(?:this|that|it)\s+(?:the\s+)?(.+?)\s*$",
        re.IGNORECASE,
    ),
                                                                    
                                                                   
    re.compile(
        r"^\s*(?:save|store|keep|file|upload)\s+my\s+(.+?)\s*$",
        re.IGNORECASE,
    ),
                                                        
    re.compile(
        r"^\s*this\s+is\s+(?:my\s+|the\s+|a\s+)?(.+?)\s*$",
        re.IGNORECASE,
    ),
                                                     
    re.compile(
        r"^\s*(?:named|labeled|labelled|called)\s+(.+?)\s*$",
        re.IGNORECASE,
    ),
]


_USE_ORIGINAL_NAME_VETO = re.compile(
                                                            
                                                                       
    r"\b(?:use|keep|leave|save\s+with|stick\s+with|stay\s+with|with)\s+"
    r"(?:the\s+|their\s+|its\s+|each(?:'s)?\s+|every\s+|all\s+)?"
    r"(?:original|same|own|default)\s+"
    r"(?:file\s*)?names?\b"
    r"|\b(?:original|default)\s+(?:file\s*)?names?\b",
    re.IGNORECASE,
)


def extract_naming_intent_from_accompanying_text(
    text: Optional[str],
) -> Optional[str]:


    if not text:
        return None
    stripped = text.strip()
    if not stripped:
        return None
                                                                     
                          
    if "?" in stripped:
        return None
                                                          
    if "\n" in stripped:
        return None
                                                                      
                                            
    if len(stripped) > 120:
        return None
                                                              
                                                                     
    if message_signals_explicit_different_intent(stripped):
        return None
                                                                   
                                                                     
    if _USE_ORIGINAL_NAME_VETO.search(stripped):
        return None

                                                                       
    stripped_via_existing = _strip_naming_command(stripped)
    if stripped_via_existing and stripped_via_existing != stripped:
        candidate = _normalize_asset_name(stripped_via_existing)
        if candidate and _USE_ORIGINAL_NAME_VETO.search(candidate):
            return None
        return candidate or None

                                                                   
    for pat in _ACCOMPANYING_TEXT_NAMING_PATTERNS:
        m = pat.match(stripped)
        if m:
            raw = m.group(1).strip().rstrip(".")
                                                                        
                                           
            if _USE_ORIGINAL_NAME_VETO.search(raw):
                return None
            candidate = _normalize_asset_name(raw)
            return candidate or None

    return None


def _read_file_semantic_inputs(
    vault_id: str, file_id: str,
) -> Optional[dict]:


    try:
        from taxonomy import DOC_TYPE_TO_FAMILY
    except Exception:
        DOC_TYPE_TO_FAMILY = {}

    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT uf.file_name, uf.saved_name, uf.asset_type,
                   vdm.doc_type
              FROM uploaded_files uf
              LEFT JOIN vault_document_metadata vdm
                     ON vdm.vault_id = uf.vault_id
                    AND vdm.uploaded_file_id = uf.id
             WHERE uf.vault_id = %s
               AND uf.id = %s
             LIMIT 1
            """,
            (vault_id, file_id),
        )
        row = cur.fetchone()
        if not row:
            return None

        cur.execute(
            """
            SELECT DISTINCT tag
              FROM vault_asset_tags
             WHERE vault_id = %s
               AND uploaded_file_id = %s
            """,
            (vault_id, file_id),
        )
        tag_rows = cur.fetchall() or []
    finally:
        conn.close()

    doc_type = row.get("doc_type")
    return {
        "saved_name":  row.get("saved_name"),
        "file_name":   row.get("file_name"),
        "asset_type":  row.get("asset_type"),
        "doc_type":    doc_type,
        "doc_family":  DOC_TYPE_TO_FAMILY.get(doc_type) if doc_type else None,
        "tags":        [str(r["tag"]) for r in tag_rows if r.get("tag")],
    }


def enqueue_semantic_profile_for_file(
    vault_id: str, file_id: str,
) -> None:


    try:
        from semantic_embedder import (
            is_enabled,
            build_semantic_profile_text,
            enqueue_uploaded_file_embedding,
        )
        if not is_enabled():
            return
        inputs = _read_file_semantic_inputs(vault_id, file_id)
        if not inputs:
            return
        text = build_semantic_profile_text(**inputs)
        if not text:
            return
        enqueue_uploaded_file_embedding(
            client, vault_id, file_id, "semantic_profile", text,
        )
    except Exception as exc:
        print(
            "[CHAT-DEBUG] enqueue_semantic_profile_failed "
            f"file_id={file_id} error_type={type(exc).__name__} repr={exc!r}",
            flush=True,
        )


def _persist_naming_text_classification(
    *,
    vault_id: str,
    file_id: str,
    naming_text: str,
    file_name: Optional[str] = None,
    content_type: Optional[str] = None,
) -> Optional[str]:


    try:
        from document_understanding import classify_naming_text_doc_type
    except Exception:
        return None

    doc_type = classify_naming_text_doc_type(naming_text)
    if not doc_type:
        return None

                                                                  
    try:
        conn = get_db()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO vault_document_metadata
                        (vault_id, uploaded_file_id,
                         doc_type, metadata_json, confidence)
                    VALUES (%s, %s, %s, %s::jsonb, %s)
                    ON CONFLICT (vault_id, uploaded_file_id) DO UPDATE SET
                        doc_type      = EXCLUDED.doc_type,
                        metadata_json = EXCLUDED.metadata_json,
                        confidence    = EXCLUDED.confidence,
                        updated_at    = NOW()
                    """,
                    (
                        vault_id, file_id,
                        doc_type, json.dumps({}), 0.85,
                    ),
                )
            conn.commit()
        finally:
            conn.close()
    except Exception as exc:
        print(
            "[CHAT-DEBUG] persist_naming_classification_db_failed "
            f"file_id={file_id} doc_type={doc_type!r} "
            f"error_type={type(exc).__name__} repr={exc!r}",
            flush=True,
        )
        return None

                                                                    
    try:
        from asset_tagger import tag_uploaded_file_safe
        tag_uploaded_file_safe(
            vault_id, file_id,
            file_name=file_name,
            content_type=content_type,
            doc_type=doc_type,
            doc_metadata={},
            replace=True,
        )
    except Exception:
        pass

                                                                      
    try:
        from taxonomy import DOC_TYPE_TO_ASSET_TYPE
        if doc_type in DOC_TYPE_TO_ASSET_TYPE:
            ct_lower = (content_type or "").lower()
            new_asset_type = (
                "id_image" if ct_lower.startswith("image/")
                else DOC_TYPE_TO_ASSET_TYPE[doc_type]
            )
            conn = get_db()
            try:
                with conn.cursor() as cur:
                    cur.execute(
                        """
                        UPDATE uploaded_files
                           SET asset_type = %s
                         WHERE id = %s AND vault_id = %s
                        """,
                        (new_asset_type, file_id, vault_id),
                    )
                conn.commit()
            finally:
                conn.close()
    except Exception:
        pass

                                                                    
    try:
        enqueue_semantic_profile_for_file(vault_id, file_id)
    except Exception:
        pass

    return doc_type


EXACT_NAME_WEIGHT      = 1.0
DOC_TYPE_WEIGHT        = 0.7
DOC_FAMILY_WEIGHT      = 0.5
TAG_WEIGHT_PER_MATCH   = 0.3
TAG_WEIGHT_MAX         = 0.6
                                                                    
                                                                  
ENTITY_WEIGHT_PER_MATCH = 0.5
ENTITY_WEIGHT_MAX       = 0.9
VECTOR_WEIGHT          = 0.6
HYBRID_MIN_SCORE       = 0.25


def score_hybrid_match(
    *,
    query_lower: str,
    candidate_saved_name: Optional[str],
    candidate_doc_type: Optional[str],
    candidate_doc_family: Optional[str],
    candidate_tags: Optional[list[str]] = None,
    query_doc_type: Optional[str] = None,
    query_tags: Optional[list[str]] = None,
    vector_score: float = 0.0,
    entity_matches: int = 0,
    matched_entity_values: Optional[list[str]] = None,
) -> dict:


    score = 0.0
    reasons: list[str] = []

                                                                     
    saved = (candidate_saved_name or "").strip().lower()
    if saved and (saved == query_lower or saved in query_lower):
        score += EXACT_NAME_WEIGHT
        reasons.append("exact_name")
    elif saved and query_lower in saved:
                                                                 
        score += EXACT_NAME_WEIGHT * 0.5
        reasons.append("name_substring")

                                                                   
    if query_doc_type and candidate_doc_type:
        if candidate_doc_type == query_doc_type:
            score += DOC_TYPE_WEIGHT
            reasons.append("doc_type")
        else:
                                                                 
                                                             
            try:
                from taxonomy import DOC_TYPE_TO_FAMILY
                if (
                    DOC_TYPE_TO_FAMILY.get(query_doc_type)
                    and DOC_TYPE_TO_FAMILY.get(query_doc_type)
                       == DOC_TYPE_TO_FAMILY.get(candidate_doc_type)
                ):
                    score += DOC_FAMILY_WEIGHT
                    reasons.append("doc_family")
            except Exception:
                pass

                                                 
    if candidate_doc_family:
        fam = candidate_doc_family.lower()
        if fam in query_lower:
            score += DOC_FAMILY_WEIGHT
            reasons.append("family_keyword")

                                                                     
    if query_tags and candidate_tags:
        ct = {t.lower() for t in candidate_tags if t}
        qt = {t.lower() for t in query_tags if t}
        overlap = ct & qt
        if overlap:
            boost = min(len(overlap) * TAG_WEIGHT_PER_MATCH, TAG_WEIGHT_MAX)
            score += boost
            reasons.append(f"tags:{','.join(sorted(overlap))}")

                                                                     
    if entity_matches and entity_matches > 0:
        boost = min(
            entity_matches * ENTITY_WEIGHT_PER_MATCH,
            ENTITY_WEIGHT_MAX,
        )
        score += boost
        if matched_entity_values:
            reasons.append(
                f"entities:{','.join(sorted(set(matched_entity_values)))}"
            )
        else:
            reasons.append(f"entities:{entity_matches}")

                                                                   
    if vector_score and vector_score > 0:
        score += max(0.0, min(1.0, vector_score)) * VECTOR_WEIGHT
        reasons.append(f"vector:{vector_score:.2f}")

    return {"score": round(score, 4), "reasons": reasons}


_US_STATE_NAMES: tuple[str, ...] = (
    "alabama","alaska","arizona","arkansas","california","colorado",
    "connecticut","delaware","florida","georgia","hawaii","idaho",
    "illinois","indiana","iowa","kansas","kentucky","louisiana",
    "maine","maryland","massachusetts","michigan","minnesota",
    "mississippi","missouri","montana","nebraska","nevada",
    "new hampshire","new jersey","new mexico","new york",
    "north carolina","north dakota","ohio","oklahoma","oregon",
    "pennsylvania","rhode island","south carolina","south dakota",
    "tennessee","texas","utah","vermont","virginia","washington",
    "west virginia","wisconsin","wyoming",
)


_QUERY_FREEFORM_STOPWORDS: frozenset[str] = frozenset({
    "the","a","an","my","your","our","their","this","that","these",
    "those","last","next","past","previous","recent","old","new",
    "year","month","week","day","time","one","two","three","four",
    "five","ten","first","second","third","yesterday","today",
    "tomorrow",
})


def extract_query_entities(
    query: str,
    *,
    now=None,
) -> list[dict]:


    out: list[dict] = []
    seen: set[tuple[str, str]] = set()
    if not query:
        return out
    q = query.lower()

    def _emit(kind: str, value: str, candidate_keys: list[str]) -> None:
        key = (kind, value)
        if key in seen:
            return
        seen.add(key)
        out.append({
            "value": value,
            "kind": kind,
            "candidate_keys": candidate_keys,
        })

                         
    try:
        from document_understanding import _COUNTRY_NAMES
    except Exception:
        _COUNTRY_NAMES = ()
    country_keys = [
        "country", "nationality", "origin", "destination", "location",
    ]
    for country in _COUNTRY_NAMES:
        if re.search(r"\b" + re.escape(country) + r"\b", q):
            _emit("country", country, country_keys)

                         
    state_keys = [
        "country", "state", "location", "origin", "destination",
    ]
    for state in _US_STATE_NAMES:
        if re.search(r"\b" + re.escape(state) + r"\b", q):
            _emit("state", state, state_keys)

                                      
    for m in re.finditer(r"\b(?:19|20)\d{2}\b", q):
        _emit("year", m.group(0), [])

                              
    if now is None:
        try:
            from datetime import datetime
            now = datetime.now()
        except Exception:
            now = None
    if now is not None:
        try:
            current_year = int(now.year)
        except Exception:
            current_year = None
        if current_year is not None:
            for phrase, offset in (
                ("last year", -1),
                ("this year",  0),
                ("next year",  1),
            ):
                if phrase in q:
                    _emit("year", str(current_year + offset), [])

                                                                   
    freeform_keys = [
        "merchant", "vendor", "hotel", "location", "city",
        "airline", "institution", "provider", "issuer",
        "store", "place", "name",
    ]
    for m in re.finditer(
        r"\b(?:from|in|at)\s+"
        r"([a-z][a-z']+(?:\s+[a-z][a-z']+)?)",
        q,
    ):
        token = m.group(1).strip()
        if not token:
            continue
                                               
        parts = token.split()
        if any(p in _QUERY_FREEFORM_STOPWORDS for p in parts):
            continue
                            
        if token.replace(" ", "").isdigit():
            continue
                                                           
        if any((kind in ("country", "state") and val == token)
               for (kind, val) in seen):
            continue
        _emit("freeform", token, freeform_keys)

    return out


def extract_query_tags(query: str) -> list[str]:


    from taxonomy import ALLOWED_TAGS, DOC_TYPE_TO_TAGS
    if not query:
        return []
    q = query.lower()
    hits: list[str] = []
    seen: set[str] = set()

    def _add(tag: str) -> None:
        if tag in ALLOWED_TAGS and tag not in seen:
            seen.add(tag)
            hits.append(tag)

                                  
    for tag in ALLOWED_TAGS:
        for term in (tag, tag + "s"):
            if re.search(r"\b" + re.escape(term) + r"\b", q):
                _add(tag)
                break

                                                                
    try:
        from document_understanding import classify_naming_text_doc_type
        dt = classify_naming_text_doc_type(query)
        if dt and dt in DOC_TYPE_TO_TAGS:
            for tag in DOC_TYPE_TO_TAGS[dt]:
                _add(tag)
    except Exception:
        pass

    return hits


def list_logins_tool(vault_id: str) -> str:


    from locales import ENGLISH_FORMATTER as _fmt

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT service
              FROM vault_items
             WHERE vault_id  = %s
               AND item_type = 'login'
             GROUP BY service
             ORDER BY service ASC
            """,
            (vault_id,),
        )
        rows = cursor.fetchall()
    finally:
        conn.close()

    if not rows:
        return _fmt.no_saved_logins()

    items = [f"- {str(r['service']).title()}" for r in rows]
    return _fmt.saved_logins_header() + "\n" + "\n".join(items)


def _normalize_asset_name(name: Optional[str]) -> str:
    return normalize_service(name)


def _title_case_asset(name: str) -> str:
    if not name:
        return "File"
    return " ".join(word.capitalize() for word in name.split())


def _is_generic_upload_instruction(message: str) -> bool:
    lower = _normalize_asset_name(message)
    if lower in GENERIC_UPLOAD_INSTRUCTIONS:
        return True

    generic_parts = [
        "analyze and save the uploaded file",
        "analyze and save the uploaded files",
        "analyze uploaded file",
        "save uploaded file",
        "save the uploaded file",
        "save the uploaded files",
    ]
    return any(part in lower for part in generic_parts)


def _is_image_content_type(content_type: Optional[str], file_name: str) -> bool:
                                                             
                                                                     
    try:
        from vault_image_formats import is_image_row
        if is_image_row(mime=content_type, file_name=file_name):
            return True
    except Exception:
                                                                   
                                                 
        pass
    ct = (content_type or "").lower()
    lower_name = (file_name or "").lower()

    if ct.startswith("image/"):
        return True

    return lower_name.endswith(
        (
            ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp",
            ".heic", ".heif", ".tif", ".tiff", ".avif",
        )
    )

def _is_video_content_type(content_type: Optional[str], file_name: str) -> bool:
    ct = (content_type or "").lower()
    lower_name = (file_name or "").lower()

    if ct.startswith("video/"):
        return True

    return lower_name.endswith((".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v"))


def _is_audio_content_type(content_type: Optional[str], file_name: str) -> bool:
    ct = (content_type or "").lower()
    lower_name = (file_name or "").lower()

    if ct.startswith("audio/"):
        return True

    return lower_name.endswith(
        (".mp3", ".m4a", ".wav", ".ogg", ".aac", ".flac")
    )

def _is_document_like_file(content_type: Optional[str], file_name: str) -> bool:
    lower_name = (file_name or "").lower()

    if _is_image_content_type(content_type, file_name):
        return False

    return lower_name.endswith((".pdf", ".docx", ".xlsx", ".txt", ".csv", ".json"))


def _should_analyze_upload(content_type: Optional[str], file_name: str) -> bool:
    return _is_document_like_file(content_type, file_name)


_IMPORT_BATCH_STATUSES: frozenset[str] = frozenset({
    "pending",
    "uploading",
    "completed",
    "completed_with_errors",
    "cancelled",
    "failed",
})

                                                                
_TERMINAL_IMPORT_BATCH_STATUSES: frozenset[str] = frozenset({
    "completed",
    "completed_with_errors",
    "cancelled",
    "failed",
})


def _is_terminal_import_status(status: Optional[str]) -> bool:


    if status is None:
        return False
    return status in _TERMINAL_IMPORT_BATCH_STATUSES


def _compute_complete_import_status(failed_count: int) -> str:


    if failed_count <= 0:
        return "completed"
    return "completed_with_errors"


def _account_storage_headroom(vault_id: str) -> tuple[int, int]:


    from billing import get_account_id_for_vault, get_entitlement
    account_id = get_account_id_for_vault(vault_id)
    if account_id is None:
        return get_vault_total_bytes(vault_id), int(MAX_VAULT_BYTES)
    ent = get_entitlement(account_id)
    return ent.used_bytes, ent.effective_limit_bytes


def _storage_limit_error_detail(
    used: int,
    limit: int,
    requested: int,
) -> dict:


    return {
        "code": "vault_storage_limit_exceeded",
        "message": (
            f"Vault storage limit reached. "
            f"Used {used / (1024*1024):.0f} MB of "
            f"{limit / (1024*1024):.0f} MB. "
            f"Delete data before uploading more."
        ),
        "used_bytes": used,
        "projected_bytes": used + requested,
        "limit_bytes": limit,
    }


def start_import_batch(
    *,
    vault_id: str,
    root_folder_name: Optional[str],
    total_files: int,
    total_bytes_planned: int,
) -> dict:


    if total_files < 0 or total_bytes_planned < 0:
        raise HTTPException(
            status_code=400,
            detail={
                "code": "invalid_import_batch_request",
                "message": "total_files and total_bytes_planned must be ≥ 0.",
            },
        )

    if total_bytes_planned > 0:
        used, limit = _account_storage_headroom(vault_id)
        if used + total_bytes_planned > limit:
            raise HTTPException(
                status_code=413,
                detail=_storage_limit_error_detail(
                    used, limit, total_bytes_planned,
                ),
            )

    import_id = str(uuid.uuid4())
    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            INSERT INTO import_batches (
                import_id, vault_id, root_folder_name,
                total_files, total_bytes_planned, status
            )
            VALUES (%s, %s, %s, %s, %s, 'pending')
            RETURNING import_id, vault_id, root_folder_name,
                      total_files, uploaded_count, failed_count,
                      skipped_duplicate_count, status,
                      total_bytes_planned, total_bytes_uploaded,
                      created_at, updated_at, completed_at
            """,
            (
                import_id,
                vault_id,
                root_folder_name,
                int(total_files),
                int(total_bytes_planned),
            ),
        )
        row = cursor.fetchone()
        conn.commit()
    finally:
        conn.close()
    return _serialize_import_batch(row)


def get_import_batch(
    *,
    vault_id: str,
    import_id: str,
) -> Optional[dict]:


    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT import_id, vault_id, root_folder_name,
                   total_files, uploaded_count, failed_count,
                   skipped_duplicate_count, status,
                   total_bytes_planned, total_bytes_uploaded,
                   created_at, updated_at, completed_at
            FROM import_batches
            WHERE import_id = %s AND vault_id = %s
            """,
            (import_id, vault_id),
        )
        row = cursor.fetchone()
    finally:
        conn.close()
    return _serialize_import_batch(row) if row else None


def _validate_import_id_for_upload(
    vault_id: str,
    import_id: Optional[str],
) -> Optional[dict]:


    if not import_id:
        return None
    row = get_import_batch(vault_id=vault_id, import_id=import_id)
    if row is None:
        raise HTTPException(
            status_code=404,
            detail={
                "code": "import_batch_not_found",
                "message": "Import batch not found for this vault.",
            },
        )
    if _is_terminal_import_status(row.get("status")):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "import_batch_terminal",
                "message": (
                    "This import batch is already "
                    f"{row.get('status')!r}; no new files can be "
                    "added."
                ),
                "status": row.get("status"),
            },
        )
    return row


def _bump_import_batch_on_upload_success(
    *,
    vault_id: str,
    import_id: Optional[str],
    file_size: int,
) -> None:


    if not import_id:
        return
    try:
        conn = get_db()
        try:
            cursor = conn.cursor()
            cursor.execute(
                """
                UPDATE import_batches
                SET uploaded_count = uploaded_count + 1,
                    total_bytes_uploaded = total_bytes_uploaded + %s,
                    status = CASE
                        WHEN status = 'pending' THEN 'uploading'
                        ELSE status
                    END,
                    updated_at = NOW()
                WHERE import_id = %s AND vault_id = %s
                  AND status IN ('pending', 'uploading')
                """,
                (int(file_size), import_id, vault_id),
            )
            conn.commit()
        finally:
            conn.close()
    except Exception:
        logger.warning(
            "import_batches count bump failed import=%s file_size=%s "
            "— upload succeeded; aggregate counts will be off",
            import_id, file_size,
        )


def _bump_import_batch_on_skipped_duplicate(
    *,
    vault_id: str,
    import_id: Optional[str],
) -> None:


    if not import_id:
        return
    try:
        conn = get_db()
        try:
            cursor = conn.cursor()
            cursor.execute(
                """
                UPDATE import_batches
                SET skipped_duplicate_count =
                        skipped_duplicate_count + 1,
                    status = CASE
                        WHEN status = 'pending' THEN 'uploading'
                        ELSE status
                    END,
                    updated_at = NOW()
                WHERE import_id = %s AND vault_id = %s
                  AND status IN ('pending', 'uploading')
                """,
                (import_id, vault_id),
            )
            conn.commit()
        finally:
            conn.close()
    except Exception:
        logger.warning(
            "import_batches skipped-duplicate bump failed import=%s "
            "— upload was deduped; aggregate counts will be off",
            import_id,
        )


def _format_duplicate_skipped_message(existing: dict) -> str:


    rp = (existing.get("relative_path") or "").strip()
    name = (
        existing.get("saved_name")
        or existing.get("file_name")
        or "this file"
    )
    if rp:
        return f"Already in your vault at {rp}."
    return f"Already in your vault as {name}."


def _format_duplicate_found_message(existing: dict) -> str:


    rp = (existing.get("relative_path") or "").strip()
    name = (
        existing.get("saved_name")
        or existing.get("file_name")
        or "this file"
    )
    if rp:
        return (
            f"This file already exists in your vault as "
            f"{name} ({rp}). Skip or keep both?"
        )
    return (
        f"This file already exists in your vault as {name}. "
        f"Skip or keep both?"
    )


def _format_name_conflict_message(existing: dict) -> str:


    rp = (existing.get("relative_path") or "").strip()
    if rp:
        return (
            "A file with this name already exists in this folder, "
            "but the content is different. "
            f"(Existing at {rp})"
        )
    return (
        "A file with this name already exists in your vault, "
        "but the content is different."
    )


def _bump_import_batch_on_failure(
    *,
    vault_id: str,
    import_id: Optional[str],
) -> None:


    if not import_id:
        return
    try:
        conn = get_db()
        try:
            cursor = conn.cursor()
            cursor.execute(
                """
                UPDATE import_batches
                SET failed_count = failed_count + 1,
                    updated_at = NOW()
                WHERE import_id = %s AND vault_id = %s
                """,
                (import_id, vault_id),
            )
            conn.commit()
        finally:
            conn.close()
    except Exception:
        logger.warning(
            "import_batches failure bump failed import=%s "
            "— aggregate counts will be off", import_id,
        )


def cancel_import_batch(
    *,
    vault_id: str,
    import_id: str,
) -> dict:


    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            UPDATE import_batches
            SET status = 'cancelled',
                completed_at = NOW(),
                updated_at = NOW()
            WHERE import_id = %s AND vault_id = %s
              AND status IN ('pending', 'uploading')
            RETURNING import_id, vault_id, root_folder_name,
                      total_files, uploaded_count, failed_count,
                      skipped_duplicate_count, status,
                      total_bytes_planned, total_bytes_uploaded,
                      created_at, updated_at, completed_at
            """,
            (import_id, vault_id),
        )
        updated = cursor.fetchone()
        conn.commit()
    finally:
        conn.close()

    if updated is not None:
        return _serialize_import_batch(updated)

                                                                     
    existing = get_import_batch(vault_id=vault_id, import_id=import_id)
    if existing is None:
        raise HTTPException(
            status_code=404,
            detail={
                "code": "import_batch_not_found",
                "message": "Import batch not found for this vault.",
            },
        )
    return existing                                       


def complete_import_batch(
    *,
    vault_id: str,
    import_id: str,
    failed_count_delta: int = 0,
    skipped_duplicate_count_delta: int = 0,
) -> dict:


    existing = get_import_batch(vault_id=vault_id, import_id=import_id)
    if existing is None:
        raise HTTPException(
            status_code=404,
            detail={
                "code": "import_batch_not_found",
                "message": "Import batch not found for this vault.",
            },
        )
    if _is_terminal_import_status(existing.get("status")):
        return existing                             

    final_failed_count = max(
        0, int(existing.get("failed_count") or 0) + int(failed_count_delta),
    )
    final_status = _compute_complete_import_status(final_failed_count)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            UPDATE import_batches
            SET failed_count = failed_count + %s,
                skipped_duplicate_count =
                    skipped_duplicate_count + %s,
                status = %s,
                completed_at = NOW(),
                updated_at = NOW()
            WHERE import_id = %s AND vault_id = %s
              AND status IN ('pending', 'uploading')
            RETURNING import_id, vault_id, root_folder_name,
                      total_files, uploaded_count, failed_count,
                      skipped_duplicate_count, status,
                      total_bytes_planned, total_bytes_uploaded,
                      created_at, updated_at, completed_at
            """,
            (
                max(0, int(failed_count_delta)),
                max(0, int(skipped_duplicate_count_delta)),
                final_status,
                import_id,
                vault_id,
            ),
        )
        updated = cursor.fetchone()
        conn.commit()
    finally:
        conn.close()

    if updated is not None:
        return _serialize_import_batch(updated)
                                                                
                                                           
    return get_import_batch(vault_id=vault_id, import_id=import_id) or existing


def _serialize_import_batch(row: Optional[dict]) -> dict:


    if not row:
        return {}
    out: dict = {}
    for key, value in dict(row).items():
        if hasattr(value, "isoformat"):
            out[key] = value.isoformat()
        else:
            out[key] = str(value) if hasattr(value, "hex") else value
    return out


_RELATIVE_PATH_MAX_LEN = 1024


_EXTENSION_FILTER_GROUPS: dict[str, tuple[str, ...]] = {
    "zip": ("zip", "tar", "gz", "tgz", "7z", "rar"),
    "python": ("py",),
    "javascript": ("js", "mjs", "cjs"),
    "html": ("html", "htm"),
    "css": ("css",),
    "json": ("json",),
    "yaml": ("yaml", "yml"),
    "shell": ("sh", "bash", "zsh"),
    "powershell": ("ps1",),
    "pdf": ("pdf",),
    "image": ("jpg", "jpeg", "png", "gif", "heic", "webp", "bmp", "tiff"),
    "video": ("mp4", "mov", "webm", "mkv", "avi", "m4v"),
    "audio": ("mp3", "m4a", "wav", "aac", "ogg", "flac"),
    "spreadsheet": ("xlsx", "xls", "csv"),
    "document": ("docx", "doc", "pdf", "rtf"),
    "code": ("py", "js", "ts", "html", "css", "json", "sh", "ps1",
             "java", "rb", "go", "rs", "cpp", "c", "h"),
    "archive": ("zip", "tar", "gz", "tgz", "7z", "rar"),
    "presentation": ("pptx", "ppt"),
}

                                                                    
_EXTENSION_FILTER_TOKENS: dict[str, str] = {
                                       
    "zip": "zip",
    "zips": "zip",
    "archive": "archive",
    "archives": "archive",
    "python": "python",
    "py": "python",
    "javascript": "javascript",
    "js": "javascript",
    "html": "html",
    "css": "css",
    "json": "json",
    "yaml": "yaml",
    "yml": "yaml",
    "shell": "shell",
    "bash": "shell",
    "powershell": "powershell",
    "ps1": "powershell",
    "pdf": "pdf",
    "pdfs": "pdf",
    "image": "image",
    "images": "image",
    "photo": "image",
    "photos": "image",
    "video": "video",
    "videos": "video",
    "audio": "audio",
    "music": "audio",
    "voice": "audio",
    "spreadsheet": "spreadsheet",
    "spreadsheets": "spreadsheet",
    "excel": "spreadsheet",
    "csv": "spreadsheet",
    "code": "code",
    "script": "code",
    "scripts": "code",
    "document": "document",
    "documents": "document",
    "doc": "document",
    "docs": "document",
    "presentation": "presentation",
    "presentations": "presentation",
    "powerpoint": "presentation",
}


def _extract_extension_filter(query: Optional[str]) -> Optional[tuple[str, ...]]:


    group_key = _extract_extension_group_key(query)
    if group_key is None:
        return None
    return _EXTENSION_FILTER_GROUPS[group_key]


def _file_matches_extensions(
    file_name: Optional[str],
    saved_name: Optional[str],
    extensions: tuple[str, ...],
) -> bool:


    if not extensions:
        return True
    for candidate in (file_name, saved_name):
        if not candidate:
            continue
        lowered = candidate.lower()
        for ext in extensions:
            ext_lower = ext.lower().lstrip(".")
            if lowered.endswith("." + ext_lower):
                return True
    return False


def _is_under_folder(
    relative_path: Optional[str],
    folder_prefix: Optional[str],
) -> bool:


    if not folder_prefix:
        return True
    if not relative_path:
        return False
    prefix_segments = [
        s for s in folder_prefix.strip("/").lower().split("/") if s
    ]
    if not prefix_segments:
        return True
                                                                   
    rp_segments = [
        s for s in relative_path.lower().split("/") if s
    ][:-1]
    n = len(prefix_segments)
    if len(rp_segments) < n:
        return False
    for i in range(len(rp_segments) - n + 1):
        if rp_segments[i:i + n] == prefix_segments:
            return True
    return False


def _format_disambiguation_reply(
    candidates: list[dict],
    requested_name: Optional[str],
) -> str:


    n = len(candidates)
    if n <= 1:
        return ""                                    
    name_phrase = (
        f' named "{requested_name}"' if requested_name else ""
    )
    lines = [f"I found {n} files{name_phrase}. Which one do you want?"]
    for c in candidates[:10]:
        rp = (c.get("relative_path") or "").strip()
        name = (
            c.get("saved_name")
            or c.get("file_name")
            or "unknown"
        )
        if rp:
            lines.append(f"  • {rp}")
        else:
            lines.append(f"  • {name}")
    if n > 10:
        lines.append(f"  …and {n - 10} more")
    return "\n".join(lines)


def build_folder_tree(
    *,
    vault_id: str,
    path: Optional[str] = None,
) -> dict:


    sanitised_path = _sanitize_relative_path(path) if path else None
    at_root = not sanitised_path

    rows = list_uploaded_files(vault_id)

    folders_count: dict[str, int] = {}
    files_at_path: list[dict] = []

    if at_root:
        for row in rows:
            rp = (row.get("relative_path") or "").strip()
            if not rp:
                files_at_path.append(_serialize_file_row(row))
                continue
                                                        
            head, sep, _rest = rp.partition("/")
            if not sep:
                                                                   
                                                                     
                files_at_path.append(_serialize_file_row(row))
                continue
            folders_count[head] = folders_count.get(head, 0) + 1
    else:
                                                                    
                                                                 
        prefix_lower = sanitised_path.lower() + "/"
        for row in rows:
            rp = (row.get("relative_path") or "").strip()
            if not rp:
                continue
            if not rp.lower().startswith(prefix_lower):
                continue
            rest = rp[len(prefix_lower):]
            if not rest:
                continue
            head, sep, _rest = rest.partition("/")
            if not sep:
                                                                   
                files_at_path.append(_serialize_file_row(row))
            else:
                folders_count[head] = folders_count.get(head, 0) + 1

    folders = [
        {"name": name, "file_count": count}
        for name, count in sorted(
            folders_count.items(), key=lambda item: item[0].lower(),
        )
    ]

    return {
        "path": sanitised_path or "",
        "breadcrumbs": (
            sanitised_path.split("/") if sanitised_path else []
        ),
        "folders": folders,
        "files": files_at_path,
    }


def _serialize_file_row(row: dict) -> dict:


    out = {}
    for key, value in row.items():
        if hasattr(value, "isoformat"):
            out[key] = value.isoformat()
        else:
            out[key] = value
    return out


_FOLDER_HINT_PREPOSITIONS: tuple[str, ...] = (
    "inside",
    "in",
    "under",
    "from",
    "within",
)


def _list_all_folder_names(vault_id: str) -> list[str]:


    rows = list_uploaded_files(vault_id)
    seen: set[str] = set()
    for row in rows:
        rp = (row.get("relative_path") or "").strip()
        if not rp:
            continue
        segments = rp.split("/")
                                                    
        for seg in segments[:-1]:
            seg = seg.strip()
            if seg:
                seen.add(seg)
    return sorted(seen, key=len, reverse=True)


def _extract_folder_hint(
    query: Optional[str],
    *,
    known_folders: Optional[list[str]] = None,
) -> Optional[str]:


    if not query:
        return None
    if known_folders is None or not known_folders:
        return None
    lowered = query.lower()

                                                                   
    for folder_name in known_folders:
        fn_lower = folder_name.lower()
        for prep in _FOLDER_HINT_PREPOSITIONS:
                                                                   
                                                              
            import re as _re
            patterns = (
                _re.compile(
                    rf"\b{_re.escape(prep)}\s+{_re.escape(fn_lower)}\b"),
                _re.compile(
                    rf"\b{_re.escape(prep)}\s+my\s+{_re.escape(fn_lower)}\b"),
                _re.compile(
                    rf"\b{_re.escape(prep)}\s+the\s+{_re.escape(fn_lower)}\b"),
            )
            for pattern in patterns:
                if pattern.search(lowered):
                    return folder_name
    return None


_LIST_BROWSE_TOKENS: frozenset[str] = frozenset({
    "show",
    "list",
    "find",
    "send",
    "give",
    "browse",
    "open",
    "display",
    "view",
})


def _is_list_or_browse_phrasing(query: Optional[str]) -> bool:


    if not query:
        return False
    lowered = query.strip().lower()
    if not lowered:
        return False
                                                               
                                            
    PLURAL_TYPE_TOKENS = {
        "photos", "pictures", "images", "videos", "audio", "music",
        "voice", "recordings", "scripts", "code", "files", "pdfs",
        "documents", "docs", "spreadsheets", "presentations",
        "archives", "zips",
    }
    import re as _re
    tokens = set(_re.findall(r"[a-z]+", lowered))
    if tokens & PLURAL_TYPE_TOKENS:
        return True
    if tokens & _LIST_BROWSE_TOKENS:
                                                        
        if "all" in tokens or "every" in tokens:
            return True
                                                                    
                                          
        if tokens & PLURAL_TYPE_TOKENS:
            return True
                                                    
        if "files" in tokens:
            return True
    return False


_LITERAL_FILENAME_RE = None                       

def _literal_filename_pattern():
    global _LITERAL_FILENAME_RE
    if _LITERAL_FILENAME_RE is None:
        import re as _re
                                                               
                                                              
        _LITERAL_FILENAME_RE = _re.compile(
            r"\b([A-Za-z][A-Za-z0-9_.\-]{0,60}\.[A-Za-z]{1,8})\b"
        )
    return _LITERAL_FILENAME_RE


def _has_literal_filename(query: Optional[str]) -> bool:


    if not query:
        return False
    return bool(_literal_filename_pattern().search(query))


def _extract_literal_filename(query: Optional[str]) -> Optional[str]:


    if not query:
        return None
    match = _literal_filename_pattern().search(query)
    if not match:
        return None
    return match.group(1).strip()


_EXTENSION_GROUP_LABELS_BY_KEY: dict[str, str] = {
    "zip": "zip files",
    "python": "Python scripts",
    "javascript": "JavaScript files",
    "html": "HTML files",
    "css": "CSS files",
    "json": "JSON files",
    "yaml": "YAML files",
    "shell": "shell scripts",
    "powershell": "PowerShell scripts",
    "pdf": "PDFs",
    "image": "images",
    "video": "videos",
    "audio": "audio files",
    "spreadsheet": "spreadsheets",
    "document": "documents",
    "code": "code files",
    "archive": "archives",
    "presentation": "presentations",
}


_EXTENSION_FILTER_SKIP_PREPOSITIONS: frozenset[str] = frozenset({
    "inside", "in", "under", "from", "within",
})

                                                                   
_EXTENSION_FILTER_SKIP_FILLERS: frozenset[str] = frozenset({
    "my", "the",
})


def _extract_extension_group_key(query: Optional[str]) -> Optional[str]:


    if not query:
        return None
    lowered = query.strip().lower()
    if not lowered:
        return None
    import re as _re
    for match in _re.finditer(r"[a-z0-9]+", lowered):
        token = match.group(0)
        group = _EXTENSION_FILTER_TOKENS.get(token)
        if group is None:
            continue
        start = match.start()
        if start > 0 and lowered[start - 1] == ".":
                                                                 
                                             
            continue
                                                            
                                                                   
        prev_words = _re.findall(r"\w+", lowered[:start])
        if prev_words:
            last = prev_words[-1]
            if last in _EXTENSION_FILTER_SKIP_PREPOSITIONS:
                continue
            if (
                last in _EXTENSION_FILTER_SKIP_FILLERS
                and len(prev_words) >= 2
                and prev_words[-2] in _EXTENSION_FILTER_SKIP_PREPOSITIONS
            ):
                continue
        return group
    return None


def _extension_group_label(
    extensions: Optional[tuple[str, ...]],
    *,
    query: Optional[str] = None,
) -> str:


    if not extensions:
        return "files"
    if query:
        group_key = _extract_extension_group_key(query)
        if group_key is not None:
            return _EXTENSION_GROUP_LABELS_BY_KEY.get(group_key, "files")
    for key, ext in _EXTENSION_FILTER_GROUPS.items():
        if ext == extensions:
            return _EXTENSION_GROUP_LABELS_BY_KEY.get(key, "files")
    return "files"


def search_files_with_filters(
    vault_id: str,
    *,
    folder_filter: Optional[str] = None,
    extension_filter: Optional[tuple[str, ...]] = None,
    name_query: Optional[str] = None,
    limit: int = 20,
) -> list[dict]:


    rows = list_uploaded_files(vault_id)
    result: list[dict] = []
    name_q_lower = (name_query or "").strip().lower()

    for row in rows:
        if folder_filter and not _is_under_folder(
            row.get("relative_path"), folder_filter,
        ):
            continue
        if extension_filter and not _file_matches_extensions(
            row.get("file_name"),
            row.get("saved_name"),
            extension_filter,
        ):
            continue
        if name_q_lower:
            haystack = " ".join([
                str(row.get("saved_name") or ""),
                str(row.get("file_name") or ""),
            ]).lower()
            if name_q_lower not in haystack:
                continue
        result.append(_serialize_file_row(row))
        if len(result) >= limit:
            break
    return result


def _format_no_match_reply(
    folder_filter: Optional[str],
    extension_filter: Optional[tuple[str, ...]],
    *,
    requested_name: Optional[str] = None,
    query: Optional[str] = None,
) -> str:


    if extension_filter and not requested_name:
        label = _extension_group_label(extension_filter, query=query)
        if folder_filter:
            return (
                f"I couldn't find any {label} inside "
                f"{folder_filter}."
            )
        return f"I couldn't find any {label} in your vault."

    name_phrase = (
        f" {requested_name}" if requested_name else " that file"
    )
    if folder_filter:
        return (
            f"I couldn't find{name_phrase} inside {folder_filter}."
        )
    return f"I couldn't find{name_phrase} in your vault."


_FILE_LIST_CARD_CAP = 25


def _format_bytes_human(size: int) -> str:
    """Human file size — matches the frontend formatter so a row
    always carries a display-ready string in addition to raw bytes.
    Ranges: <1 KB = "N B", <1 MB = "N.N KB", <1 GB = "N.N MB",
    otherwise "N.NN GB"."""
    if size < 0:
        size = 0
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"
    if size < 1024 * 1024 * 1024:
        return f"{size / (1024 * 1024):.1f} MB"
    return f"{size / (1024 * 1024 * 1024):.2f} GB"


def _serialize_file_for_list_card(row: dict) -> dict:


    out = {
        "file_id": row.get("id"),
        "file_name": row.get("file_name"),
        "saved_name": row.get("saved_name"),
        "mime_type": row.get("content_type"),
        "asset_type": row.get("asset_type") or "file",
        # Downloadable is always true for tracked vault files. Kept
        # explicit so the frontend renders a Download button (distinct
        # from View) without inferring from mime_type.
        "downloadable": True,
    }
    rp = row.get("relative_path")
    if rp:
        out["relative_path"] = rp
    size = row.get("file_size")
    if isinstance(size, int):
        out["size_bytes"] = size
        out["size_display"] = _format_bytes_human(size)
    # Ship the upload timestamp so the frontend can render a
    # clean localized "uploaded on …" line.
    ts = row.get("uploaded_at") or row.get("created_at")
    if ts is not None:
        try:
            out["uploaded_at"] = (
                ts.isoformat() if hasattr(ts, "isoformat") else str(ts)
            )
        except Exception:
            pass
    return out


def _build_file_search_envelope(
    *,
    query: str,
    results: list[dict],
    message: str,
    pending_count: int = 0,
    stale_count: int = 0,
    pending_embedding_count: int = 0,
    semantic_available: bool = True,
) -> str:


    safe_results: list[dict] = []
    for r in results:
        if not isinstance(r, dict):
            continue
                                                                  
                                                                 
        sim = r.get("similarity")
        sim_out = None
        try:
            if sim is not None:
                sim_out = round(float(sim), 4)
        except Exception:
            sim_out = None
        safe_results.append({
            "file_id":       str(r.get("file_id") or ""),
            "file_name":     str(r.get("file_name") or ""),
            "saved_name":    r.get("saved_name") or None,
            "relative_path": r.get("relative_path") or None,
            "mime_type":     r.get("mime_type") or None,
            "asset_type":    r.get("asset_type") or "file",
            "match_type":    str(r.get("match_type") or "filename"),
            "match_reason":  str(r.get("match_reason") or ""),
            "match_confidence": (
                str(r.get("confidence") or "weak").lower()
            ),
                                                                   
                                                              
            "is_stale":      bool(r.get("is_stale") or False),
                                                                  
            "similarity":    sim_out,
                                                           
                                                           
            "is_archive_match": bool(r.get("is_archive_match") or False),
                                                               
                                                              
            "archive_match_details": (
                r.get("archive_match_details")
                if isinstance(r.get("archive_match_details"), dict)
                else None
            ),
            "purpose":       r.get("purpose") or None,
            "purpose_label": r.get("purpose_label") or None,
        })
                                                               
                                                
    try:
        from vault_chat_result_cards import (
            SCHEMA_VERSION as _FILE_SEARCH_SCHEMA_VERSION,
            COPY_VERSION as _FILE_SEARCH_COPY_VERSION,
        )
    except Exception:
        _FILE_SEARCH_SCHEMA_VERSION = "file_search_results.v2"
        _FILE_SEARCH_COPY_VERSION = "sharp_empty_state_2026_06_28"
    payload = {
        "type":            "file_search_results",
        "schema_version":  _FILE_SEARCH_SCHEMA_VERSION,
        "copy_version":    _FILE_SEARCH_COPY_VERSION,
        "query":           str(query or ""),
        "count":           len(safe_results),
        "results":         safe_results,
        "message":         str(message or ""),
        "pending_count":   int(pending_count or 0),
        "stale_count":     int(stale_count or 0),
                                                              
        "pending_embedding_count": int(pending_embedding_count or 0),
        "semantic_available":      bool(semantic_available),
                                                                  
                                                            
        "is_complete":     True,
    }
    return json.dumps(payload)


_RELATIONSHIP_TYPE_PRIORITY: dict[str, int] = {
    "duplicate":               0,
    "front_back_pair":         1,
    "near_duplicate":          2,
    "same_person":             3,
    "same_financial_account":  4,
    "same_company":            5,
    "same_trip":               6,
    "supporting_document":     7,
    "archive_contains_signal": 8,
    "same_document_family":    9,
    "semantic_related":       10,
    "same_import_batch":      11,
    "same_folder":            12,
}


def _confidence_label_for(value: float) -> str:


    try:
        v = float(value)
    except Exception:
        return "weak"
    if v >= 0.75:
        return "strong"
    if v >= 0.50:
        return "medium"
    return "weak"


def _safe_relationship_evidence(value) -> dict:


    allowed = {
        "shared_entity_names",
        "shared_topics",
        "shared_categories",
        "shared_doc_types",
        "embedding_cosine",
        "filename_pattern",
        "content_sha256_match",
        "import_id",
        "folder_path",
        "shared_email_domains",
        "archive_file_id",
        "inner_file_count",
    }
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except Exception:
            return {}
    if not isinstance(value, dict):
        return {}
    out: dict[str, Any] = {}
    for k, v in value.items():
        if k in allowed:
            out[k] = v
    return out


def _sort_relationship_rows(
    rows: list[dict], anchor_id: str,
) -> list[dict]:


    def _key(r: dict) -> tuple:
        confidence = float(r.get("confidence") or 0.0)
        type_pri = _RELATIONSHIP_TYPE_PRIORITY.get(
            str(r.get("relationship_type") or ""), 99,
        )
        updated = r.get("updated_at")
        try:
            updated_ts = updated.timestamp() if updated else 0.0
        except Exception:
            updated_ts = 0.0
                                                              
                                                       
        return (-confidence, type_pri, -updated_ts)
    return sorted(rows, key=_key)


def _build_related_files_graph_envelope(
    *,
    anchor: dict,
    rows: list[dict],
    file_metadata_by_id: dict,
    message: str,
) -> str:


    safe_anchor = _safe_file_object(anchor)
    safe_rels: list[dict] = []
    for r in rows:
        if not isinstance(r, dict):
            continue
                                                
        a_id = str(r.get("file_a_id") or "")
        b_id = str(r.get("file_b_id") or "")
        anchor_id = str(safe_anchor.get("file_id") or "")
        other_id = b_id if a_id == anchor_id else a_id
                                                           
                                                            
        if other_id not in file_metadata_by_id:
            continue
        other_meta = file_metadata_by_id[other_id]
        other_obj = _safe_file_object(other_meta, fallback_id=other_id)
        if not other_obj.get("file_id"):
            continue
        reasons = r.get("reasons_jsonb")
        if isinstance(reasons, str):
            try:
                reasons = json.loads(reasons)
            except Exception:
                reasons = []
        if not isinstance(reasons, list):
            reasons = []
                                         
        safe_reasons = [
            str(x).strip() for x in reasons
            if isinstance(x, str) and str(x).strip()
        ][:4]
        confidence = float(r.get("confidence") or 0.0)
        safe_rels.append({
            "file":              other_obj,
            "relationship_type": str(r.get("relationship_type") or ""),
            "confidence":        round(confidence, 4),
            "confidence_label":  _confidence_label_for(confidence),
            "reasons":           safe_reasons,
            "evidence":          _safe_relationship_evidence(
                r.get("evidence_jsonb"),
            ),
        })

    payload = {
        "type":          "related_files_graph",
        "anchor":        safe_anchor,
        "count":         len(safe_rels),
        "message":       str(message or ""),
        "relationships": safe_rels,
    }
    return json.dumps(payload)


def _load_safe_file_metadata(
    vault_id: str, file_ids: list[str],
) -> dict:


    if not file_ids:
        return {}
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT id::text AS file_id,
                   file_name,
                   saved_name,
                   relative_path,
                   content_type,
                   asset_type
            FROM uploaded_files
            WHERE vault_id = %s
              AND id = ANY(%s::uuid[])
            """,
            (vault_id, list(file_ids)),
        )
        rows = list(cur.fetchall() or [])
    finally:
        conn.close()
    out: dict[str, dict] = {}
    for r in rows:
        fid = str(r.get("file_id") or "")
        if fid:
            out[fid] = dict(r)
    return out


def _load_safe_anchor_row(
    vault_id: str, file_id: str,
) -> Optional[dict]:


    if not vault_id or not file_id:
        return None
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT id::text AS id,
                   file_name,
                   saved_name,
                   relative_path,
                   content_type,
                   asset_type
            FROM uploaded_files
            WHERE vault_id = %s
              AND id = %s
            """,
            (vault_id, file_id),
        )
        row = cur.fetchone()
    finally:
        conn.close()
    return dict(row) if row else None


def _compose_related_files_envelope(
    *, vault_id: str, anchor_row: dict,
) -> dict:


    from vault_relationship_graph import get_relationships_for_file

    anchor_id = str(anchor_row.get("id") or "")
    if not anchor_id:
        return {
            "type":          "related_files_graph",
            "anchor":        {},
            "count":         0,
            "message":       "",
            "relationships": [],
        }

    rows = get_relationships_for_file(
        vault_id, anchor_id, limit=200,
    )
    sorted_rows = _sort_relationship_rows(rows, anchor_id)

    other_ids: list[str] = []
    seen: set[str] = set()
    for r in sorted_rows:
        a_id = str(r.get("file_a_id") or "")
        b_id = str(r.get("file_b_id") or "")
        other = b_id if a_id == anchor_id else a_id
        if other and other not in seen:
            seen.add(other)
            other_ids.append(other)
    file_metadata_by_id = _load_safe_file_metadata(
        vault_id, other_ids,
    )

    safe_anchor = {
        "file_id":       anchor_id,
        "file_name":     anchor_row.get("file_name") or "",
        "saved_name":    anchor_row.get("saved_name") or None,
        "relative_path": anchor_row.get("relative_path") or None,
        "mime_type":     anchor_row.get("content_type") or None,
        "asset_type":    anchor_row.get("asset_type") or "file",
    }

    anchor_label = (
        safe_anchor["saved_name"]
        or safe_anchor["file_name"]
        or "this file"
    )

    kept_rows: list[dict] = []
    for r in sorted_rows:
        a_id = str(r.get("file_a_id") or "")
        b_id = str(r.get("file_b_id") or "")
        other = b_id if a_id == anchor_id else a_id
        if other in file_metadata_by_id:
            kept_rows.append(r)
    n = len(kept_rows)

    if n == 0:
        message = (
            f"I don't see strong related files for "
            f"\"{anchor_label}\" yet. As more files are "
            "analyzed, related items may appear."
        )
    else:
        message = (
            f"I found {n} file{'s' if n != 1 else ''} "
            f"related to \"{anchor_label}\"."
        )

    env_json = _build_related_files_graph_envelope(
        anchor=safe_anchor,
        rows=kept_rows,
        file_metadata_by_id=file_metadata_by_id,
        message=message,
    )
    return json.loads(env_json)


_CLUSTER_TYPE_PRIORITY: dict[str, int] = {
    "identity":        0,
    "travel":          1,
    "finance":         2,
    "tax":             3,
    "application":     4,
    "company":         5,
    "duplicates":      6,
    "archive_content": 7,
    "same_person":     8,
    "same_folder":     9,
}


_CLUSTER_TITLE_BY_TYPE: dict[str, str] = {
    "identity":        "Identity documents",
    "travel":          "Travel documents",
    "finance":         "Financial documents",
    "tax":             "Tax documents",
    "application":     "Application + supporting documents",
    "company":         "Company / employer documents",
    "duplicates":      "Duplicate files",
    "archive_content": "Archive contents",
    "same_person":     "Same person",
    "same_folder":     "Same folder",
}


def _select_cluster_type(edge_types: set[str]) -> str:


    has_person = "same_person" in edge_types
    has_front_back = "front_back_pair" in edge_types
    has_supporting = "supporting_document" in edge_types
    if has_person and (has_front_back or has_supporting):
        return "identity"
    if "same_trip" in edge_types:
        return "travel"
    if "same_financial_account" in edge_types:
        return "finance"
    if "same_company" in edge_types and has_supporting:
        return "tax"
    if has_supporting:
        return "application"
    if "same_company" in edge_types:
        return "company"
    if "duplicate" in edge_types or "near_duplicate" in edge_types:
        return "duplicates"
    if "archive_contains_signal" in edge_types:
        return "archive_content"
    if has_person or has_front_back:
        return "same_person"
    return "same_folder"


def _cluster_relationship_edges(
    edges: list[dict],
) -> list[dict]:


    if not edges:
        return []

                                       
    parent: dict[str, str] = {}

    def _find(x: str) -> str:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def _union(a: str, b: str) -> None:
        ra, rb = _find(a), _find(b)
        if ra != rb:
            parent[ra] = rb

    safe_edges: list[dict] = []
    for e in edges:
        if not isinstance(e, dict):
            continue
        a = str(e.get("file_a_id") or "").strip()
        b = str(e.get("file_b_id") or "").strip()
        if not a or not b or a == b:
            continue
        if a not in parent:
            parent[a] = a
        if b not in parent:
            parent[b] = b
        _union(a, b)
        safe_edges.append(e)

                             
    groups: dict[str, list[dict]] = {}
    for e in safe_edges:
        root = _find(str(e["file_a_id"]))
        groups.setdefault(root, []).append(e)

    clusters: list[dict] = []
    for root, group_edges in groups.items():
        file_ids: set[str] = set()
        edge_types: set[str] = set()
        reasons: list[str] = []
        seen_reasons: set[str] = set()
        entity_names: list[str] = []
        seen_entities: set[str] = set()
        strong_count = 0
        max_conf = 0.0
        latest_ts = 0.0
        for e in group_edges:
            file_ids.add(str(e["file_a_id"]))
            file_ids.add(str(e["file_b_id"]))
            rt = str(e.get("relationship_type") or "")
            if rt:
                edge_types.add(rt)
            try:
                c = float(e.get("confidence") or 0.0)
            except Exception:
                c = 0.0
            if c >= 0.75:
                strong_count += 1
            if c > max_conf:
                max_conf = c
                                                           
            raw_reasons = e.get("reasons_jsonb")
            if isinstance(raw_reasons, str):
                try:
                    raw_reasons = json.loads(raw_reasons)
                except Exception:
                    raw_reasons = []
            if isinstance(raw_reasons, list):
                for r in raw_reasons:
                    if isinstance(r, str):
                        s = r.strip()
                        if s and s not in seen_reasons:
                            seen_reasons.add(s)
                            reasons.append(s)
                                                           
                                                               
            raw_ev = e.get("evidence_jsonb")
            if isinstance(raw_ev, str):
                try:
                    raw_ev = json.loads(raw_ev)
                except Exception:
                    raw_ev = {}
            ev = _safe_relationship_evidence(raw_ev)
            for name in (ev.get("shared_entity_names") or []):
                s = str(name).strip()
                if s and s not in seen_entities:
                    seen_entities.add(s)
                    entity_names.append(s)
                      
            ts = e.get("updated_at")
            try:
                if ts is not None:
                    latest_ts = max(latest_ts, ts.timestamp())
            except Exception:
                pass

        cluster_type = _select_cluster_type(edge_types)
        clusters.append({
            "file_ids":                 sorted(file_ids),
            "edge_types":               sorted(edge_types),
            "relationship_count":       len(group_edges),
            "strong_relationship_count": strong_count,
            "max_confidence":           round(max_conf, 4),
            "latest_updated_ts":        latest_ts,
            "main_reasons":             reasons[:3],
            "cluster_type":             cluster_type,
            "entity_names":             entity_names[:3],
        })
    return clusters


def _filter_meaningful_clusters(clusters: list[dict]) -> list[dict]:


    out: list[dict] = []
    for c in clusters:
        edge_types = set(c.get("edge_types") or [])
        if edge_types == {"same_folder"}:
            continue
        out.append(c)
    return out


def _sort_clusters(clusters: list[dict]) -> list[dict]:


    def _key(c: dict) -> tuple:
        strong = int(c.get("strong_relationship_count") or 0)
        files = len(c.get("file_ids") or [])
        type_pri = _CLUSTER_TYPE_PRIORITY.get(
            str(c.get("cluster_type") or ""), 99,
        )
        ts = float(c.get("latest_updated_ts") or 0.0)
        return (-strong, -files, type_pri, -ts)
    return sorted(clusters, key=_key)


def _confidence_label_for_cluster(c: dict) -> str:


    return _confidence_label_for(float(c.get("max_confidence") or 0.0))


def _cluster_title(cluster: dict) -> str:


    cluster_type = str(cluster.get("cluster_type") or "same_folder")
    base = _CLUSTER_TITLE_BY_TYPE.get(cluster_type, "Connected files")
    names = cluster.get("entity_names") or []
    if not names:
        return base
    primary = str(names[0]).strip()
    if not primary:
        return base
    if cluster_type == "identity":
        return f"{primary} identity documents"
    if cluster_type == "same_person":
        return f"{primary} documents"
    if cluster_type == "company":
        return f"{primary} documents"
    if cluster_type == "tax":
        return f"{primary} tax documents"
    return base


def _build_vault_clusters_envelope_dict(
    *,
    clusters: list[dict],
    file_metadata_by_id: dict,
    message: str,
    has_pending_analysis: bool = False,
) -> dict:


    cluster_files_cap = 50
    safe_rows: list[dict] = []
    for i, c in enumerate(clusters):
        file_ids = list(c.get("file_ids") or [])
                                                         
                                                            
        kept_ids = [fid for fid in file_ids if fid in file_metadata_by_id]
        if not kept_ids:
            continue
        reps: list[dict] = []
        for fid in kept_ids[:cluster_files_cap]:
            row = file_metadata_by_id.get(fid)
            obj = _safe_file_object(row, fallback_id=fid)
            if obj.get("file_id"):
                reps.append(obj)
        if not reps:
            continue
        warnings: list[str] = []
        if has_pending_analysis:
            warnings.append(
                "Some connections may still be refreshing.",
            )
        safe_rows.append({
            "cluster_id":               f"cluster-{i + 1}",
            "cluster_type":             str(c.get("cluster_type") or ""),
            "title":                    _cluster_title(c),
            "confidence":               _confidence_label_for_cluster(c),
            "file_count":               len(kept_ids),
            "relationship_count":       int(c.get("relationship_count") or 0),
            "strong_relationship_count":
                int(c.get("strong_relationship_count") or 0),
            "main_reasons":             list(c.get("main_reasons") or [])[:3],
            "representative_files":     reps,
            "related_file_ids":         kept_ids,
            "warnings":                 warnings,
        })

    return {
        "type":     "vault_relationship_clusters",
        "message":  str(message or ""),
        "count":    len(safe_rows),
        "clusters": safe_rows,
    }


def _load_all_relationships_for_vault(
    vault_id: str, *, limit: int = 5000,
) -> list[dict]:


    if not vault_id:
        return []
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT file_a_id::text AS file_a_id,
                   file_b_id::text AS file_b_id,
                   relationship_type,
                   confidence,
                   reasons_jsonb,
                   evidence_jsonb,
                   updated_at
            FROM vault_file_relationships
            WHERE vault_id = %s
            ORDER BY confidence DESC, updated_at DESC
            LIMIT %s
            """,
            (vault_id, int(limit)),
        )
        return list(cur.fetchall() or [])
    finally:
        conn.close()


def _compose_vault_clusters_envelope(vault_id: str) -> dict:


    edges = _load_all_relationships_for_vault(vault_id)
    raw_clusters = _cluster_relationship_edges(edges)
    meaningful = _filter_meaningful_clusters(raw_clusters)
    ranked = _sort_clusters(meaningful)

                                                              
    all_ids: list[str] = []
    seen: set[str] = set()
    for c in ranked:
        for fid in c.get("file_ids") or []:
            if fid and fid not in seen:
                seen.add(fid)
                all_ids.append(fid)
    file_metadata_by_id = _load_safe_file_metadata(vault_id, all_ids)

                                                          
    has_pending = False
    try:
        conn = get_db()
        try:
            cur = conn.cursor()
            cur.execute(
                """
                SELECT 1 FROM vault_analysis_jobs
                WHERE vault_id = %s
                  AND status IN ('pending', 'processing')
                LIMIT 1
                """,
                (vault_id,),
            )
            has_pending = bool(cur.fetchone())
        finally:
            conn.close()
    except Exception:
        has_pending = False

    envelope = _build_vault_clusters_envelope_dict(
        clusters=ranked,
        file_metadata_by_id=file_metadata_by_id,
        message="",
        has_pending_analysis=has_pending,
    )

    count = envelope["count"]
    if count == 0:
        if has_pending:
            envelope["message"] = (
                "I don't see connected groups yet — your "
                "vault analysis is still in progress. Try "
                "again in a moment."
            )
        else:
            envelope["message"] = (
                "I don't see connected groups in your vault "
                "yet. As more files are analyzed, document "
                "groups may appear."
            )
    else:
        plural = "s" if count != 1 else ""
        envelope["message"] = (
            f"I found {count} connected group{plural} in your vault."
        )
    return envelope


def _safe_file_object(
    row: dict, *, fallback_id: Optional[str] = None,
) -> dict:


    if not isinstance(row, dict):
        return {}
    file_id = str(
        row.get("file_id") or row.get("id") or fallback_id or "",
    )
    if not file_id:
        return {}
    return {
        "file_id":       file_id,
        "file_name":     str(row.get("file_name") or ""),
        "saved_name":    row.get("saved_name") or None,
        "relative_path": row.get("relative_path") or None,
        "mime_type":     row.get("mime_type")
                         or row.get("content_type") or None,
        "asset_type":    row.get("asset_type") or "file",
    }


def _sort_files_for_pagination(rows: list[dict]) -> list[dict]:
    """Stable order for chat file listings: newest-first, tie-break on
    id. Guarantees repeated "show me all my files" and subsequent
    "show more" pages never skip or duplicate a row, even if two
    files share a created_at timestamp.
    """
    def _key(r: dict):
        ts = r.get("created_at")
        try:
            iso = ts.isoformat() if hasattr(ts, "isoformat") else str(ts or "")
        except Exception:
            iso = ""
        return (iso, str(r.get("id") or ""))

    return sorted(rows, key=_key, reverse=True)


def _build_vault_file_list_envelope(
    rows: list[dict],
    title: str,
    *,
    message: Optional[str] = None,
    requested_name: Optional[str] = None,
    offset: int = 0,
    page_size: Optional[int] = None,
) -> str:
    """Build a paginated vault_file_list envelope.

    ``rows`` is the FULL, stably-ordered result set (already ordered
    by `list_uploaded_files` — newest-first). ``offset`` and
    ``page_size`` produce a deterministic slice; the envelope carries
    ``offset``, ``page_size``, ``next_offset`` and ``has_more`` so the
    frontend can render a real "Show more" button and re-issue a
    follow-up chat message.

    Stable ordering is enforced by ``_sort_files_for_pagination`` —
    the same input list always yields the same page N.
    """
    size = int(page_size or _FILE_LIST_CARD_CAP)
    if size <= 0:
        size = _FILE_LIST_CARD_CAP
    start = max(0, int(offset or 0))

    ordered = _sort_files_for_pagination(rows)
    total = len(ordered)
    end = min(total, start + size)
    page_rows = ordered[start:end]
    files = [_serialize_file_for_list_card(r) for r in page_rows]

    next_offset = end
    has_more = end < total
    more_count = max(0, total - end)

    payload: dict[str, object] = {
        "type": "vault_file_list",
        "title": title,
        "count": len(page_rows),
        "total_count": total,
        "offset": start,
        "page_size": size,
        "next_offset": next_offset,
        "has_more": has_more,
        "files": files,
    }
    if message:
        payload["message"] = message
    if more_count > 0:
        payload["more_count"] = more_count
    if requested_name:
        payload["requested_name"] = requested_name
    return json.dumps(payload)


def _build_vault_inventory_envelope(
    *,
    summary: dict,
    message: str,
    analysis_coverage: Optional[dict] = None,
) -> str:


    recent = summary.get("recent_files") or []
    payload: dict[str, object] = {
        "type":         "vault_inventory",
        "message":      message,
        "total_files":  int(summary.get("total_files") or 0),
        "total_bytes":  int(summary.get("total_bytes") or 0),
        "folder_count": int(summary.get("folder_count") or 0),
        "top_folders":  list(summary.get("top_folders") or []),
        "type_counts":  dict(summary.get("type_counts") or {}),
        "recent_files": list(recent),
    }
    if analysis_coverage is not None:
        payload["analysis_coverage"] = dict(analysis_coverage)
    return json.dumps(payload)


def _build_travel_readiness_envelope(
    *, report: dict, message: str,
) -> str:


    payload = {
        "type":          "travel_readiness",
        "message":       message,
        "confidence":    str(report.get("confidence") or "blocked"),
        "found":         list(report.get("found") or []),
        "missing":       list(report.get("missing") or []),
        "expired":       list(report.get("expired") or []),
        "expiring_soon": list(report.get("expiring_soon") or []),
    }
    return json.dumps(payload)


def _build_credential_files_envelope(
    *,
    matches: list[dict],
    message: str,
    scanned_count: int = 0,
    not_scanned_count: int = 0,
    is_partial: bool = False,
    sections: Optional[dict] = None,                                     
    has_content_matches: Optional[bool] = None,                            
) -> str:


    actions: list[dict] = []
    if int(not_scanned_count) > 0:
        actions.append({
            "type":       "scan_remaining",
            "label":      "Scan remaining files",
            "file_count": int(not_scanned_count),
        })
    payload: dict[str, object] = {
        "type":               "credential_files",
        "message":            message,
        "count":              len(matches),
        "files":              list(matches),
        "scanned_count":      int(scanned_count),
        "not_scanned_count":  int(not_scanned_count),
        "is_partial":         bool(is_partial),
        "actions":            actions,
    }
    return json.dumps(payload)


def _maybe_build_deep_answer_envelope(
    *,
    vault_id: str,
    intent: str,
    query: str,
    key: Optional[bytes],
    coverage_threshold: float = 0.95,
) -> Optional[str]:


    try:
        from vault_deep_answer import (
            is_deep_intent,
            start_or_resume_job,
            step_deep_answer,
            safe_job_snapshot,
            build_coverage_report,
        )
        from vault_analysis import analysis_coverage_for_vault
        from routes.deep_answer_routes import (
            _coverage_loader as _da_coverage_loader,
            _drain_fns as _da_drain_fns,
            _final_results_fn as _da_final_results_fn,
            _enqueue_missing_fn as _da_enqueue_missing_fn,
        )
    except Exception:
        logger.exception("deep-answer gating import failed")
        return None

    if not is_deep_intent(intent):
        return None
    if not key:
                                                                  
                                                                    
        return None

    raw_coverage = analysis_coverage_for_vault(vault_id) or {}
    coverage = build_coverage_report(raw_coverage)
    total = int(coverage.get("total") or 0)
    if total == 0 or coverage.get("scan_complete"):
        return None
    scanned = int(coverage.get("scanned") or 0)
    fraction = (scanned / total) if total else 1.0
    if fraction >= coverage_threshold:
        return None

    try:
        job = start_or_resume_job(
            vault_id=vault_id,
            intent=intent,
            query=query or "",
        )
        step_deep_answer(
            job=job,
            vault_id=vault_id,
            key=key,
            coverage_loader=_da_coverage_loader(vault_id),
            drain_fns=_da_drain_fns(),
            final_results_fn=_da_final_results_fn,
            enqueue_missing_fn=_da_enqueue_missing_fn,
        )
        snapshot = safe_job_snapshot(job)
                                                              
                                                             
        intent_headline = {
            "search_files_for_credentials":
                "Scanning your vault for files that contain saved credentials",
            "search_files_about":
                "Scanning your vault for files about this topic",
            "list_by_tag":
                "Scanning your vault for files matching this category",
            "travel_readiness":
                "Scanning your vault for travel documents",
            "related_files":
                "Scanning your vault for related files",
            "vault_clusters":
                "Scanning your vault for connected groups of files",
        }.get(intent, "Scanning your vault")
        if scanned == 0:
            user_message = f"{intent_headline}… Preparing scan…"
        else:
            user_message = (
                f"{intent_headline}… "
                f"{scanned} of {total} files read so far."
            )
        payload = {
            "type":     "deep_answer_progress",
            "message":  user_message,
            "job_id":   snapshot.get("job_id"),
            "intent":   snapshot.get("intent"),
            "status":   snapshot.get("status"),
            "progress": snapshot.get("progress") or {},
            "results":  snapshot.get("results"),
        }
        return json.dumps(payload)
    except Exception:
        logger.exception(
            "deep-answer gating failed vault=%s intent=%s",
            vault_id, intent,
        )
        return None


def _build_credential_extraction_review_envelope(
    *,
    file_id: Optional[str],
    file_name: Optional[str],
    saved_name: Optional[str],
    relative_path: Optional[str],
    records: list[dict],
    message: str,
    text_available: bool,
) -> str:


    safe_records: list[dict] = []
    for raw in records[:50]:
        if not isinstance(raw, dict):
            continue
        fields = raw.get("fields") if isinstance(raw.get("fields"), dict) else {}
        service = (raw.get("service") or "general").strip() or "general"
        username = fields.get("username") if isinstance(fields, dict) else None
        email = fields.get("email") if isinstance(fields, dict) else None
        if isinstance(username, str):
            username = username.strip() or None
        else:
            username = None
        if isinstance(email, str):
            email = email.strip() or None
        else:
            email = None
        safe_records.append({
            "service":           service,
            "username":          username,
            "email":             email,
            "password_present":  bool(fields.get("password")) if isinstance(fields, dict) else False,
            "pin_present":       bool(fields.get("pin")) if isinstance(fields, dict) else False,
            "note_present":      bool(fields.get("note")) if isinstance(fields, dict) else False,
        })
    payload = {
        "type":           "credential_extraction_review",
        "message":        message,
        "file": {
            "file_id":       str(file_id) if file_id else None,
            "file_name":     file_name,
            "saved_name":    saved_name,
            "relative_path": relative_path,
        },
        "count":          len(safe_records),
        "records":        safe_records,
        "text_available": bool(text_available),
    }
    return json.dumps(payload)


def _handle_extract_logins_from_file(
    *,
    vault_id: str,
    key: Optional[bytes],
    asset_name: Optional[str],
) -> str:


    file_row = _resolve_file_for_analysis(
        vault_id,
        asset_name=asset_name,
        decrypted_message=asset_name or "",
        key=key,
    )
    if not file_row:
        return _build_credential_extraction_review_envelope(
            file_id=None,
            file_name=None,
            saved_name=None,
            relative_path=None,
            records=[],
            message=(
                "I couldn't find a file to extract from. Open the "
                "file first (or ask me to find it), then ask again."
            ),
            text_available=False,
        )

    plaintext = file_row.get("extracted_text")
    file_id = str(file_row.get("id") or "")
    file_name = file_row.get("file_name")
    saved_name = file_row.get("saved_name")
    relative_path = file_row.get("relative_path")

    if not plaintext:
        return _build_credential_extraction_review_envelope(
            file_id=file_id,
            file_name=file_name,
            saved_name=saved_name,
            relative_path=relative_path,
            records=[],
            message=(
                f"I haven't extracted the text from "
                f"{saved_name or file_name or 'this file'} yet. Run "
                "vault analysis to scan the remaining files, then ask "
                "again."
            ),
            text_available=False,
        )

    try:
        from extractor import extract_multiple_credentials
        records = extract_multiple_credentials(plaintext) or []
    except Exception:
        logger.exception("extract_multiple_credentials failed")
        records = []

    if not records:
        return _build_credential_extraction_review_envelope(
            file_id=file_id,
            file_name=file_name,
            saved_name=saved_name,
            relative_path=relative_path,
            records=[],
            message=(
                f"I scanned {saved_name or file_name or 'that file'} "
                "but couldn't pick out distinct login records. The "
                "file may use a layout I don't recognize — open it to "
                "verify."
            ),
            text_available=True,
        )

    label = saved_name or file_name or "this file"
    n = len(records)
    message = (
        f"I found {n} login record{'s' if n != 1 else ''} in {label}. "
        "Review the list below — nothing is saved until you confirm."
    )
    return _build_credential_extraction_review_envelope(
        file_id=file_id,
        file_name=file_name,
        saved_name=saved_name,
        relative_path=relative_path,
        records=records,
        message=message,
        text_available=True,
    )


def _build_file_disambiguation_envelope(
    *,
    files: list[dict],
    title: str,
    message: str,
    context_kind: str,
) -> str:


    safe_files: list[dict] = []
    for raw in files:
        if not isinstance(raw, dict):
            continue
        raw_reasons = raw.get("reasons")
        if isinstance(raw_reasons, list):
            safe_reasons = [
                str(r) for r in raw_reasons
                if isinstance(r, str) and r.strip()
            ]
        else:
            safe_reasons = []
        entry: dict = {
            "file_id":       str(raw.get("file_id") or ""),
            "file_name":     str(raw.get("file_name") or ""),
            "saved_name":    raw.get("saved_name") or None,
            "relative_path": raw.get("relative_path") or None,
            "mime_type":     raw.get("mime_type") or None,
            "asset_type":    raw.get("asset_type") or "file",
            "confidence":    (raw.get("confidence") or "").lower() or None,
            "reasons":       safe_reasons,
        }
                                                                 
                                                                
        if raw.get("best_match") is True:
            entry["best_match"] = True
        density = raw.get("credential_density")
        if isinstance(density, (int, float)) and density > 0:
            entry["credential_density"] = round(float(density), 4)
        blocks = raw.get("credential_block_count")
        if isinstance(blocks, int) and blocks > 0:
            entry["credential_block_count"] = blocks
        services = raw.get("service_count")
        if isinstance(services, int) and services > 0:
            entry["service_count"] = services
        if raw.get("mostly_credentials") is True:
            entry["mostly_credentials"] = True
                                                                  
                                                                 
        purpose = raw.get("purpose")
        if isinstance(purpose, str) and purpose:
            entry["purpose"] = purpose
        purpose_label = raw.get("purpose_label")
        if isinstance(purpose_label, str) and purpose_label.strip():
            entry["purpose_label"] = purpose_label.strip()
        safe_files.append(entry)

    payload = {
        "type":         "file_disambiguation",
        "title":        title or "Which file do you mean?",
        "message":      message or "",
        "context_kind": str(context_kind or ""),
        "count":        len(safe_files),
        "files":        safe_files,
    }
    return json.dumps(payload)


def _build_vault_brain_answer_envelope(
    *,
    message: str,
    intent: str,
    evidence_rows: list[dict],
    no_evidence: bool,
    coverage_note: str,
    retrieval_mode: str,
    breadth: str = "",
    brain_coverage: Optional[dict] = None,
    continuation_available: bool = False,
) -> str:


    safe_rows: list[dict] = []
    for row in evidence_rows or []:
        if not isinstance(row, dict):
            continue
        snippet = str(row.get("snippet") or "")
                                                                  
                                                                     
        if len(snippet) > 600:
            snippet = snippet[:600].rstrip() + "…"
        try:
            score = float(row.get("score") or 0.0)
        except (TypeError, ValueError):
            score = 0.0
        try:
            chunk_index = int(row.get("chunk_index") or 0)
        except (TypeError, ValueError):
            chunk_index = 0
        safe_rows.append({
            "file_id":           str(row.get("file_id") or ""),
            "file_name":         str(row.get("file_name") or ""),
            "snippet":           snippet,
            "extraction_source": str(row.get("extraction_source") or ""),
            "score":             round(score, 4),
            "chunk_index":       chunk_index,
        })
                                                              
                                                                     
    safe_brain_coverage: dict = {}
    if isinstance(brain_coverage, dict):
                                                                
                                                               
        ALLOWED_KEYS = {
            "vault_id", "total_files", "files_with_extracted_text",
            "files_with_chunks", "files_with_any_chunks",
            "files_with_embedded_chunks",
            "files_with_all_chunks_embedded",
            "files_with_partial_embeddings",
            "files_missing_chunks", "files_missing_embeddings",
            "chunks_total", "chunks_embedded", "chunks_unembedded",
            "pending_chunking_jobs", "pending_embedding_jobs",
            "failed_chunking_jobs", "failed_embedding_jobs",
            "unsupported_files", "coverage_percentage",
            "chunk_coverage_percentage",
            "is_complete", "has_failures",
        }
        for k, v in brain_coverage.items():
            if k in ALLOWED_KEYS:
                safe_brain_coverage[str(k)] = v
    payload = {
        "type":                   "vault_brain_answer",
        "message":                str(message or ""),
        "intent":                 str(intent or ""),
        "no_evidence":            bool(no_evidence),
        "coverage_note":          str(coverage_note or ""),
        "retrieval_mode":         str(retrieval_mode or ""),
        "breadth":                str(breadth or ""),
        "coverage":               safe_brain_coverage,
        "continuation_available": bool(continuation_available),
        "count":                  len(safe_rows),
        "evidence":               safe_rows,
    }
    return json.dumps(payload)


def _build_related_files_envelope(
    *,
    anchor: dict,
    results: list[dict],
    message: str,
) -> str:


    payload = {
        "type":    "related_files",
        "message": message,
        "anchor":  dict(anchor),
        "count":   len(results),
        "results": list(results),
    }
    return json.dumps(payload)


def _list_uploaded_files_for_credential_search(
    vault_id: str, key: Optional[bytes],
) -> list[dict]:


    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT id, file_name, content_type, file_size, detected_type,
                   detected_service, saved_name, asset_type, needs_naming,
                   created_at, relative_path, content_sha256,
                   extracted_text, extracted_text_encrypted
            FROM uploaded_files
            WHERE vault_id = %s
              AND upload_status = 'complete'
            ORDER BY created_at DESC
            """,
            (vault_id,),
        )
        rows = cursor.fetchall() or []
    finally:
        conn.close()

    result: list[dict] = []
    for row in rows:
        stored_text = row.get("extracted_text")
        encrypted = bool(row.get("extracted_text_encrypted"))
        plaintext: Optional[str]
        if stored_text and encrypted and key:
            try:
                plaintext = decrypt_message(stored_text, key)
            except Exception:
                plaintext = None
        elif stored_text and not encrypted:
            plaintext = stored_text
        else:
            plaintext = None
                                                                    
                                                                  
        row_dict = dict(row)
        row_dict["extracted_text"] = plaintext
        row_dict.pop("extracted_text_encrypted", None)
        result.append(row_dict)
    return result


def _load_one_file_for_analysis(
    vault_id: str,
    file_id: str,
    key: Optional[bytes],
) -> Optional[dict]:


    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT id, file_name, content_type, file_size, detected_type,
                   detected_service, saved_name, asset_type,
                   created_at, relative_path, content_sha256,
                   extracted_text, extracted_text_encrypted
            FROM uploaded_files
            WHERE vault_id = %s
              AND id = %s
              AND upload_status = 'complete'
            """,
            (vault_id, file_id),
        )
        row = cursor.fetchone()
    finally:
        conn.close()

    if not row:
        return None

    stored_text = row.get("extracted_text")
    encrypted = bool(row.get("extracted_text_encrypted"))
    plaintext: Optional[str]
    if stored_text and encrypted and key:
        try:
            plaintext = decrypt_message(stored_text, key)
        except Exception:
            plaintext = None
    elif stored_text and not encrypted:
        plaintext = stored_text
    else:
        plaintext = None

    row_dict = dict(row)
    row_dict["extracted_text"] = plaintext
    row_dict.pop("extracted_text_encrypted", None)
    return row_dict


def _resolve_file_for_analysis(
    vault_id: str,
    *,
    asset_name: Optional[str],
    decrypted_message: str,
    key: Optional[bytes],
) -> Optional[dict]:


    candidate = (asset_name or "").strip()

                                                             
    if candidate:
        row = _find_file_by_exact_name(vault_id, candidate)
        if row:
            return _load_one_file_for_analysis(
                vault_id, str(row["id"]), key,
            )

                                                                   
    if candidate:
        try:
            import asyncio
            try:
                loop = asyncio.get_event_loop()
            except RuntimeError:
                loop = None
            if loop and loop.is_running():
                                                                  
                                                                   
                hits = []
            else:
                hits = []
        except Exception:
            hits = []
                                                                
                                                                   
    try:
        from vault_chat_memory import get_last_file_search_results
        last = get_last_file_search_results(vault_id)
    except Exception:
        last = None
    if last and last.get("results"):
        first_id = (last["results"][0] or {}).get("file_id")
        if first_id:
            return _load_one_file_for_analysis(vault_id, str(first_id), key)

    return None


def _find_file_by_exact_name(
    vault_id: str, name_token: str,
) -> Optional[dict]:


    if not name_token or not name_token.strip():
        return None
    token = name_token.strip().lower()
    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT id, file_name, saved_name
            FROM uploaded_files
            WHERE vault_id = %s
              AND upload_status = 'complete'
              AND (
                LOWER(COALESCE(saved_name, '')) = %s
                OR LOWER(file_name) = %s
              )
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (vault_id, token, token),
        )
        return cursor.fetchone()
    finally:
        conn.close()


def _list_document_metadata_for_vault(vault_id: str) -> list[dict]:


    try:
        conn = get_db()
        try:
            cursor = conn.cursor(cursor_factory=RealDictCursor)
            cursor.execute(
                """
                SELECT uploaded_file_id, doc_type, metadata_json
                FROM vault_document_metadata
                WHERE vault_id = %s
                """,
                (vault_id,),
            )
            return cursor.fetchall() or []
        finally:
            conn.close()
    except Exception:
        logger.warning(
            "list_document_metadata_for_vault failed for vault=%s",
            vault_id,
        )
        return []


def _file_list_title(
    folder_filter: Optional[str],
    extension_filter: Optional[tuple[str, ...]],
    *,
    query: Optional[str] = None,
) -> str:


    type_label = (
        _extension_group_label(extension_filter, query=query)
        if extension_filter
        else "Files"
    )
                                                                     
    display_label = type_label[:1].upper() + type_label[1:]
    if folder_filter:
        return f"{display_label} inside {folder_filter}"
    return display_label


def _file_list_message(
    count: int,
    folder_filter: Optional[str],
    extension_filter: Optional[tuple[str, ...]],
    *,
    query: Optional[str] = None,
) -> str:


    type_label = (
        _extension_group_label(extension_filter, query=query)
        if extension_filter
        else "files"
    )
    if folder_filter:
        return f"I found {count} {type_label} inside {folder_filter}."
    return f"I found {count} {type_label}."


def _format_file_list_reply(
    rows: list[dict],
    folder_filter: Optional[str],
    extension_filter: Optional[tuple[str, ...]],
    *,
    query: Optional[str] = None,
) -> str:


    n = len(rows)
    if n == 0:
        return _format_no_match_reply(
            folder_filter, extension_filter, query=query,
        )

    type_label = (
        _extension_group_label(extension_filter, query=query)
        if extension_filter
        else "files"
    )
    if folder_filter:
        header = f"I found {n} {type_label} inside {folder_filter}:"
    else:
        header = f"I found {n} {type_label}:"

    lines = [header]
    for row in rows[:25]:
        rp = (row.get("relative_path") or "").strip()
        name = (
            row.get("saved_name") or row.get("file_name") or "file"
        )
        lines.append(f"  • {rp or name}")
    if n > 25:
        lines.append(f"  …and {n - 25} more")
    return "\n".join(lines)


def _try_folder_aware_file_retrieval(
    *,
    vault_id: str,
    decrypted_message: Optional[str],
    asset_name_from_llm: Optional[str] = None,
) -> Optional[str]:


    if not decrypted_message:
        return None

    extension_filter = _extract_extension_filter(decrypted_message)
    known_folders = _list_all_folder_names(vault_id)
    folder_filter = _extract_folder_hint(
        decrypted_message, known_folders=known_folders,
    )

    llm_name = (asset_name_from_llm or "").strip()
    literal_name = _extract_literal_filename(decrypted_message)
    has_specific_name = bool(llm_name) or literal_name is not None

                          
    if (
        extension_filter is None
        and folder_filter is None
        and not has_specific_name
    ):
        return None

    list_mode = (
        (extension_filter is not None and not has_specific_name)
        or (
            folder_filter is not None
            and not has_specific_name
            and _is_list_or_browse_phrasing(decrypted_message)
        )
    )

    if list_mode:
                                                                    
                                                                    
        rows = search_files_with_filters(
            vault_id,
            folder_filter=folder_filter,
            extension_filter=extension_filter,
            limit=200,
        )
        if not rows:
            return _format_no_match_reply(
                folder_filter, extension_filter,
                query=decrypted_message,
            )
                                                                      
                                                                 
        return _build_vault_file_list_envelope(
            rows,
            _file_list_title(
                folder_filter, extension_filter,
                query=decrypted_message,
            ),
            message=_file_list_message(
                len(rows), folder_filter, extension_filter,
                query=decrypted_message,
            ),
        )

                                                        
    name_q = literal_name or llm_name or decrypted_message
    candidates = search_files_with_filters(
        vault_id,
        folder_filter=folder_filter,
        extension_filter=extension_filter,
        name_query=name_q,
        limit=10,
    )

    if not candidates:
                                                                
                                                                  
        if extension_filter is None and folder_filter is None:
            return None
        return _format_no_match_reply(
            folder_filter,
            extension_filter,
            requested_name=(llm_name or literal_name or None),
            query=decrypted_message,
        )
    if len(candidates) == 1:
        return _build_structured_asset_reply(
            candidates[0],
            llm_name or literal_name or "file",
        )
                                                                     
                                                               
    requested = llm_name or literal_name or None
    title = (
        f'Files named "{requested}"' if requested else "Matching files"
    )
    message = (
        f'I found {len(candidates)} files named "{requested}". Which one?'
        if requested
        else f"I found {len(candidates)} matching files. Which one?"
    )
    return _build_vault_file_list_envelope(
        candidates,
        title,
        message=message,
        requested_name=requested,
    )


def _sanitize_relative_path(raw: Optional[str]) -> Optional[str]:


    if raw is None:
        return None
    if not isinstance(raw, str):
        return None
    candidate = raw.strip()
    if not candidate:
        return None

    if len(candidate) > _RELATIVE_PATH_MAX_LEN:
        return None

                                                                     
    normalized = candidate.replace("\\", "/")

                                                        
    if normalized.startswith("/"):
        return None
    if (
        len(normalized) >= 2
        and normalized[1] == ":"
        and normalized[0].isalpha()
    ):
                                                 
        return None

                                                                 
    for forbidden in ("\x00", "\r", "\n"):
        if forbidden in normalized:
            return None

                                                                     
    segments: list[str] = []
    for segment in normalized.split("/"):
        if not segment:
            continue
        if segment in (".", ".."):
            return None
        segments.append(segment)

    if not segments:
        return None

    return "/".join(segments)


def _lookup_doc_type_for_file(vault_id: str, file_id: str) -> Optional[str]:


    try:
        conn = get_db()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    """SELECT doc_type FROM vault_document_metadata
                       WHERE vault_id=%s AND uploaded_file_id=%s""",
                    (vault_id, file_id),
                )
                row = cur.fetchone()
        finally:
            conn.close()
        return row[0] if row else None
    except Exception as e:
        logger.warning(
            "_lookup_doc_type_for_file failed vault=%s file=%s: %s",
            vault_id, file_id, e,
        )
        return None


def _infer_asset_type(
    file_name: str,
    content_type: Optional[str],
    extracted_text: Optional[str],
    requested_name: Optional[str] = None,
    *,
    vault_id: Optional[str] = None,
    file_id: Optional[str] = None,
) -> str:


    from taxonomy import DOC_TYPE_TO_ASSET_TYPE

    ct_lower = (content_type or "").lower()

                                              
    if vault_id and file_id:
        doc_type = _lookup_doc_type_for_file(vault_id, file_id)
        if doc_type and doc_type in DOC_TYPE_TO_ASSET_TYPE:
            if ct_lower.startswith("image/"):
                return "id_image"
            return DOC_TYPE_TO_ASSET_TYPE[doc_type]

                                               
    lower_name = (requested_name or file_name or "").lower()
    text = (extracted_text or "").lower()

    if any(
        x in lower_name
        for x in [
            "driver license",
            "drivers license",
            "driver's license",
            "driver licence",
            "drivers licence",
            "photo id",
            "passport",
            "id card",
            "identity card",
            "national id",
        ]
    ):
        return "id_image"

    if ct_lower.startswith("image/"):
        return "image"

    if file_name.lower().endswith(".pdf"):
        return "pdf"

    if file_name.lower().endswith(".docx"):
        return "docx"

    if file_name.lower().endswith(".xlsx"):
        return "spreadsheet"

    if any(
        word in text
        for word in [
            "passport",
            "driver license",
            "drivers license",
            "driver's license",
            "national id",
            "identity card",
        ]
    ):
        return "id_document"

    return "file"


def _default_asset_type_for_upload(
    file_name: str,
    content_type: Optional[str],
    extracted_text: Optional[str],
) -> str:
    if _is_image_content_type(content_type, file_name):
        return "image"

    if _is_video_content_type(content_type, file_name):
        return "video"

    if _is_audio_content_type(content_type, file_name):
        return "audio"

    return _infer_asset_type(file_name, content_type, extracted_text)


def _asset_noun_for_message(asset_type: Optional[str]) -> str:


    if asset_type == "image":
        return "image"
    if asset_type == "video":
        return "video"
    if asset_type == "audio":
        return "recording"
    return "file"


def _build_upload_message(
    result: dict,
    *,
    suppress_naming_prompt: bool = False,
    auto_saved_name: Optional[str] = None,
    is_batch_upload: bool = False,
) -> str:


    filename = result["filename"]
    asset_type = result.get("asset_type")
    noun = _asset_noun_for_message(asset_type)

    if is_batch_upload:
                                                                     
                                                                  
        return f"Saved {noun} {filename}."

    if auto_saved_name:
                                                                   
                                                                    
        return f"Saved this {noun} as {_title_case_asset(auto_saved_name)}."

    if result.get("autosaved_secret"):
        message = (
            f"Saved {filename} and automatically extracted "
            f"{result.get('saved_count', 0)} {result.get('detected_type', 'secret')} secret(s)."
        )
        if result.get("skipped_count", 0) > 0:
            message += f" Skipped {result['skipped_count']} invalid item(s)."
        return message

    if suppress_naming_prompt:
                                                                  
                                                                    
        return f"Saved {noun} {filename}."

    if asset_type == "image":
        return (
            f"Saved image {filename}.\n"
            f"Tell me what you want to call it, like 'photo id', 'passport', 'selfie', or 'map'."
        )

    if asset_type == "video":
        return (
            f"Saved video {filename}. "
            f"What should I save this video as? "
            f"(e.g. 'private video', 'memory', or 'evidence')"
        )

    if asset_type == "audio":
        return (
            f"Saved recording {filename}. "
            f"What should I save this recording as? "
            f"(e.g. 'voice note', 'meeting', or 'idea')"
        )

    return (
        f"Saved file {filename}.\n"
        f"Tell me what you want to call it, like 'passport', 'school transcript', 'contract', or 'map'."
    )


def _build_structured_asset_reply(asset: dict, asset_name: str) -> str:
    pretty_name = _title_case_asset(asset.get("saved_name") or asset_name or "file")
    content_type = asset.get("content_type") or "application/octet-stream"
    message_type = "vault_image" if str(content_type).lower().startswith("image/") else "vault_file"

    payload = {
        "type": message_type,
        "message": f"I found your {pretty_name}.",
        "file_id": asset["id"],
        "file_name": asset["file_name"],
        "content_type": content_type,
        "asset_type": asset.get("asset_type") or "file",
    }
                                                                   
                                                                
    relative_path = asset.get("relative_path")
    if relative_path:
        payload["relative_path"] = relative_path
    return json.dumps(payload)


def ensure_vault_exists(vault_id: str):


    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)

        cursor.execute(
            """
            SELECT 1
            FROM vaults
            WHERE vault_id = %s
            LIMIT 1
            """,
            (vault_id,),
        )
        if not cursor.fetchone():
            raise HTTPException(status_code=404, detail="Vault not found")
    finally:
        conn.close()

def get_vault_total_bytes(vault_id: str) -> int:
    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT total_bytes
            FROM vaults
            WHERE vault_id = %s
            """,
            (vault_id,),
        )
        row = cursor.fetchone()
        return int(row["total_bytes"]) if row else 0
    finally:
        conn.close()


def _lookup_vault_name_safe(vault_id: str) -> str:


    try:
        conn = get_db()
        try:
            cur = conn.cursor(cursor_factory=RealDictCursor)
            cur.execute(
                """
                SELECT vault_name
                FROM vaults
                WHERE vault_id = %s
                LIMIT 1
                """,
                (vault_id,),
            )
            row = cur.fetchone()
            return str((row or {}).get("vault_name") or "")
        finally:
            conn.close()
    except Exception:
        return ""


def bump_vault_total_bytes(vault_id: str, delta_bytes: int):


    conn = get_db()
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            UPDATE vaults
            SET total_bytes = GREATEST(total_bytes + %s, 0)
            WHERE vault_id = %s
            """,
            (delta_bytes, vault_id),
        )
        conn.commit()
    finally:
        conn.close()
                                                                      
                                                                
    if delta_bytes != 0:
        try:
            from billing import (
                bump_account_total_bytes,
                get_account_id_for_vault,
            )
            _account_id = get_account_id_for_vault(vault_id)
            if _account_id is not None:
                bump_account_total_bytes(_account_id, int(delta_bytes))
        except Exception:
            logger.warning(
                "account_storage_totals bump failed vault=%s delta=%s "
                "— nightly reconciliation will catch this",
                vault_id, delta_bytes,
            )


def _is_valid_secret_payload(payload: Optional[dict]) -> bool:
    if not payload or not isinstance(payload, dict):
        return False

    secret_type = payload.get("secret_type")
    service = _normalize_service_name(payload.get("service"))
    fields = payload.get("fields")

    if not secret_type or service == "general":
        return False

    if not isinstance(fields, dict) or not fields:
        return False

    return any(str(v).strip() for v in fields.values() if v is not None)


def _is_high_confidence_login(payload: Optional[dict]) -> bool:


    if not _is_valid_secret_payload(payload):
        return False
    fields = payload.get("fields") or {}

    def _has(name: str) -> bool:
        v = fields.get(name)
        return bool(v is not None and str(v).strip())

    has_identifier = _has("username") or _has("email")
    has_secret = _has("password") or _has("pin") or _has("token") or _has("api_key")
    return has_identifier and has_secret


_EXPLICIT_LOGIN_EXTRACTION = re.compile(
    r"""(?xi)
    \b
    (?:
        # Bare verbs of intent. "check/look/search" used to require
        # "for" but that broke "check these documents for login
        # details" (the verb is separated from "for" by the object).
        # The triad gate downstream prevents false-positive saves
        # from a loose match here, so we err on letting the regex
        # match the user's phrasing.
        find | extract | scan | check | look | search |
        pull | grab |
        # "save" alone is too broad (matches "save my passport" /
        # "save these files" / the chat-side save_login path). Only
        # the "save any/the/all" idioms map to "extract logins".
        save \s+ (?:any|the|all)
    )
    \b
    # Up to ~40 chars of stuffing between verb and noun (e.g.
    # "scan this folder for saved").
    .{0,40}?
    \b
    (?:
        log\s?ins? | passwords? | credentials? | accounts? | secrets? |
        login \s+ (?:details?|info|credentials?|data)
    )
    \b
    """,
)


def _user_requested_login_extraction(text: Optional[str]) -> bool:


    if not text:
        return False
    stripped = text.strip()
    if not stripped:
        return False
    return bool(_EXPLICIT_LOGIN_EXTRACTION.search(stripped))


def _maybe_extract_login_payload(
    message: str,
    fallback_service: Optional[str] = None,
) -> Optional[dict]:


    extracted = extract_credentials(message)
    if not extracted.get("has_potential_secret"):
        return None

    fields = {}
    if extracted.get("username"):
        fields["username"] = extracted["username"]
    if extracted.get("email"):
        fields["email"] = extracted["email"]
    if extracted.get("password"):
        fields["password"] = extracted["password"]
    if extracted.get("pin"):
        fields["pin"] = extracted["pin"]
    if extracted.get("note"):
        fields["note"] = extracted["note"]

    if not fields:
        return None

    service = _normalize_service_name(extracted.get("service"))
    if service == "general":
                                                                      
                                                           
        service = _normalize_service_name(fallback_service)
    if service == "general":
        return None

    payload = {
        "secret_type": "login",
        "service": service,
        "fields": fields,
    }
    return payload if _is_valid_secret_payload(payload) else None


def _extract_text_from_bytes(file_name: str, file_bytes: bytes) -> Optional[str]:
    lower_name = file_name.lower()

    try:
        if lower_name.endswith(".txt") or lower_name.endswith(".csv"):
            return file_bytes.decode("utf-8", errors="ignore").strip()

        if lower_name.endswith(".json"):
            try:
                obj = json.loads(file_bytes.decode("utf-8", errors="ignore"))
                return json.dumps(obj, indent=2)
            except Exception:
                return file_bytes.decode("utf-8", errors="ignore").strip()

        if lower_name.endswith(".pdf") and PyPDF2 is not None:
            reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            pages = []
            for page in reader.pages:
                pages.append((page.extract_text() or "").strip())
            return "\n".join([p for p in pages if p]).strip()

        if lower_name.endswith(".docx") and docx is not None:
            document = docx.Document(io.BytesIO(file_bytes))
            return "\n".join([p.text for p in document.paragraphs if p.text]).strip()

        if lower_name.endswith(".xlsx"):
            workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                    if values:
                        rows.append(" | ".join(values))
            return "\n".join(rows).strip()
    except Exception:
        logger.warning("Failed to extract text from uploaded file")

    return None


def _classify_secret_from_text(text: Optional[str], file_name: str) -> dict:
    if not text:
        return {
            "detected_type": "file",
            "detected_service": "general",
            "payloads": [],
        }

    lower = text.lower()
    login_payloads = extract_multiple_credentials(text)
                                                                     
                                                                      
    high_confidence_logins = [
        p for p in login_payloads if _is_high_confidence_login(p)
    ]
    if high_confidence_logins:
        return {
            "detected_type": "login",
            "detected_service": high_confidence_logins[0]["service"],
            "payloads": high_confidence_logins,
        }

    if any(word in lower for word in ["passport", "driver license", "driver's license", "national id", "identity card"]):
        payload = {
            "secret_type": "id",
            "service": "id document",
            "fields": {"document_text": text[:4000]},
        }
        return {
            "detected_type": "id",
            "detected_service": "id document",
            "payloads": [payload] if _is_valid_secret_payload(payload) else [],
        }

    if any(word in lower for word in ["card number", "cvv", "expiry", "expiration", "debit card", "credit card"]):
        payload = {
            "secret_type": "card",
            "service": "payment card",
            "fields": {"card_text": text[:4000]},
        }
        return {
            "detected_type": "card",
            "detected_service": "payment card",
            "payloads": [payload] if _is_valid_secret_payload(payload) else [],
        }

    if any(word in lower for word in ["bank account", "routing", "account number", "swift", "iban"]):
        payload = {
            "secret_type": "bank",
            "service": "bank account",
            "fields": {"bank_text": text[:4000]},
        }
        return {
            "detected_type": "bank",
            "detected_service": "bank account",
            "payloads": [payload] if _is_valid_secret_payload(payload) else [],
        }

    if any(word in lower for word in ["backup code", "backup codes"]):
        payload = {
            "secret_type": "backup",
            "service": "backup codes",
            "fields": {"codes_text": text[:4000]},
        }
        return {
            "detected_type": "backup",
            "detected_service": "backup codes",
            "payloads": [payload] if _is_valid_secret_payload(payload) else [],
        }

    if any(word in lower for word in ["seed phrase", "recovery phrase", "mnemonic phrase"]):
        payload = {
            "secret_type": "seed",
            "service": "seed phrase",
            "fields": {"seed_text": text[:4000]},
        }
        return {
            "detected_type": "seed",
            "detected_service": "seed phrase",
            "payloads": [payload] if _is_valid_secret_payload(payload) else [],
        }

    return {
        "detected_type": "file",
        "detected_service": _normalize_service_name(file_name),
        "payloads": [],
    }


def save_named_uploaded_asset(
    vault_id: str,
    file_id: str,
    saved_name: str,
    key: Optional[bytes] = None,
):
    normalized_name = _normalize_asset_name(saved_name)
    if normalized_name == "general":
        raise ValueError("Missing asset name")

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
                                                                          
                                                                        
        cursor.execute(
            """
            SELECT id, file_name, content_type, detected_service,
                   extracted_text, extracted_text_encrypted
            FROM uploaded_files
            WHERE id = %s AND vault_id = %s
              AND upload_status = 'complete'
            LIMIT 1
            """,
            (file_id, vault_id),
        )
        row = cursor.fetchone()

        if not row:
            raise ValueError("Uploaded file not found")

        stored_text = row.get("extracted_text")
        if row.get("extracted_text_encrypted") and stored_text and key:
            try:
                inference_text = decrypt_message(stored_text, key)
            except Exception:
                inference_text = None
        else:
                                                                    
                                                                
            inference_text = stored_text if not row.get("extracted_text_encrypted") else None

        asset_type = _infer_asset_type(
            file_name=row["file_name"],
            content_type=row.get("content_type"),
            extracted_text=inference_text,
            requested_name=normalized_name,
            vault_id=vault_id,
            file_id=file_id,
        )

        cursor.execute(
            """
            UPDATE uploaded_files
            SET saved_name = %s,
                asset_type = %s,
                needs_naming = FALSE
            WHERE id = %s
            """,
            (normalized_name, asset_type, file_id),
        )
        conn.commit()

                                                                    
        try:
            from asset_tagger import tag_uploaded_file_safe
            tag_uploaded_file_safe(
                vault_id, file_id,
                file_name=row["file_name"],
                saved_name=normalized_name,
                asset_type=asset_type,
                content_type=row.get("content_type"),
                detected_service=row.get("detected_service"),
                replace=True,
            )
        except Exception:
            pass

                                                                         
        try:
            from document_understanding import extract_and_store_document_safe
            extract_and_store_document_safe(
                vault_id, file_id,
                file_name=row["file_name"],
                saved_name=normalized_name,
                asset_type=asset_type,
                content_type=row.get("content_type"),
                detected_service=row.get("detected_service"),
                extracted_text=inference_text,
                extracted_text_encrypted=False if inference_text else
                    bool(row.get("extracted_text_encrypted")),
                replace=True,
            )
        except Exception:
            pass

                                                                    
        try:
            enqueue_semantic_profile_for_file(vault_id, file_id)
        except Exception:
            pass

                                                                  
        try:
            from relationship_builder import build_relationships_for_file_safe
            build_relationships_for_file_safe(
                vault_id, file_id, replace=True,
            )
        except Exception:
            pass

                                                                  
        try:
            from expiry_engine import build_expiry_alerts_for_file_safe
            build_expiry_alerts_for_file_safe(
                vault_id, file_id, replace=True,
            )
        except Exception:
            pass

                                                                    
        try:
            from document_entities import build_entities_for_file_safe
            build_entities_for_file_safe(
                vault_id, file_id, replace=True,
            )
        except Exception:
            pass

        return {
            "saved_name": normalized_name,
            "asset_type": asset_type,
            "file_name": row["file_name"],
        }
    finally:
        conn.close()


def retrieve_saved_asset(vault_id: str, asset_name: str) -> Optional[dict]:
    normalized_name = _normalize_asset_name(asset_name)
    if normalized_name == "general":
        return None

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)

                                                           
        cursor.execute(
            """
            SELECT id, file_name, content_type, file_size, saved_name, asset_type, created_at
            FROM uploaded_files
            WHERE vault_id = %s
              AND LOWER(saved_name) = LOWER(%s)
              AND upload_status = 'complete'
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (vault_id, normalized_name),
        )
        row = cursor.fetchone()
        if row:
            return row

        cursor.execute(
            """
            SELECT id, file_name, content_type, file_size, saved_name, asset_type, created_at
            FROM uploaded_files
            WHERE vault_id = %s
              AND LOWER(saved_name) LIKE LOWER(%s)
              AND upload_status = 'complete'
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (vault_id, f"%{normalized_name}%"),
        )
        row = cursor.fetchone()
        if row:
            return row

        words = [w for w in normalized_name.split() if len(w) >= 2]
        if words:
            like_clauses = " OR ".join(["LOWER(saved_name) LIKE LOWER(%s)"] * len(words))
            params = [vault_id] + [f"%{w}%" for w in words]
            cursor.execute(
                f"""
                SELECT id, file_name, content_type, file_size, saved_name, asset_type, created_at
                FROM uploaded_files
                WHERE vault_id = %s
                  AND ({like_clauses})
                  AND upload_status = 'complete'
                ORDER BY created_at DESC
                LIMIT 1
                """,
                params,
            )
            return cursor.fetchone()

        return None
    finally:
        conn.close()


def get_pending_named_file(vault_id: str) -> Optional[dict]:
    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)


        cursor.execute(
            """
            SELECT id, file_name, content_type, created_at,
                   needs_naming, upload_status
            FROM uploaded_files
            WHERE vault_id = %s
              AND needs_naming = TRUE
              AND upload_status = 'complete'
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (vault_id,),
        )
        row = cursor.fetchone()
        try:
            if row is None:



                cursor.execute(
                    """
                    SELECT COUNT(*)::INT AS uploads,
                           SUM(CASE WHEN needs_naming = TRUE THEN 1 ELSE 0 END)::INT AS pending_naming,
                           SUM(CASE WHEN upload_status <> 'complete' THEN 1 ELSE 0 END)::INT AS in_flight
                    FROM uploaded_files
                    WHERE vault_id = %s
                    """,
                    (vault_id,),
                )
                stats = cursor.fetchone() or {}
                print(
                    "[CHAT-DEBUG] pending_file_lookup vault="
                    f"{(vault_id or '')[:8]}... row=None "
                    f"uploads={stats.get('uploads')} "
                    f"pending_naming={stats.get('pending_naming')} "
                    f"in_flight={stats.get('in_flight')}",
                    flush=True,
                )
            else:
                print(
                    "[CHAT-DEBUG] pending_file_lookup vault="
                    f"{(vault_id or '')[:8]}... row_id="
                    f"{row.get('id')!r} needs_naming="
                    f"{row.get('needs_naming')} upload_status="
                    f"{row.get('upload_status')!r}",
                    flush=True,
                )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] pending_file_lookup diagnostic failed",
            )
        return row
    finally:
        conn.close()


async def handle_tool_call(
    tool_name: str, args: dict, vault_id: str, key: bytes,
    token_id: str = "",
):
    try:
                                                            
                                                                 
        try:
            from vault_tool_result_cache import (
                maybe_get_cached, maybe_store,
            )
            cached_result = maybe_get_cached(
                vault_id=vault_id, token_id=token_id or "",
                tool_name=tool_name, args=args or {},
            )
            if cached_result is not None:
                return cached_result
        except Exception:
                                              
            logger.exception("[CHAT-TRACE] tool cache lookup failed")
            maybe_store = None                

        if tool_name == "save_secret":
            result = save_secret_tool(vault_id, args, key)
            _set_last_service(vault_id, args.get("service"))
                                                              
                                          
            try:
                from vault_tool_result_cache import (
                    invalidate_for_event,
                )
                invalidate_for_event(
                    vault_id=vault_id, event="credential_saved",
                )
            except Exception:
                logger.exception(
                    "[CHAT-TRACE] save_secret invalidation failed",
                )
            return result

        if tool_name == "retrieve_secret":
            result = retrieve_secret_tool(vault_id, args, key)
            _set_last_service(vault_id, args.get("service"))
            return result

        if tool_name == "list_secrets":
            result = list_secrets_tool(vault_id)
            if maybe_store is not None:
                try:
                    maybe_store(
                        vault_id=vault_id, token_id=token_id or "",
                        tool_name="list_secrets",
                        args=args or {}, result=result or "",
                    )
                except Exception:
                    pass
            return result

                                                                 
        try:
            from vault_knowledge_tools import VAULT_KNOWLEDGE_DISPATCH
        except Exception:
            VAULT_KNOWLEDGE_DISPATCH = {}
        if tool_name in VAULT_KNOWLEDGE_DISPATCH:
            fn = VAULT_KNOWLEDGE_DISPATCH[tool_name]
            try:
                safe_args = dict(args or {})
                                                               
                                                             
                safe_args.pop("vault_id", None)
                safe_args.pop("key", None)
                result = await asyncio.to_thread(
                    fn, vault_id=vault_id, key=key, **safe_args,
                )
                                                                 
                                                                 
                if maybe_store is not None and isinstance(result, str):
                    try:
                        maybe_store(
                            vault_id=vault_id,
                            token_id=token_id or "",
                            tool_name=tool_name,
                            args=safe_args, result=result,
                        )
                                                            
                                                              
                        if tool_name == (
                            "save_generated_credential_after_confirmation"
                        ):
                            from vault_tool_result_cache import (
                                invalidate_for_event,
                            )
                            invalidate_for_event(
                                vault_id=vault_id,
                                event="credential_saved",
                            )
                    except Exception:
                        pass
                return result
            except Exception:
                logger.exception(
                    "[CHAT-TRACE] vault_knowledge tool failed name=%s",
                    tool_name,
                )
                return json.dumps({"error": "unavailable"})

        return f"Unknown tool: {tool_name}"
    except Exception:
        logger.error("Tool execution failed")
        return "Tool execution failed."


def _peek_existing_login_fields(
    vault_id: str, service: str, key: bytes,
) -> dict:


    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT encrypted_data
            FROM vault_items
            WHERE vault_id = %s
              AND item_type = 'login'
              AND LOWER(service) = LOWER(%s)
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (vault_id, service),
        )
        row = cursor.fetchone()
    finally:
        conn.close()

    if not row or not row.get("encrypted_data"):
        return {}
    try:
        decoded = json.loads(decrypt_message(row["encrypted_data"], key))
        return decoded if isinstance(decoded, dict) else {}
    except Exception:
        return {}


class SaveSecretError(Exception):
    pass


class SaveSecretInvalidPayloadError(SaveSecretError):


    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


class SaveSecretStorageLimitError(SaveSecretError):


    def __init__(self, *, used_bytes: int, projected_bytes: int, limit_bytes: int):
        super().__init__("vault_storage_limit_exceeded")
        self.used_bytes = int(used_bytes)
        self.projected_bytes = int(projected_bytes)
        self.limit_bytes = int(limit_bytes)


_PLACEHOLDER_VALUE_PATTERNS: tuple = (
    # LLM template placeholders that must never be committed as a
    # real credential. 2026-07-22 incident: the assistant's chat
    # freeform reply text like "Username: <your Instagram
    # username>" is user-facing chat only, but if any code path
    # ever routed that literal string into save_secret_tool it
    # would corrupt the vault. Defensively reject.
    re.compile(r"^\s*<[^>]{2,64}>\s*$"),
    re.compile(r"^\s*\[[^\]]{2,64}\]\s*$"),
    re.compile(r"^\s*\{\{[^}]{2,64}\}\}\s*$"),
    re.compile(r"^\s*your\s+(?:instagram|facebook|twitter|"
               r"linkedin|gmail|email|username|password|login|"
               r"account)\s+(?:username|password|email|handle)?"
               r"\s*$", re.IGNORECASE),
    re.compile(r"^\s*enter\s+(?:your\s+)?"
               r"(?:username|password|email|login)\s*$",
               re.IGNORECASE),
    re.compile(r"^\s*REPLACE(?:_ME)?\s*$", re.IGNORECASE),
    re.compile(r"^\s*TODO\s*$", re.IGNORECASE),
    re.compile(r"^\s*<REDACTED>\s*$", re.IGNORECASE),
)


def _looks_like_placeholder_value(value) -> bool:
    """True when ``value`` looks like an LLM template placeholder
    rather than a real credential the user typed. Belt-and-braces
    check: the LLM's chat text should never reach this function,
    but we refuse to commit obvious placeholders as a defensive
    guarantee (2026-07-22 chat deep-fix)."""
    if not isinstance(value, str):
        return False
    if not value.strip():
        return False
    for pat in _PLACEHOLDER_VALUE_PATTERNS:
        if pat.match(value):
            return True
    return False


def _classify_save_login_payload(args) -> Optional[str]:


    if not isinstance(args, dict):
        return "missing_args"
    secret_type = args.get("secret_type")
    if not secret_type:
        return "missing_secret_type"
    service = _normalize_service_name(args.get("service"))
    if service == "general":
        return "missing_service"
    fields = args.get("fields", {})
    if not isinstance(fields, dict):
        return "fields_not_dict"
    if not fields:
        return "empty_fields"
    # 2026-07-22 chat deep-fix (item 5): reject placeholder-shaped
    # field values so a template like "<your Instagram username>"
    # can never be persisted as a real credential.
    for _fname, _fval in fields.items():
        if _looks_like_placeholder_value(_fval):
            return "placeholder_field_value"
    return None


def save_secret_tool(vault_id: str, args: dict, key: bytes,
                     *, generated: bool = False):


    ensure_vault_exists(vault_id)

    reason = _classify_save_login_payload(args)
    if reason is not None:
        raise SaveSecretInvalidPayloadError(reason)

    secret_type = args.get("secret_type")
    service = _normalize_service_name(args.get("service"))
    new_fields = args.get("fields", {})

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)

        cursor.execute(
            """
            SELECT id, encrypted_data
            FROM vault_items
            WHERE vault_id = %s
              AND item_type = %s
              AND LOWER(service) = LOWER(%s)
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (vault_id, secret_type, service),
        )
        existing = cursor.fetchone()

                                                                
        _SCHEMA_B_MARKERS = frozenset({
            "category", "title", "fields", "item_id",
        })
        old_fields: dict = {}
        old_notes: Optional[str] = None
        previous_size = 0

        if existing and existing.get("encrypted_data"):
            previous_size = len(existing["encrypted_data"].encode("utf-8"))
            try:
                decoded = json.loads(
                    decrypt_message(existing["encrypted_data"], key),
                )
                if isinstance(decoded, dict):
                    if any(k in decoded for k in _SCHEMA_B_MARKERS):
                                                                   
                                                        
                        inner = decoded.get("fields")
                        old_fields = inner if isinstance(inner, dict) else {}
                        if isinstance(decoded.get("notes"), str):
                            old_notes = decoded.get("notes")
                    else:
                                                                 
                                                             
                        old_fields = decoded
            except Exception:
                                                                
                                                             
                old_fields = {}
                old_notes = None

                                                             
        fields = {**old_fields, **new_fields}

                                                                 
        caller_notes = args.get("notes")
        if isinstance(caller_notes, str):
            envelope_notes = caller_notes
        elif old_notes is not None:
            envelope_notes = old_notes
        else:
            envelope_notes = ""
        envelope_payload = {
            "category": secret_type,
            "title":    service,
            "fields":   fields,
            "notes":    envelope_notes,
        }
        encrypted_data = encrypt_message(
            json.dumps(envelope_payload, ensure_ascii=False), key,
        )
        new_size = len(encrypted_data.encode("utf-8"))

                                                                    
        from billing import get_account_id_for_vault, get_entitlement
        _billing_account_id = get_account_id_for_vault(vault_id)
        if _billing_account_id is None:
                                                                      
                                                                      
            _billing_limit = int(MAX_VAULT_BYTES)
            _billing_used = get_vault_total_bytes(vault_id)
        else:
            _ent = get_entitlement(_billing_account_id)
            _billing_limit = _ent.effective_limit_bytes
            _billing_used = _ent.used_bytes

        projected_total = _billing_used - previous_size + new_size
        if projected_total > _billing_limit:
            raise SaveSecretStorageLimitError(
                used_bytes=_billing_used,
                projected_bytes=projected_total,
                limit_bytes=_billing_limit,
            )

        if existing:
            cursor.execute(
                """
                UPDATE vault_items
                SET encrypted_data = %s,
                    created_at = CURRENT_TIMESTAMP
                WHERE id = %s
                """,
                (encrypted_data, existing["id"]),
            )
        else:
            cursor.execute(
                """
                INSERT INTO vault_items (vault_id, item_type, service, encrypted_data)
                VALUES (%s, %s, %s, %s)
                RETURNING id
                """,
                (vault_id, secret_type, service, encrypted_data),
            )
            _new_item_row = cursor.fetchone()
            _new_item_id = (
                _new_item_row["id"] if isinstance(_new_item_row, dict)
                else (_new_item_row[0] if _new_item_row else None)
            )
                                                                                
            if _new_item_id is not None:
                from semantic_embedder import enqueue_vault_item_embedding
                enqueue_vault_item_embedding(client, vault_id, _new_item_id, "item_service", service)
                enqueue_vault_item_embedding(client, vault_id, _new_item_id, "item_type", secret_type)
                                                                                             
                try:
                    from asset_tagger import tag_vault_item_safe
                    tag_vault_item_safe(vault_id, _new_item_id,
                                        service=service, item_type=secret_type)
                except Exception:
                    pass

        conn.commit()

        delta = new_size - previous_size
        if delta != 0:
                                                           
                                                                       
            bump_vault_total_bytes(vault_id, delta)

                                                                      
        try:
            _audit_pw = fields.get("password")
            _audit_id = (
                existing["id"] if existing
                else locals().get("_new_item_id")
            )
            if _audit_pw and _audit_id is not None:
                from password_audit import upsert_password_audit_safe
                upsert_password_audit_safe(
                    vault_id, int(_audit_id), _audit_pw, generated=generated,
                )
        except Exception:
            pass

                                                                       
        try:
            _rel_id = (
                existing["id"] if existing
                else locals().get("_new_item_id")
            )
            if _rel_id is not None:
                from relationship_builder import build_relationships_for_item_safe
                build_relationships_for_item_safe(
                    vault_id, int(_rel_id), replace=True,
                )
        except Exception:
            pass

    finally:
        conn.close()

    _set_last_service(vault_id, service)
    remember_service(vault_id, service)

                                                                 
    try:
        from vault_intelligence_updater import on_credential_changed
        on_credential_changed(vault_id)
    except Exception:
        pass

    pretty_service = service.title()

    warning = None

    if secret_type == "login":
        warning = password_strength_warning(fields.get("password"))

    if warning:
        return (
            f"Saved your {pretty_service} login.\n\n"
            f"{warning} I can generate a stronger password for you and save it."
        )

    return f"Saved your {pretty_service} login."


def retrieve_secret_tool(vault_id: str, args: dict, key: bytes):
    service = _normalize_service_name(args.get("service"))

    if service == "general":
        raise ValueError("Missing service name")

    from vault_item_visibility import SYSTEM_ITEM_TYPE_SQL_TUPLE

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT id, item_type, service, encrypted_data, created_at
            FROM vault_items
            WHERE vault_id = %s
              AND LOWER(service) = LOWER(%s)
              AND item_type NOT IN %s
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (vault_id, service, SYSTEM_ITEM_TYPE_SQL_TUPLE),
        )
        result = cursor.fetchone()
    finally:
        conn.close()

    if not result:
        return f"I could not find a saved login for {service.title()}."

    try:
        decrypted_data = decrypt_message(result["encrypted_data"], key)
        fields = json.loads(decrypted_data)

                                                                          
        try:
            _pw = fields.get("password") if isinstance(fields, dict) else None
            if _pw and result.get("id") is not None:
                from password_audit import upsert_password_audit_safe
                upsert_password_audit_safe(
                    vault_id, int(result["id"]), _pw, generated=False,
                )
        except Exception:
            pass

                                                                      
        from locales import ENGLISH_FORMATTER as _fmt
        pretty_service = result["service"].title()
        lines = [_fmt.found_login(pretty_service)]

        if fields.get("username"):
            lines.append(f"{_fmt.field_label('username')}: {fields['username']}")

        if fields.get("email"):
            lines.append(f"{_fmt.field_label('email')}: {fields['email']}")

        if fields.get("password"):
            lines.append(f"{_fmt.field_label('password')}: {fields['password']}")

        if fields.get("pin"):
            lines.append(f"{_fmt.field_label('pin')}: {fields['pin']}")

        if fields.get("note"):
            lines.append(f"{_fmt.field_label('note')}: {fields['note']}")

        _set_last_service(vault_id, result["service"])
        remember_service(vault_id, result["service"])
        return "\n".join(lines)

    except Exception:
        from locales import ENGLISH_FORMATTER as _fmt
        return _fmt.found_login_decrypt_failed(service.title())

def list_secrets_tool(vault_id: str):
    from vault_item_visibility import SYSTEM_ITEM_TYPE_SQL_TUPLE

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT item_type, service
            FROM vault_items
            WHERE vault_id = %s
              AND item_type NOT IN %s
            GROUP BY item_type, service
            ORDER BY service ASC
            """,
            (vault_id, SYSTEM_ITEM_TYPE_SQL_TUPLE),
        )
        results = cursor.fetchall()
    finally:
        conn.close()

                                                        
    from locales import ENGLISH_FORMATTER as _fmt
    if not results:
        return _fmt.vault_empty()

    items = []
    for r in results:
        item_type = str(r["item_type"]).title()
        service = str(r["service"]).title()
        items.append(f"- {service} ({item_type})")

    return _fmt.vault_list_header() + "\n" + "\n".join(items)


def kickoff_analysis_for_uploaded_file(
    *,
    vault_id: str,
    file_id: str,
    file_name: Optional[str],
    content_type: Optional[str],
) -> None:


    try:
        from vault_analysis import (
            default_stages_for_file,
            enqueue_analysis_job,
            mark_file_analysis_pending,
            mark_file_analysis_unsupported,
        )
        stages = default_stages_for_file(
            file_name=file_name, content_type=content_type,
        )
        if not stages:
            mark_file_analysis_unsupported(file_id, vault_id=vault_id)
        else:
            mark_file_analysis_pending(file_id, vault_id=vault_id)
            for stage in stages:
                enqueue_analysis_job(
                    vault_id=vault_id,
                    file_id=file_id,
                    stage=stage,
                )
    except Exception:
        logger.exception(
            "[analysis-kickoff] failed file_id=%s stages_skipped", file_id,
        )

                                                         
    try:
        from vault_intelligence_updater import on_file_uploaded
        on_file_uploaded(vault_id)
    except Exception:
        pass


def _run_post_upload_processing(
    *,
    vault_id: str,
    file_id: str,
    file_name: str,
    asset_type: str,
    content_type: Optional[str],
    detected_service: Optional[str],
    stored_extracted_text: Optional[str],
    extracted_text_encrypted: bool,
) -> None:


    _bg_t0 = time.perf_counter()
    _ulog("post.start", file_id=file_id, asset_type=asset_type)
    try:
        from semantic_embedder import enqueue_uploaded_file_embedding
        enqueue_uploaded_file_embedding(client, vault_id, file_id, "file_name", file_name)
        enqueue_uploaded_file_embedding(client, vault_id, file_id, "asset_type", asset_type)
        enqueue_uploaded_file_embedding(client, vault_id, file_id, "detected_service", detected_service)
    except Exception:
        logger.exception("[upload-bg] embedding failed file_id=%s", file_id)
    _ulog("post.embed", file_id=file_id, elapsed_ms=int((time.perf_counter() - _bg_t0) * 1000))

    try:
        from asset_tagger import tag_uploaded_file_safe
        tag_uploaded_file_safe(
            vault_id, file_id,
            file_name=file_name, asset_type=asset_type,
            content_type=content_type, detected_service=detected_service,
        )
    except Exception:
        logger.exception("[upload-bg] asset tagging failed file_id=%s", file_id)
    _ulog("post.tag", file_id=file_id, elapsed_ms=int((time.perf_counter() - _bg_t0) * 1000))

    try:
        from document_understanding import extract_and_store_document_safe
        extract_and_store_document_safe(
            vault_id, file_id,
            file_name=file_name, asset_type=asset_type, content_type=content_type,
            detected_service=detected_service,
            extracted_text=stored_extracted_text,
            extracted_text_encrypted=extracted_text_encrypted,
        )
    except Exception:
        logger.exception("[upload-bg] doc understanding failed file_id=%s", file_id)
    _ulog("post.doc", file_id=file_id, elapsed_ms=int((time.perf_counter() - _bg_t0) * 1000))

    try:
        from relationship_builder import build_relationships_for_file_safe
        build_relationships_for_file_safe(vault_id, file_id, replace=True)
    except Exception:
        logger.exception("[upload-bg] relationships failed file_id=%s", file_id)
    _ulog("post.rel", file_id=file_id, elapsed_ms=int((time.perf_counter() - _bg_t0) * 1000))

    try:
        from expiry_engine import build_expiry_alerts_for_file_safe
        build_expiry_alerts_for_file_safe(vault_id, file_id, replace=True)
    except Exception:
        logger.exception("[upload-bg] expiry alerts failed file_id=%s", file_id)
    _ulog("post.exp", file_id=file_id, elapsed_ms=int((time.perf_counter() - _bg_t0) * 1000))

    try:
        from document_entities import build_entities_for_file_safe
        build_entities_for_file_safe(vault_id, file_id, replace=True)
    except Exception:
        logger.exception("[upload-bg] entities failed file_id=%s", file_id)
    _ulog("post.done", file_id=file_id, elapsed_ms=int((time.perf_counter() - _bg_t0) * 1000))


def save_uploaded_file(
    vault_id: str,
    file_name: str,
    content_type: Optional[str],
    file_bytes: bytes,
    key: bytes,
    background_tasks: Optional[BackgroundTasks] = None,
    is_batch_upload: bool = False,
    auto_save_login_credentials: bool = False,
    relative_path: Optional[str] = None,
    import_id: Optional[str] = None,


    content_sha256: Optional[str] = None,
    duplicate_action: str = "prompt",
    # ZK ciphertext-first mode: when both ciphertext blobs are
    # supplied, the row is INSERTed with legacy readable columns
    # (file_name, content_type, detected_type, detected_service,
    # saved_name, asset_type) as NULL. Plaintext `file_name` /
    # `content_type` still travel in RAM inside this handler for
    # OCR/classification/sanitize helpers, but never touch the DB
    # on the ZK path.
    filename_ciphertext_bytes: Optional[bytes] = None,
    content_type_ciphertext_bytes: Optional[bytes] = None,
):


    ensure_vault_exists(vault_id)
    _t0 = time.perf_counter()
    _ulog("upload.enter", file_name=file_name, content_type=content_type,
          size=len(file_bytes), is_batch=is_batch_upload,
          has_import_id=bool(import_id))

                                                                      
    _validate_import_id_for_upload(vault_id, import_id)

    file_size = len(file_bytes)

                                                                       
    from billing import get_account_id_for_vault, get_entitlement
    _billing_account_id = get_account_id_for_vault(vault_id)
    if _billing_account_id is None:
        _billing_limit = int(MAX_VAULT_BYTES)
        _billing_used = get_vault_total_bytes(vault_id)
    else:
        _ent = get_entitlement(_billing_account_id)
        _billing_limit = _ent.effective_limit_bytes
        _billing_used = _ent.used_bytes

    if _billing_used + file_size > _billing_limit:
        raise HTTPException(
            status_code=413,
            detail={
                "code": "vault_storage_limit_exceeded",
                "message": (
                    f"Vault storage limit reached. "
                    f"Used {_billing_used / (1024*1024):.0f} MB of "
                    f"{_billing_limit / (1024*1024):.0f} MB. "
                    f"Delete data before uploading more."
                ),
                "used_bytes": _billing_used,
                "projected_bytes": _billing_used + file_size,
                "limit_bytes": _billing_limit,
            },
        )

                                                                
    safe_hash = normalize_content_sha256(content_sha256)
    action_mode = normalize_duplicate_action(duplicate_action)
    duplicate_of_file_id: Optional[str] = None
    existing_match: Optional[dict] = None

    if safe_hash is not None:
        all_rows = list_uploaded_files(vault_id)
        existing_match = find_existing_file_by_hash(
            all_rows,
            content_sha256=safe_hash,
            file_size=file_size,
        )

        if existing_match is not None:
            if action_mode == "skip":
                                                                   
                                                                 
                _bump_import_batch_on_skipped_duplicate(
                    vault_id=vault_id,
                    import_id=import_id,
                )
                _ulog(
                    "upload.duplicate_skipped",
                    file_name=file_name,
                    size=file_size,
                    hash_prefix=short_hash_for_log(safe_hash),
                    existing_file_id=str(existing_match.get("id")),
                )
                return {
                    "status": "skipped_duplicate",
                    "duplicate_of_file_id": str(existing_match.get("id")),
                    "existing_file_name": existing_match.get("file_name"),
                    "existing_saved_name": existing_match.get("saved_name"),
                    "existing_relative_path": existing_match.get("relative_path"),
                    "existing_uploaded_at": (
                        existing_match.get("created_at").isoformat()
                        if hasattr(
                            existing_match.get("created_at"), "isoformat"
                        )
                        else existing_match.get("created_at")
                    ),
                    "filename": file_name,
                    "size": file_size,
                    "content_type": content_type,
                    "message": _format_duplicate_skipped_message(
                        existing_match
                    ),
                }

            if action_mode == "keep_both":
                                                                    
                                                                    
                duplicate_of_file_id = str(existing_match.get("id"))
                _ulog(
                    "upload.duplicate_keep_both",
                    file_name=file_name,
                    hash_prefix=short_hash_for_log(safe_hash),
                    existing_file_id=duplicate_of_file_id,
                )
                                                         

            else:
                                                                   
                                                                     
                _ulog(
                    "upload.duplicate_found",
                    file_name=file_name,
                    size=file_size,
                    hash_prefix=short_hash_for_log(safe_hash),
                    existing_file_id=str(existing_match.get("id")),
                )
                raise HTTPException(
                    status_code=409,
                    detail={
                        "code": "duplicate_found",
                        "message": _format_duplicate_found_message(
                            existing_match
                        ),
                        "existing_file_id": str(existing_match.get("id")),
                        "existing_file_name": existing_match.get(
                            "file_name"
                        ),
                        "existing_saved_name": existing_match.get(
                            "saved_name"
                        ),
                        "existing_relative_path": existing_match.get(
                            "relative_path"
                        ),
                        "existing_uploaded_at": (
                            existing_match.get("created_at").isoformat()
                            if hasattr(
                                existing_match.get("created_at"),
                                "isoformat",
                            )
                            else existing_match.get("created_at")
                        ),
                        "incoming_file_name": file_name,
                        "incoming_size": file_size,
                    },
                )

                                                                       
    default_saved_name: Optional[str] = None
    if is_batch_upload:
        try:
            normalized_default = _normalize_asset_name(file_name or "")
            if normalized_default and normalized_default != "general":
                default_saved_name = normalized_default
        except Exception:
            default_saved_name = None

    name_conflict_existing: Optional[dict] = None
    incoming_renamed_from: Optional[str] = None
    proposed_versioned_name: Optional[str] = None
    name_conflict_version_seed: int = 1

    if duplicate_of_file_id is None:
                                                                
                                                  
        rows_for_conflict = (
            list_uploaded_files(vault_id)
            if safe_hash is None
            else list_uploaded_files(vault_id)
        )
        check_saved_name = default_saved_name
        check_file_name = file_name

        name_conflict_existing = find_name_conflict(
            rows_for_conflict,
            saved_name=check_saved_name,
            file_name=check_file_name,
            relative_path=relative_path,
            incoming_content_sha256=safe_hash,
        )

        if name_conflict_existing is not None:
            proposed_versioned_name = next_available_versioned_name(
                rows_for_conflict,
                base_name=(
                    check_saved_name or check_file_name or "file"
                ),
                relative_path=relative_path,
            )
            name_conflict_version_seed = int(
                name_conflict_existing.get("version_number") or 1
            )

                                                                    
            if action_mode == "prompt" and not is_batch_upload:
                _ulog(
                    "upload.name_conflict_prompt",
                    file_name=file_name,
                    hash_prefix=short_hash_for_log(safe_hash),
                    existing_file_id=str(
                        name_conflict_existing.get("id")
                    ),
                )
                raise HTTPException(
                    status_code=409,
                    detail={
                        "code": "name_conflict",
                        "message": _format_name_conflict_message(
                            name_conflict_existing
                        ),
                        "existing_file_id": str(
                            name_conflict_existing.get("id")
                        ),
                        "existing_file_name":
                            name_conflict_existing.get("file_name"),
                        "existing_saved_name":
                            name_conflict_existing.get("saved_name"),
                        "existing_relative_path":
                            name_conflict_existing.get("relative_path"),
                        "existing_uploaded_at": (
                            name_conflict_existing.get(
                                "created_at"
                            ).isoformat()
                            if hasattr(
                                name_conflict_existing.get(
                                    "created_at"
                                ),
                                "isoformat",
                            )
                            else name_conflict_existing.get(
                                "created_at"
                            )
                        ),
                        "incoming_file_name": file_name,
                        "incoming_size": file_size,
                        "proposed_versioned_name":
                            proposed_versioned_name,
                    },
                )

                                                              
            incoming_renamed_from = (
                check_saved_name or check_file_name
            )
            if is_batch_upload:
                default_saved_name = proposed_versioned_name
            _ulog(
                "upload.name_conflict_auto_versioned",
                file_name=file_name,
                hash_prefix=short_hash_for_log(safe_hash),
                existing_file_id=str(
                    name_conflict_existing.get("id")
                ),
                proposed=proposed_versioned_name,
            )

    encrypted_file_data = encrypt_bytes(file_bytes, key)
    _ulog("upload.encrypted", file_name=file_name, elapsed_ms=int((time.perf_counter() - _t0) * 1000))

    should_analyze = _should_analyze_upload(content_type, file_name)
    extracted_text = _extract_text_from_bytes(file_name, file_bytes) if should_analyze else None
    _ulog("upload.text_extracted", file_name=file_name, had_text=bool(extracted_text),
          elapsed_ms=int((time.perf_counter() - _t0) * 1000))

    classification = (
        _classify_secret_from_text(extracted_text, file_name)
        if should_analyze
        else {
            "detected_type": "file",
            "detected_service": "general",
            "payloads": [],
        }
    )

    file_id = str(uuid.uuid4())

    payloads = classification.get("payloads") or []
    saved_count = 0
    skipped_count = 0
    skipped_login_count = 0                                            

    if should_analyze:
        for payload in payloads:
                                                                         
                                                                      
            is_login_payload = (payload or {}).get("secret_type") == "login"
            if is_login_payload and not auto_save_login_credentials:
                skipped_login_count += 1
                continue
            try:
                if _is_valid_secret_payload(payload):
                    save_secret_tool(vault_id, payload, key)
                    saved_count += 1
                else:
                    skipped_count += 1
            except Exception:
                skipped_count += 1
                logger.warning("Skipped invalid extracted payload from uploaded file")

    if skipped_login_count > 0:
                                                                      
                                                                   
        _ulog(
            "upload.login_autosave_skipped",
            file_name=file_name,
            skipped_count=skipped_login_count,
            reason="auto_save_login_credentials=False",
        )

    autosaved_secret = saved_count > 0
    asset_type = _default_asset_type_for_upload(file_name, content_type, extracted_text)
                                                                     
                                                                    
    needs_naming = default_saved_name is None

                                                                  
    version_number = 1
    if duplicate_of_file_id is not None:
        version_base = (
            default_saved_name or file_name or "file"
        )
        vault_rows = list_uploaded_files(vault_id)
        versioned = next_available_versioned_name(
            vault_rows,
            base_name=version_base,
            relative_path=relative_path,
        )
        if is_batch_upload:
            default_saved_name = versioned
            needs_naming = False
        version_number = max(
            int(existing_match.get("version_number") or 1) + 1
            if existing_match is not None else 2,
            2,
        )
    elif name_conflict_existing is not None:
                                                                    
                                                                    
        version_number = max(name_conflict_version_seed + 1, 2)
        needs_naming = default_saved_name is None

                                                                             
    if extracted_text:
        stored_extracted_text = encrypt_message(extracted_text, key)
        extracted_text_encrypted = True
    else:
        stored_extracted_text = None
        extracted_text_encrypted = False

    is_zk_upload = filename_ciphertext_bytes is not None

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        if is_zk_upload:
            # ZK ciphertext-first: readable legacy columns are NULL
            # from the initial INSERT. Server-inferred fields
            # (detected_type, detected_service, asset_type) that
            # would normally come from `classification` are also
            # left NULL — document understanding runs later in
            # the request and its output (if any) is encrypted and
            # persisted via /vault/ciphertext/uploaded-files.
            cursor.execute(
                """
                INSERT INTO uploaded_files (
                    id, vault_id, file_name, content_type,
                    file_size, encrypted_file_data, extracted_text,
                    extracted_text_encrypted,
                    detected_type, detected_service, autosaved_secret,
                    saved_name, asset_type, needs_naming, relative_path,
                    import_id, content_sha256, duplicate_of_file_id,
                    version_number,
                    file_name_ciphertext, content_type_ciphertext
                )
                VALUES (%s, %s, NULL, NULL,
                        %s, %s, %s, %s,
                        NULL, NULL, %s,
                        NULL, NULL, %s, %s,
                        %s, %s, %s,
                        %s,
                        %s, %s)
                """,
                (
                    file_id,
                    vault_id,
                    file_size,
                    encrypted_file_data,
                    stored_extracted_text,
                    extracted_text_encrypted,
                    autosaved_secret,
                    needs_naming,
                    relative_path,
                    import_id,
                    safe_hash,
                    duplicate_of_file_id,
                    version_number,
                    filename_ciphertext_bytes,
                    content_type_ciphertext_bytes,
                ),
            )
        else:
            cursor.execute(
                """
                INSERT INTO uploaded_files (
                    id, vault_id, file_name, content_type,
                    file_size, encrypted_file_data, extracted_text,
                    extracted_text_encrypted,
                    detected_type, detected_service, autosaved_secret,
                    saved_name, asset_type, needs_naming, relative_path,
                    import_id, content_sha256, duplicate_of_file_id,
                    version_number
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    file_id,
                    vault_id,
                    file_name,
                    content_type,
                    file_size,
                    encrypted_file_data,
                    stored_extracted_text,
                    extracted_text_encrypted,
                    classification["detected_type"],
                    classification["detected_service"],
                    autosaved_secret,
                    default_saved_name,
                    asset_type,
                    needs_naming,
                    relative_path,
                    import_id,
                    safe_hash,
                    duplicate_of_file_id,
                    version_number,
                ),
            )
        conn.commit()
    finally:
        conn.close()
    _ulog("upload.inserted", file_id=file_id,
          has_relative_path=bool(relative_path),
          has_import_id=bool(import_id),
          elapsed_ms=int((time.perf_counter() - _t0) * 1000))

                                                                    
    kickoff_analysis_for_uploaded_file(
        vault_id=vault_id,
        file_id=file_id,
        file_name=file_name,
        content_type=content_type,
    )

                                                                   
    _bump_import_batch_on_upload_success(
        vault_id=vault_id,
        import_id=import_id,
        file_size=file_size,
    )

    bump_vault_total_bytes(vault_id, file_size)
                                                                 
                                                                     
    if _billing_account_id is not None:
        try:
            from billing import bump_account_total_bytes
            bump_account_total_bytes(_billing_account_id, file_size)
        except Exception:
            logger.warning(
                "bump_account_total_bytes failed account=%s delta=%s "
                "after save_uploaded_file — reconciliation will catch it",
                _billing_account_id, file_size,
            )

                                                                     
    if background_tasks is not None:
        background_tasks.add_task(
            _run_post_upload_processing,
            vault_id=vault_id,
            file_id=file_id,
            file_name=file_name,
            asset_type=asset_type,
            content_type=content_type,
            detected_service=classification.get("detected_service"),
            stored_extracted_text=stored_extracted_text,
            extracted_text_encrypted=extracted_text_encrypted,
        )
    else:
        _run_post_upload_processing(
            vault_id=vault_id,
            file_id=file_id,
            file_name=file_name,
            asset_type=asset_type,
            content_type=content_type,
            detected_service=classification.get("detected_service"),
            stored_extracted_text=stored_extracted_text,
            extracted_text_encrypted=extracted_text_encrypted,
        )

    return {
        "status": "uploaded",
        "file_id": file_id,
        "filename": file_name,
        "content_type": content_type,
        "size": file_size,
        "text_extracted": bool(extracted_text),
        "detected_type": classification["detected_type"],
        "detected_service": classification["detected_service"],
        "autosaved_secret": autosaved_secret,
        "saved_count": saved_count,
        "skipped_count": skipped_count,
        "asset_type": asset_type,
        "needs_naming": needs_naming,
                                                             
                                                                        
        "saved_name": default_saved_name,
        "was_analyzed": should_analyze,
                                                                       
                                                                       
        "duplicate_of_file_id": duplicate_of_file_id,
        "version_number": version_number,
                                                                    
                                                                     
        "renamed": incoming_renamed_from is not None,
        "original_saved_name": incoming_renamed_from,
    }


def list_uploaded_files(vault_id: str):
    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT id, file_name, content_type, file_size, detected_type,
                   detected_service, autosaved_secret, saved_name, asset_type,
                   needs_naming, created_at, relative_path
            FROM uploaded_files
            WHERE vault_id = %s
              AND upload_status = 'complete'
            ORDER BY created_at DESC
            """,
            (vault_id,),
        )
        return cursor.fetchall() or []
    finally:
        conn.close()


def _get_uploaded_files_for_chat(vault_id: str, file_ids: list[str]) -> list[dict]:
    if not file_ids:
        return []

                                                                             
    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT id, file_name, content_type, file_size,
                   detected_type, detected_service, autosaved_secret,
                   saved_name, asset_type, needs_naming, created_at
            FROM uploaded_files
            WHERE vault_id = %s
              AND id = ANY(%s)
              AND upload_status = 'complete'
            ORDER BY created_at DESC
            """,
            (vault_id, file_ids),
        )
        return cursor.fetchall() or []
    finally:
        conn.close()


def _extract_requested_service(message: str, vault_id: str) -> str:
    extracted = extract_credentials(message)
    service = _normalize_service_name(extracted.get("service"))

    if service != "general":
        return service

    simple_patterns = [
        r"(?:what is|what's|show|show me|get|find|retrieve|fetch)\s+my\s+(.+?)\s+(?:login|logins|password|credentials|account|accounts)\b",
        r"(?:login|logins|password|credentials|account|accounts)\s+for\s+(.+)$",
    ]

    lower = message.lower().strip()
    for pattern in simple_patterns:
        match = re.search(pattern, lower)
        if match:
            candidate = _normalize_service_name(match.group(1))
            if candidate != "general":
                return candidate

    last_service = _get_last_service(vault_id)
    if last_service:
        return last_service

    return "general"


def _vault_functions_for_message(message: str) -> list:


    if not message:
        return list(VAULT_FUNCTIONS)
    try:
        from vault_brain_intent import (
            looks_like_vault_files_question,
            is_credential_files_query,
        )
        if is_credential_files_query(message) or \
           looks_like_vault_files_question(message):
            return [
                fn for fn in VAULT_FUNCTIONS
                if fn.get("function", {}).get("name") != "list_secrets"
            ]
    except Exception:
                                                                  
                                                                    
        logger.warning("vault_functions filter failed; using full list")

                                                               
    try:
        if _looks_like_vault_overview_question(message):
            return [
                fn for fn in VAULT_FUNCTIONS
                if fn.get("function", {}).get("name") != "list_secrets"
            ]
    except Exception:
        logger.warning(
            "vault_functions overview-shape filter failed; "
            "using full list",
        )
    return list(VAULT_FUNCTIONS)


_VAULT_OVERVIEW_PATTERNS = (
                                                                   
    r"\breview(?:ing|s|ed)?\b.*\b(everything|all|vault|files?)\b",
                                                        
    r"\bscan(?:ning|s|ned)?\b.*\b(everything|all|vault|files?)\b",
                                                                           
    r"\b(?:go|read|look)\s+through\b.*\b(everything|all|vault|files?)\b",
                                
    r"\banaly[sz]e\b.*\b(everything|vault|files?)\b",
                                                                   
    r"\b(?:what\s+do\s+you\s+know|tell\s+me)\b.*\bmy\s+vault\b",
                                                        
    r"\bwhat\b.*\b(?:needs?\s+attention|pending|outstanding)\b",
                                                         
    r"\brev[ie]+w\b.*\b(everything|all|vault|files?)\b",
)

                                                                       
_IDENTITY_PATTERNS = (
    r"^\s*(?:hi|hello|hey|yo|hiya|howdy)\b",
    r"^\s*good\s+(?:morning|afternoon|evening|night|day)\b",
    r"\bwho\s+are\s+you\b",
    r"\bwhat\s+are\s+you\b",
    r"\bwhat\s+can\s+you\s+do\b",
)


def _looks_like_vault_overview_question(message: str) -> bool:


    if not isinstance(message, str) or not message.strip():
        return False
    text = message.strip().lower()
    import re as _re
    for pat in _VAULT_OVERVIEW_PATTERNS:
        if _re.search(pat, text):
            return True
    for pat in _IDENTITY_PATTERNS:
        if _re.search(pat, text):
            return True
    return False


async def ai_stream(
    messages, vault_id: str, key: bytes,
    last_user_message: str = "",
    token_id: str = "",
):
    try:
                                                             
                                                                
        from vault_ai_provider import (
            get_chat_client, active_provider_name,
        )
        from vault_config import ai as _ai_cfg
        provider_name = active_provider_name()

                                                                 
        try:
            from vault_history_compressor import (
                compress_messages as _compress,
            )
            _compression = _compress(messages or [])
            if _compression.compression_summary.get("compressed"):
                                                            
                                                            
                if messages and isinstance(messages[-1], dict) \
                        and messages[-1].get("role") == "user":
                    last_user = messages[-1]
                    rebuilt = _compression.assembled()
                                                                 
                                                    
                    if rebuilt and rebuilt[-1] is last_user:
                        rebuilt.pop()
                    rebuilt.append(last_user)
                    messages = rebuilt
                logger.info(
                    "[CHAT-TRACE] history_compressed kept=%d "
                    "summarized=%d memory_chars=%d",
                    _compression.compression_summary.get("kept_count", 0),
                    _compression.compression_summary.get("summarized_count", 0),
                    _compression.compression_summary.get("memory_chars", 0),
                )
        except Exception:
            logger.exception(
                "[CHAT-TRACE] history compression failed",
            )

                                                                  
        try:
            from vault_planner import plan_user_message
            _planner = await plan_user_message(
                message=last_user_message or "",
                recent_history=[
                    m for m in (messages or [])[-5:]
                    if isinstance(m, dict)
                    and m.get("role") in ("user", "assistant")
                ],
            )
        except Exception:
            logger.exception("[CHAT-TRACE] planner failed")
            _planner = None

                                                             
        allowed_tools = _vault_functions_for_message(last_user_message)

                                                                  
        _skip_tools = False
        _planner_used = False
        if _planner is not None:
            try:
                if _planner.can_skip_tools():
                    _skip_tools = True
                    _planner_used = True
                elif _planner.planned_tools:
                    _planner_tool_names = set(_planner.planned_tools)
                    _planner_pruned = [
                        fn for fn in allowed_tools
                        if fn.get("function", {}).get("name")
                        in _planner_tool_names
                    ]
                    if _planner_pruned:
                        allowed_tools = _planner_pruned
                        _planner_used = True
            except Exception:
                logger.exception(
                    "[CHAT-TRACE] planner-driven tool prune failed",
                )

                                                             
        if not _planner_used and not _skip_tools:
            try:
                from vault_tool_router import filter_function_schemas
                partitioned = filter_function_schemas(
                    allowed_tools, last_user_message or "",
                )
                if partitioned:
                    allowed_tools = partitioned
            except Exception:
                                                                 
                                                      
                logger.exception("[CHAT-TRACE] tool router failed")

        if _skip_tools:
            allowed_tools = []

        try:
            blocked = [
                fn["function"]["name"] for fn in VAULT_FUNCTIONS
                if fn not in allowed_tools
            ]
        except Exception:
            blocked = []
                                                              
                                                          
        _forced_tool_name: Optional[str] = None
        if _planner is not None:
            if _planner.intent == "credential_creation_draft":
                _forced_tool_name = "generate_credential_draft"
            elif _planner.intent == "credential_save_confirmation":
                                                                  
                                                                    
                _allow_force_save = True
                try:
                    from vault_active_context import (
                        CONTEXT_CREDENTIAL_DRAFT,
                        get_active_context,
                    )
                    from vault_chat_context_gate import (
                        is_explicit_credential_save_phrase,
                    )
                    _planner_active_ctx = get_active_context(vault_id)
                    if (
                        _planner_active_ctx is not None
                        and _planner_active_ctx
                        != CONTEXT_CREDENTIAL_DRAFT
                        and not is_explicit_credential_save_phrase(
                            last_user_message or "",
                        )
                    ):
                        _allow_force_save = False
                        print(
                            "[CHAT-TRACE] credential_save_force_"
                            "suppressed vault="
                            f"{(vault_id or '')[:8]}... "
                            f"active_ctx={_planner_active_ctx}",
                            flush=True,
                        )
                except Exception:
                    _allow_force_save = True
                if _allow_force_save:
                    _forced_tool_name = (
                        "save_generated_credential_after_confirmation"
                    )
            elif _planner.intent == "document_search":
                                                                
                                                                  
                _forced_tool_name = "find_in_vault"

                                                   
        if _forced_tool_name and allowed_tools:
            if not any(
                fn.get("function", {}).get("name") == _forced_tool_name
                for fn in allowed_tools
            ):
                _forced_schema = next(
                    (
                        fn for fn in VAULT_FUNCTIONS
                        if fn.get("function", {}).get("name")
                           == _forced_tool_name
                    ),
                    None,
                )
                if _forced_schema is not None:
                    allowed_tools = list(allowed_tools) + [_forced_schema]

                                                               
        _planner_intent_str = (
            _planner.intent if _planner is not None else "none"
        )
        _planner_tools_str = (
            ",".join(_planner.planned_tools)
            if _planner is not None and _planner.planned_tools
            else "none"
        )
        _trace_line = (
            f"[CHAT-TRACE] ai_stream provider={provider_name} "
            f"vault={(vault_id or '')[:8]}... "
            f"tool_count={len(allowed_tools)} "
            f"blocked_tools={','.join(blocked) or 'none'} "
            f"msg_len={len(last_user_message or '')} "
            f"planner_intent={_planner_intent_str} "
            f"planner_used={_planner_used} "
            f"skip_tools={_skip_tools} "
            f"planned_tools={_planner_tools_str} "
            f"forced_tool={_forced_tool_name or 'none'}"
        )
        print(_trace_line, flush=True)
        logger.info("%s", _trace_line)
        stream_client = get_chat_client()
                                                               
                                                                 
        _output_cap = int(
            os.getenv("VAULTAI_CHAT_OUTPUT_TOKENS", "500"),
        )
                                                             
                                                                   
        try:
            from vault_active_context import get_active_context
            from vault_chat_context_gate import context_hint_system_message
            _ai_active_ctx = get_active_context(vault_id)
            _ctx_hint = context_hint_system_message(
                active_context=_ai_active_ctx,
                has_pending_draft=False,
                planner_intent=(
                    _planner.intent if _planner is not None else None
                ),
            )
        except Exception:
            _ctx_hint = None
        _stream_messages = (
            [{"role": "system", "content": _ctx_hint}] + list(messages)
            if _ctx_hint else messages
        )
        if _ctx_hint:
            print(
                "[CHAT-TRACE] context_hint_injected "
                f"vault={(vault_id or '')[:8]}... "
                f"active_ctx={_ai_active_ctx}",
                flush=True,
            )

                                                                
        _create_kwargs: dict = dict(
            model=_ai_cfg().chat_model,
            messages=_stream_messages,
            stream=True,
            max_tokens=_output_cap,
        )
        if not _skip_tools and allowed_tools:
            _create_kwargs["tools"] = allowed_tools
                                                                
                                                                
            if _forced_tool_name and any(
                fn.get("function", {}).get("name") == _forced_tool_name
                for fn in allowed_tools
            ):
                _create_kwargs["tool_choice"] = {
                    "type": "function",
                    "function": {"name": _forced_tool_name},
                }
            else:
                _create_kwargs["tool_choice"] = "auto"
        response = await stream_client.chat.completions.create(
            **_create_kwargs,
        )

        tool_name = None
        tool_args_buffer = ""
        tool_call_id: Optional[str] = None
        any_content_streamed = False

                                               
        from vault_chat_safety_sanitizer import (
            sanitize_user_facing_text,
        )
        from vault_chat_save_guard import (
            guard_response_text,
            planner_intent_is_save,
            tool_call_was_successful_save,
        )
        from vault_chat_truth_guard import (
            truth_guard as _truth_guard,
            GUARDED_TOOL_NAMES as _TRUTH_GUARDED_TOOLS,
        )
        _reply_buf: list[str] = []
        _save_tool_succeeded_this_turn = False
                                                                 
                                                            
        _truth_guard_tool_name: Optional[str] = None
        _truth_guard_tool_result: Optional[str] = None

        async for chunk in response:
            delta = chunk.choices[0].delta

            if delta.tool_calls:
                tc = delta.tool_calls[0]

                if tc.function and tc.function.name:
                    tool_name = tc.function.name

                if getattr(tc, "id", None):
                    tool_call_id = tc.id

                if tc.function and tc.function.arguments:
                    tool_args_buffer += tc.function.arguments

                continue

            if delta.content:
                any_content_streamed = True
                                                              
                                                                  
                _reply_buf.append(
                    sanitize_user_facing_text(delta.content),
                )

                                                                    
        if tool_name and tool_args_buffer:
            try:
                args = json.loads(tool_args_buffer)
            except json.JSONDecodeError:
                args = {}
            logger.info("Tool called: %s", tool_name)
                                                                    
                                                                     
            _entered_find_in_vault = (tool_name == "find_in_vault")
            if _entered_find_in_vault:
                print(
                    "[CHAT-TRACE] find_in_vault_entered=true "
                    f"vault={(vault_id or '')[:8]}... "
                    f"planner_intent={_planner_intent_str} "
                    f"forced_tool={_forced_tool_name or 'none'}",
                    flush=True,
                )
            _tool_exception_class = ""
            try:
                result = await handle_tool_call(
                    tool_name, args, vault_id, key,
                    token_id=token_id,
                )
            except Exception as _exc:
                _tool_exception_class = type(_exc).__name__
                logger.exception(
                    "[CHAT-TRACE] tool_call_failed name=%s vault=%s "
                    "exception_class=%s",
                    tool_name, (vault_id or "")[:8] + "…",
                    _tool_exception_class,
                )
                result = json.dumps({
                    "error": "unavailable",
                    "exception_class": _tool_exception_class,
                })

                                                           
            if tool_call_was_successful_save(tool_name, result):
                _save_tool_succeeded_this_turn = True
                print(
                    f"[CHAT-TRACE] save_tool_succeeded name={tool_name}",
                    flush=True,
                )

                                                             
            if tool_name in _TRUTH_GUARDED_TOOLS:
                _truth_guard_tool_name = tool_name
                _truth_guard_tool_result = (
                    result if isinstance(result, str)
                    else json.dumps(result)
                )

                                                                 
            from vault_chat_safety_sanitizer import (
                sanitize_tool_result,
            )
            sanitized_for_log = sanitize_tool_result(
                tool_name, result,
            )
            logger.info(
                "[CHAT-TRACE] tool_result_sanitized name=%s "
                "raw_len=%d sanitized_len=%d",
                tool_name,
                len(result or ""),
                len(sanitized_for_log or ""),
            )

                                                                        
            _envelope_shipped = False
            if tool_name == "find_in_vault":
                                                                    
                 
                _find_result_dict: Optional[dict] = None
                _envelope_error_slug = ""
                _envelope_exception_class = ""
                try:
                    _find_result_dict = (
                        json.loads(result)
                        if isinstance(result, str) else result
                    )
                except Exception as _exc:
                    _envelope_error_slug = "tool_result_not_json"
                    _envelope_exception_class = type(_exc).__name__
                    _find_result_dict = None

                _result_is_error_band = bool(
                    isinstance(_find_result_dict, dict)
                    and _find_result_dict.get("error")
                )
                                                                   
                                                  
                if not _envelope_exception_class and isinstance(
                    _find_result_dict, dict,
                ):
                    _envelope_exception_class = str(
                        _find_result_dict.get("exception_class") or ""
                    )
                if not _envelope_exception_class and _tool_exception_class:
                    _envelope_exception_class = _tool_exception_class
                _result_has_hits_shape = bool(
                    isinstance(_find_result_dict, dict)
                    and (
                        "hits" in _find_result_dict
                        or "complete" in _find_result_dict
                    )
                )

                _envelope_str: Optional[str] = None
                try:
                    from vault_chat_result_cards import (
                        build_find_in_vault_envelope,
                        build_system_error_envelope,
                    )
                    if _result_has_hits_shape and not _result_is_error_band:
                        _envelope_str = build_find_in_vault_envelope(
                            query=last_user_message or "",
                            find_result=_find_result_dict or {},
                            key=key,
                            vault_id=vault_id,
                        )
                    else:
                                                                      
                                                                       
                        _err_slug = (
                            str(
                                _find_result_dict.get("error")
                                if isinstance(_find_result_dict, dict)
                                else _envelope_error_slug
                            )
                            or "tool_raised"
                        )
                                                                       
                                                                        
                        _envelope_str = build_system_error_envelope(
                            query=last_user_message or "",
                            query_kind=(
                                str(
                                    _find_result_dict.get("query_kind")
                                    or ""
                                ) if isinstance(_find_result_dict, dict)
                                else ""
                            ) or "generic",
                            error_slug=_err_slug,
                            exception_class=_envelope_exception_class
                                or None,
                        )
                except Exception as _exc:
                                                                      
                                                                        
                    logger.exception(
                        "[CHAT-TRACE] find_in_vault_envelope_build_raised "
                        "vault=%s exception_class=%s",
                        (vault_id or "")[:8] + "…",
                        type(_exc).__name__,
                    )
                    _envelope_str = json.dumps({
                        "type": "file_search_results",
                        "query": last_user_message or "",
                        "count": 0, "results": [],
                        "message": (
                            "Something went wrong while searching your "
                            "vault. Please try again."
                        ),
                        "is_complete": False,
                        "incomplete_reason": "envelope_build_raised",
                        "query_kind": "generic",
                        "error_band": "system_error",
                        "exception_class": type(_exc).__name__,
                    })

                if _envelope_str:
                    _reply_buf.clear()
                    _reply_buf.append(_envelope_str)
                    any_content_streamed = True
                    _envelope_shipped = True
                    _truth_guard_tool_name = None
                    _truth_guard_tool_result = None
                                                                
                                                             
                    try:
                        from vault_active_context import (
                            set_active_context,
                            CONTEXT_FILE_SEARCH_RESULTS,
                        )
                        set_active_context(
                            vault_id, CONTEXT_FILE_SEARCH_RESULTS,
                        )
                    except Exception:
                                                                 
                                                             
                        pass
                                                                  
                                                                   
                    _tele_qk = (
                        str(_find_result_dict.get("query_kind"))
                        if isinstance(_find_result_dict, dict)
                        else "unknown"
                    )
                    _tele_count = (
                        len(_find_result_dict.get("hits") or [])
                        if isinstance(_find_result_dict, dict)
                        else 0
                    )
                    _tele_complete = (
                        bool(_find_result_dict.get("complete", True))
                        if isinstance(_find_result_dict, dict)
                        and not _result_is_error_band
                        else False
                    )
                    print(
                        "[CHAT-TRACE] find_in_vault_envelope_shipped "
                        f"vault={(vault_id or '')[:8]}... "
                        f"query_kind={_tele_qk} "
                        f"count={_tele_count} "
                        f"complete={_tele_complete} "
                        f"error_band="
                        f"{'system_error' if _result_is_error_band else 'none'} "
                        f"exception_class="
                        f"{_envelope_exception_class or 'none'} "
                        f"bytes={len(_envelope_str)}",
                        flush=True,
                    )

                                                                  
            followup_messages = list(messages) + [
                {
                    "role": "assistant",
                    "content": None,
                    "tool_calls": [
                        {
                            "id":   tool_call_id or "call_0",
                            "type": "function",
                            "function": {
                                "name":      tool_name,
                                "arguments": tool_args_buffer,
                            },
                        }
                    ],
                },
                {
                    "role": "tool",
                    "tool_call_id": tool_call_id or "call_0",
                    "content": result if isinstance(result, str)
                               else json.dumps(result),
                },
            ]
            if not _envelope_shipped:
                try:
                    followup = await stream_client.chat.completions.create(
                        model=_ai_cfg().chat_model,
                        messages=followup_messages,
                        stream=True,
                        max_tokens=_output_cap,
                    )
                    async for fchunk in followup:
                        fdelta = fchunk.choices[0].delta
                        if fdelta.content:
                            any_content_streamed = True
                            _reply_buf.append(
                                sanitize_user_facing_text(fdelta.content),
                            )
                except Exception:
                    logger.exception(
                        "[CHAT-TRACE] followup_stream_failed name=%s "
                        "vault=%s", tool_name,
                        (vault_id or "")[:8] + "…",
                    )
                                                                        
                                                                   
                    _reply_buf.append(sanitized_for_log)
                    any_content_streamed = True

                                                              
        if not any_content_streamed:
            from vault_chat_safety_sanitizer import (
                SENTENCE_GENERIC_TOOL_FAILED,
            )
            _reply_buf.append(SENTENCE_GENERIC_TOOL_FAILED)

                                                                    
        _full_reply = "".join(_reply_buf)
        # 2026-07-21: gate the descriptive-prose branch of the save-
        # claim guard by whether the current turn's planner intent
        # was actually to save something. On a capability_question
        # or identity_question turn the LLM's natural answer often
        # contains descriptive phrases like "credentials are stored"
        # while explaining what the vault does — those are not
        # save-claim lies. Strong first-person claims ("I've saved
        # X") are still caught regardless. See vault_chat_save_guard.
        _intent_was_save = planner_intent_is_save(
            planner=_planner,
            allowed_tool_names={
                fn.get("function", {}).get("name")
                for fn in allowed_tools
                if isinstance(fn.get("function", {}).get("name"), str)
            },
        )
        _corrected_reply, _triggered_slug = guard_response_text(
            reply_text=_full_reply,
            save_tool_succeeded=_save_tool_succeeded_this_turn,
            intent_was_save=_intent_was_save,
        )
        if _triggered_slug is not None:
                                                               
                                                            
            print(
                f"[CHAT-TRACE] save_claim_blocked vault="
                f"{(vault_id or '')[:8]}... slug={_triggered_slug} "
                f"reply_len={len(_full_reply)}",
                flush=True,
            )
            logger.warning(
                "[CHAT-TRACE] save_claim_blocked vault=%s slug=%s "
                "reply_len=%d",
                (vault_id or "")[:8] + "...",
                _triggered_slug,
                len(_full_reply),
            )

                                                                    
        _truth_corrected, _truth_slug = _truth_guard(
            query=last_user_message,
            reply_text=_corrected_reply,
            tool_name=_truth_guard_tool_name,
            tool_result=_truth_guard_tool_result,
        )
        if _truth_slug is not None:
                                                              
                                                           
            try:
                _hit_count = 0
                if _truth_guard_tool_result:
                    _parsed = json.loads(_truth_guard_tool_result)
                    _hits = _parsed.get("hits") if isinstance(_parsed, dict) else None
                    if isinstance(_hits, list):
                        _hit_count = len(_hits)
            except Exception:
                _hit_count = -1
            print(
                f"[CHAT-TRACE] truth_guard_blocked vault="
                f"{(vault_id or '')[:8]}... slug={_truth_slug} "
                f"hit_count={_hit_count} "
                f"reply_len={len(_corrected_reply)}",
                flush=True,
            )
            logger.warning(
                "[CHAT-TRACE] truth_guard_blocked vault=%s slug=%s "
                "hit_count=%d reply_len=%d",
                (vault_id or "")[:8] + "...",
                _truth_slug,
                _hit_count,
                len(_corrected_reply),
            )
            _corrected_reply = _truth_corrected

                                                  
        try:
            from vault_chat_style_sanitizer import sanitize_reply_style
            _looks_like_envelope = (
                isinstance(_corrected_reply, str)
                and _corrected_reply.strip().startswith("{")
                and "\"type\":" in _corrected_reply[:120]
            )
            if not _looks_like_envelope:
                _styled, _style_slugs = sanitize_reply_style(
                    _corrected_reply,
                )
                if _style_slugs:
                    print(
                        "[CHAT-TRACE] style_sanitized vault="
                        f"{(vault_id or '')[:8]}... "
                        f"slugs={','.join(sorted(_style_slugs))} "
                        f"reply_len={len(_styled)}",
                        flush=True,
                    )
                if _styled:
                    _corrected_reply = _styled
        except Exception:
                                                                    
            logger.exception(
                "[CHAT-TRACE] style_sanitizer_raised vault=%s",
                (vault_id or "")[:8] + "…",
            )

        yield _corrected_reply.encode("utf-8")

    except Exception:
        logger.exception("OpenAI stream error")
        from vault_chat_safety_sanitizer import (
            SENTENCE_GENERIC_TOOL_FAILED,
        )
        yield SENTENCE_GENERIC_TOOL_FAILED.encode("utf-8")


@app.get("/")
async def root():
    return {"status": "VaultAI backend running"}


@app.get("/debug/whoami")
async def debug_whoami(request: Request):


    from device_gate import _is_dev_environment
    return {
        "ok":                 True,
        "your_origin":        request.headers.get("origin"),
        "your_user_agent":    request.headers.get("user-agent"),
        "your_x_device_id":   request.headers.get("x-device-id"),
        "your_referer":       request.headers.get("referer"),
        "cors_regex_active":  CORS_ALLOWED_ORIGIN_REGEX,
        "cors_dev_mode":      _is_dev_environment(),
        "server_time_iso":    __import__("datetime")
                              .datetime.utcnow().isoformat() + "Z",
    }


@app.get("/health")
async def health():


    try:
        conn = get_db()
        try:
            cur = conn.cursor()
            cur.execute("SELECT 1")
            cur.fetchone()
            conn.commit()
        finally:
            conn.close()
        return {
            "status": "ok",
            "db":     "connected",
        }
    except Exception as exc:
                                                                  
                                                                   
        logger.warning(
            "[HEALTH] DB readiness probe failed: %s",
            type(exc).__name__,
        )
        return JSONResponse(
            status_code=503,
            content={
                "status": "unhealthy",
                "db":     "unavailable",
                "error":  type(exc).__name__,
            },
        )


@app.post("/get-or-create-vault-meta")
async def get_or_create_vault_meta_endpoint(
    request: Request,
    payload: VaultMetaRequest,
    principal=Depends(verify_trusted_device),
):
                                                                      
                                                                        
    vault_id = principal["vault_id"]

    pin_salt = get_or_create_pin_salt(vault_id)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT vault_id, vault_name, pin_salt, pin_verifier,
                   total_bytes, created_at, kdf_iterations
            FROM vaults
            WHERE vault_id = %s
            LIMIT 1
            """,
            (vault_id,),
        )
        row = cursor.fetchone()

        if not row:
            raise HTTPException(status_code=404, detail="Vault not found")

        kdf_iterations = int(row.get("kdf_iterations") or KDF_LEGACY_ITERATIONS)

                                                                       
        logger.info(
            "[PIN-DEBUG] /get-or-create-vault-meta vault=...%s vault_name=%r "
            "pin_salt_len=%d kdf_iterations=%d pin_verifier_present_before=%s "
            "pin_len=%d",
            (vault_id or "")[-8:], row["vault_name"],
            len(pin_salt or ""), kdf_iterations,
            bool(row.get("pin_verifier")),
            len(payload.pin or ""),
        )

        if not row.get("pin_verifier"):
                                                                      
                                                                           
            if not payload.pin:
                raise HTTPException(status_code=400, detail="PIN required for first vault setup")
            if not payload.pin.isdigit():
                raise HTTPException(
                    status_code=400, detail="PIN can only contain digits.",
                )
            if len(payload.pin) > 64:
                raise HTTPException(
                    status_code=400, detail="PIN must be 64 digits or fewer.",
                )
            if len(payload.pin) < MIN_PIN_LENGTH_NEW_VAULT:
                raise HTTPException(
                    status_code=400, detail="PIN must be at least 6 digits.",
                )

                                                                            
            pin_verifier = create_pin_verifier(payload.pin, pin_salt, kdf_iterations)

            cursor.execute(
                """
                UPDATE vaults
                SET pin_verifier = %s
                WHERE vault_id = %s
                """,
                (pin_verifier, vault_id),
            )
            conn.commit()
            row["pin_verifier"] = pin_verifier
            logger.info(
                "[PIN-DEBUG] /get-or-create-vault-meta WROTE_VERIFIER "
                "vault=...%s kdf_iterations=%d "
                "pin_salt_len=%d",
                (vault_id or "")[-8:],
                kdf_iterations, len(pin_salt or ""),
            )

        return {
            "vault_name": row["vault_name"],
            "vault_id": str(row["vault_id"]),
            "pin_salt": pin_salt,
            "kdf_iterations": kdf_iterations,
            "total_bytes": int(row["total_bytes"] or 0),
            "created_at": row["created_at"].isoformat() if row.get("created_at") else None,
        }
    finally:
        conn.close()

@app.get("/vault-meta")
async def vault_meta_endpoint(
    request: Request,
    principal=Depends(verify_trusted_device),
):
    vault_id = principal["vault_id"]

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT vault_id, vault_name, pin_salt, total_bytes,
                   created_at, kdf_iterations
            FROM vaults
            WHERE vault_id = %s
            LIMIT 1
            """,
            (vault_id,),
        )
        row = cursor.fetchone()

        if not row:
            raise HTTPException(status_code=404, detail="Vault not found")

        return {
            "vault_name": row["vault_name"],
            "vault_id": str(row["vault_id"]),
            "pin_salt": row["pin_salt"],
            "kdf_iterations": int(row.get("kdf_iterations") or KDF_LEGACY_ITERATIONS),
            "total_bytes": int(row["total_bytes"] or 0),
            "created_at": row["created_at"].isoformat() if row.get("created_at") else None,
        }
    finally:
        conn.close()

@app.post("/verify-pin")
async def verify_pin_endpoint(
    payload: VerifyPinRequest,
    request: Request,
    principal=Depends(verify_trusted_device),
):


    print("[PIN-DEBUG] ROUTE_ENTRY", flush=True)
    vault_id = principal["vault_id"]

    try:
        from rate_limit_sensitive import enforce_pin_verify_rate_limit
        enforce_pin_verify_rate_limit(request, vault_id=vault_id)
    except HTTPException:
        import security_event_log as _sec_log
        _sec_log.emit(
            reason=_sec_log.REASON_RATE_LIMITED_PIN_VERIFY,
            route=_sec_log.ROUTE_GROUP_PIN,
            subject=_sec_log.short_hash(vault_id),
        )
        raise
    print(
        f"[PIN-DEBUG] AUTH_OK vault=...{str(vault_id)[-6:]}",
        flush=True,
    )
    print("[PIN-DEBUG] DEVICE_GATE_OK", flush=True)
    print(
        f"[PIN-DEBUG] VERIFY_ROUTE_ENTRY vault=...{str(vault_id)[-6:]} "
        f"vault_name={payload.vault_name!r} "
        f"pin_len={len(payload.pin or '')}",
        flush=True,
    )

    try:
        print("[PIN-DEBUG] VERIFY_START", flush=True)
                                                                           
                                                    
        verify_vault_pin(vault_id, payload.pin)
        print(
            f"[PIN-DEBUG] VERIFY_ROUTE_OK vault=...{str(vault_id)[-6:]}",
            flush=True,
        )
        return {"valid": True}
    except HTTPException as exc:
        print(
            f"[PIN-DEBUG] HTTP_FAIL status={exc.status_code} "
            f"detail={exc.detail!r}",
            flush=True,
        )
                                                                  
        if exc.status_code in (401, 423):
            detail = exc.detail if isinstance(exc.detail, dict) else {"message": str(exc.detail)}
            print(
                f"[PIN-DEBUG] VERIFY_ROUTE_REJECT vault=...{str(vault_id)[-6:]} "
                f"status={exc.status_code} code={detail.get('code')!r} "
                f"attempts_left={detail.get('attempts_left')!r}",
                flush=True,
            )
            return {
                "valid": False,
                "code": detail.get("code"),
                "message": detail.get("message"),
                "attempts_left": detail.get("attempts_left"),
            }
        print(
            f"[PIN-DEBUG] VERIFY_ROUTE_HTTPEXC vault=...{str(vault_id)[-6:]} "
            f"status={exc.status_code}",
            flush=True,
        )
        raise
    except Exception as exc:
        print(
            f"[PIN-DEBUG] VERIFY_ROUTE_EXC vault=...{str(vault_id)[-6:]} "
            f"error_type={type(exc).__name__} repr={exc!r}",
            flush=True,
        )
        raise

from datetime import datetime, timedelta, timezone


def _create_notification(
    vault_id: str,
    kind: str,
    title: str,
    body: str,
    metadata: Optional[dict] = None,
) -> None:


    try:
        from vault_core import is_vault_zk_adopted
        # ZK/adopted vault: the backend must NOT persist readable
        # notification title / body / user-derived metadata. Only
        # the structural `kind` enum is readable. The next unlock
        # of the client will encrypt any queued notification payload
        # locally and POST to /vault/ciphertext/notifications.
        # If no client finalization ever happens, the notification
        # is silently dropped — the correct privacy tradeoff for a
        # ZK vault whose user has not returned yet.
        if is_vault_zk_adopted(vault_id):
            conn = get_db()
            try:
                cursor = conn.cursor()
                # Insert a shell row with structural `kind` only.
                # The client's lazy-migration loop will pick this
                # row up on next unlock, encrypt the queued
                # server-side proposal, and finalize ciphertext.
                # Until then, the plaintext title/body do NOT touch
                # the DB.
                cursor.execute(
                    """
                    INSERT INTO notifications
                        (vault_id, kind, title, body, metadata)
                    VALUES (%s, %s, NULL, NULL, NULL)
                    """,
                    (vault_id, kind),
                )
                conn.commit()
            finally:
                conn.close()
            return

        conn = get_db()
        try:
            cursor = conn.cursor()
            cursor.execute(
                """
                INSERT INTO notifications (vault_id, kind, title, body, metadata)
                VALUES (%s, %s, %s, %s, %s)
                """,
                (vault_id, kind, title, body,
                 json.dumps(metadata) if metadata else None),
            )
            conn.commit()
        finally:
            conn.close()
    except Exception:
        logger.exception("Failed to insert notification (kind=%s)", kind)


class MarkNotificationReadRequest(BaseModel):
    notification_id: Optional[int] = None                              


class BackfillArchiveSignalsRequest(BaseModel):


    vault_name: str
    pin: str


@app.post("/vault-analysis/backfill-archive-signals")
@limiter.limit("12/hour")
async def vault_analysis_backfill_archive_signals_endpoint(
    request: Request,
    payload: BackfillArchiveSignalsRequest,
    principal=Depends(verify_trusted_device),
):


    vault_id = principal["vault_id"]

                                                            
    verify_vault_pin(vault_id, payload.pin)

    try:
        import vault_understanding as vu
        report = vu.backfill_archive_signals(vault_id)
    except Exception:
        logger.exception(
            "[backfill-archive] endpoint failed vault=%s", vault_id,
        )
        raise HTTPException(
            status_code=500,
            detail="backfill failed; check server logs",
        )

    return report


class RelatedFilesRequest(BaseModel):


    pin: str


@app.post("/files/{file_id}/related")
@limiter.limit("120/hour")
async def related_files_for_file_endpoint(
    request: Request,
    file_id: str,
    payload: RelatedFilesRequest,
    principal=Depends(verify_trusted_device),
):


    vault_id = principal["vault_id"]

                                                            
    verify_vault_pin(vault_id, payload.pin)

                                                            
    anchor_row = _load_safe_anchor_row(vault_id, file_id)
    if anchor_row is None:
        raise HTTPException(
            status_code=404,
            detail="file not found in this vault",
        )

    try:
        envelope = _compose_related_files_envelope(
            vault_id=vault_id, anchor_row=anchor_row,
        )
    except Exception:
        logger.exception(
            "[related-files] endpoint failed vault=%s file=%s",
            vault_id, file_id,
        )
        raise HTTPException(
            status_code=500,
            detail="related-files lookup failed",
        )

    return envelope


@app.get("/notifications")
@limiter.limit("120/hour")
async def notifications_list_endpoint(
    request: Request,
    principal=Depends(verify_trusted_device),
):

    vault_id = principal["vault_id"]

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT id, kind, title, body, metadata, read_at, created_at
            FROM notifications
            WHERE vault_id = %s
            ORDER BY created_at DESC
            LIMIT 100
            """,
            (vault_id,),
        )
        rows = cursor.fetchall() or []

        cursor.execute(
            """
            SELECT COUNT(*) AS c FROM notifications
            WHERE vault_id = %s AND read_at IS NULL
            """,
            (vault_id,),
        )
        unread = int((cursor.fetchone() or {}).get("c") or 0)
    finally:
        conn.close()

    return {
        "unread_count": unread,
        "notifications": [
            {
                "id": r["id"],
                "kind": r["kind"],
                "title": r["title"],
                "body": r["body"],
                "metadata": json.loads(r["metadata"]) if r["metadata"] else None,
                "read_at": r["read_at"].isoformat() if r["read_at"] else None,
                "created_at": r["created_at"].isoformat() if r["created_at"] else None,
            }
            for r in rows
        ],
    }


@app.post("/notifications/mark-read")
@limiter.limit("120/hour")
async def notifications_mark_read_endpoint(
    request: Request,
    payload: MarkNotificationReadRequest,
    principal=Depends(verify_trusted_device),
):

    vault_id = principal["vault_id"]

    conn = get_db()
    try:
        cursor = conn.cursor()
        if payload.notification_id is None:
            cursor.execute(
                """
                UPDATE notifications
                SET read_at = NOW()
                WHERE vault_id = %s AND read_at IS NULL
                """,
                (vault_id,),
            )
        else:
            cursor.execute(
                """
                UPDATE notifications
                SET read_at = NOW()
                WHERE vault_id = %s AND id = %s AND read_at IS NULL
                """,
                (vault_id, payload.notification_id),
            )
        conn.commit()
    finally:
        conn.close()

    return {"ok": True}


from vault_core import (
    derive_pairing_wrap_key,
    generate_pairing_code,
    hash_pairing_code,
    PAIRING_CODE_TTL_MINUTES,
)


@app.post("/beneficiary/create")
@limiter.limit("10/hour")
async def beneficiary_create_endpoint(
    request: Request,
    payload: CreateBeneficiaryRequest,
    principal=Depends(verify_trusted_device),
):


    vault_id = principal["vault_id"]

    label = (payload.label or "").strip()
    if not label or len(label) > 60:
        raise HTTPException(
            status_code=400,
            detail="Beneficiary label is required and must be 1-60 characters.",
        )

    passer_vault_key = verify_vault_pin(vault_id, payload.pin)

    pairing_code = generate_pairing_code()
    code_hash = hash_pairing_code(pairing_code)
    wrap_key = derive_pairing_wrap_key(pairing_code)
    wrapped = encrypt_bytes(passer_vault_key, wrap_key)
    expires_at = datetime.now(timezone.utc) + timedelta(minutes=PAIRING_CODE_TTL_MINUTES)

    # ZK/adopted vault: never persist the readable passer_label.
    # Client is expected to POST the ciphertext label to
    # /vault/ciphertext/beneficiary-links immediately after
    # /beneficiary/create returns. Until finalized the row has
    # passer_label = NULL. Legacy un-adopted vaults keep the
    # plaintext write.
    from vault_core import is_vault_zk_adopted
    _zk = is_vault_zk_adopted(vault_id)
    _stored_label = None if _zk else label

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            INSERT INTO beneficiary_links
              (passer_vault_id, passer_label,
               pairing_code_hash, pairing_expires_at,
               wrapped_vault_key, status)
            VALUES (%s, %s, %s, %s, %s, 'pairing_pending')
            RETURNING id
            """,
            (vault_id, _stored_label, code_hash, expires_at, wrapped),
        )
        row = cursor.fetchone()
        conn.commit()
    finally:
        conn.close()

    return {
        "link_id": row["id"],
        "label": label,
        "pairing_code": pairing_code,
        "expires_at": expires_at.isoformat(),
        "ttl_minutes": PAIRING_CODE_TTL_MINUTES,
    }


@app.post("/beneficiary/link")
@limiter.limit("10/hour")
async def beneficiary_link_endpoint(
    request: Request,
    payload: LinkBeneficiaryRequest,
    principal=Depends(verify_trusted_device),
):


    vault_id = principal["vault_id"]

    pairing_code = (payload.pairing_code or "").strip()
    if not pairing_code:
        raise HTTPException(status_code=400, detail="Pairing code is required.")

    beneficiary_vault_key = verify_vault_pin(vault_id, payload.pin)

    code_hash = hash_pairing_code(pairing_code)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT id, passer_vault_id, passer_label,
                   pairing_expires_at, wrapped_vault_key, status
            FROM beneficiary_links
            WHERE pairing_code_hash = %s
            LIMIT 1
            """,
            (code_hash,),
        )
        row = cursor.fetchone()

        if not row:
            raise HTTPException(status_code=404, detail="Invalid or expired pairing code.")

        if row["status"] != "pairing_pending":
            raise HTTPException(status_code=400, detail="Pairing code already used.")

        cursor.execute("SELECT NOW() AS now")
        now = (cursor.fetchone() or {}).get("now")
        if row["pairing_expires_at"] and now and row["pairing_expires_at"] < now:
            raise HTTPException(status_code=410, detail="Pairing code expired.")

        if str(row["passer_vault_id"]) == vault_id:
            raise HTTPException(status_code=400, detail="Cannot link a vault to itself.")

                                                                                                  
        wrap_key = derive_pairing_wrap_key(pairing_code)
        try:
            passer_vault_key = decrypt_bytes(row["wrapped_vault_key"], wrap_key)
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid pairing code.")

        new_wrapped = encrypt_bytes(passer_vault_key, beneficiary_vault_key)

        cursor.execute(
            """
            UPDATE beneficiary_links
            SET beneficiary_vault_id = %s,
                pairing_code_hash = NULL,
                pairing_expires_at = NULL,
                wrapped_vault_key = %s,
                status = 'linked',
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
            """,
            (vault_id, new_wrapped, row["id"]),
        )
        conn.commit()
    finally:
        conn.close()

    return {
        "link_id": row["id"],
        "passer_label": row["passer_label"],
        "status": "linked",
    }


@app.post("/beneficiary/list-mine")
@limiter.limit("60/hour")
async def beneficiary_list_mine_endpoint(
    request: Request,
    payload: BeneficiaryListRequest,
    principal=Depends(verify_trusted_device),
):

    vault_id = principal["vault_id"]
    verify_vault_pin(vault_id, payload.pin)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        # LEFT JOIN inheritance_credentials so a row appears with
        # ``credentials_saved = False`` when no active package
        # exists — the client uses this to decide whether to show
        # "Add credentials" or "Update / Delete credentials".
        # ``pairing_state`` is the operator-visible state
        # introduced in migration 0028; back-compat default is
        # ``paired_no_credentials`` for rows that predate the
        # migration.
        cursor.execute(
            """
            SELECT bl.id,
                   bl.passer_label,
                   bl.status,
                   bl.beneficiary_vault_id IS NOT NULL AS is_linked,
                   bl.pairing_expires_at,
                   bl.transfer_requested_at,
                   bl.transfer_executes_at,
                   bl.created_at,
                   COALESCE(bl.pairing_state, 'paired_no_credentials')
                       AS pairing_state,
                   bl.access_requested_at,
                   bl.cooldown_ends_at,
                   bl.decision_at,
                   (ic.id IS NOT NULL) AS credentials_saved,
                   ic.crypto_version   AS credential_crypto_version,
                   ic.updated_at       AS credential_updated_at
              FROM beneficiary_links bl
              LEFT JOIN inheritance_credentials ic
                     ON ic.beneficiary_link_id = bl.id
                    AND ic.deleted_at IS NULL
             WHERE bl.passer_vault_id = %s
             ORDER BY bl.created_at DESC
            """,
            (vault_id,),
        )
        rows = cursor.fetchall() or []
    finally:
        conn.close()

    def iso(v):
        return v.isoformat() if v else None

    return {
        "beneficiaries": [
            {
                "id": r["id"],
                "label": r["passer_label"],
                "status": r["status"],
                "pairing_state": r.get("pairing_state") or "paired_no_credentials",
                "is_linked": bool(r["is_linked"]),
                "pairing_expires_at": iso(r["pairing_expires_at"]),
                "transfer_requested_at": iso(r["transfer_requested_at"]),
                "transfer_executes_at": iso(r["transfer_executes_at"]),
                "created_at": iso(r["created_at"]),
                "credentials_saved": bool(r.get("credentials_saved")),
                "credential_crypto_version": (
                    int(r["credential_crypto_version"])
                    if r.get("credential_crypto_version") is not None
                    else None
                ),
                "credential_updated_at": iso(r.get("credential_updated_at")),
                "access_requested_at": iso(r.get("access_requested_at")),
                "cooldown_ends_at": iso(r.get("cooldown_ends_at")),
                "decision_at": iso(r.get("decision_at")),
            }
            for r in rows
        ]
    }


@app.post("/beneficiary/list-inheritances")
@limiter.limit("60/hour")
async def beneficiary_list_inheritances_endpoint(
    request: Request,
    payload: BeneficiaryListRequest,
    principal=Depends(verify_trusted_device),
):

    vault_id = principal["vault_id"]
    verify_vault_pin(vault_id, payload.pin)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        # Include the escrow state (Phase 2) so the beneficiary UI
        # can decide whether to show Request access / Cancel / Claim
        # / Reveal. ``credentials_saved`` is set by the LEFT JOIN
        # on ``inheritance_credentials`` (soft-delete respected).
        #
        # 2026-07-23: also expose the owner's chosen ``vault_name``
        # so the beneficiary UI can render a human-readable owner
        # identity even when ``passer_label`` is NULL (which happens
        # for every ZK-created pairing — the plaintext label is
        # encrypted under the OWNER's metadataKey and stored in
        # ``passer_label_ciphertext``, which the beneficiary cannot
        # decrypt). The vault_name is already server-visible product
        # metadata (returned by /auth/me to the owner) and the
        # beneficiary already holds a pairing code they exchanged
        # with the owner, so surfacing it here does not leak new
        # information. It is nullable — a pre-migration-0031 ZK
        # account that never adopted a real vault_name returns NULL
        # and the frontend falls back to 'Unknown'.
        cursor.execute(
            """
            SELECT bl.id,
                   bl.passer_label,
                   ov.vault_name AS owner_vault_name,
                   bl.status,
                   bl.transfer_requested_at,
                   bl.transfer_executes_at,
                   bl.created_at,
                   COALESCE(bl.pairing_state, 'paired_no_credentials')
                       AS pairing_state,
                   bl.access_requested_at,
                   bl.cooldown_ends_at,
                   bl.decision_at,
                   (ic.id IS NOT NULL) AS credentials_saved
              FROM beneficiary_links bl
              LEFT JOIN inheritance_credentials ic
                     ON ic.beneficiary_link_id = bl.id
                    AND ic.deleted_at IS NULL
              LEFT JOIN vaults ov
                     ON ov.vault_id = bl.passer_vault_id
             WHERE bl.beneficiary_vault_id = %s
             ORDER BY bl.created_at DESC
            """,
            (vault_id,),
        )
        rows = cursor.fetchall() or []
    finally:
        conn.close()

    def iso(v):
        return v.isoformat() if v else None

    from datetime import datetime, timezone as _tz
    server_now_iso = datetime.now(_tz.utc).isoformat()

    return {
        "server_now": server_now_iso,
        "inheritances": [
            {
                "id": r["id"],
                "passer_label": r["passer_label"],
                # 2026-07-23: owner's chosen vault_name, nullable.
                # Frontend uses this as a fallback when passer_label
                # is NULL (every ZK-created pairing).
                "owner_vault_name": r.get("owner_vault_name"),
                "status": r["status"],
                "pairing_state": r.get("pairing_state") or "paired_no_credentials",
                "credentials_saved": bool(r.get("credentials_saved")),
                "access_requested_at": iso(r.get("access_requested_at")),
                "cooldown_ends_at": iso(r.get("cooldown_ends_at")),
                "decision_at": iso(r.get("decision_at")),
                "transfer_requested_at": iso(r["transfer_requested_at"]),
                "transfer_executes_at": iso(r["transfer_executes_at"]),
                "created_at": iso(r["created_at"]),
            }
            for r in rows
        ],
    }


TRANSFER_COOLDOWN_DAYS = int(os.getenv("TRANSFER_COOLDOWN_DAYS", "30"))


@app.post("/beneficiary/request-transfer")
@limiter.limit("10/hour")
async def beneficiary_request_transfer_endpoint(
    request: Request,
    payload: TransferActionRequest,
    principal=Depends(verify_trusted_device),
):

    vault_id = principal["vault_id"]
    verify_vault_pin(vault_id, payload.pin)

    requested_at = datetime.now(timezone.utc)
    executes_at = requested_at + timedelta(days=TRANSFER_COOLDOWN_DAYS)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            UPDATE beneficiary_links
            SET status = 'transfer_pending',
                transfer_requested_at = %s,
                transfer_executes_at = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
              AND beneficiary_vault_id = %s
              AND status = 'linked'
            RETURNING id, passer_vault_id, passer_label
            """,
            (requested_at, executes_at, payload.link_id, vault_id),
        )
        row = cursor.fetchone()
        conn.commit()
    finally:
        conn.close()

    if not row:
        raise HTTPException(
            status_code=404,
            detail="Inheritance link not found or not in a state where transfer can be requested.",
        )

    _create_notification(
        vault_id=str(row["passer_vault_id"]),
        kind="transfer_requested",
        title=f'Transfer requested for "{row["passer_label"]}"',
        body=(
            f'Your beneficiary linked to "{row["passer_label"]}" has '
            f"requested transfer of its contents. If you're alive and well, "
            f"open the Inheritance page and click \"Cancel transfer\" before "
            f"{executes_at.strftime('%Y-%m-%d %H:%M UTC')} "
            f"({TRANSFER_COOLDOWN_DAYS} days from now). After that, the "
            f"transfer executes automatically."
        ),
        metadata={"link_id": row["id"], "executes_at": executes_at.isoformat()},
    )

    return {
        "status": "transfer_pending",
        "transfer_requested_at": requested_at.isoformat(),
        "transfer_executes_at": executes_at.isoformat(),
    }


@app.post("/beneficiary/cancel-transfer")
@limiter.limit("20/hour")
async def beneficiary_cancel_transfer_endpoint(
    request: Request,
    payload: TransferActionRequest,
    principal=Depends(verify_trusted_device),
):

    vault_id = principal["vault_id"]
    verify_vault_pin(vault_id, payload.pin)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            UPDATE beneficiary_links
            SET status = 'linked',
                transfer_requested_at = NULL,
                transfer_executes_at = NULL,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
              AND passer_vault_id = %s
              AND status = 'transfer_pending'
            RETURNING id, passer_label, beneficiary_vault_id
            """,
            (payload.link_id, vault_id),
        )
        row = cursor.fetchone()
        conn.commit()
    finally:
        conn.close()

    if not row:
        raise HTTPException(
            status_code=404,
            detail="No pending transfer found for this link.",
        )

    if row.get("beneficiary_vault_id"):
        _create_notification(
            vault_id=str(row["beneficiary_vault_id"]),
            kind="transfer_cancelled",
            title=f'Transfer cancelled for "{row["passer_label"]}"',
            body=(
                f'The owner of "{row["passer_label"]}" cancelled the pending '
                f"transfer. The inheritance link is still active — you can "
                f"request transfer again later if needed."
            ),
            metadata={"link_id": row["id"]},
        )

    return {"status": "linked"}


VAULT_FREEZE_DAYS = int(os.getenv("VAULT_FREEZE_DAYS", "90"))


def _generate_inherited_vault_name(
    base_name: str,
) -> str:


    base = (base_name or "Inherited vault").strip()[:50] or "Inherited vault"
    candidate = f"{base} (inherited)"
    suffix = 1

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        while True:
            cursor.execute(
                """
                SELECT 1 FROM vaults
                WHERE LOWER(vault_name) = LOWER(%s)
                LIMIT 1
                """,
                (candidate,),
            )
            if cursor.fetchone() is None:
                return candidate
            suffix += 1
            candidate = f"{base} (inherited {suffix})"
            if suffix > 99:               
                return f"{base} ({uuid.uuid4().hex[:6]})"
    finally:
        conn.close()


@app.post("/beneficiary/claim-transfer")
@limiter.limit("5/hour")
async def beneficiary_claim_transfer_endpoint(
    request: Request,
    payload: ClaimTransferRequest,
    principal=Depends(verify_trusted_device),
):


    vault_id = principal["vault_id"]

    beneficiary_vault_key = verify_vault_pin(vault_id, payload.pin)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)

        cursor.execute(
            """
            SELECT id, passer_vault_id, passer_label,
                   beneficiary_vault_id,
                   wrapped_vault_key, status, transfer_executes_at
            FROM beneficiary_links
            WHERE id = %s
            LIMIT 1
            """,
            (payload.link_id,),
        )
        link = cursor.fetchone()
        if not link:
            raise HTTPException(status_code=404, detail="Inheritance link not found.")
        if str(link["beneficiary_vault_id"]) != vault_id:
            raise HTTPException(status_code=403, detail="This inheritance does not belong to you.")
        if link["status"] != "transfer_pending":
            raise HTTPException(status_code=400, detail="No pending transfer to claim.")

                                                                      
        cursor.execute("SELECT (NOW() AT TIME ZONE 'UTC') AS now")
        now = (cursor.fetchone() or {}).get("now")
        if not link["transfer_executes_at"] or (now and link["transfer_executes_at"] > now):
            raise HTTPException(
                status_code=400,
                detail="Cooldown not elapsed yet — transfer not ready to claim.",
            )

                                                                                          
        try:
            passer_vault_key = decrypt_bytes(link["wrapped_vault_key"], beneficiary_vault_key)
        except Exception:
            raise HTTPException(status_code=500, detail="Failed to unwrap inherited key.")

        passer_vault_id = str(link["passer_vault_id"])
        passer_label = link["passer_label"]

                                                                    
        requested_name = (payload.inherited_vault_name or "").strip()
        if requested_name:
            cursor.execute(
                "SELECT 1 FROM vaults WHERE LOWER(vault_name) = LOWER(%s) LIMIT 1",
                (requested_name,),
            )
            if cursor.fetchone() is not None:
                raise HTTPException(
                    status_code=409,
                    detail=f"A vault named \"{requested_name}\" already exists. Pick another.",
                )
            new_vault_name = requested_name
        else:
            new_vault_name = _generate_inherited_vault_name(passer_label)

                                                                     
        cursor.execute(
            "SELECT account_id FROM vaults WHERE vault_id = %s LIMIT 1",
            (vault_id,),
        )
        ben_acct_row = cursor.fetchone()
        if not ben_acct_row or not ben_acct_row.get("account_id"):
            raise HTTPException(
                status_code=500,
                detail="Beneficiary vault has no account binding; transfer aborted.",
            )
        new_account_id = ben_acct_row["account_id"]

                                                                  
        new_pin_salt = generate_pin_salt()
        new_kdf_iterations = KDF_TARGET_ITERATIONS
        new_inherited_key = derive_key(
            payload.pin, new_pin_salt, iterations=new_kdf_iterations,
        )
        new_pin_verifier = encrypt_message(
            "vaultai_pin_ok", new_inherited_key,
        )

        new_vault_id = str(uuid.uuid4())
        cursor.execute(
            """
            INSERT INTO vaults
              (vault_id, vault_name, account_id, pin_salt, pin_verifier,
               kdf_iterations, inherited_from_label, total_bytes)
            VALUES (%s, %s, %s, %s, %s, %s, %s, 0)
            """,
            (new_vault_id, new_vault_name, new_account_id,
             new_pin_salt, new_pin_verifier,
             new_kdf_iterations, passer_label),
        )

                                                 
        cursor.execute(
            """
            SELECT item_type, service, encrypted_data
            FROM vault_items
            WHERE vault_id = %s
            """,
            (passer_vault_id,),
        )
        item_rows = cursor.fetchall() or []
        new_total_bytes = 0
        for item in item_rows:
            plaintext = decrypt_message(item["encrypted_data"], passer_vault_key)
            new_blob = encrypt_message(plaintext, new_inherited_key)
            new_total_bytes += len(new_blob.encode("utf-8"))
            cursor.execute(
                """
                INSERT INTO vault_items
                  (vault_id, item_type, service, encrypted_data)
                VALUES (%s, %s, %s, %s)
                RETURNING id
                """,
                (new_vault_id, item["item_type"], item["service"], new_blob),
            )
            _inh_row = cursor.fetchone()
            _inh_item_id = (
                _inh_row["id"] if isinstance(_inh_row, dict)
                else (_inh_row[0] if _inh_row else None)
            )
                                                                            
            if _inh_item_id is not None:
                from semantic_embedder import enqueue_vault_item_embedding
                enqueue_vault_item_embedding(client, new_vault_id, _inh_item_id, "item_service", item["service"])
                enqueue_vault_item_embedding(client, new_vault_id, _inh_item_id, "item_type", item["item_type"])
                                                                                 
                try:
                    from asset_tagger import tag_vault_item_safe
                    tag_vault_item_safe(new_vault_id, _inh_item_id,
                                        service=item["service"], item_type=item["item_type"])
                except Exception:
                    pass

                                                                  
        cursor.execute(
            """
            SELECT id, file_name, content_type, file_size,
                   encrypted_file_data, extracted_text, extracted_text_encrypted,
                   detected_type, detected_service, autosaved_secret,
                   saved_name, asset_type, needs_naming,
                   storage_mode, chunk_size, chunk_count, upload_status
            FROM uploaded_files
            WHERE vault_id = %s
              AND upload_status = 'complete'
            """,
            (passer_vault_id,),
        )
        file_rows = cursor.fetchall() or []
        for f in file_rows:
            storage_mode = (f["storage_mode"] or "inline")
            new_file_id = str(uuid.uuid4())

                                                                         
            new_extracted = f["extracted_text"]
            new_extracted_encrypted = bool(f["extracted_text_encrypted"])
            if new_extracted_encrypted and new_extracted:
                try:
                    plain_text = decrypt_message(new_extracted, passer_vault_key)
                    new_extracted = encrypt_message(plain_text, new_inherited_key)
                except Exception:
                                                                                
                    pass

            if storage_mode == "inline":
                                                             
                raw_bytes = decrypt_bytes(f["encrypted_file_data"], passer_vault_key)
                new_file_blob = encrypt_bytes(raw_bytes, new_inherited_key)
                new_total_bytes += len(new_file_blob.encode("utf-8"))

                cursor.execute(
                    """
                    INSERT INTO uploaded_files
                      (id, vault_id, file_name, content_type, file_size,
                       encrypted_file_data, extracted_text, extracted_text_encrypted,
                       detected_type, detected_service, autosaved_secret,
                       saved_name, asset_type, needs_naming,
                       storage_mode, upload_status)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                            'inline', 'complete')
                    """,
                    (
                        new_file_id,
                        new_vault_id,
                        f["file_name"], f["content_type"], f["file_size"],
                        new_file_blob, new_extracted, new_extracted_encrypted,
                        f["detected_type"], f["detected_service"], f["autosaved_secret"],
                        f["saved_name"], f["asset_type"], f["needs_naming"],
                    ),
                )
                                                                                        
                from semantic_embedder import enqueue_uploaded_file_embedding
                enqueue_uploaded_file_embedding(client, new_vault_id, new_file_id, "file_name", f["file_name"])
                enqueue_uploaded_file_embedding(client, new_vault_id, new_file_id, "saved_name", f["saved_name"])
                enqueue_uploaded_file_embedding(client, new_vault_id, new_file_id, "asset_type", f["asset_type"])
                enqueue_uploaded_file_embedding(client, new_vault_id, new_file_id, "detected_service", f["detected_service"])
                                                                                          
                try:
                    from asset_tagger import tag_uploaded_file_safe
                    tag_uploaded_file_safe(
                        new_vault_id, new_file_id,
                        file_name=f["file_name"], saved_name=f["saved_name"],
                        asset_type=f["asset_type"], content_type=f["content_type"],
                        detected_service=f["detected_service"],
                    )
                except Exception:
                    pass
                                                                                
                try:
                    from document_understanding import extract_and_store_document_safe
                    extract_and_store_document_safe(
                        new_vault_id, new_file_id,
                        file_name=f["file_name"], saved_name=f["saved_name"],
                        asset_type=f["asset_type"],
                        content_type=f["content_type"],
                        detected_service=f["detected_service"],
                        extracted_text=None,
                        extracted_text_encrypted=True,
                    )
                except Exception:
                    pass
                continue

            if storage_mode != "chunks":
                                                                      
                                                                
                raise HTTPException(
                    status_code=500,
                    detail=(
                        f"Unknown storage_mode {storage_mode!r} for file "
                        f"{f['file_name']!r}. Inheritance aborted."
                    ),
                )

                                      
            chunk_size = int(f["chunk_size"] or 0)
            chunk_count = int(f["chunk_count"] or 0)
            file_size = int(f["file_size"] or 0)
            if chunk_count <= 0 or chunk_size <= 0:
                raise HTTPException(
                    status_code=500,
                    detail=(
                        f"Chunked file {f['file_name']!r} has invalid "
                        f"chunk_size/chunk_count. Inheritance aborted."
                    ),
                )

                                                                   
            cursor.execute(
                """
                INSERT INTO uploaded_files
                  (id, vault_id, file_name, content_type, file_size,
                   encrypted_file_data, extracted_text, extracted_text_encrypted,
                   detected_type, detected_service, autosaved_secret,
                   saved_name, asset_type, needs_naming,
                   storage_mode, chunk_size, chunk_count, upload_status)
                VALUES (%s, %s, %s, %s, %s, NULL, %s, %s, %s, %s, %s, %s, %s, %s,
                        'chunks', %s, %s, 'complete')
                """,
                (
                    new_file_id,
                    new_vault_id,
                    f["file_name"], f["content_type"], file_size,
                    new_extracted, new_extracted_encrypted,
                    f["detected_type"], f["detected_service"], f["autosaved_secret"],
                    f["saved_name"], f["asset_type"], f["needs_naming"],
                    chunk_size, chunk_count,
                ),
            )
                                                                                     
            from semantic_embedder import enqueue_uploaded_file_embedding
            enqueue_uploaded_file_embedding(client, new_vault_id, new_file_id, "file_name", f["file_name"])
            enqueue_uploaded_file_embedding(client, new_vault_id, new_file_id, "saved_name", f["saved_name"])
            enqueue_uploaded_file_embedding(client, new_vault_id, new_file_id, "asset_type", f["asset_type"])
            enqueue_uploaded_file_embedding(client, new_vault_id, new_file_id, "detected_service", f["detected_service"])
                                                                                       
            try:
                from asset_tagger import tag_uploaded_file_safe
                tag_uploaded_file_safe(
                    new_vault_id, new_file_id,
                    file_name=f["file_name"], saved_name=f["saved_name"],
                    asset_type=f["asset_type"], content_type=f["content_type"],
                    detected_service=f["detected_service"],
                )
            except Exception:
                pass
                                                                             
            try:
                from document_understanding import extract_and_store_document_safe
                extract_and_store_document_safe(
                    new_vault_id, new_file_id,
                    file_name=f["file_name"], saved_name=f["saved_name"],
                    asset_type=f["asset_type"],
                    content_type=f["content_type"],
                    detected_service=f["detected_service"],
                    extracted_text=None,
                    extracted_text_encrypted=True,
                )
            except Exception:
                pass

                                                                      
            for chunk_index in range(chunk_count):
                cursor.execute(
                    """
                    SELECT chunk_bytes
                    FROM uploaded_file_chunks
                    WHERE file_id = %s AND chunk_index = %s
                    """,
                    (f["id"], chunk_index),
                )
                crow = cursor.fetchone()
                if not crow:
                    raise HTTPException(
                        status_code=500,
                        detail=(
                            f"Chunk {chunk_index} missing for inherited file "
                            f"{f['file_name']!r}. Inheritance aborted."
                        ),
                    )

                try:
                    plain = decrypt_chunk(
                        bytes(crow["chunk_bytes"]),
                        passer_vault_key,
                        chunk_index,
                    )
                except InvalidTag:
                    raise HTTPException(
                        status_code=500,
                        detail=(
                            f"Chunk {chunk_index} of file {f['file_name']!r} "
                            f"failed authentication. Inheritance aborted."
                        ),
                    )

                new_frame = encrypt_chunk(plain, new_inherited_key, chunk_index)

                cursor.execute(
                    """
                    INSERT INTO uploaded_file_chunks (file_id, chunk_index, chunk_bytes)
                    VALUES (%s, %s, %s)
                    """,
                    (new_file_id, chunk_index, psycopg2.Binary(new_frame)),
                )

                                                                           
            new_total_bytes += file_size

                                             
        cursor.execute(
            """
            UPDATE vaults
            SET total_bytes = %s
            WHERE vault_id = %s
            """,
            (new_total_bytes, new_vault_id),
        )

                                             
        frozen_until = datetime.now(timezone.utc) + timedelta(days=VAULT_FREEZE_DAYS)
        cursor.execute(
            """
            UPDATE vaults
            SET must_reset = TRUE,
                frozen_until = %s,
                failed_pin_attempts = 0
            WHERE vault_id = %s
            """,
            (frozen_until, passer_vault_id),
        )

                                                
        cursor.execute(
            """
            UPDATE beneficiary_links
            SET status = 'transferred',
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
            """,
            (link["id"],),
        )

        conn.commit()
    except HTTPException:
        conn.rollback()
        raise
    except Exception:
        conn.rollback()
        logger.exception("Transfer claim failed")
        raise HTTPException(status_code=500, detail="Transfer claim failed.")
    finally:
        conn.close()

    logger.warning(
        "Transfer claimed by vault=%s: %d items + %d files copied from "
        "passer vault=%s into new vault=%s (%s)",
        vault_id, len(item_rows), len(file_rows), passer_vault_id, new_vault_name, new_vault_id,
    )

                                           
    try:
        from relationship_builder import rebuild_vault_relationships_safe
        rebuild_vault_relationships_safe(new_vault_id)
    except Exception:
        pass

                                                                        
    try:
        from expiry_engine import rebuild_expiry_alerts_safe
        rebuild_expiry_alerts_safe(new_vault_id)
    except Exception:
        pass

                                                              
    try:
        from document_entities import rebuild_entities_for_vault_safe
        rebuild_entities_for_vault_safe(new_vault_id)
    except Exception:
        pass

    _create_notification(
        vault_id=passer_vault_id,
        kind="transfer_completed",
        title=f'Your vault "{passer_label}" has been transferred',
        body=(
            f"The 30-day cooldown elapsed without cancellation, and your "
            f'beneficiary has claimed "{passer_label}". Your vault has been '
            f"frozen for {VAULT_FREEZE_DAYS} days. If this was a mistake, "
            f"sign in and contact support."
        ),
        metadata={"link_id": link["id"], "passer_label": passer_label},
    )

    return {
        "claimed": True,
        "inherited_vault_name": new_vault_name,
        "pin_salt": new_pin_salt,
        "kdf_iterations": new_kdf_iterations,
        "items_copied": len(item_rows),
        "files_copied": len(file_rows),
    }


@app.get("/list-my-vaults")
async def list_my_vaults_endpoint(
    request: Request,
    principal=Depends(verify_trusted_device),
):


    vault_id = principal["vault_id"]

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT vault_name, inherited_from_label, must_reset,
                   frozen_until, total_bytes, created_at,
                   pin_verifier IS NOT NULL AS has_pin
            FROM vaults
            WHERE account_id = (
                SELECT account_id FROM vaults WHERE vault_id = %s
            )
            ORDER BY created_at ASC
            """,
            (vault_id,),
        )
        rows = cursor.fetchall() or []
    finally:
        conn.close()

    return {
        "vaults": [
            {
                "vault_name": r["vault_name"],
                "inherited_from_label": r["inherited_from_label"],
                "must_reset": bool(r["must_reset"]),
                "frozen_until": r["frozen_until"].isoformat() if r["frozen_until"] else None,
                "total_bytes": int(r["total_bytes"] or 0),
                "has_pin": bool(r["has_pin"]),
                "created_at": r["created_at"].isoformat() if r.get("created_at") else None,
            }
            for r in rows
        ]
    }


@app.post("/beneficiary/delete")
@limiter.limit("20/hour")
async def beneficiary_delete_endpoint(
    request: Request,
    payload: DeleteBeneficiaryRequest,
    principal=Depends(verify_trusted_device),
):

    vault_id = principal["vault_id"]
    verify_vault_pin(vault_id, payload.pin)

    conn = get_db()
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            DELETE FROM beneficiary_links
            WHERE id = %s
              AND passer_vault_id = %s
            """,
            (payload.link_id, vault_id),
        )
        affected = cursor.rowcount
        conn.commit()
    finally:
        conn.close()

    if not affected:
        raise HTTPException(status_code=404, detail="Beneficiary link not found.")

    return {"deleted": True}


@app.post("/rotate-vault-kdf")
@limiter.limit("5/minute")
async def rotate_vault_kdf_endpoint(
    request: Request,
    payload: RotateVaultKdfRequest,
    principal=Depends(verify_trusted_device),
):


    vault_id = principal["vault_id"]

                                                                    
    old_key = verify_vault_pin(vault_id, payload.pin)

    result = rotate_vault_kdf_if_needed(
        vault_id, payload.pin, old_key,
    )

    return {
        "rotated": result["rotated"],
        "pin_salt": result["pin_salt"],
        "kdf_iterations": result["kdf_iterations"],
    }


@app.post("/wipe-orphan-data")
@limiter.limit("5/minute")
async def wipe_orphan_data_endpoint(
    request: Request,
    payload: WipeOrphanRequest,
    principal=Depends(verify_trusted_device),
):


    if not payload.confirm:
        raise HTTPException(status_code=400, detail="confirm=true required")
    return {"status": "noop", "items_deleted": 0, "files_deleted": 0}


@app.post("/vault-stats")
async def vault_stats_endpoint(
    request: Request,
    payload: VaultStatsRequest,
    principal=Depends(verify_trusted_device),
):
    vault_id = principal["vault_id"]

    verify_vault_pin(vault_id, payload.pin)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)

        cursor.execute(
            """
            SELECT COALESCE(total_bytes, 0) AS total_bytes
            FROM vaults
            WHERE vault_id = %s
            """,
            (vault_id,),
        )
        vault_row = cursor.fetchone()
        storage_used_bytes = int(vault_row["total_bytes"]) if vault_row else 0

        cursor.execute(
            """
            SELECT
                COUNT(*) FILTER (WHERE item_type = 'login') AS login_count,
                COUNT(*) FILTER (WHERE item_type = 'card') AS card_count,
                COUNT(*) FILTER (WHERE item_type = 'id') AS id_count
            FROM vault_items
            WHERE vault_id = %s
            """,
            (vault_id,),
        )
        item_counts = cursor.fetchone() or {}

        cursor.execute(
            """
            SELECT COUNT(*) AS file_count
            FROM uploaded_files
            WHERE vault_id = %s
              AND upload_status = 'complete'
            """,
            (vault_id,),
        )
        file_row = cursor.fetchone() or {}

                                                                 
        cursor.execute(
            """
            SELECT COALESCE(SUM(file_size), 0)::BIGINT AS pending_bytes
            FROM uploaded_files
            WHERE vault_id = %s
              AND upload_status = 'uploading'
            """,
            (vault_id,),
        )
        pending_row = cursor.fetchone() or {}

                                                                   
        from billing import (
            get_account_id_for_vault as _get_acct_id,
            get_effective_storage_limit as _get_eff_limit,
        )
        _account_id = _get_acct_id(vault_id)
        storage_limit_bytes = _get_eff_limit(_account_id)

        return {
            "login_count": int(item_counts.get("login_count") or 0),
            "card_count": int(item_counts.get("card_count") or 0),
            "id_count": int(item_counts.get("id_count") or 0),
            "file_count": int(file_row.get("file_count") or 0),
            "storage_used_bytes": storage_used_bytes,
            "storage_limit_bytes": storage_limit_bytes,
            "upload_safety_cap_bytes": MAX_UPLOAD_BYTES,
            "storage_pending_bytes": int(pending_row.get("pending_bytes") or 0),
        }
    finally:
        conn.close()


@app.post("/vault-name-available")
@limiter.limit("10/minute")
async def check_vault_name_available(
    request: Request,
    payload: VaultNameCheck,
    principal=Depends(verify_trusted_device),
):
                                                                            
                                                                  
    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT 1
            FROM vaults
            WHERE LOWER(vault_name) = LOWER(%s)
            """,
            (payload.vault_name,),
        )
        available = cursor.fetchone() is None
    finally:
        conn.close()

    return {"available": available}

@app.post("/upload-file")
@limiter.limit("10/minute")
async def upload_file_endpoint(
    request: Request,
    background_tasks: BackgroundTasks,
    vault_name: str = Form(...),
    pin: str = Form(...),
    content_type: Optional[str] = Form(None),
    accompanying_text: Optional[str] = Form(None),
                                                              
                                                                  
    is_batch_upload: bool = Form(False),
                                                             
                                                                   
    relative_path: Optional[str] = Form(None),
                                                                       
                                                                     
    import_id: Optional[str] = Form(None),
                                                                  
                                                                   
    content_sha256: Optional[str] = Form(None),
    duplicate_action: Optional[str] = Form(None),
    # ZK ciphertext-first mode. When supplied, save_uploaded_file
    # writes the row with readable metadata columns NULL from the
    # initial INSERT (no cleanup dependency).
    filename_ciphertext: Optional[str] = Form(None),
    content_type_ciphertext: Optional[str] = Form(None),
    file: UploadFile = File(...),
    principal = Depends(verify_trusted_device),
):


    vault_id = principal["vault_id"]
    try:
        key = get_verified_vault_key(vault_id, pin)
        file_bytes = await file.read()
        resolved_content_type = content_type or file.content_type

        auto_save_login_credentials = _user_requested_login_extraction(
            accompanying_text,
        )

        safe_relative_path = _sanitize_relative_path(relative_path)

        # Decode ZK ciphertext fields (base64url) — HTTP 400 on
        # malformed input; skipping cleanly when both are absent.
        import base64 as _b64
        import binascii as _binascii
        def _b64u(val: Optional[str]) -> Optional[bytes]:
            if not val:
                return None
            try:
                return _b64.urlsafe_b64decode(val + "=" * (-len(val) % 4))
            except (_binascii.Error, ValueError):
                raise HTTPException(
                    status_code=400,
                    detail="ciphertext must be base64url",
                )
        _fname_ct_bytes = _b64u(filename_ciphertext)
        _ctype_ct_bytes = _b64u(content_type_ciphertext)

        result = save_uploaded_file(
            vault_id=vault_id,
            file_name=file.filename or "uploaded_file",
            content_type=resolved_content_type,
            file_bytes=file_bytes,
            key=key,
            background_tasks=background_tasks,
            is_batch_upload=is_batch_upload,
            auto_save_login_credentials=auto_save_login_credentials,
            relative_path=safe_relative_path,
            import_id=import_id,
            content_sha256=content_sha256,
            duplicate_action=duplicate_action or "prompt",
            filename_ciphertext_bytes=_fname_ct_bytes,
            content_type_ciphertext_bytes=_ctype_ct_bytes,
        )

                                                                       
        if result.get("status") == "skipped_duplicate":
            return result

                                                                   
        text_for_classification = (accompanying_text or "").strip()
        auto_saved_name: Optional[str] = None
        auto_classified_doc_type: Optional[str] = None
        if text_for_classification:
            candidate_name = extract_naming_intent_from_accompanying_text(
                text_for_classification,
            )
            if candidate_name:
                try:
                    named = save_named_uploaded_asset(
                        vault_id=vault_id,
                        file_id=result["file_id"],
                        saved_name=candidate_name,
                        key=key,
                    )
                    auto_saved_name = named.get("saved_name") or candidate_name
                except Exception as exc:
                                                                      
                    print(
                        "[CHAT-DEBUG] upload_auto_name_failed "
                        f"file_id={result.get('file_id')} "
                        f"candidate_name={candidate_name!r} "
                        f"error_type={type(exc).__name__} repr={exc!r}",
                        flush=True,
                    )

                                                                  
            try:
                auto_classified_doc_type = (
                    _persist_naming_text_classification(
                        vault_id=vault_id,
                        file_id=result["file_id"],
                        naming_text=text_for_classification,
                        file_name=result.get("filename"),
                        content_type=resolved_content_type,
                    )
                )
            except Exception as exc:
                                                                    
                print(
                    "[CHAT-DEBUG] upload_naming_classification_failed "
                    f"file_id={result.get('file_id')} "
                    f"error_type={type(exc).__name__} repr={exc!r}",
                    flush=True,
                )

        result["auto_named"] = auto_saved_name is not None
        result["saved_name"] = auto_saved_name
        result["auto_doc_type"] = auto_classified_doc_type

                                                               
        try:
            from vault_tool_result_cache import invalidate_for_event
            invalidate_for_event(
                vault_id=vault_id, event="file_uploaded",
            )
        except Exception:
            pass
                                                                   
                                                                   
        result["message"] = _build_upload_message(
            result,
            suppress_naming_prompt=bool(text_for_classification),
            auto_saved_name=auto_saved_name,
            is_batch_upload=is_batch_upload,
        )

        # 2026-07-24 chat-brain rebuild — bind THIS upload to the
        # session + turn that produced it so a later text-only
        # reply cannot revive an unrelated unnamed file. Only
        # bind when the upload still needs naming (auto-named
        # files were already committed).
        try:
            if result.get("file_id") and not auto_saved_name:
                from vault_chat_upload_binding import bind_upload
                _upload_session_id = None
                try:
                    _upload_session_id = (
                        principal.get("token_id")
                        if isinstance(principal, dict) else None
                    ) or None
                except Exception:
                    _upload_session_id = None
                _upload_turn_id = ""
                try:
                    _upload_turn_id = (
                        request.headers.get("X-Chat-Turn-Id")
                        or request.headers.get("x-chat-turn-id")
                        or ""
                    )
                except Exception:
                    _upload_turn_id = ""
                if not _upload_turn_id:
                    _upload_turn_id = str(result["file_id"])
                bind_upload(
                    vault_id=vault_id,
                    session_id=_upload_session_id,
                    turn_id=_upload_turn_id,
                    uploaded_file_id=str(result["file_id"]),
                    filename=str(result.get("filename")
                                 or result.get("file_name") or ""),
                    content_type=str(resolved_content_type or ""),
                )
        except Exception:
            logger.exception(
                "[UPLOAD-BINDING] bind_call_failed "
                "vault=%s file_id=%s",
                (vault_id or "")[:8] + "…",
                str(result.get("file_id") or "")[:8] + "…",
            )

        print(
            "[CHAT-DEBUG] upload_accompanying_text "
            f"text_present={bool(text_for_classification)} "
            f"text_len={len(text_for_classification)} "
            f"auto_named={result['auto_named']} "
            f"saved_name={auto_saved_name!r} "
            f"auto_doc_type={auto_classified_doc_type!r} "
            f"asset_type={result.get('asset_type')!r}",
            flush=True,
        )

        return result
    except HTTPException:
        raise
    except Exception:
        logger.error("Upload failed")
        raise HTTPException(status_code=400, detail="Upload failed")


@app.post("/name-file")
async def name_file_endpoint(
    request: Request,
    vault_name: str = Form(...),
    file_id: str = Form(...),
    saved_name: str = Form(...),
    pin: str = Form(...),
    principal = Depends(verify_trusted_device),
):
    vault_id = principal["vault_id"]

    key = verify_vault_pin(vault_id, pin)

    result = save_named_uploaded_asset(
        vault_id=vault_id,
        file_id=file_id,
        saved_name=saved_name,
        key=key,
    )
    return {
        "status": "named",
        **result,
        "message": f"Saved {result['file_name']} as {result['saved_name']}.",
    }


@app.post("/list-files")
async def list_files_endpoint(
    request: Request,
    payload: ListFilesRequest,
    principal = Depends(verify_trusted_device),
):
    vault_id = principal["vault_id"]

    verify_vault_pin(vault_id, payload.pin)

    files = list_uploaded_files(vault_id)
    return {"files": files}


@app.post("/download-file")
async def download_file_endpoint(
    request: Request,
    payload: DownloadFileRequest,
    principal = Depends(verify_trusted_device),
):
    vault_id = principal["vault_id"]
    key = get_verified_vault_key(vault_id, payload.pin)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute(
            """
            SELECT file_name, content_type, encrypted_file_data,
                   storage_mode, upload_status
            FROM uploaded_files
            WHERE id = %s AND vault_id = %s
            LIMIT 1
            """,
            (payload.file_id, vault_id),
        )
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="File not found")

                                                                           
        if (row.get("upload_status") or "complete") != "complete":
            raise HTTPException(status_code=409, detail="File upload not finalized yet")
        if (row.get("storage_mode") or "inline") != "inline":
            raise HTTPException(
                status_code=400,
                detail="This file requires the chunked download endpoint (Phase 9.8.C).",
            )

        raw_bytes = decrypt_bytes(row["encrypted_file_data"], key)
        encoded = base64.b64encode(raw_bytes).decode("utf-8")

        return {
            "file_id": payload.file_id,
            "file_name": row["file_name"],
            "content_type": row.get("content_type"),
            "base64_data": encoded,
        }
    finally:
        conn.close()


@app.delete("/delete-file")
async def delete_file_endpoint(
    request: Request,
    payload: DeleteFileRequest,
    principal=Depends(verify_trusted_device),
):
    vault_id = principal["vault_id"]

    verify_vault_pin(vault_id, payload.pin)

    conn = get_db()
    try:
        cursor = conn.cursor(cursor_factory=RealDictCursor)

        cursor.execute(
            """
            DELETE FROM uploaded_files
            WHERE id = %s AND vault_id = %s
            RETURNING file_name, file_size
            """,
            (payload.file_id, vault_id),
        )

        deleted = cursor.fetchone()
        conn.commit()

        if not deleted:
            raise HTTPException(status_code=404, detail="File not found")

        bump_vault_total_bytes(vault_id, -int(deleted["file_size"] or 0))

                                                      
        try:
            from vault_relationship_graph import (
                cleanup_relationships_for_file,
            )
            cleanup_relationships_for_file(vault_id, payload.file_id)
        except Exception:
            logger.exception(
                "[delete-file] relationship cleanup failed "
                "vault=%s file=%s", vault_id, payload.file_id,
            )

        return {
            "status": "deleted",
            "file_name": deleted["file_name"],
            "message": f"Deleted {deleted['file_name']}."
        }

    finally:
        conn.close()


def _load_semantic_chat_min_score() -> float:


    raw = os.getenv("VAULTAI_SEMANTIC_CHAT_MIN_SCORE", "0.25")
    try:
        v = float(raw)
    except (TypeError, ValueError):
        return 0.25
    if v != v or v < 0.0 or v > 1.0:                       
        return 0.25
    return v


SEMANTIC_CHAT_MIN_SCORE = _load_semantic_chat_min_score()


async def _run_semantic_search_for_chat(
    vault_id: str, query: str, limit: int = 10,
) -> list[dict]:


    from semantic_embedder import is_enabled, embed_text
    if not is_enabled() or not (query or "").strip():
        return []
    try:
        vec = await embed_text(client, query)
    except Exception as e:
        logger.warning("chat semantic embed failed: %s", e)
        return []
    if vec is None:
        return []
    conn = get_db()
    rows: list = []
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT si.source_kind,
                   si.uploaded_file_id,
                   si.vault_item_id,
                   1 - (si.embedding <=> %s::vector) AS score
            FROM semantic_index si
            WHERE si.vault_id = %s
            ORDER BY si.embedding <=> %s::vector
            LIMIT %s;
            """,
            (vec, vault_id, vec, limit),
        )
        rows = cur.fetchall() or []
    except Exception as e:
        logger.warning("chat semantic query failed: %s", e)
        rows = []
    finally:
        conn.close()
    return [
        {
            "source_kind": r["source_kind"],
            "uploaded_file_id": r["uploaded_file_id"],
            "vault_item_id": r["vault_item_id"],
            "score": float(r["score"]),
        }
        for r in rows
    ]


def _lookup_file_by_id(vault_id: str, file_id: str) -> Optional[dict]:
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT id, file_name, content_type, file_size, saved_name,
                   asset_type, created_at
            FROM uploaded_files
            WHERE vault_id = %s AND id = %s
              AND upload_status = 'complete'
            """,
            (vault_id, file_id),
        )
        return cur.fetchone()
    finally:
        conn.close()


def _lookup_item_service_by_id(
    vault_id: str, item_id: int,
) -> Optional[str]:
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT service FROM vault_items
            WHERE vault_id = %s AND id = %s
            """,
            (vault_id, item_id),
        )
        row = cur.fetchone()
        return row["service"] if row else None
    finally:
        conn.close()


async def hybrid_retrieve_files(
    vault_id: str,
    query: str,
    *,
    limit: int = 10,
) -> list[dict]:


    query_norm = (query or "").strip()
    if not query_norm:
        return []
    query_lower = query_norm.lower()

                                          
    try:
        from document_understanding import classify_naming_text_doc_type
        query_doc_type = classify_naming_text_doc_type(query_norm)
    except Exception:
        query_doc_type = None
    query_tags = extract_query_tags(query_norm)
    query_entities = extract_query_entities(query_norm)

    candidates: dict[str, dict] = {}
                                                                  
                                                                   
    entity_hits: dict[str, dict] = {}

                                                                    
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            WITH base AS (
              SELECT uf.id, uf.file_name, uf.saved_name, uf.asset_type,
                     uf.content_type, uf.created_at,
                     vdm.doc_type
                FROM uploaded_files uf
                LEFT JOIN vault_document_metadata vdm
                       ON vdm.vault_id = uf.vault_id
                      AND vdm.uploaded_file_id = uf.id
               WHERE uf.vault_id = %s
                 AND uf.upload_status = 'complete'
                 AND (
                       LOWER(COALESCE(uf.saved_name, '')) = %s
                    OR LOWER(COALESCE(uf.saved_name, '')) LIKE %s
                    OR %s LIKE '%%' || LOWER(COALESCE(uf.saved_name, '')) || '%%'
                    OR EXISTS (
                         SELECT 1 FROM vault_asset_tags t
                          WHERE t.vault_id = uf.vault_id
                            AND t.uploaded_file_id = uf.id
                            AND t.tag = ANY(%s)
                       )
                 )
            )
            SELECT b.*,
                   COALESCE(
                     ARRAY(
                       SELECT DISTINCT t.tag FROM vault_asset_tags t
                        WHERE t.vault_id = %s
                          AND t.uploaded_file_id = b.id
                     ),
                     ARRAY[]::text[]
                   ) AS tags
              FROM base b
             ORDER BY b.created_at DESC
             LIMIT 50
            """,
            (
                vault_id,
                query_lower,
                f"%{query_lower}%",
                query_lower,
                query_tags or [],
                vault_id,
            ),
        )
        for row in cur.fetchall() or []:
            candidates[row["id"]] = {
                "id":          row["id"],
                "file_name":   row.get("file_name"),
                "saved_name":  row.get("saved_name"),
                "asset_type":  row.get("asset_type"),
                "content_type": row.get("content_type"),
                "doc_type":    row.get("doc_type"),
                "tags":        list(row.get("tags") or []),
                "vector_score": 0.0,
            }
    except Exception as e:
        logger.warning("hybrid_retrieve_files SQL pass failed: %s", e)
    finally:
        conn.close()

                                                                    
    if query_entities:
        clauses: list[str] = []
        params: list = [vault_id]
        for ent in query_entities:
            value = ent["value"]
            if ent["kind"] == "year":
                clauses.append("entity_value LIKE %s")
                params.append(f"{value}%")
            else:
                clauses.append("LOWER(entity_value) = LOWER(%s)")
                params.append(value)
        sql_pred = " OR ".join(clauses)
        econn = get_db()
        try:
            ecur = econn.cursor(cursor_factory=RealDictCursor)
            ecur.execute(
                f"""
                SELECT source_file_id,
                       COUNT(*)                 AS match_count,
                       ARRAY_AGG(DISTINCT entity_value) AS matched_values
                  FROM vault_document_entities
                 WHERE vault_id = %s
                   AND ({sql_pred})
                 GROUP BY source_file_id
                 LIMIT 100
                """,
                params,
            )
            for row in ecur.fetchall() or []:
                fid = row["source_file_id"]
                entity_hits[fid] = {
                    "count": int(row["match_count"] or 0),
                    "values": list(row["matched_values"] or []),
                }
        except Exception as e:
            logger.warning("hybrid_retrieve_files entity pass failed: %s", e)
        finally:
            econn.close()

                                                                      
        missing_entity_fids = [
            fid for fid in entity_hits.keys() if fid not in candidates
        ]
        if missing_entity_fids:
            econn2 = get_db()
            try:
                ecur2 = econn2.cursor(cursor_factory=RealDictCursor)
                ecur2.execute(
                    """
                    SELECT uf.id, uf.file_name, uf.saved_name,
                           uf.asset_type, uf.content_type,
                           vdm.doc_type,
                           COALESCE(
                             ARRAY(
                               SELECT DISTINCT t.tag FROM vault_asset_tags t
                                WHERE t.vault_id = uf.vault_id
                                  AND t.uploaded_file_id = uf.id
                             ),
                             ARRAY[]::text[]
                           ) AS tags
                      FROM uploaded_files uf
                      LEFT JOIN vault_document_metadata vdm
                             ON vdm.vault_id = uf.vault_id
                            AND vdm.uploaded_file_id = uf.id
                     WHERE uf.vault_id = %s
                       AND uf.id = ANY(%s)
                       AND uf.upload_status = 'complete'
                    """,
                    (vault_id, missing_entity_fids),
                )
                for row in ecur2.fetchall() or []:
                    candidates[row["id"]] = {
                        "id":          row["id"],
                        "file_name":   row.get("file_name"),
                        "saved_name":  row.get("saved_name"),
                        "asset_type":  row.get("asset_type"),
                        "content_type": row.get("content_type"),
                        "doc_type":    row.get("doc_type"),
                        "tags":        list(row.get("tags") or []),
                        "vector_score": 0.0,
                    }
            except Exception as e:
                logger.warning(
                    "hybrid_retrieve_files entity-only enrich failed: %s", e,
                )
            finally:
                econn2.close()

                                                                    
    try:
        from semantic_embedder import is_enabled
        embedder_on = is_enabled()
    except Exception:
        embedder_on = False
    if embedder_on:
        try:
            from semantic_embedder import embed_text
            vec = await embed_text(client, query_norm)
        except Exception as e:
            logger.warning("hybrid retrieve embed failed: %s", e)
            vec = None
        if vec is not None:
            vconn = get_db()
            try:
                vcur = vconn.cursor(cursor_factory=RealDictCursor)
                vcur.execute(
                    """
                    SELECT si.uploaded_file_id,
                           1 - (si.embedding <=> %s::vector) AS score,
                           uf.file_name, uf.saved_name, uf.asset_type,
                           uf.content_type,
                           vdm.doc_type
                      FROM semantic_index si
                      JOIN uploaded_files uf ON uf.id = si.uploaded_file_id
                      LEFT JOIN vault_document_metadata vdm
                             ON vdm.vault_id = uf.vault_id
                            AND vdm.uploaded_file_id = uf.id
                     WHERE si.vault_id = %s
                       AND si.source_kind = 'semantic_profile'
                       AND uf.upload_status = 'complete'
                     ORDER BY si.embedding <=> %s::vector
                     LIMIT %s
                    """,
                    (vec, vault_id, vec, max(limit * 3, 30)),
                )
                for row in vcur.fetchall() or []:
                    fid = row["uploaded_file_id"]
                    if fid in candidates:
                        candidates[fid]["vector_score"] = max(
                            candidates[fid]["vector_score"],
                            float(row["score"]),
                        )
                    else:
                        candidates[fid] = {
                            "id":          fid,
                            "file_name":   row.get("file_name"),
                            "saved_name":  row.get("saved_name"),
                            "asset_type":  row.get("asset_type"),
                            "content_type": row.get("content_type"),
                            "doc_type":    row.get("doc_type"),
                            "tags":        [],
                            "vector_score": float(row["score"]),
                        }
            except Exception as e:
                logger.warning("hybrid_retrieve_files vector pass failed: %s", e)
            finally:
                vconn.close()

                                                                    
    try:
        from taxonomy import DOC_TYPE_TO_FAMILY
    except Exception:
        DOC_TYPE_TO_FAMILY = {}

    scored: list[dict] = []
    for c in candidates.values():
        doc_family = (
            DOC_TYPE_TO_FAMILY.get(c["doc_type"]) if c.get("doc_type") else None
        )
        ent = entity_hits.get(c["id"], {})
        result = score_hybrid_match(
            query_lower=query_lower,
            candidate_saved_name=c.get("saved_name"),
            candidate_doc_type=c.get("doc_type"),
            candidate_doc_family=doc_family,
            candidate_tags=c.get("tags") or [],
            query_doc_type=query_doc_type,
            query_tags=query_tags,
            vector_score=c.get("vector_score", 0.0),
            entity_matches=int(ent.get("count", 0)),
            matched_entity_values=ent.get("values"),
        )
        if result["score"] < HYBRID_MIN_SCORE:
            continue
        scored.append({**c, **result})

    scored.sort(key=lambda r: r["score"], reverse=True)
    return scored[:limit]


async def _handle_search_memory_intent(
    vault_id: str, query: str, key: bytes,
) -> Optional[str]:


    try:
        file_hits = await hybrid_retrieve_files(
            vault_id, query, limit=10,
        )
    except Exception as e:
        logger.warning("hybrid_retrieve_files raised in search_memory: %s", e)
        file_hits = []

    if file_hits:
        top = file_hits[0]
        print(
            "[CHAT-DEBUG] hybrid_retrieve_top "
            f"file_id={top.get('id')} "
            f"saved_name={top.get('saved_name')!r} "
            f"doc_type={top.get('doc_type')!r} "
            f"score={top.get('score')!r} "
            f"reasons={top.get('reasons')!r}",
            flush=True,
        )
        asset = _lookup_file_by_id(vault_id, top["id"])
        if asset:
            return _build_structured_asset_reply(
                asset,
                asset.get("saved_name") or asset.get("file_name") or "file",
            )

                                                                   
    matches = await _run_semantic_search_for_chat(vault_id, query, limit=10)
    if not matches:
        return None
    top = matches[0]
    if top["score"] < SEMANTIC_CHAT_MIN_SCORE:
        return None
    if top["uploaded_file_id"]:
        asset = _lookup_file_by_id(vault_id, top["uploaded_file_id"])
        if asset:
            return _build_structured_asset_reply(
                asset, asset.get("saved_name") or asset.get("file_name") or "file"
            )
    if top["vault_item_id"] is not None:
        service = _lookup_item_service_by_id(vault_id, top["vault_item_id"])
        if service:
            return retrieve_secret_tool(vault_id, {"service": service}, key)
    return None
                                                                 

ALLOWED_TAGS_FOR_CHAT = {
    "travel", "finance", "legal", "medical", "business", "education",
    "identity", "government", "security", "personal", "media",
    "receipt", "tax",
}

LIST_BY_TAG_RENDER_MAX = 15                                          


def _handle_list_by_tag(
    vault_id: str, tag: str,
) -> Optional[str]:


    tag = (tag or "").strip().lower()
    if tag not in ALLOWED_TAGS_FOR_CHAT:
        return None

    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT uf.file_name, uf.saved_name, uf.asset_type, uf.content_type
            FROM vault_asset_tags t
            JOIN uploaded_files uf ON uf.id = t.uploaded_file_id
            WHERE t.vault_id = %s AND t.tag = %s
              AND t.uploaded_file_id IS NOT NULL
              AND uf.upload_status = 'complete'
            ORDER BY uf.created_at DESC
            LIMIT 25
            """,
            (vault_id, tag),
        )
        file_rows = cur.fetchall() or []

        cur.execute(
            """
            SELECT vi.service, vi.item_type
            FROM vault_asset_tags t
            JOIN vault_items vi ON vi.id = t.vault_item_id
            WHERE t.vault_id = %s AND t.tag = %s
              AND t.vault_item_id IS NOT NULL
            ORDER BY vi.created_at DESC
            LIMIT 25
            """,
            (vault_id, tag),
        )
        item_rows = cur.fetchall() or []
    finally:
        conn.close()

                                                        
    files_unique: list[dict] = []
    seen_file_names: set[str] = set()
    for r in file_rows:
        display = (r["saved_name"] or r["file_name"] or "untitled").strip()
        k = display.lower()
        if k in seen_file_names:
            continue
        seen_file_names.add(k)
        files_unique.append(r)

    items_unique: list[dict] = []
    seen_services: set[str] = set()
    for r in item_rows:
        service = (r["service"] or "untitled").strip()
        k = service.lower()
        if k in seen_services:
            continue
        seen_services.add(k)
        items_unique.append(r)

    if not files_unique and not items_unique:
        return None

                                                        
    from locales import ENGLISH_FORMATTER as _fmt
    lines: list[str] = [_fmt.tag_list_header(tag)]

    if files_unique:
        lines.append("")
        lines.append(_fmt.files_section_header())
        shown_files = files_unique[:LIST_BY_TAG_RENDER_MAX]
        for f in shown_files:
            name = f["saved_name"] or f["file_name"] or "untitled"
            kind = (f["asset_type"] or "file").title()
            lines.append(f"- {_title_case_asset(name)} ({kind})")
        extra_files = len(files_unique) - len(shown_files)
        if extra_files > 0:
            lines.append(_fmt.and_more_files(extra_files))

    if items_unique:
        lines.append("")
        lines.append(_fmt.logins_section_header())
        shown_items = items_unique[:LIST_BY_TAG_RENDER_MAX]
        for it in shown_items:
            service = (it["service"] or "untitled").title()
            type_lbl = (it["item_type"] or "login").title()
            lines.append(f"- {service} ({type_lbl})")
        extra_items = len(items_unique) - len(shown_items)
        if extra_items > 0:
            lines.append(_fmt.and_more_logins(extra_items))

    return "\n".join(lines).rstrip()
                                                                 

ALLOWED_DOC_TYPES_FOR_CHAT = {
    "passport", "visa", "boarding_pass", "ticket", "hotel_itinerary",
    "receipt", "invoice", "contract", "agreement", "degree",
    "certificate", "insurance", "driver_license", "id_card",
    "tax_document", "medical_record",
}
ALLOWED_DOC_FIELDS_FOR_CHAT = {
    "country", "expiry_date", "issue_date", "purchase_date", "due_date",
    "effective_date", "renewal_date", "departure_date", "event_date",
    "record_date", "check_in", "check_out", "amount", "currency",
    "merchant", "vendor", "parties", "airline", "flight_number",
    "origin", "destination", "hotel", "location", "provider",
    "institution", "year", "field", "issuer", "tax_year", "form_type",
    "passport_number_last4", "id_number_last4",
    "license_number_last4", "policy_number_last4", "license_class",
}
DOC_LIST_RENDER_MAX = 10

_DOC_FIELD_ORDER = (
    "country", "expiry_date", "issue_date", "purchase_date", "due_date",
    "effective_date", "renewal_date", "check_in", "check_out",
    "departure_date", "event_date", "record_date",
    "airline", "flight_number", "origin", "destination",
    "hotel", "location", "provider", "institution", "field", "year",
    "issuer", "merchant", "vendor", "amount", "currency",
    "tax_year", "form_type", "license_class",
    "passport_number_last4", "id_number_last4",
    "license_number_last4", "policy_number_last4",
    "entry_type", "parties",
)


def _doc_type_pretty(dt: str) -> str:
    return dt.replace("_", " ").title()


def _format_expiry_status(iso_date: Optional[str]) -> Optional[str]:


    from locales import ENGLISH_FORMATTER
    return ENGLISH_FORMATTER.expiry_status_absolute(iso_date)


def _render_doc_entry(doc_type: str, md: dict, fallback_name: str) -> str:


    parts: list[str] = []
    seen_currency = False
    for key in _DOC_FIELD_ORDER:
        if key not in md:
            continue
        val = md[key]
        if val in (None, "", []):
            continue
        if key == "expiry_date":
            status = _format_expiry_status(str(val))
            parts.append(
                f"Expiry: {val}" + (f" ({status})" if status else "")
            )
        elif key == "amount":
            currency = md.get("currency", "")
            parts.append(f"Amount: {currency} {val}".strip())
            seen_currency = True
        elif key == "currency":
            if not seen_currency:
                parts.append(f"Currency: {val}")
        elif key == "parties" and isinstance(val, list):
            parts.append("Parties: " + ", ".join(str(x) for x in val[:5]))
        else:
            label = key.replace("_last4", " (last 4)").replace("_", " ").title()
            parts.append(f"{label}: {val}")
    if parts:
        return "- " + " · ".join(parts)
    return f"- {fallback_name or _doc_type_pretty(doc_type)}"


def _handle_understand_document(
    vault_id: str,
    doc_type: Optional[str],
    doc_field: Optional[str],
) -> Optional[str]:


    if doc_type:
        doc_type = doc_type.strip().lower()
        if doc_type not in ALLOWED_DOC_TYPES_FOR_CHAT:
            doc_type = None
    if doc_field:
        doc_field = doc_field.strip().lower()
        if doc_field not in ALLOWED_DOC_FIELDS_FOR_CHAT:
            doc_field = None

    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        if doc_type:
            cur.execute(
                """
                SELECT vdm.doc_type, vdm.metadata_json,
                       uf.file_name, uf.saved_name
                FROM vault_document_metadata vdm
                JOIN uploaded_files uf ON uf.id = vdm.uploaded_file_id
                WHERE vdm.vault_id = %s
                  AND vdm.doc_type = %s
                ORDER BY vdm.updated_at DESC
                LIMIT 50
                """,
                (vault_id, doc_type),
            )
        elif doc_field:
            cur.execute(
                """
                SELECT vdm.doc_type, vdm.metadata_json,
                       uf.file_name, uf.saved_name
                FROM vault_document_metadata vdm
                JOIN uploaded_files uf ON uf.id = vdm.uploaded_file_id
                WHERE vdm.vault_id = %s
                  AND vdm.metadata_json ? %s
                ORDER BY vdm.doc_type, vdm.updated_at DESC
                LIMIT 50
                """,
                (vault_id, doc_field),
            )
        else:
            cur.execute(
                """
                SELECT vdm.doc_type, vdm.metadata_json,
                       uf.file_name, uf.saved_name
                FROM vault_document_metadata vdm
                JOIN uploaded_files uf ON uf.id = vdm.uploaded_file_id
                WHERE vdm.vault_id = %s
                ORDER BY vdm.doc_type, vdm.updated_at DESC
                LIMIT 100
                """,
                (vault_id,),
            )
        rows = cur.fetchall() or []
    finally:
        conn.close()

    if not rows:
        return None

    grouped: dict[str, list] = {}
    for r in rows:
        grouped.setdefault(r["doc_type"], []).append(r)

    lines: list[str] = []
    for dt in sorted(grouped.keys()):
        items = grouped[dt][:DOC_LIST_RENDER_MAX]
        more = len(grouped[dt]) - len(items)
        lines.append(f"{_doc_type_pretty(dt)}:")
        for r in items:
            md_raw = r["metadata_json"]
            if isinstance(md_raw, dict):
                md = md_raw
            else:
                try:
                    md = json.loads(md_raw or "{}")
                except Exception:
                    md = {}
            fallback = r.get("saved_name") or r.get("file_name") or ""
            lines.append(_render_doc_entry(dt, md, fallback))
        if more > 0:
            lines.append(
                f"  ...and {more} more "
                f"{'document' if more == 1 else 'documents'}."
            )
        lines.append("")

    return "\n".join(lines).rstrip()
                                                                 

ALLOWED_EXPIRY_DATE_FIELDS = (
    "expiry_date", "departure_date", "check_out", "due_date", "event_date",
)
ALLOWED_RENEWAL_DATE_FIELDS = ("renewal_date",)
_ALL_DATE_FIELDS = ALLOWED_EXPIRY_DATE_FIELDS + ALLOWED_RENEWAL_DATE_FIELDS


def _days_until_iso(iso_date: str) -> Optional[int]:


    try:
        from datetime import date
        parts = (iso_date or "").split("-")
        if len(parts) != 3:
            return None
        y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
        return (date(y, m, d) - date.today()).days
    except Exception:
        return None


def _expiry_status(days_delta: int) -> str:


    if days_delta < 0:
        return "expired"
    if days_delta <= 7:
        return "critical"
    if days_delta <= 30:
        return "warning"
    if days_delta <= 90:
        return "upcoming"
    return "valid"


def _describe_date(field: str, days_delta: int) -> str:


    from locales import ENGLISH_FORMATTER
    return ENGLISH_FORMATTER.date_relative(field, days_delta)


def _handle_expiry_alerts(
    vault_id: str,
) -> Optional[str]:


    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT vdm.doc_type, vdm.metadata_json,
                   uf.file_name, uf.saved_name
            FROM vault_document_metadata vdm
            JOIN uploaded_files uf ON uf.id = vdm.uploaded_file_id
            WHERE vdm.vault_id = %s
            """,
            (vault_id,),
        )
        rows = cur.fetchall() or []
    finally:
        conn.close()

                                                                     
    expired: list[tuple[int, str, str]] = []
    critical: list[tuple[int, str, str]] = []
    warning: list[tuple[int, str, str]] = []
    upcoming: list[tuple[int, str, str]] = []
    renewals: list[tuple[int, str, str]] = []

    for r in rows:
        md_raw = r["metadata_json"]
        if isinstance(md_raw, dict):
            md = md_raw
        else:
            try:
                md = json.loads(md_raw or "{}")
            except Exception:
                md = {}
        if not md:
            continue
        title = (r.get("saved_name") or r.get("file_name")
                 or _doc_type_pretty(r["doc_type"]))
        title = str(title).strip()

        for field in _ALL_DATE_FIELDS:
            iso = md.get(field)
            if not iso:
                continue
            days = _days_until_iso(str(iso))
            if days is None:
                continue
            status = _expiry_status(days)
            if status == "valid":
                continue
            descriptor = _describe_date(field, days)
            entry = (days, title, descriptor)
            if field in ALLOWED_RENEWAL_DATE_FIELDS:
                renewals.append(entry)
            elif status == "expired":
                expired.append(entry)
            elif status == "critical":
                critical.append(entry)
            elif status == "warning":
                warning.append(entry)
            elif status == "upcoming":
                upcoming.append(entry)

                                                                      
    expired.sort(key=lambda x: x[0], reverse=True)
    critical.sort(key=lambda x: x[0])
    warning.sort(key=lambda x: x[0])
    upcoming.sort(key=lambda x: x[0])
    renewals.sort(key=lambda x: x[0])

    sections: list[tuple[str, list]] = [
        ("Expired",   expired),
        ("Critical",  critical),
        ("Warning",   warning),
        ("Upcoming",  upcoming),
        ("Renewals",  renewals),
    ]
    if not any(items for _, items in sections):
        return None

    lines: list[str] = []
    for header, items in sections:
        if not items:
            continue
        lines.append(f"{header}:")
        for _days, title, descriptor in items:
            lines.append(f"- {title} - {descriptor}")
        lines.append("")
    return "\n".join(lines).rstrip()
                                                                 

_EXPIRY_TYPE_PRETTY = {
    "passport":         "Passport",
    "visa":             "Visa",
    "id_card":          "ID card",
    "driver_license":   "Driver license",
    "insurance":        "Insurance",
    "tax":              "Tax",
    "contract":         "Contract",
    "subscription":     "Subscription",
    "inheritance":      "Inheritance cooldown",
    "custom":           "Reminder",
}


def _handle_expiry_status(
    vault_id: str,
    expiry_filter: Optional[str] = None,
) -> Optional[str]:


    try:
        from expiry_engine import fetch_active_alerts, is_enabled
        if not is_enabled():
            return None
        filt = (expiry_filter or "").strip().lower() or None
        rows = fetch_active_alerts(
            vault_id, expiry_type_filter=filt, limit=100,
        )
    except Exception:
        return None
    if not rows:
        return None

                                                                   
    from locales import ENGLISH_FORMATTER as _fmt
    by_sev: dict[str, list[dict]] = {"critical": [], "warning": [], "info": []}
    for r in rows:
        sev = r.get("severity") or "info"
        if sev in by_sev:
            by_sev[sev].append(r)

    label_cache: dict = {}

    def _label_for(r: dict) -> str:


        kind = r.get("source_kind")
        pretty = _EXPIRY_TYPE_PRETTY.get(
            r.get("expiry_type") or "", (r.get("expiry_type") or "").title()
        )
        if kind == "uploaded_file" and r.get("source_file_id"):
            cache_key = ("file", r["source_file_id"])
            if cache_key not in label_cache:
                try:
                    conn = get_db()
                    try:
                        cur = conn.cursor()
                        cur.execute(
                            """SELECT saved_name, file_name FROM uploaded_files
                               WHERE id=%s AND vault_id=%s""",
                            (r["source_file_id"], vault_id),
                        )
                        row = cur.fetchone()
                    finally:
                        conn.close()
                    if row:
                        name = (row[0] or row[1] or "").strip()
                        label_cache[cache_key] = name or pretty
                    else:
                        label_cache[cache_key] = pretty
                except Exception:
                    label_cache[cache_key] = pretty
            return label_cache[cache_key] or pretty
        if kind == "vault_item" and r.get("source_item_id") is not None:
            cache_key = ("item", r["source_item_id"])
            if cache_key not in label_cache:
                try:
                    conn = get_db()
                    try:
                        cur = conn.cursor()
                        cur.execute(
                            """SELECT service FROM vault_items
                               WHERE id=%s AND vault_id=%s""",
                            (r["source_item_id"], vault_id),
                        )
                        row = cur.fetchone()
                    finally:
                        conn.close()
                    if row and row[0]:
                        label_cache[cache_key] = row[0].strip().title()
                    else:
                        label_cache[cache_key] = pretty
                except Exception:
                    label_cache[cache_key] = pretty
            return label_cache[cache_key] or pretty
        if kind == "inheritance":
            return "Inheritance cooldown"
        if kind == "memory":
            return "Reminder"
        return pretty

                                                            
    try:
        from vault_language_engine import (
            build_language_context, choose_reply_language,
        )
        from locales import get_formatter
        ctx = build_language_context(vault_id)
        reply_lang = choose_reply_language(
            ctx.get("preferred_locale"),
            chat_language_override=ctx.get("chat_language"),
        )
        reply_fmt = get_formatter(reply_lang)
    except Exception:
        reply_lang = "en"
        reply_fmt = _fmt                    

    lines: list[str] = []
    sev_order = ("critical", "warning", "info")
    for sev in sev_order:
        items = by_sev.get(sev) or []
        if not items:
            continue
                                                                 
                                                                    
        seen: set = set()
        unique: list[dict] = []
        for r in items:
            label = _label_for(r)
            k = (label.lower(), r.get("expiry_date") or "")
            if k in seen:
                continue
            seen.add(k)
            unique.append(r)
        if not unique:
            continue
        if lines:
            lines.append("")
                                                                       
        lines.append(f"{reply_fmt.severity_label(sev)}:")
        for r in unique:
                                                                     
                                                                       
            doc_type = r.get("expiry_type") or ""
            try:
                                                                   
                                                                  
                from datetime import date as _date
                iso = (r.get("expiry_date") or "")[:10]
                parts = iso.split("-")
                if len(parts) == 3:
                    y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                    days = (_date(y, m, d) - _date.today()).days
                    if doc_type in (
                        "passport","visa","id_card","driver_license",
                        "boarding_pass","hotel_itinerary","ticket",
                        "invoice","receipt","contract","agreement",
                        "degree","certificate","insurance",
                        "tax_document","medical_record",
                    ):
                                                   
                        loc_label = reply_fmt.doc_type_label(doc_type)
                    else:
                        loc_label = _label_for(r)
                    lines.append(
                        f"- {reply_fmt.expiry_phrase(loc_label, days)}"
                    )
                    continue
            except Exception:
                pass
                                                                      
            label = _label_for(r)
            descriptor = _fmt.expiry_status_absolute(r.get("expiry_date"))
            if not descriptor:
                continue
            d = descriptor[0].lower() + descriptor[1:] if descriptor else ""
            lines.append(f"- {label} {d}")
    if not lines:
        return None
    return "\n".join(lines)
                                                                 

_ENTITY_KEY_LABELS_EN = {
    "document_type":           "Type",
    "country":                 "Country",
    "nationality":             "Nationality",
    "issue_date":              "Issue date",
    "expiry_date":             "Expiry",
    "passport_number_last4":   "Passport number (last 4)",
    "id_number_last4":         "ID number (last 4)",
    "license_number_last4":    "License number (last 4)",
    "policy_number_last4":     "Policy number (last 4)",
    "license_class":           "License class",
    "entry_type":              "Entry type",
    "visa_type":               "Visa type",
    "airline":                 "Airline",
    "flight_number":           "Flight",
    "departure_date":          "Departure",
    "origin":                  "Origin",
    "destination":             "Destination",
    "hotel":                   "Hotel",
    "location":                "Location",
    "check_in":                "Check-in",
    "check_out":               "Check-out",
    "vendor":                  "Vendor",
    "event_date":              "Event date",
    "merchant":                "Merchant",
    "amount":                  "Amount",
    "currency":                "Currency",
    "invoice_number":          "Invoice number",
    "due_date":                "Due date",
    "purchase_date":           "Purchase date",
    "parties":                 "Parties",
    "effective_date":          "Effective",
    "renewal_date":            "Renewal",
    "provider":                "Provider",
    "record_type":             "Record type",
    "record_date":             "Record date",
    "institution":             "Institution",
    "year":                    "Year",
    "field":                   "Field",
    "graduation_date":         "Graduation",
    "issuer":                  "Issuer",
    "tax_year":                "Tax year",
    "form_type":               "Form",
    "filing_date":             "Filing date",
}


def _entity_key_label(key: str) -> str:
    return _ENTITY_KEY_LABELS_EN.get(
        key, key.replace("_", " ").title() if key else "Field",
    )


_ENTITY_TYPE_RENDER_ORDER = (
    "identity", "travel", "government", "tax", "finance",
    "insurance", "legal", "business", "medical", "education",
    "personal", "other",
)


def _document_type_pretty(value: Optional[str]) -> str:

    if not value:
        return "Document"
    return value.replace("_", " ").title()


def _handle_document_details(
    vault_id: str,
    entity_doc_type: Optional[str] = None,
    entity_type: Optional[str] = None,
    entity_key: Optional[str] = None,
) -> Optional[str]:


    try:
        from document_entities import (
            fetch_entities, is_enabled, ALLOWED_ENTITY_TYPES,
        )
        if not is_enabled():
            return None

                                                                
        dt = (entity_doc_type or "").strip().lower() or None
        et = (entity_type or "").strip().lower() or None
        ek = (entity_key or "").strip().lower() or None
        if et and et not in ALLOWED_ENTITY_TYPES:
            et = None

        rows = fetch_entities(
            vault_id,
            doc_type=dt, entity_type=et, entity_key=ek, limit=300,
        )
    except Exception:
        return None
    if not rows:
        return None

                                                                    
    if ek:
                                                                   
                                               
        per_file: dict[str, dict] = {}
        for r in rows:
            fid = r["source_file_id"]
            per_file.setdefault(fid, {"doc_type": None, "values": []})
            if r["entity_key"] == "document_type":
                per_file[fid]["doc_type"] = r["entity_value"]
            else:
                per_file[fid]["values"].append(r["entity_value"])
                                                               
                                                                       
        try:
            from document_entities import fetch_entities as _fe
            dt_rows = _fe(
                vault_id, doc_type=dt, entity_key="document_type",
                limit=300,
            )
            for r in dt_rows:
                fid = r["source_file_id"]
                if fid in per_file and per_file[fid]["doc_type"] is None:
                    per_file[fid]["doc_type"] = r["entity_value"]
        except Exception:
            pass

        label = _entity_key_label(ek)
        lines = [f"{label}:"]
        for fid, info in per_file.items():
            for v in info["values"]:
                pretty = _document_type_pretty(info["doc_type"])
                lines.append(f"- {pretty}: {v}")
        if len(lines) == 1:
            return None
        return "\n".join(lines)

                                                                    
    by_file: dict[str, list[dict]] = {}
    for r in rows:
        by_file.setdefault(r["source_file_id"], []).append(r)

    lines: list[str] = []
    file_count = 0
    DOC_LIMIT = 8
    for fid, file_rows in by_file.items():
        if file_count >= DOC_LIMIT:
            extra = len(by_file) - file_count
            if extra > 0:
                lines.append("")
                lines.append(f"...and {extra} more document"
                             f"{'s' if extra != 1 else ''}.")
            break
        file_count += 1

                                             
        dt_entity = next(
            (r for r in file_rows if r["entity_key"] == "document_type"),
            None,
        )
        heading = _document_type_pretty(
            dt_entity["entity_value"] if dt_entity else None,
        )
        if lines:
            lines.append("")
        lines.append(f"{heading}:")

                                                                    
        ordered_type_index = {t: i for i, t in enumerate(_ENTITY_TYPE_RENDER_ORDER)}
        def _rkey(r):
            return (
                ordered_type_index.get(r["entity_type"], 99),
                r["entity_key"],
            )
        for r in sorted(file_rows, key=_rkey):
            if r["entity_key"] == "document_type":
                continue
            lines.append(
                f"- {_entity_key_label(r['entity_key'])}: {r['entity_value']}"
            )
    if not lines:
        return None
    return "\n".join(lines)
                                                                 

_MEMORY_RENDER_ORDER = (
    "company", "project", "goal", "identity",
    "relationship", "family",
    "travel", "location", "date", "life_event",
    "preference", "note",
)


def _handle_remember_fact(
    vault_id: str,
    memory_type: Optional[str],
    memory_key: Optional[str],
    memory_value: Optional[str],
    memory_event_date: Optional[str] = None,
) -> Optional[str]:


    from ai_memory import (
        is_enabled, ALLOWED_MEMORY_TYPES, is_forbidden_memory_value,
        slugify_memory_key, update_memory_safe,
    )
    from memory_recall import label_for_type
    from vault_core import is_vault_zk_adopted
    if not is_enabled():
        return None
    mt = (memory_type or "").strip().lower()
    if mt not in ALLOWED_MEMORY_TYPES:
        return None
    val = (memory_value or "").strip()
    if not val:
        return None
    if is_forbidden_memory_value(val):

        return ("I shouldn't remember secret-looking data. "
                "Please save secrets via the vault instead.")
    key = slugify_memory_key(memory_key) or slugify_memory_key(val)
    if not key:
        return None

    # ZK/adopted-vault boundary: the backend MUST NOT persist
    # readable memory_key / memory_value for a ZK vault. Emit a
    # ``memory_proposal`` sentinel prefix that the Flutter chat
    # SSE handler recognizes, encrypts client-side under memoryKey,
    # computes memory_lookup_hash locally, and POSTs to
    # /vault/ciphertext/vault-ai-memory. If the client fails to
    # finalize (network drop, tab close), the memory is not saved
    # — the correct privacy tradeoff. No server-side plaintext
    # persistence fallback.
    if is_vault_zk_adopted(vault_id):
        _proposal_json = json.dumps(
            {
                "memory_type": mt,
                "memory_key": key,
                "memory_value": val,
                "memory_event_date": memory_event_date,
            },
            separators=(",", ":"),
        )
        label = label_for_type(mt)
        return (
            f"<<VAULTAI_MEMORY_PROPOSAL>>{_proposal_json}<<END>>\n\n"
            f"Got it.\n\n{label}: {val}\n\n"
            "I'll remember that for this vault (encrypted locally)."
        )

    result = update_memory_safe(
        vault_id, mt, key, val,
        event_date=memory_event_date,
    )
    if not result or not result.get("ok"):
        return None
                                                             
    try:
        from relationship_builder import build_relationships_for_memory_safe
        build_relationships_for_memory_safe(vault_id, mt, key)
    except Exception:
        pass

                                                                   
    try:
        from expiry_engine import build_expiry_alerts_for_memory_safe
        build_expiry_alerts_for_memory_safe(vault_id, mt, key)
    except Exception:
        pass

    label = label_for_type(mt)
    status = result.get("status")
    if status == "updated":
        prior = result.get("prior_value") or "(previous value)"
        return (
            f"Updated.\n\n{label}: {val}\n\n"
            f"I had \"{prior}\" before; updated to \"{val}\". "
            f"The previous value is kept in history."
        )
    if status == "unchanged":
        return (
            f"Already remembered.\n\n{label}: {val}"
        )
              
    return (
        f"Got it.\n\n{label}: {val}\n\n"
        "I'll remember that for this vault."
    )


def _handle_recall_memory(
    vault_id: str,
    memory_type: Optional[str],
    memory_query: Optional[str] = None,
    memory_anchor: Optional[str] = None,
    memory_direction: Optional[str] = None,
) -> Optional[str]:


    from ai_memory import is_enabled, ALLOWED_MEMORY_TYPES
    from memory_recall import (
        recall_ranked, render_grouped_by_type, render_rows, label_for_type,
    )
    from locales import ENGLISH_FORMATTER as _fmt
    if not is_enabled():
        return None
    mt = (memory_type or "").strip().lower() or None
    if mt and mt not in ALLOWED_MEMORY_TYPES:
        mt = None

    query = (memory_query or "").strip() or None
    anchor = (memory_anchor or "").strip() or None
    direction = (memory_direction or "").strip().lower() or None
    if direction and direction not in ("before", "after", "around"):
        direction = None

                                                                    
    if query or anchor or direction:
        rows = recall_ranked(
            vault_id,
            memory_type=mt, query_text=query,
            anchor_text=anchor, direction=direction,
            limit=15,
        )
        if not rows:
            return "I don't have anything matching that in memory yet."
        lines = [_fmt.remember_header()]
                                                               
        grouped: dict[str, list[dict]] = {}
        for r in rows:
            grouped.setdefault(r.get("memory_type") or "note", []).append(r)
        any_event_date = any(r.get("event_date") for r in rows)
        for t in _MEMORY_RENDER_ORDER:
            items = grouped.get(t)
            if not items:
                continue
            lines.append("")
            lines.append(label_for_type(t))
            lines.extend(render_rows(items, show_dates=any_event_date))
        return "\n".join(lines).rstrip()

                                                                      
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        if mt:
            cur.execute(
                """SELECT memory_type, memory_value FROM vault_ai_memory
                   WHERE vault_id=%s AND memory_type=%s
                     AND superseded_at IS NULL
                   ORDER BY updated_at DESC LIMIT 50""",
                (vault_id, mt),
            )
        else:
            cur.execute(
                """SELECT memory_type, memory_value FROM vault_ai_memory
                   WHERE vault_id=%s
                     AND superseded_at IS NULL
                   ORDER BY memory_type, updated_at DESC LIMIT 200""",
                (vault_id,),
            )
        rows = cur.fetchall() or []
    finally:
        conn.close()
    if not rows:
        return "I don't have anything saved in memory yet."
    grouped: dict[str, list[str]] = {}
    for r in rows:
        grouped.setdefault(r["memory_type"], []).append(r["memory_value"])
    lines = [_fmt.remember_header()]
    for t in _MEMORY_RENDER_ORDER:
        items = grouped.get(t)
        if not items:
            continue
        seen: set[str] = set()
        unique: list[str] = []
        for v in items:
            k = v.lower()
            if k in seen:
                continue
            seen.add(k)
            unique.append(v)
        lines.append("")
        lines.append(label_for_type(t))
        for v in unique:
            lines.append(f"- {v}")
    return "\n".join(lines).rstrip()
                                                                        

def _flag_enabled(env_var: str, *, default_true: bool = True) -> bool:


    raw = os.getenv(env_var)
    if raw is None:
        return default_true
    return raw.strip().lower() not in ("false", "0", "no", "off")


def _fetch_vault_name_for_prompt(vault_id: str) -> Optional[str]:
    """Server-authoritative lookup of the user-chosen vault name.

    Reads ``vaults.vault_name`` for the AUTHENTICATED vault_id
    (never trusts a client-supplied value on this turn). The vault
    name is the user-chosen identity for BOTH the vault and its AI
    keeper — one string, one meaning. Returns None on lookup
    failure, missing row, or NULL column; callers fall back to
    ``tools.VAULT_NAME_FALLBACK``.

    Failure is intentionally swallowed: a DB hiccup during prompt
    build must not fail the chat request. The prompt just uses the
    fallback name that turn.
    """
    try:
        conn = get_db()
        try:
            cur = conn.cursor()
            cur.execute(
                "SELECT vault_name FROM vaults WHERE vault_id = %s",
                (vault_id,),
            )
            row = cur.fetchone()
            if not row:
                return None
            value = row[0]
            if value is None:
                return None
            return str(value)
        finally:
            conn.close()
    except Exception:
        logger.exception("[PROMPT] vault_name lookup failed")
        return None


def _build_chat_prompt_context(
    *, vault_id: str, request: Request
) -> dict:


    has_memory = _flag_enabled("VAULTAI_AI_MEMORY_ENABLED", default_true=True)
    has_relationships = _flag_enabled(
        "VAULTAI_RELATIONSHIPS_ENABLED", default_true=True,
    )
    has_expiry = _flag_enabled(
        "VAULTAI_EXPIRY_DASHBOARD_ENABLED", default_true=True,
    )
    has_documents = _flag_enabled(
        "VAULTAI_DOCUMENT_UNDERSTANDING_ENABLED", default_true=True,
    )
    has_security_center = _flag_enabled(
        "VAULTAI_SECURITY_CENTER_ENABLED", default_true=True,
    )
    has_semantic = _flag_enabled(
        "VAULTAI_SEMANTIC_SEARCH_ENABLED", default_true=False,
    )

    capabilities = []
    if has_memory:
        capabilities.append("memory timeline")
    if has_relationships:
        capabilities.append("relationship graph")
    if has_expiry:
        capabilities.append("expiry tracking")
        capabilities.append("travel readiness")
    if has_documents:
        capabilities.append("document understanding")
    capabilities.extend(["files", "credentials", "inheritance"])
    if has_security_center:
        capabilities.append("security posture")
    if has_semantic:
        capabilities.append("semantic search")
    capabilities.extend(
        ["multilingual conversation", "vault reasoning", "natural interaction"],
    )

    enabled_features_str = ", ".join(capabilities) if capabilities else (
        "core vault only"
    )

                                                                      
    raw_locale = (request.headers.get("accept-language") or "").strip()
    locale_hint = raw_locale.split(",")[0].strip()[:16] or "auto"

    # SERVER-AUTHORITATIVE vault_name lookup. The chat request may
    # also carry a vault_name field on the wire, but the identity
    # slot in the prompt trusts ONLY the value from the authenticated
    # ``vaults`` row keyed by principal["vault_id"]. On any failure
    # the prompt builder returns None and
    # ``tools.build_vault_runtime_context`` substitutes the neutral
    # literal ``VaultAI`` — never a hash, handle, UUID, template
    # token, or the random 32-hex placeholder that used to leak
    # into the UI before 2026-07-20.
    #
    # 2026-07-22 chat deep-fix note: this slot is the AI keeper's
    # own name in the LLM's system prompt. It is the USER-CHOSEN
    # vault name (e.g. "Brain", "My Safe", "Family Vault") — a
    # per-vault value, NOT a hardcoded global constant. It must
    # NEVER be sourced from the user's DISPLAY NAME
    # (``users.display_name``), which is the human owner's own
    # label; the prior "Chosen is thinking..." regression came
    # from the frontend binding the ChatMessageList vault-name
    # parameter to ``app.displayName`` instead of to
    # ``app.vaultName``. Both fields have a legitimate purpose:
    # display_name identifies the human owner in UI headers /
    # greetings; vault_name is the vault's own identity used both
    # for signing in and for the AI keeper.
    vault_name = _fetch_vault_name_for_prompt(vault_id)

    return {
        "VAULT_NAME": vault_name,
        "VAULT_STATE": "unlocked",
        "LOCALE": locale_hint,
        "ENABLED_FEATURES": enabled_features_str,
        "HAS_MEMORY": "yes" if has_memory else "no",
        "HAS_RELATIONSHIPS": "yes" if has_relationships else "no",
        "HAS_EXPIRY": "yes" if has_expiry else "no",
    }


@app.post("/chat")
@limiter.limit("5/minute")
async def chat_endpoint(
    request: Request,
    req: ChatRequest,
    principal=Depends(verify_trusted_device),
):
    vault_id = principal["vault_id"]
    logger.debug("Chat request from vault: %s", vault_id)


    try:
        from rate_limit_sensitive import enforce_chat_rate_limit
        enforce_chat_rate_limit(request, vault_id=vault_id)
    except HTTPException:
        import security_event_log as _sec_log
        _sec_log.emit(
            reason=_sec_log.REASON_RATE_LIMITED_CHAT,
            route=_sec_log.ROUTE_GROUP_CHAT,
            subject=_sec_log.short_hash(vault_id),
        )
        raise

                                                                   
    print(
        f"[CHAT-DEBUG] entry vault_id={vault_id} vault_name={req.vault_name!r} "
        f"pin_len={len(req.pin or '')} encrypted_len={len(req.encrypted_message or '')} "
        f"uploaded_file_count={len(req.uploaded_file_ids or [])}",
        flush=True,
    )

    if not req.encrypted_message:
        print("[CHAT-DEBUG] missing_encrypted_message", flush=True)
        raise HTTPException(status_code=400, detail="Missing encrypted message")

    # 2026-07-22 (2) protocol-version enforcement + KDF version gate.
    #
    # (1) The client's declared `crypto_protocol_version` in the
    #     REQUEST BODY (not a header) is the AUTHORITATIVE boundary
    #     for whether the KDF fields are mandatory. >= 2 => required;
    #     missing => 400 missing_kdf_generation_fields (typed).
    #     Header-based enforcement was rejected on review — a
    #     spoofable header is not a valid crypto-protocol boundary.
    #
    # (2) Once the fields are declared, [check_kdf_generation_fresh]
    #     compares them to the current DB row and 409s on mismatch.
    #     See vault_kdf_generation.py for the full compare-and-swap
    #     semantics.
    from vault_kdf_generation import (
        check_kdf_generation_fresh,
        client_requires_kdf_fields,
        missing_kdf_fields_error,
        salt_fingerprint,
        key_fingerprint,
    )
    # X-App-Release stays as a diagnostic-only log field. It is
    # NEVER used to decide whether the crypto gate runs.
    _app_release = None
    try:
        _app_release = (
            request.headers.get("X-App-Release")
            or request.headers.get("x-app-release")
        )
    except Exception:
        _app_release = None
    _requires_kdf = client_requires_kdf_fields(
        req.crypto_protocol_version,
    )
    if _requires_kdf and (
        not req.kdf_salt_used or req.kdf_iterations_used is None
    ):
        # Modern-protocol client dropped the KDF fields — do NOT
        # fall into the legacy pass-through. Fail fast with a typed
        # code so the client can surface a specific message (and
        # the exact production regression cannot recur silently).
        print(
            "[CHAT-DEBUG] modern_protocol_missing_kdf "
            f"crypto_protocol_version={req.crypto_protocol_version} "
            f"app_release_diag={(_app_release or '-')[:12]}",
            flush=True,
        )
        raise missing_kdf_fields_error()

    # Diagnostic log — non-secret fingerprints only. Used to
    # correlate client-side [chat.body.diag] entries with server
    # decisions during incident triage. app_release is logged for
    # release-correlation ONLY; it does not gate anything.
    _chat_request_id = request.headers.get("X-Chat-Request-Id") or "-"
    _received_salt_fp = salt_fingerprint(req.kdf_salt_used)
    print(
        "[CHAT-DEBUG] kdf_gate_pre "
        f"chat_request_id={_chat_request_id} vault_id={vault_id} "
        f"crypto_protocol_version={req.crypto_protocol_version} "
        f"requires_kdf={_requires_kdf} "
        f"kdf_salt_used_present={bool(req.kdf_salt_used)} "
        f"received_salt_fp12={_received_salt_fp} "
        f"received_iterations={req.kdf_iterations_used} "
        f"app_release_diag={(_app_release or '-')[:12]}",
        flush=True,
    )

    _conn_for_kdf_check = get_db()
    try:
        # Read current DB fingerprint for the diag log BEFORE the
        # gate call (gate reads the same row; second read here is
        # informational only — never used for the actual gating).
        _cur = _conn_for_kdf_check.cursor()
        _cur.execute(
            "SELECT pin_salt, kdf_iterations FROM vaults "
            "WHERE vault_id = %s LIMIT 1",
            (vault_id,),
        )
        _row = _cur.fetchone()
        _db_salt = _row[0] if _row else None
        _db_iter = int(_row[1]) if _row and _row[1] is not None else None
        _db_salt_fp = salt_fingerprint(_db_salt)

        _gate_result: str
        try:
            check_kdf_generation_fresh(
                _conn_for_kdf_check,
                vault_id,
                req.kdf_salt_used,
                req.kdf_iterations_used,
            )
            _gate_result = (
                "legacy_pass"
                if (not req.kdf_salt_used
                    and req.kdf_iterations_used is None)
                else "match"
            )
        except HTTPException as _gate_exc:
            _gate_result = (
                "stale"
                if _gate_exc.status_code == 409
                else f"other_{_gate_exc.status_code}"
            )
            print(
                "[CHAT-DEBUG] kdf_gate_post "
                f"chat_request_id={_chat_request_id} "
                f"db_salt_fp12={_db_salt_fp} db_iter={_db_iter} "
                f"gate_result={_gate_result}",
                flush=True,
            )
            raise
        print(
            "[CHAT-DEBUG] kdf_gate_post "
            f"chat_request_id={_chat_request_id} "
            f"db_salt_fp12={_db_salt_fp} db_iter={_db_iter} "
            f"gate_result={_gate_result}",
            flush=True,
        )
    finally:
        _conn_for_kdf_check.close()

    try:
        print("[CHAT-DEBUG] verify_pin_start", flush=True)
        try:
            key = get_verified_vault_key(vault_id, req.pin)
        except HTTPException as e:
            print(
                f"[CHAT-DEBUG] verify_pin_failed status={e.status_code} detail={e.detail!r}",
                flush=True,
            )
            raise
        except Exception as e:
            print(
                f"[CHAT-DEBUG] verify_pin_failed error_type={type(e).__name__} repr={e!r}",
                flush=True,
            )
            raise
        _server_key_fp = key_fingerprint(key)
        print(
            "[CHAT-DEBUG] verify_pin_ok "
            f"chat_request_id={_chat_request_id} "
            f"server_key_fp12={_server_key_fp}",
            flush=True,
        )

        print(
            "[CHAT-DEBUG] decrypt_start "
            f"chat_request_id={_chat_request_id} "
            f"encrypted_len={len(req.encrypted_message)}",
            flush=True,
        )
        try:
            decrypted_message = decrypt_message(req.encrypted_message, key)
        except HTTPException as e:
            # Correlate with the client's [chat.body.diag] entry so
            # an incident triage can prove which key-vs-metadata
            # pairing the request carried without exposing any
            # secret material.
            print(
                "[CHAT-DEBUG] decrypt_failed "
                f"chat_request_id={_chat_request_id} "
                f"status={e.status_code} detail={e.detail!r} "
                f"server_key_fp12={_server_key_fp} "
                f"received_salt_fp12={_received_salt_fp} "
                f"gate_result={_gate_result}",
                flush=True,
            )
            raise
        except Exception as e:
            print(
                f"[CHAT-DEBUG] decrypt_failed "
                f"chat_request_id={_chat_request_id} "
                f"error_type={type(e).__name__} repr={e!r}",
                flush=True,
            )
            raise
        print(
            "[CHAT-DEBUG] decrypt_ok "
            f"chat_request_id={_chat_request_id}",
            flush=True,
        )



        try:
            import vault_multilingual as _mling
            _detected_lang = _mling.detect_language(
                decrypted_message or ""
            )
            _app_locale_hint = req.app_locale
            if not _app_locale_hint:
                try:
                    _app_locale_hint = (
                        request.headers.get("X-App-Locale")
                        or request.headers.get("x-app-locale")
                    )
                except Exception:
                    _app_locale_hint = None
            _header_locale_hint = None
            try:
                _header_locale_hint = (
                    request.headers.get("accept-language") or ""
                ).split(",")[0].strip() or None
            except Exception:
                _header_locale_hint = None
            _reply_language = _mling.resolve_reply_language(
                detected_from_message=_detected_lang,
                app_locale_hint=_app_locale_hint,
                header_locale_hint=_header_locale_hint,
            )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] reply-language resolution failed; "
                "falling back to English"
            )
            _reply_language = "en"


        # Selection-hint pin. When the user taps a row on a chat card
        # the frontend sends {"kind":"login"|"file","id":"<uuid>"} in
        # the request body — we pin the active entity BEFORE routing
        # so downstream builders can prefer id over query and never
        # confuse two rows that share a title.
        try:
            _hint = req.selection_hint
            if isinstance(_hint, dict):
                _hint_kind = str(_hint.get("kind") or "").strip().lower()
                _hint_id = str(_hint.get("id") or "").strip()
                if _hint_kind and _hint_id and len(_hint_id) <= 64:
                    _hint_session_id = str(
                        (principal or {}).get("token_id") or "",
                    ) or None
                    from vault_chat_active_entity import (
                        set_active_entity as _hint_set_ae,
                        ENTITY_LOGIN as _HINT_ENTITY_LOGIN,
                        ENTITY_FILE as _HINT_ENTITY_FILE,
                        ACTION_SHOW as _HINT_ACT_SHOW,
                        ACTION_OPEN as _HINT_ACT_OPEN,
                        ACTION_VIEW as _HINT_ACT_VIEW,
                        ACTION_COPY as _HINT_ACT_COPY,
                        ACTION_RENAME as _HINT_ACT_RENAME,
                        ACTION_DELETE as _HINT_ACT_DELETE,
                        ACTION_EDIT as _HINT_ACT_EDIT,
                        ACTION_DOWNLOAD as _HINT_ACT_DL,
                    )
                    if _hint_kind == "login":
                        _hint_set_ae(
                            vault_id,
                            entity_type=_HINT_ENTITY_LOGIN,
                            entity_ref={"id": _hint_id},
                            display_label="",
                            allowed_actions=(
                                _HINT_ACT_SHOW, _HINT_ACT_OPEN,
                                _HINT_ACT_VIEW, _HINT_ACT_COPY,
                                _HINT_ACT_RENAME, _HINT_ACT_DELETE,
                                _HINT_ACT_EDIT,
                            ),
                            session_id=_hint_session_id,
                        )
                        print(
                            "[CHAT-TRACE] selection_hint pinned "
                            f"kind=login id={_hint_id[:8]}...",
                            flush=True,
                        )
                    elif _hint_kind == "file":
                        _hint_set_ae(
                            vault_id,
                            entity_type=_HINT_ENTITY_FILE,
                            entity_ref={"file_id": _hint_id},
                            display_label="",
                            allowed_actions=(
                                _HINT_ACT_SHOW, _HINT_ACT_OPEN,
                                _HINT_ACT_VIEW, _HINT_ACT_DELETE,
                                _HINT_ACT_DL,
                            ),
                            session_id=_hint_session_id,
                        )
                        print(
                            "[CHAT-TRACE] selection_hint pinned "
                            f"kind=file id={_hint_id[:8]}...",
                            flush=True,
                        )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] selection_hint_pin_failed"
            )


        def encrypted_reply(text: str):
            print(
                f"[CHAT-DEBUG] encrypted_reply_start reply_text_len={len(text or '')}",
                flush=True,
            )
            resp = StreamingResponse(
                _stream_single_message(text, key),
                media_type="text/event-stream",
                headers={
                    "Cache-Control": "no-cache",
                    "Connection": "keep-alive",
                    "X-Accel-Buffering": "no",
                },
            )
            print("[CHAT-DEBUG] encrypted_reply_ok", flush=True)
            return resp

                                                               
        _direct_ai_tools_enabled = os.getenv(
            "VAULTAI_DIRECT_AI_TOOLS_ENABLED", "true",
        ).strip().lower() in ("1", "true", "yes", "on")

                                                                   
        def _route_to_ai_planner_stream():
            safe_message = redact_message(decrypted_message)
            prompt_context = _build_chat_prompt_context(
                vault_id=vault_id,
                request=request,
            )
            from tools import (
                STATIC_VAULT_SYSTEM_PROMPT,
                build_vault_runtime_context,
            )
            try:



                _effective_locale = _reply_language or str(
                    prompt_context.get("LOCALE", "auto")
                )
                _dynamic_context = build_vault_runtime_context(
                    # vault_name is the AUTHENTICATED value from the
                    # vaults row (or None). tools.build_vault_runtime_context
                    # substitutes "VaultAI" when None. Never pass
                    # req.vault_name here — that field carries client
                    # state and may be stale, missing, or a legacy
                    # value; the authoritative source is the DB row.
                    vault_name=prompt_context.get("VAULT_NAME"),
                    vault_state=str(prompt_context.get("VAULT_STATE", "unlocked")),
                    locale=_effective_locale,
                    enabled_features=str(prompt_context.get("ENABLED_FEATURES", "")),
                    has_memory=str(prompt_context.get("HAS_MEMORY", "on")),
                    has_relationships=str(prompt_context.get("HAS_RELATIONSHIPS", "on")),
                    has_expiry=str(prompt_context.get("HAS_EXPIRY", "on")),
                )
            except Exception:
                logger.exception("[CHAT-DEBUG] dynamic context build failed")
                _dynamic_context = ""
            messages: list[dict] = [
                {"role": "system", "content": STATIC_VAULT_SYSTEM_PROMPT},
            ]
            if _dynamic_context:
                messages.append({
                    "role": "system", "content": _dynamic_context,
                })
            try:
                from vault_history_compressor import compress_messages
                _compressed = compress_messages(messages)
                messages = _compressed.assembled()
            except Exception:
                logger.exception("[CHAT-DEBUG] history compression failed")
            messages.append({"role": "user", "content": safe_message})

            print("[CHAT-DEBUG] openai_stream_start", flush=True)

            async def generate():
                try:
                    chunk_count = 0
                    async for chunk in ai_stream(
                        messages,
                        vault_id,
                        key,
                        last_user_message=safe_message or "",
                        token_id=str(principal.get("token_id") or ""),
                    ):
                        chunk_str = chunk.decode("utf-8")
                        encrypted_chunk = encrypt_message(chunk_str, key)
                        chunk_count += 1
                        yield f"data: {encrypted_chunk}\n\n".encode("utf-8")
                    print(
                        f"[CHAT-DEBUG] openai_stream_ok chunks={chunk_count}",
                        flush=True,
                    )
                except Exception as e:
                    print(
                        f"[CHAT-DEBUG] openai_stream_failed error_type={type(e).__name__} repr={e!r}",
                        flush=True,
                    )
                    raise

            return StreamingResponse(
                generate(),
                media_type="text/event-stream",
                headers={
                    "Cache-Control": "no-cache",
                    "Connection": "keep-alive",
                    "X-Accel-Buffering": "no",
                },
            )






        _chat_started_at = time.monotonic()
        try:
            import chat_fast_path as _cfp
            _cfp.emit_span(
                _cfp.SPAN_REQUEST_START,
                vault_id=vault_id,
                message_len_bytes=len(decrypted_message or ""),
            )
        except Exception:
            _cfp = None

        _fast_envelope: dict | None = None











        try:
            from vault_chat_pronoun_followup import detect_pronoun_followup
            from vault_chat_active_entity import (
                get_active_entity,
                entity_matches_action,
                ENTITY_LOGIN,
                ENTITY_FILE,
                ENTITY_FILE_LIST,
                ENTITY_GENERATED_LOGIN_DRAFT,
                ENTITY_CRYPTO_WALLET,
            )
            _pronoun_hit = detect_pronoun_followup(decrypted_message)
            if _pronoun_hit is not None:
                _verb = _pronoun_hit["verb"]
                _session_id_for_ctx = str(
                    (principal or {}).get("token_id") or "",
                ) or None
                _active = get_active_entity(
                    vault_id,
                    session_id=_session_id_for_ctx,
                )
                print(
                    "[CHAT-DEBUG] pronoun_followup_check "
                    f"vault={(vault_id or '')[:8]}... "
                    f"verb={_verb!r} "
                    f"active={(_active or {}).get('entity_type')!r}",
                    flush=True,
                )
                if _active and entity_matches_action(_active, _verb):
                    _etype = _active.get("entity_type")
                    _eq = _active.get("query")


                    # All login pronoun follow-ups resolve back to a detail
                    # card for the same login. The frontend uses
                    # `pending_action` to decide whether to open edit mode,
                    # show the delete confirmation, auto-copy a field, or
                    # open the website.
                    _login_followup_verbs = (
                        "show", "open", "view",
                        "edit", "delete", "copy", "save",
                    )
                    if (
                        _etype == ENTITY_LOGIN
                        and _verb in _login_followup_verbs
                        and isinstance(_eq, str) and _eq.strip()
                    ):
                        try:
                            from vault_chat_router import (
                                build_vault_chat_envelope as _fp_build_envelope_lf,
                            )
                            _followup_envelope = _fp_build_envelope_lf(
                                f"find my {_eq} login",
                            )
                            if isinstance(_followup_envelope, dict):
                                try:
                                    from vault_chat_card_data import (
                                        populate_vault_chat_card_data,
                                    )
                                    _followup_envelope = populate_vault_chat_card_data(
                                        _followup_envelope,
                                        vault_id=vault_id,
                                        key=key,
                                    )
                                except Exception:
                                    logger.exception(
                                        "[CHAT-DEBUG] followup_populate_failed",
                                    )
                                # Annotate the card with a `pending_action`
                                # the frontend translates into a UI trigger.
                                # Never plaintext — just a verb tag.
                                try:
                                    _card = _followup_envelope.get("card")
                                    if isinstance(_card, dict):
                                        _card_data = _card.get("data")
                                        if isinstance(_card_data, dict):
                                            _card_data["pending_action"] = _verb
                                except Exception:
                                    pass
                                try:
                                    if isinstance(_followup_envelope, dict):
                                        _followup_envelope["locale"] = _reply_language
                                except Exception:
                                    pass
                                return encrypted_reply(
                                    json.dumps(_followup_envelope),
                                )
                        except Exception:
                            logger.exception(
                                "[CHAT-DEBUG] active_entity_dispatch_failed "
                                "type=login vault=%s",
                                (vault_id or "")[:8] + "...",
                            )




                    # File-list pagination follow-up: "show more" /
                    # "more" / "next" resurfaces the same file list
                    # card sliced at the next offset. Uses the stored
                    # active-entity ref (offset + page_size + total)
                    # so every subsequent page is deterministic and
                    # cannot skip or duplicate rows.
                    if (
                        _etype == ENTITY_FILE_LIST
                        and _verb == "more"
                    ):
                        try:
                            _ref = _active.get("entity_ref") or {}
                            _fl_offset = int(_ref.get("offset") or 0)
                            _fl_page = int(
                                _ref.get("page_size") or 0,
                            ) or 25
                            rows = list_uploaded_files(vault_id)
                            _new_total = len(rows)
                            _fl_reply = _build_vault_file_list_envelope(
                                rows,
                                f"More files ({_fl_offset + 1}"
                                f"–{min(_fl_offset + _fl_page, _new_total)}"
                                f" of {_new_total})",
                                message="Here are the next files.",
                                offset=_fl_offset,
                                page_size=_fl_page,
                            )
                            _fl_parsed = json.loads(_fl_reply)
                            _fl_next = int(
                                _fl_parsed.get("next_offset") or _fl_offset,
                            )
                            _fl_has_more = bool(
                                _fl_parsed.get("has_more"),
                            )
                            try:
                                from vault_chat_active_entity import (
                                    set_active_entity as _fl_set_ae,
                                    ENTITY_FILE_LIST as _FL_ENTITY,
                                    ACTION_MORE as _FL_ACT_MORE,
                                )
                                _fl_set_ae(
                                    vault_id,
                                    entity_type=_FL_ENTITY,
                                    entity_ref={
                                        "offset":      _fl_next,
                                        "page_size":   _fl_page,
                                        "total_count": _new_total,
                                    },
                                    display_label="All files",
                                    allowed_actions=(
                                        (_FL_ACT_MORE,)
                                        if _fl_has_more else ()
                                    ),
                                    session_id=_session_id_for_ctx,
                                )
                            except Exception:
                                logger.exception(
                                    "[CHAT-DEBUG] file_list_more "
                                    "set_active_entity failed",
                                )
                            print(
                                "[CHAT-TRACE] file_list_more "
                                f"vault={(vault_id or '')[:8]}... "
                                f"offset={_fl_offset} next={_fl_next} "
                                f"has_more={_fl_has_more}",
                                flush=True,
                            )
                            return encrypted_reply(_fl_reply)
                        except Exception:
                            logger.exception(
                                "[CHAT-DEBUG] file_list_more_dispatch_failed "
                                "vault=%s",
                                (vault_id or "")[:8] + "...",
                            )

                    # File pronoun follow-up:
                    #   "download it"  -> resurface the ACTIVE FILE as a
                    #                     structured file card whose data
                    #                     carries pending_action="download".
                    #                     Frontend triggers _downloadVaultFileCard.
                    #   "open it" / "view it" -> same, pending_action="open".
                    # The backend NEVER hands the browser plaintext bytes
                    # here; it only re-emits the same safe file card the
                    # user just saw and tags the pending action.
                    elif (
                        _etype == ENTITY_FILE
                        and _verb in ("download", "open", "view", "show")
                    ):
                        try:
                            _ref = _active.get("entity_ref") or {}
                            _file_id = str(_ref.get("file_id") or "")
                            _display = str(
                                _active.get("display_label") or ""
                            )
                            _mime = str(_ref.get("content_type") or "")
                            _rel = str(_ref.get("relative_path") or "")
                            if _file_id:
                                _payload = {
                                    "type":         "vault_file",
                                    "message":      "",
                                    "file_id":      _file_id,
                                    "file_name":    _display,
                                    # content_type — matches the frontend
                                    # parser at main.dart:7797 which reads
                                    # this key, not mime_type.
                                    "content_type": _mime or None,
                                    "relative_path": _rel or None,
                                    "pending_action": _verb,
                                }
                                print(
                                    "[CHAT-TRACE] file_pronoun_followup "
                                    f"vault={(vault_id or '')[:8]}... "
                                    f"verb={_verb} file={_file_id[:8]}...",
                                    flush=True,
                                )
                                return encrypted_reply(json.dumps(_payload))
                        except Exception:
                            logger.exception(
                                "[CHAT-DEBUG] file_followup_dispatch_failed "
                                "vault=%s", (vault_id or "")[:8] + "...",
                            )

                    elif (
                        _etype == ENTITY_GENERATED_LOGIN_DRAFT
                        and _verb == "save"
                    ):
                        print(
                            "[CHAT-DEBUG] pronoun_followup_deferred_to_pending "
                            f"vault={(vault_id or '')[:8]}...",
                            flush=True,
                        )

                    # Crypto Vault follow-up: "open it" / "take me
                    # there" / "use it" when the last surfaced entity
                    # was the Crypto Vault card. Rebuild the same
                    # delegated show-vault envelope so the entitlement
                    # branch runs — non-upgraded users still see the
                    # upgrade CTA, entitled users get the Open button.
                    # This is the fix for the "'open it' returns
                    # nothing" bug.
                    elif (
                        _etype == ENTITY_CRYPTO_WALLET
                        and _verb in ("open", "show", "view")
                    ):
                        try:
                            from vault_chat_router import (
                                build_crypto_delegated_show_vault_envelope,
                            )
                            _followup_envelope = (
                                build_crypto_delegated_show_vault_envelope()
                            )
                            _cv_user_tier = "free"
                            try:
                                from billing import (
                                    get_account_id_for_vault as _cv_acct,
                                    get_entitlement as _cv_ent,
                                )
                                _cv_acct_id = _cv_acct(vault_id)
                                if _cv_acct_id:
                                    _cv_e = _cv_ent(_cv_acct_id)
                                    if (
                                        _cv_e.block_count > 0
                                        and _cv_e.purchased_bytes > 0
                                    ):
                                        _cv_user_tier = "upgraded"
                                    else:
                                        _cv_user_tier = "free"
                            except Exception:
                                _cv_user_tier = "free"
                            try:
                                from vault_chat_crypto_data import (
                                    populate_crypto_delegated_card_data,
                                )
                                _followup_envelope = (
                                    populate_crypto_delegated_card_data(
                                        _followup_envelope,
                                        vault_id=vault_id,
                                        client_platform=None,
                                        user_tier=_cv_user_tier,
                                    )
                                )
                            except Exception:
                                logger.exception(
                                    "[CHAT-DEBUG] crypto_followup_populate_failed",
                                )
                            try:
                                if isinstance(_followup_envelope, dict):
                                    _followup_envelope["locale"] = _reply_language
                            except Exception:
                                pass
                            print(
                                "[CHAT-TRACE] crypto_pronoun_followup "
                                f"vault={(vault_id or '')[:8]}... "
                                f"verb={_verb} tier={_cv_user_tier}",
                                flush=True,
                            )
                            return encrypted_reply(
                                json.dumps(_followup_envelope),
                            )
                        except Exception:
                            logger.exception(
                                "[CHAT-DEBUG] crypto_followup_dispatch_failed "
                                "vault=%s",
                                (vault_id or "")[:8] + "...",
                            )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] pronoun_followup_check_failed "
                "vault=%s",
                (vault_id or "")[:8] + "...",
            )

        if _cfp is not None:
            try:
                _cfp.emit_span(
                    _cfp.SPAN_FAST_ROUTER_START,
                    vault_id=vault_id,
                    started_at_monotonic=_chat_started_at,
                )
                from vault_chat_router import (
                    build_vault_chat_envelope as _fp_build_envelope,
                )
                _fast_envelope = _cfp.peek_intent_without_side_effects(
                    decrypted_message or "",
                    build_envelope=_fp_build_envelope,
                )



                if _fast_envelope is None:
                    try:
                        _translated_query = (
                            _mling.translate_to_english_router_query(
                                decrypted_message or ""
                            )
                        )
                    except Exception:
                        _translated_query = None
                    if _translated_query:
                        try:
                            _fast_envelope = (
                                _cfp.peek_intent_without_side_effects(
                                    _translated_query,
                                    build_envelope=_fp_build_envelope,
                                )
                            )
                        except Exception:
                            logger.exception(
                                "[CHAT-PERF] multilingual re-route failed"
                            )
                            _fast_envelope = None
                _cfp.emit_span(
                    _cfp.SPAN_FAST_ROUTER_DONE,
                    vault_id=vault_id,
                    started_at_monotonic=_chat_started_at,
                    intent=(
                        _fast_envelope.get("intent")
                        if isinstance(_fast_envelope, dict) else None
                    ),
                )
            except Exception:
                logger.exception(
                    "[CHAT-PERF] fast-path peek failed; drop through"
                )
                _fast_envelope = None


        _fast_pending_confirm = False
        _fast_has_active_context = False
        if _cfp is not None and _fast_envelope is not None:
            try:
                from vault_pending_draft_confirm import (
                    is_pending_draft_confirm_phrase as _fp_pending_check,
                    is_save_themed_confirm_phrase as _fp_save_themed,
                )
                _fast_pending_confirm = (
                    _fp_pending_check(decrypted_message or "")
                    or _fp_save_themed(decrypted_message or "")
                )
            except Exception:
                _fast_pending_confirm = True

            try:
                from vault_active_context import (
                    get_active_context as _fp_get_active_ctx,
                )
                _fast_has_active_context = bool(
                    _fp_get_active_ctx(vault_id)
                )
            except Exception:
                _fast_has_active_context = True


        # Short-circuit: "show me all my files" — the router tags this
        # as vault_file_list_all with an empty card. Replace it with a
        # real structured vault_file_list envelope so the frontend
        # never has to render prose for a whole-vault listing.
        # Also record the active file-list entity so a subsequent
        # "show more" chat follow-up can request page 2 without the
        # frontend having to re-issue "show me all my files".
        if (
            isinstance(_fast_envelope, dict)
            and str(_fast_envelope.get("intent") or "")
                == "vault_file_list_all"
        ):
            try:
                rows = list_uploaded_files(vault_id)
                total = len(rows)
                title = (
                    f"All {total} files in your vault"
                    if total > 0 else "Your vault is empty"
                )
                reply = _build_vault_file_list_envelope(
                    rows,
                    title,
                    message="Here are your files.",
                    offset=0,
                )
                try:
                    _reply_parsed = json.loads(reply)
                    _next_offset = int(
                        _reply_parsed.get("next_offset") or 0,
                    )
                    _page_size = int(
                        _reply_parsed.get("page_size") or 0,
                    )
                    _has_more = bool(_reply_parsed.get("has_more"))
                    _fla_session_id = str(
                        (principal or {}).get("token_id") or "",
                    ) or None
                    from vault_chat_active_entity import (
                        set_active_entity as _fla_set_ae,
                        ENTITY_FILE_LIST as _FLA_ENTITY,
                        ACTION_MORE as _FLA_ACT_MORE,
                    )
                    _fla_set_ae(
                        vault_id,
                        entity_type=_FLA_ENTITY,
                        entity_ref={
                            "offset":      _next_offset,
                            "page_size":   _page_size,
                            "total_count": total,
                        },
                        display_label="All files",
                        allowed_actions=(
                            (_FLA_ACT_MORE,) if _has_more else ()
                        ),
                        session_id=_fla_session_id,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-PERF] file_list_all set_active_entity "
                        "failed vault=%s",
                        (vault_id or "")[:8] + "...",
                    )
                return encrypted_reply(reply)
            except Exception:
                logger.exception(
                    "[CHAT-PERF] file_list_all fast-path failed vault=%s",
                    (vault_id or "")[:8] + "...",
                )

        if (
            _cfp is not None
            and _cfp.can_skip_drains(
                _fast_envelope,
                is_pending_confirm=_fast_pending_confirm,
                has_active_context=_fast_has_active_context,
            )
        ):
            try:
                from vault_chat_card_data import (
                    populate_vault_chat_card_data,
                )
                _fast_envelope = populate_vault_chat_card_data(
                    _fast_envelope,
                    vault_id=vault_id,
                    key=key,
                )
            except Exception:
                logger.exception(
                    "[CHAT-PERF] fast path populate_vault_chat_card_data "
                    "failed; response will still be sent from router "
                    "envelope"
                )
            try:
                _client_platform_fp = None
                try:
                    _client_platform_fp = (
                        request.headers.get("X-Client-Platform")
                        or request.headers.get("x-client-platform")
                    )
                    if _client_platform_fp:
                        _client_platform_fp = (
                            _client_platform_fp.strip().lower()
                        )
                except Exception:
                    _client_platform_fp = None
                from vault_chat_crypto_data import (
                    populate_crypto_delegated_card_data,
                )

                _fp_user_tier = "free"
                try:
                    from billing import (
                        get_account_id_for_vault as _fp_get_acct,
                        get_entitlement as _fp_get_ent,
                    )
                    _fp_acct = _fp_get_acct(vault_id)
                    if _fp_acct:
                        _fp_ent = _fp_get_ent(_fp_acct)
                        if (
                            _fp_ent.block_count > 0
                            and _fp_ent.purchased_bytes > 0
                        ):
                            _fp_user_tier = "upgraded"
                except Exception:
                    logger.exception(
                        "[CHAT-PERF] fast path crypto tier lookup failed "
                        "vault=%s; defaulting to free",
                        (vault_id or "")[:8] + "...",
                    )
                    _fp_user_tier = "free"
                _fast_envelope = populate_crypto_delegated_card_data(
                    _fast_envelope,
                    vault_id=vault_id,
                    client_platform=_client_platform_fp,
                    user_tier=_fp_user_tier,
                )
            except Exception:
                logger.exception(
                    "[CHAT-PERF] fast path populate_crypto_delegated_card_data "
                    "failed; response will still be sent"
                )



            try:
                import vault_faq_content_i18n as _faq_i18n
                _fast_envelope = _faq_i18n.apply_locale_to_faq_envelope(
                    _fast_envelope, _reply_language,
                )
            except Exception:
                logger.exception(
                    "[CHAT-PERF] fast path apply_locale_to_faq_envelope "
                    "failed; response will still be sent in English"
                )
            try:
                if isinstance(_fast_envelope, dict):
                    _fast_envelope["locale"] = _reply_language
            except Exception:
                pass









            try:
                _fp_intent_str = str(_fast_envelope.get("intent") or "")
                _fp_card = _fast_envelope.get("card") or {}
                _fp_query = (
                    _fp_card.get("query")
                    if isinstance(_fp_card, dict) else None
                )
                _fp_session_id = str(
                    (principal or {}).get("token_id") or ""
                ) or None
                _fp_data = (
                    _fp_card.get("data")
                    if isinstance(_fp_card, dict) else None
                )
                _fp_login_id: str = ""
                _fp_login_label: str = ""
                _fp_is_multi: bool = False
                _fp_candidates: list = []
                if isinstance(_fp_data, dict):
                    _view = str(_fp_data.get("view") or "")
                    if _view == "detail" and isinstance(
                        _fp_data.get("login"), dict,
                    ):
                        _fp_login_id = str(
                            _fp_data["login"].get("id") or ""
                        )
                        _fp_login_label = str(
                            _fp_data["login"].get("title") or ""
                        )
                    elif _view == "chooser" and isinstance(
                        _fp_data.get("logins"), list,
                    ):
                        _fp_is_multi = True
                        for _row in _fp_data["logins"][:20]:
                            if not isinstance(_row, dict):
                                continue
                            _rid = str(_row.get("id") or "")
                            _rtitle = str(_row.get("title") or "")
                            if _rid:
                                _fp_candidates.append({
                                    "id": _rid,
                                })
                            _ = _rtitle  # kept plaintext-free
                _login_intents = (
                    "vault_login_search",
                    "vault_login_reveal",
                    "vault_login_copy",
                )
                # 2026-07-12: "show me all my logins" that collapses to a
                # single detail card must pin the active entity to that
                # login so "edit it" / "copy the password" / "delete it"
                # follow-ups work — same shape as INTENT_LOGIN_SEARCH
                # with one hit. When the count is >1 the fast-path
                # renders a chooser (existing behavior below); when
                # count is 0 the not_found view carries no id to pin.
                _login_list_single = (
                    _fp_intent_str == "vault_login_list"
                    and isinstance(_fp_data, dict)
                    and _fp_data.get("view") == "detail"
                    and bool(_fp_login_id)
                )
                if _login_list_single:
                    from vault_chat_active_entity import (
                        set_active_entity as _lls_set_ae,
                        ENTITY_LOGIN as _LLS_ENTITY_LOGIN,
                        ACTION_SHOW as _LLS_ACT_SHOW,
                        ACTION_OPEN as _LLS_ACT_OPEN,
                        ACTION_VIEW as _LLS_ACT_VIEW,
                        ACTION_COPY as _LLS_ACT_COPY,
                        ACTION_RENAME as _LLS_ACT_RENAME,
                        ACTION_DELETE as _LLS_ACT_DELETE,
                        ACTION_EDIT as _LLS_ACT_EDIT,
                        ACTION_SAVE as _LLS_ACT_SAVE,
                    )
                    _lls_set_ae(
                        vault_id,
                        entity_type=_LLS_ENTITY_LOGIN,
                        entity_ref={"id": _fp_login_id},
                        display_label=(_fp_login_label or "Login"),
                        allowed_actions=(
                            _LLS_ACT_SHOW, _LLS_ACT_OPEN, _LLS_ACT_VIEW,
                            _LLS_ACT_COPY, _LLS_ACT_RENAME, _LLS_ACT_DELETE,
                            _LLS_ACT_EDIT, _LLS_ACT_SAVE,
                        ),
                        session_id=_fp_session_id,
                    )

                if _fp_intent_str in _login_intents and \
                   isinstance(_fp_query, str) and _fp_query.strip():
                    from vault_chat_active_entity import (
                        set_active_entity,
                        ENTITY_LOGIN,
                        ACTION_SHOW, ACTION_OPEN, ACTION_VIEW,
                        ACTION_COPY, ACTION_RENAME, ACTION_DELETE,
                        ACTION_EDIT, ACTION_SAVE,
                    )
                    _ref: dict = {}
                    if _fp_login_id:

                        _ref = {"id": _fp_login_id}
                    else:

                        _ref = {"query": _fp_query.strip()}
                    _label = (
                        _fp_login_label
                        or _fp_query.strip()
                    )
                    set_active_entity(
                        vault_id,
                        entity_type=ENTITY_LOGIN,
                        entity_ref=_ref,
                        display_label=_label,
                        query=_fp_query.strip(),
                        allowed_actions=(
                            ACTION_SHOW, ACTION_OPEN, ACTION_VIEW,
                            ACTION_COPY, ACTION_RENAME, ACTION_DELETE,
                            ACTION_EDIT, ACTION_SAVE,
                        ),
                        session_id=_fp_session_id,
                        is_multi=_fp_is_multi,
                        candidates=_fp_candidates,
                    )

                # When we surface a crypto-vault delegated show-vault
                # card, remember that the last active entity is the
                # Crypto Vault. A subsequent bare "open it" / "take me
                # there" / "use it" is then routed through the same
                # entitlement-aware envelope by the pronoun follow-up
                # dispatcher above.
                if (
                    _fp_intent_str == "vault_crypto_delegated"
                    and isinstance(_fp_card, dict)
                    and _fp_card.get("innerIntent") == "crypto_vault_show_vault"
                ):
                    from vault_chat_active_entity import (
                        set_active_entity as _fp_set_ae,
                        ENTITY_CRYPTO_WALLET as _FP_ENTITY_CV,
                        ACTION_OPEN as _FP_ACT_OPEN,
                        ACTION_SHOW as _FP_ACT_SHOW,
                        ACTION_VIEW as _FP_ACT_VIEW,
                        ACTION_UPGRADE as _FP_ACT_UPGRADE,
                    )
                    _fp_set_ae(
                        vault_id,
                        entity_type=_FP_ENTITY_CV,
                        entity_ref={},
                        display_label="Crypto Vault",
                        allowed_actions=(
                            _FP_ACT_OPEN, _FP_ACT_SHOW,
                            _FP_ACT_VIEW, _FP_ACT_UPGRADE,
                        ),
                        session_id=_fp_session_id,
                    )
            except Exception:
                logger.exception(
                    "[CHAT-DEBUG] set_active_entity failed "
                    "vault=%s", (vault_id or "")[:8] + "...",
                )

            _cfp.emit_span(
                _cfp.SPAN_RESPONSE_READY,
                vault_id=vault_id,
                started_at_monotonic=_chat_started_at,
                intent=str(_fast_envelope.get("intent") or ""),
                extra=f"fast_path=1 lang={_reply_language}",
            )
            return encrypted_reply(json.dumps(_fast_envelope))


        if _cfp is not None:
            _cfp.emit_span(
                _cfp.SPAN_DRAIN_SCHEDULED,
                vault_id=vault_id,
                started_at_monotonic=_chat_started_at,
            )

        try:
            from vault_analysis_worker import drain_text_extraction
            drain_text_extraction(vault_id=vault_id, key=key)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] text_extraction_drain failed vault=%s",
                vault_id,
            )

                                                                
        try:
            from vault_ocr_worker import drain_ocr
            drain_ocr(vault_id=vault_id, key=key)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] ocr_drain failed vault=%s",
                vault_id,
            )

                                                              
        try:
            from vault_audio_worker import drain_audio_transcription
            drain_audio_transcription(vault_id=vault_id, key=key)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] audio_transcription_drain failed vault=%s",
                vault_id,
            )

                                                                
        try:
            from vault_video_worker import drain_video_transcription
            drain_video_transcription(vault_id=vault_id, key=key)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] video_transcription_drain failed vault=%s",
                vault_id,
            )

                                                              
        try:
            from vault_archive_worker import drain_archive_indexing
            drain_archive_indexing(vault_id=vault_id, key=key)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] archive_indexing_drain failed vault=%s",
                vault_id,
            )

                                                              
        try:
            import vault_understanding as vu
            vu.mark_stale_understandings_for_changed_text(vault_id)
            vu.mark_stale_embeddings_for_changed_text(vault_id)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] stale_sweep failed vault=%s",
                vault_id,
            )

                                                                  
        try:
            from vault_understanding_worker import drain_file_understanding
            drain_file_understanding(vault_id=vault_id, key=key)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] file_understanding_drain failed vault=%s",
                vault_id,
            )

                                                             
        try:
            from vault_embedding_worker import drain_file_embedding
            drain_file_embedding(vault_id=vault_id, key=key)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] file_embedding_drain failed vault=%s",
                vault_id,
            )

                                                        
        try:
            from vault_relationship_graph import (
                cleanup_stale_relationships_for_vault,
            )
            cleanup_report = cleanup_stale_relationships_for_vault(
                vault_id,
            )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] relationship_cleanup failed "
                "vault=%s", vault_id,
            )
            cleanup_report = None

                                                             
        try:
            if cleanup_report and (
                cleanup_report.get("deleted_orphans", 0) > 0
                or cleanup_report.get("stale_version_rows", 0) > 0
            ):
                from vault_relationship_worker import (
                    enqueue_relationship_building,
                )
                enqueue_relationship_building(vault_id=vault_id)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] relationship_cleanup_enqueue failed "
                "vault=%s", vault_id,
            )

                                                             
        try:
            from vault_relationship_worker import (
                drain_relationship_building,
            )
            drain_relationship_building(vault_id=vault_id, key=key)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] relationship_building_drain failed "
                "vault=%s", vault_id,
            )


        print("[CHAT-DEBUG] memory_start", flush=True)
        try:
            memory = get_memory(vault_id)
        except Exception as e:
            print(
                f"[CHAT-DEBUG] memory_failed error_type={type(e).__name__} repr={e!r}",
                flush=True,
            )
            raise
        print("[CHAT-DEBUG] memory_ok", flush=True)

        # -----------------------------------------------------------
        # 2026-07-24 chat-brain rebuild — semantic decider band.
        #
        # Runs BEFORE the deterministic pending-confirm cascade so
        # a user's response to a pending action is interpreted
        # semantically against the actual live pending state,
        # rather than by "whichever store's regex matches first"
        # (the pre-rebuild bug that let a stale unnamed video
        # hijack a delete confirmation). Falls through on any
        # decider error, low-confidence guess, or unknown-tool
        # response so the existing cascade continues to handle
        # everything the brain doesn't own.
        try:
            from vault_chat_brain import run_chat_brain
            _brain_result = await run_chat_brain(
                vault_id=vault_id,
                session_id=(
                    principal.get("token_id") if isinstance(principal, dict)
                    else None
                ),
                turn_id=str(_chat_request_id or ""),
                vault_name=req.vault_name or "",
                reply_language=(_reply_language or "en"),
                user_message=decrypted_message or "",
                memory=memory,
                key=key,
                unlocked=True,
                user_tier=(
                    principal.get("user_tier") if isinstance(principal, dict)
                    else "free"
                ) or "free",
                features={},
            )
        except Exception:
            logger.exception(
                "[CHAT-BRAIN] band_crashed vault=%s",
                (vault_id or "")[:8] + "…",
            )
            _brain_result = None

        if _brain_result is not None and _brain_result.handled:
            return encrypted_reply(_brain_result.reply_text)


        try:
            from vault_pending_draft_confirm import (
                NO_DRAFT_FRIENDLY_REPLY,
                is_pending_draft_confirm_phrase,
                is_save_themed_confirm_phrase,
            )
            _pending_confirm = is_pending_draft_confirm_phrase(
                decrypted_message or "",
            )
            _pending_save_themed = is_save_themed_confirm_phrase(
                decrypted_message or "",
            )
        except Exception:
            _pending_confirm = False
            _pending_save_themed = False
            NO_DRAFT_FRIENDLY_REPLY = (
                "I don't have a pending save right now. Tell me "
                "what you'd like to save first."
            )                

        if _pending_confirm:
                                                                    
             
            _credential_draft = None
            try:
                from vault_credential_draft import consume_draft as _consume_credential_draft
                _credential_draft = _consume_credential_draft(
                    vault_id=vault_id,
                )
            except Exception:
                logger.exception(
                    "[CHAT-DEBUG] credential_draft_consume_failed "
                    "vault=%s", (vault_id or "")[:8] + "...",
                )
                _credential_draft = None

            if _credential_draft is not None:
                _draft_service = str(_credential_draft.service_name or "")
                _draft_username = str(_credential_draft.username or "")
                _draft_password = str(_credential_draft.password or "")
                _new_fields = {"password": _draft_password}
                if _draft_username:
                    _new_fields["username"] = _draft_username
                try:
                    save_secret_tool(
                        vault_id,
                        {
                            "secret_type": "login",
                            "service": _draft_service,
                            "fields": _new_fields,
                        },
                        key,
                        generated=True,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] confirm_save_credential_draft_failed "
                        "vault=%s", (vault_id or "")[:8] + "...",
                    )
                    return encrypted_reply(
                        f"I couldn't save the {_draft_service.title()} "
                        "login right now. Try again in a moment."
                    )
                                                            
                                                              
                try:
                    memory.pop("pending_login_draft", None)
                except Exception:
                    pass
                try:
                    from vault_active_context import (
                        clear_active_context as _clear_active_ctx_now,
                    )
                    _clear_active_ctx_now(vault_id)
                except Exception:
                    pass
                memory["last_generated_login"] = {
                    "service": _draft_service,
                    "generated_fields": list(_new_fields.keys()),
                    "ts": int(time.time()),
                }
                                                           
                                                                  
                logger.info(
                    "[CHAT-TRACE] pending_credential_confirmed "
                    "vault=%s service=%s draft_prefix=%s store=persistent",
                    (vault_id or "")[:8] + "...",
                    _draft_service[:16],
                    str(_credential_draft.draft_id)[:8],
                )
                return encrypted_reply(
                    f"Saved your {_draft_service.title()} login to "
                    "your vault \U0001F510"
                )

            try:
                _login_draft = memory.get("pending_login_draft")
            except Exception:
                _login_draft = None
            if (
                isinstance(_login_draft, dict)
                and _login_draft.get("service")
            ):
                                                             
                                                             
                _draft_service = str(_login_draft.get("service") or "")
                _draft_password = str(_login_draft.get("password") or "")
                _opts = list(_login_draft.get("username_options") or [])
                _picked_username = _opts[0] if _opts else None
                _draft_email_required = bool(
                    _login_draft.get("policy_email_required")
                )
                if _draft_email_required and not _login_draft.get(
                    "existing_email"
                ):
                    return encrypted_reply(
                        f"{_draft_service.title()} needs an email "
                        "address as the username — send me the "
                        "email you want to use and I'll save it "
                        "with this password."
                    )
                _new_fields: dict = {"password": _draft_password}
                if _picked_username:
                    _new_fields["username"] = _picked_username
                try:
                    save_secret_tool(
                        vault_id,
                        {
                            "secret_type": "login",
                            "service": _draft_service,
                            "fields": _new_fields,
                        },
                        key,
                        generated=True,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] confirm_save_credential_failed "
                        "vault=%s", (vault_id or "")[:8] + "...",
                    )
                    return encrypted_reply(
                        f"I couldn't save the {_draft_service.title()} "
                        "login right now. Try again in a moment."
                    )
                                                               
                                                
                try:
                    memory.pop("pending_login_draft", None)
                except Exception:
                    pass
                try:
                    from vault_active_context import (
                        clear_active_context as _clear_active_ctx_now,
                    )
                    _clear_active_ctx_now(vault_id)
                except Exception:
                    pass
                memory["last_generated_login"] = {
                    "service": _draft_service,
                    "generated_fields": list(_new_fields.keys()),
                    "ts": int(time.time()),
                }
                                                            
                                                                  
                logger.info(
                    "[CHAT-TRACE] pending_credential_confirmed "
                    "vault=%s service=%s",
                    (vault_id or "")[:8] + "...",
                    _draft_service[:16],
                )
                return encrypted_reply(
                    f"Saved your {_draft_service.title()} login to "
                    "your vault \U0001F510"
                )

            # 2026-07-22 chat deep-fix — attachment save via pure
            # confirm phrase ("save this", "save it", "put this in
            # my vault", etc.). If neither a credential draft nor a
            # pending_login_draft matched but a chat-uploaded file
            # is awaiting naming (needs_naming = TRUE in
            # uploaded_files), commit it using the file's original
            # filename as the default saved name. This is what the
            # user expects when they upload an image/video and then
            # say "save this" without specifying a name.
            # Attachments are DB-backed, so this fallback is
            # cross-worker safe with no per-process state.
            try:
                _pending_attachment = get_pending_named_file(vault_id)
            except Exception:
                logger.exception(
                    "[CHAT-DEBUG] pending_attachment_lookup_failed "
                    "vault=%s", (vault_id or "")[:8] + "...",
                )
                _pending_attachment = None
            if _pending_attachment is not None:
                _att_ct = (
                    _pending_attachment.get("content_type") or ""
                ).lower()
                if _att_ct.startswith("audio/"):
                    _att_noun = "recording"
                elif _att_ct.startswith("video/"):
                    _att_noun = "video"
                elif _att_ct.startswith("image/"):
                    _att_noun = "image"
                else:
                    _att_noun = "file"
                _att_default_name = str(
                    _pending_attachment.get("file_name") or ""
                ).strip()
                if not _att_default_name:
                    _att_default_name = _att_noun
                _att_clean_name = _normalize_asset_name(
                    _att_default_name,
                )
                if not _att_clean_name or _att_clean_name == "general":
                    _att_clean_name = _att_noun
                try:
                    _att_result = save_named_uploaded_asset(
                        vault_id=vault_id,
                        file_id=_pending_attachment["id"],
                        saved_name=_att_clean_name,
                        key=key,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] confirm_save_attachment_failed "
                        "vault=%s file_id=%s",
                        (vault_id or "")[:8] + "...",
                        (_pending_attachment.get("id") or "")[:8] + "...",
                    )
                    # False-success guard (item 7 of the deep-fix):
                    # never claim we saved when the save failed.
                    return encrypted_reply(
                        f"I couldn't save that {_att_noun} right "
                        "now. Try again in a moment."
                    )
                # Post-save verification: save_named_uploaded_asset
                # must return a dict with a saved_name field on
                # success. If either is missing, treat as failure
                # (defensive false-success guard).
                if (
                    not isinstance(_att_result, dict)
                    or not _att_result.get("saved_name")
                ):
                    logger.warning(
                        "[CHAT-DEBUG] confirm_save_attachment_shape "
                        "vault=%s result=%r",
                        (vault_id or "")[:8] + "...",
                        _att_result,
                    )
                    return encrypted_reply(
                        f"I couldn't save that {_att_noun} right "
                        "now. Try again in a moment."
                    )
                logger.info(
                    "[CHAT-TRACE] pending_attachment_confirmed "
                    "vault=%s file_id=%s content_type=%s",
                    (vault_id or "")[:8] + "...",
                    (_pending_attachment.get("id") or "")[:8] + "...",
                    _att_ct or "?",
                )
                return encrypted_reply(
                    f"Saved this {_att_noun} as "
                    f"{_title_case_asset(_att_result['saved_name'])}."
                )



        try:
            from vault_chat_router import (
                build_vault_chat_envelope as _vcr_build_envelope,
            )
            try:
                from vault_active_context import (
                    get_active_context as _vcr_get_active_context,
                )
                _vcr_active_ctx = _vcr_get_active_context(vault_id)
            except Exception:
                _vcr_active_ctx = None
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] vault_chat_router_import_failed vault=%s",
                (vault_id or "")[:8] + "...",
            )
            _vcr_build_envelope = None
            _vcr_active_ctx = None



        if _vcr_build_envelope is not None and not _vcr_active_ctx:
            try:
                _vcr_envelope = _vcr_build_envelope(
                    decrypted_message or "",
                )
            except Exception:
                logger.exception(
                    "[CHAT-DEBUG] vault_chat_router_classify_failed "
                    "vault=%s", (vault_id or "")[:8] + "...",
                )
                _vcr_envelope = None
            if isinstance(_vcr_envelope, dict):


                try:
                    from vault_chat_card_data import (
                        populate_vault_chat_card_data,
                    )
                    _vcr_envelope = populate_vault_chat_card_data(
                        _vcr_envelope,
                        vault_id=vault_id,
                        key=key,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] vault_chat_router_populate_failed "
                        "vault=%s", (vault_id or "")[:8] + "...",
                    )



                try:
                    _client_platform = None
                    try:
                        _client_platform = (
                            request.headers.get("X-Client-Platform")
                            or request.headers.get("x-client-platform")
                        )
                        if _client_platform:
                            _client_platform = _client_platform.strip().lower()
                    except Exception:
                        _client_platform = None
                    from vault_chat_crypto_data import (
                        populate_crypto_delegated_card_data,
                    )

                    _vcr_user_tier = "free"
                    try:
                        from billing import (
                            get_account_id_for_vault as _vcr_get_acct,
                            get_entitlement as _vcr_get_ent,
                        )
                        _vcr_acct = _vcr_get_acct(vault_id)
                        if _vcr_acct:
                            _vcr_ent = _vcr_get_ent(_vcr_acct)
                            if (
                                _vcr_ent.block_count > 0
                                and _vcr_ent.purchased_bytes > 0
                            ):
                                _vcr_user_tier = "upgraded"
                    except Exception:
                        logger.exception(
                            "[CHAT-DEBUG] slow-path crypto tier lookup "
                            "failed vault=%s; defaulting to free",
                            (vault_id or "")[:8] + "...",
                        )
                        _vcr_user_tier = "free"
                    _vcr_envelope = populate_crypto_delegated_card_data(
                        _vcr_envelope,
                        vault_id=vault_id,
                        client_platform=_client_platform,
                        user_tier=_vcr_user_tier,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] vault_chat_crypto_populate_failed "
                        "vault=%s", (vault_id or "")[:8] + "...",
                    )


                try:
                    import vault_faq_content_i18n as _faq_i18n_2
                    _vcr_envelope = _faq_i18n_2.apply_locale_to_faq_envelope(
                        _vcr_envelope, _reply_language,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] vault_chat_faq_localize_failed "
                        "vault=%s", (vault_id or "")[:8] + "...",
                    )
                _vcr_card = _vcr_envelope.get("card") or {}
                _vcr_data = _vcr_card.get("data") or {}
                logger.info(
                    "[CHAT-TRACE] vault_chat_router vault=%s "
                    "intent=%s card_type=%s data_available=%s "
                    "envelope_type=%s envelope_schema=%s "
                    "msg_len=%d streaming=false",
                    (vault_id or "")[:8] + "...",
                    str(_vcr_envelope.get("intent") or ""),
                    str((_vcr_card or {}).get("cardType") or ""),
                    bool(_vcr_data.get("available")),
                    str(_vcr_envelope.get("type") or ""),
                    str(_vcr_envelope.get("schema") or ""),
                    len(decrypted_message or ""),
                )




                try:
                    _vcr_intent_str = str(_vcr_envelope.get("intent") or "")
                    _vcr_q = _vcr_card.get("query") if isinstance(_vcr_card, dict) else None
                    _vcr_session_id = str((principal or {}).get("token_id") or "") or None
                    _vcr_login_intents = (
                        "vault_login_search",
                        "vault_login_reveal",
                        "vault_login_copy",
                    )
                    _vcr_card = _vcr_envelope.get("card") or {}
                    _vcr_data = (
                        _vcr_card.get("data")
                        if isinstance(_vcr_card, dict) else None
                    )
                    _vcr_login_id = ""
                    _vcr_label = ""
                    _vcr_is_multi = False
                    _vcr_candidates: list = []
                    if isinstance(_vcr_data, dict):
                        _v = str(_vcr_data.get("view") or "")
                        if _v == "detail" and isinstance(
                            _vcr_data.get("login"), dict,
                        ):
                            _vcr_login_id = str(
                                _vcr_data["login"].get("id") or ""
                            )
                            _vcr_label = str(
                                _vcr_data["login"].get("title") or ""
                            )
                        elif _v == "chooser" and isinstance(
                            _vcr_data.get("logins"), list,
                        ):
                            _vcr_is_multi = True
                            for _r in _vcr_data["logins"][:20]:
                                if not isinstance(_r, dict):
                                    continue
                                _rid = str(_r.get("id") or "")
                                if _rid:
                                    _vcr_candidates.append({"id": _rid})
                    if _vcr_intent_str in _vcr_login_intents and \
                       isinstance(_vcr_q, str) and _vcr_q.strip():
                        from vault_chat_active_entity import (
                            set_active_entity,
                            ENTITY_LOGIN,
                            ACTION_SHOW, ACTION_OPEN, ACTION_VIEW,
                            ACTION_COPY, ACTION_RENAME, ACTION_DELETE,
                            ACTION_EDIT, ACTION_SAVE,
                        )
                        _vcr_ref = (
                            {"id": _vcr_login_id}
                            if _vcr_login_id
                            else {"query": _vcr_q.strip()}
                        )
                        set_active_entity(
                            vault_id,
                            entity_type=ENTITY_LOGIN,
                            entity_ref=_vcr_ref,
                            display_label=(
                                _vcr_label or _vcr_q.strip()
                            ),
                            query=_vcr_q.strip(),
                            allowed_actions=(
                                ACTION_SHOW, ACTION_OPEN, ACTION_VIEW,
                                ACTION_COPY, ACTION_RENAME, ACTION_DELETE,
                                ACTION_EDIT, ACTION_SAVE,
                            ),
                            session_id=_vcr_session_id,
                            is_multi=_vcr_is_multi,
                            candidates=_vcr_candidates,
                        )
                    # Slow-path symmetry with the fast-path: pin the
                    # single-login-list result to the active entity so
                    # follow-ups like "edit it" / "delete it" fire on
                    # the exact login just rendered.
                    elif (
                        _vcr_intent_str == "vault_login_list"
                        and isinstance(_vcr_data, dict)
                        and _vcr_data.get("view") == "detail"
                        and bool(_vcr_login_id)
                    ):
                        from vault_chat_active_entity import (
                            set_active_entity as _sl_set_ae,
                            ENTITY_LOGIN as _SL_ENTITY_LOGIN,
                            ACTION_SHOW as _SL_ACT_SHOW,
                            ACTION_OPEN as _SL_ACT_OPEN,
                            ACTION_VIEW as _SL_ACT_VIEW,
                            ACTION_COPY as _SL_ACT_COPY,
                            ACTION_RENAME as _SL_ACT_RENAME,
                            ACTION_DELETE as _SL_ACT_DELETE,
                            ACTION_EDIT as _SL_ACT_EDIT,
                            ACTION_SAVE as _SL_ACT_SAVE,
                        )
                        _sl_set_ae(
                            vault_id,
                            entity_type=_SL_ENTITY_LOGIN,
                            entity_ref={"id": _vcr_login_id},
                            display_label=(_vcr_label or "Login"),
                            allowed_actions=(
                                _SL_ACT_SHOW, _SL_ACT_OPEN, _SL_ACT_VIEW,
                                _SL_ACT_COPY, _SL_ACT_RENAME, _SL_ACT_DELETE,
                                _SL_ACT_EDIT, _SL_ACT_SAVE,
                            ),
                            session_id=_vcr_session_id,
                        )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] set_active_entity "
                        "(slow path) failed vault=%s",
                        (vault_id or "")[:8] + "...",
                    )

                return encrypted_reply(json.dumps(_vcr_envelope))

        try:
            from vault_config import (
                crypto_default_network,
                crypto_wallet_engine_enabled,
            )
            from wallet_engine_chat import (
                compose_wallet_engine_chat_envelope,
                parse_wallet_engine_chat_message,
                safe_telemetry_label,
            )
            if crypto_wallet_engine_enabled():
                _wallet_parsed = parse_wallet_engine_chat_message(
                    decrypted_message,
                    default_network=crypto_default_network(),
                )
                if _wallet_parsed is not None:
                                                                      
                                                                      
                    from evm_networks import (
                        NETWORK_ETHEREUM_MAINNET,
                        is_receive_enabled,
                        is_send_enabled,
                    )
                    from vault_config import (
                        ethereum_mainnet_erc20_receive_enabled,
                    )
                    _wallet_envelope = compose_wallet_engine_chat_envelope(
                        _wallet_parsed, engine_enabled=True,
                        mainnet_receive_enabled=is_receive_enabled(
                            NETWORK_ETHEREUM_MAINNET,
                        ),
                        mainnet_send_enabled=is_send_enabled(
                            NETWORK_ETHEREUM_MAINNET,
                        ),
                        mainnet_erc20_receive_enabled=(
                            ethereum_mainnet_erc20_receive_enabled()
                        ),
                    )
                                                                   
                                               
                    logger.info(
                        "[CHAT-TRACE] wallet_engine_chat vault=%s %s",
                        (vault_id or "")[:8] + "...",
                        safe_telemetry_label(_wallet_parsed),
                    )
                    return encrypted_reply(json.dumps(_wallet_envelope))
        except Exception:
                                                                    
                                                              
            logger.exception(
                "[CHAT-DEBUG] wallet_engine_chat_failed vault=%s",
                (vault_id or "")[:8] + "...",
            )


        try:
            from vault_secure_item_save import (
                BAND_DELETE_PROMPT as _SECURE_BAND_DELETE_PROMPT,
                BAND_DELETE_CANCELLED as _SECURE_BAND_DELETE_CANCELLED,
                BAND_DELETE_CONFIRMATION_PENDING as _SECURE_BAND_DELETE_CONF,
                BAND_DELETED as _SECURE_BAND_DELETED,
                BAND_DRAFTED as _SECURE_BAND_DRAFTED,
                BAND_FOLLOWUP_CLARIFY as _SECURE_BAND_FOLLOWUP_CLARIFY,
                BAND_NEEDS_CLARIFICATION as _SECURE_BAND_NEEDS_CLAR,
                BAND_NOT_FOUND as _SECURE_BAND_NOT_FOUND,
                BAND_RETRIEVED as _SECURE_BAND_RETRIEVED,
                BAND_SAVED as _SECURE_BAND_SAVED,
                BAND_TIER_REQUIRED as _SECURE_BAND_TIER_REQUIRED,
                BAND_TITLE_UPDATED as _SECURE_BAND_TITLE_UPDATED,
                BAND_VAULT_LOCKED as _SECURE_BAND_VAULT_LOCKED,
                BAND_WARN_DRAFTED as _SECURE_BAND_WARN_DRAFTED,
                route_secure_item_message,
            )
            from vault_active_context import (
                CONTEXT_SECURE_ITEM_DELETE_CONFIRMATION as _SECURE_DELETE_CTX,
                CONTEXT_SECURE_ITEM_DRAFT as _SECURE_CTX,
                CONTEXT_SECURE_ITEM_RESULTS as _SECURE_RESULTS_CTX,
                clear_active_context as _clear_active_ctx,
                set_active_context as _set_active_ctx,
            )
                                                                   
                                                                  
            _user_tier = "free"
            try:
                from billing import (
                    get_account_id_for_vault, get_entitlement,
                )
                _acct_id = get_account_id_for_vault(vault_id)
                if _acct_id:
                    _ent = get_entitlement(_acct_id)
                    if _ent.block_count > 0 and _ent.purchased_bytes > 0:
                        _user_tier = "upgraded"
            except Exception:
                logger.exception(
                    "[CHAT-DEBUG] tier_lookup_failed vault=%s",
                    vault_id,
                )
                _user_tier = "free"

            _secure_result = route_secure_item_message(
                vault_id=vault_id, key=key,
                user_message=decrypted_message or "",
                user_tier=_user_tier,
            )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] secure_item_router_failed vault=%s",
                vault_id,
            )
            _secure_result = {"band": "no_intent", "message": ""}

        _secure_band = (_secure_result or {}).get("band")
        if _secure_band in (
            _SECURE_BAND_DRAFTED,
            _SECURE_BAND_WARN_DRAFTED,
            _SECURE_BAND_TITLE_UPDATED,
        ):
            try:
                _set_active_ctx(vault_id, _SECURE_CTX)
            except Exception:
                pass
            print(
                "[CHAT-TRACE] secure_item_router band="
                f"{_secure_band} vault={(vault_id or '')[:8]}...",
                flush=True,
            )
            return encrypted_reply(_secure_result["message"])
        if _secure_band == _SECURE_BAND_TIER_REQUIRED:
            print(
                "[CHAT-TRACE] secure_item_router band=tier_required "
                f"vault={(vault_id or '')[:8]}...",
                flush=True,
            )
            return encrypted_reply(_secure_result["message"])
        if _secure_band == _SECURE_BAND_SAVED:
                                                              
                                            
            try:
                _clear_active_ctx(vault_id)
            except Exception:
                pass
            print(
                "[CHAT-TRACE] secure_item_router band=saved "
                f"vault={(vault_id or '')[:8]}...",
                flush=True,
            )
            return encrypted_reply(_secure_result["message"])
        if _secure_band in (
            _SECURE_BAND_NEEDS_CLAR,
            _SECURE_BAND_VAULT_LOCKED,
            _SECURE_BAND_RETRIEVED,
            _SECURE_BAND_NOT_FOUND,
        ):
                                                              
                                                        
            if _secure_band == _SECURE_BAND_RETRIEVED:
                try:
                    _set_active_ctx(vault_id, _SECURE_RESULTS_CTX)
                except Exception:
                    pass
            print(
                "[CHAT-TRACE] secure_item_router band="
                f"{_secure_band} vault={(vault_id or '')[:8]}...",
                flush=True,
            )
            return encrypted_reply(_secure_result["message"])
        if _secure_band in (
            _SECURE_BAND_FOLLOWUP_CLARIFY,
            _SECURE_BAND_DELETE_PROMPT,
        ):
                                                                   
                                                              
            try:
                _set_active_ctx(vault_id, _SECURE_RESULTS_CTX)
            except Exception:
                pass
            print(
                "[CHAT-TRACE] secure_item_router band="
                f"{_secure_band} vault={(vault_id or '')[:8]}...",
                flush=True,
            )
            return encrypted_reply(_secure_result["message"])
                                                  
        if _secure_band == _SECURE_BAND_DELETE_CONF:
                                                            
                                                              
            try:
                _set_active_ctx(vault_id, _SECURE_DELETE_CTX)
            except Exception:
                pass
            print(
                "[CHAT-TRACE] secure_item_router band=delete_confirmation_pending "
                f"vault={(vault_id or '')[:8]}...",
                flush=True,
            )
            return encrypted_reply(_secure_result["message"])
        if _secure_band in (
            _SECURE_BAND_DELETED,
            _SECURE_BAND_DELETE_CANCELLED,
        ):
                                                                 
                                                                  
            try:
                _clear_active_ctx(vault_id)
            except Exception:
                pass
            print(
                "[CHAT-TRACE] secure_item_router band="
                f"{_secure_band} vault={(vault_id or '')[:8]}...",
                flush=True,
            )
            return encrypted_reply(_secure_result["message"])
                                                                
                       
        if _pending_save_themed:
            logger.info(
                "[CHAT-TRACE] no_pending_draft_save_themed vault=%s",
                (vault_id or "")[:8] + "...",
            )
            return encrypted_reply(NO_DRAFT_FRIENDLY_REPLY)


        try:
            from vault_crypto_locked_chat import (
                BAND_DEFLECTED as _CRYPTO_BAND_DEFLECTED,
                route_crypto_question,
            )
                                                                   
                                                              
            _crypto_result = route_crypto_question(
                user_message=decrypted_message or "",
                user_tier=locals().get("_user_tier", "free"),
            )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] crypto_deflector_failed vault=%s",
                vault_id,
            )
            _crypto_result = {"band": "no_intent", "message": ""}

        if (_crypto_result or {}).get("band") == _CRYPTO_BAND_DEFLECTED:
            print(
                "[CHAT-TRACE] crypto_deflected vault="
                f"{(vault_id or '')[:8]}... "
                f"intent={_crypto_result.get('intent')}",
                flush=True,
            )
            return encrypted_reply(_crypto_result["message"])


        try:
            _draft = memory.get("pending_login_draft")
        except Exception:
            _draft = None
                                                             
                                                              
        _has_pending_draft = (
            isinstance(_draft, dict) and bool(_draft.get("service"))
        )
        try:
            from vault_active_context import get_active_context
            from vault_chat_context_gate import (
                AMBIGUITY_CLARIFICATION_TEXT,
                should_request_save_ambiguity_clarification,
                should_suppress_pending_draft_save_shortcut,
            )
            _active_ctx = get_active_context(vault_id)
        except Exception:
            _active_ctx = None
            should_request_save_ambiguity_clarification = None                
            should_suppress_pending_draft_save_shortcut = None                
            AMBIGUITY_CLARIFICATION_TEXT = ""                

                                                                  
        if (
            _has_pending_draft
            and should_request_save_ambiguity_clarification is not None
            and should_request_save_ambiguity_clarification(
                active_context=_active_ctx,
                user_message=decrypted_message or "",
                has_pending_draft=True,
            )
        ):
            print(
                "[CHAT-TRACE] save_ambiguity_clarification vault="
                f"{(vault_id or '')[:8]}... "
                f"active_ctx={_active_ctx}",
                flush=True,
            )
            return encrypted_reply(AMBIGUITY_CLARIFICATION_TEXT)

                                                                 
        _suppress_save_shortcut = bool(
            _has_pending_draft
            and should_suppress_pending_draft_save_shortcut is not None
            and should_suppress_pending_draft_save_shortcut(
                active_context=_active_ctx,
                user_message=decrypted_message or "",
                has_pending_draft=True,
            )
        )

        if (
            isinstance(_draft, dict)
            and _draft.get("service")
            and not _suppress_save_shortcut
        ):
            _draft_msg = (decrypted_message or "").strip().lower()
            _pick_idx: Optional[int] = None
            _save_signal = False
            _ordinal_map = {
                "first": 1, "1st": 1, "one": 1,
                "second": 2, "2nd": 2, "two": 2,
                "third": 3, "3rd": 3, "three": 3,
            }
                                                           
            _m_num = re.match(
                r"^\s*(?:#|option\s*)?\s*([123])\s*\.?\s*$",
                _draft_msg,
            )
            if _m_num:
                _pick_idx = int(_m_num.group(1))
                                                                      
            for word, idx in _ordinal_map.items():
                if re.search(rf"\b(?:the\s+)?{word}\s+(one|option)?\b",
                             _draft_msg):
                    _pick_idx = idx
                    break
                                                                      
            if re.search(
                r"\bsave\s+(it|that|this)\b", _draft_msg,
            ):
                _save_signal = True
                if _pick_idx is None:
                    _pick_idx = 1
                                                             
            if _pick_idx is not None and re.search(
                r"\bsave\b", _draft_msg,
            ):
                _save_signal = True

            if _pick_idx is not None:
                _opts = list(_draft.get("username_options") or [])
                if 1 <= _pick_idx <= len(_opts):
                    _picked_username = _opts[_pick_idx - 1]
                else:
                    _picked_username = _opts[0] if _opts else None
                _draft_service = str(_draft.get("service") or "")
                _draft_password = str(_draft.get("password") or "")
                _draft_email_required = bool(
                    _draft.get("policy_email_required")
                )
                if _draft_email_required and not _draft.get(
                    "existing_email"
                ):
                    return encrypted_reply(
                        f"{_draft_service.title()} needs an email "
                        "address as the username — send me the email "
                        "you want to use and I'll save it with this "
                        "password."
                    )
                _new_fields: dict = {"password": _draft_password}
                if _picked_username:
                    _new_fields["username"] = _picked_username
                try:
                    save_secret_tool(
                        vault_id,
                        {
                            "secret_type": "login",
                            "service": _draft_service,
                            "fields": _new_fields,
                        },
                        key,
                        generated=True,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] save_picked_username_failed "
                        "vault=%s", vault_id,
                    )
                    return encrypted_reply(
                        f"I couldn't save the {_draft_service.title()} "
                        "login right now. Try again in a moment."
                    )
                                                                     
                try:
                    memory.pop("pending_login_draft", None)
                except Exception:
                    pass
                memory["last_generated_login"] = {
                    "service": _draft_service,
                    "generated_fields": list(_new_fields.keys()),
                    "ts": int(time.time()),
                }
                                                                  
                                                                  
                _save_lines = [
                    f"Saved your {_draft_service.title()} login. "
                    "The username and password are stored securely "
                    "in your vault.",
                ]
                logger.info(
                    "[CHAT-TRACE] pending_login_pick vault=%s "
                    "service=%s pick_idx=%d save_signal=%s",
                    (vault_id or "")[:8] + "...",
                    _draft_service[:16],
                    _pick_idx,
                    _save_signal,
                )
                return encrypted_reply("\n".join(_save_lines))

                                                                    
        if _direct_ai_tools_enabled:
            logger.info(
                "[CHAT-TRACE] direct_ai_tools_path vault=%s msg_len=%d",
                (vault_id or "")[:8] + "...",
                len(decrypted_message or ""),
            )
            return _route_to_ai_planner_stream()

                                                                    
        if not _direct_ai_tools_enabled:
            try:
                from vault_person_document_lookup import (
                    detect_person_document_query,
                    handle_person_document_query,
                )
                _person_intent = detect_person_document_query(
                    decrypted_message or "",
                )
                if _person_intent is not None:
                    _reply = handle_person_document_query(
                        vault_id=vault_id,
                        key=key,
                        person_name=_person_intent.person_name,
                        doc_kind_hint=_person_intent.doc_kind_hint,
                    )
                    logger.info(
                        "[CHAT-TRACE] person_document_lookup "
                        "(LEGACY canned path) vault=%s person=%r "
                        "doc_kind=%r reply_len=%d",
                        (vault_id or "")[:8] + "...",
                        _person_intent.person_name[:32],
                        _person_intent.doc_kind_hint,
                        len(_reply or ""),
                    )
                    return encrypted_reply(_reply)
            except Exception:
                logger.exception(
                    "[CHAT-DEBUG] person_document_lookup raised "
                    "vault=%s", vault_id,
                )

                                                    
        try:
            from vault_chat_planner_pipeline import run_pipeline as _p2_run
            from semantic_embedder import embed_text as _p2_embed_text

            async def _p2_embed(t: str):
                try:
                    return await _p2_embed_text(client, t)
                except Exception:
                    return None

            def _p2_coverage(v: str) -> dict:
                try:
                    from vault_analysis import analysis_coverage_for_vault
                    from vault_deep_answer import build_coverage_report
                    raw = analysis_coverage_for_vault(v) or {}
                    return build_coverage_report(raw)
                except Exception:
                    return {}

            def _p2_daemon_active() -> bool:
                try:
                    from vault_analysis_daemon import get_status
                    return bool(get_status().running)
                except Exception:
                    return False

            _p2_decision = await _p2_run(
                vault_id=vault_id,
                message=decrypted_message or "",
                embed_fn=_p2_embed,
                coverage_loader=_p2_coverage,
                daemon_active_probe=_p2_daemon_active,
            )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] Phase 2 pipeline raised vault=%s", vault_id,
            )
            _p2_decision = None

                                                                 
        try:
            from vault_result_context import get_last_assistant_result
            _trace_last = get_last_assistant_result(vault_id)
            _trace_lar_present = _trace_last is not None
            _trace_lar_type = (
                _trace_last.result_type if _trace_last else None
            )
        except Exception:
            _trace_lar_present = False
            _trace_lar_type = None
        logger.info(
            "[CHAT-TRACE] pipeline=phase2_followup vault=%s "
            "last_result_present=%s last_result_type=%s "
            "p2_handled=%s p2_delegate=%s p2_reason=%s "
            "msg_len=%d",
            (vault_id or "")[:8] + "…",
            _trace_lar_present,
            _trace_lar_type,
            bool(_p2_decision and _p2_decision.handled),
            bool(_p2_decision and _p2_decision.delegate_to_regex_resolver),
            (_p2_decision.reason if _p2_decision else "raised"),
            len(decrypted_message or ""),
        )

        if _p2_decision is not None and _p2_decision.handled:
            try:
                from vault_config import vault_brain_chat_enabled
                if not vault_brain_chat_enabled():
                    _p2_brain_decision = None
                    logger.info(
                        "[CHAT-TRACE] brain_pipeline_disabled vault=%s "
                        "flag=VAULTAI_EXPERIMENTAL_VAULT_BRAIN_CHAT_ENABLED=false",
                        (vault_id or "")[:8] + "...",
                    )
                else:
                    from vault_brain_chat_pipeline import (
                        run_brain_chat_pipeline as _p2_brain_run,
                    )
                    from brain_embedder import make_query_embed_fn as _p2_brain_embed
                    from vault_analysis import (
                        analysis_coverage_for_vault as _p2_brain_cov,
                    )
                    try:
                        _p2_brain_coverage = _p2_brain_cov(vault_id) or {}
                    except Exception:
                        _p2_brain_coverage = {}
                    _p2_brain_decision = await _p2_brain_run(
                        vault_id=vault_id,
                        message=decrypted_message or "",
                        key=key,
                        embed_fn=_p2_brain_embed(client),
                        coverage=_p2_brain_coverage,
                        pending_file_present=(
                            get_pending_named_file(vault_id) is not None
                        ),
                        has_uploaded_files_in_turn=bool(req.uploaded_file_ids),
                        grounded_answer_client=client,
                    )
                if _p2_brain_decision is not None and _p2_brain_decision.handled:
                    envelope = _build_vault_brain_answer_envelope(
                        message=_p2_brain_decision.reply_body,
                        intent=_p2_brain_decision.intent,
                        evidence_rows=[
                            r.to_dict() for r in _p2_brain_decision.evidence_rows
                        ],
                        no_evidence=_p2_brain_decision.no_evidence,
                        coverage_note=_p2_brain_decision.coverage_note,
                        retrieval_mode=_p2_brain_decision.retrieval_mode,
                        breadth=_p2_brain_decision.breadth,
                        brain_coverage=_p2_brain_decision.brain_coverage_dict,
                        continuation_available=(
                            _p2_brain_decision.continuation_available
                        ),
                    )
                    return encrypted_reply(envelope)
            except Exception:
                logger.exception(
                    "[CHAT-DEBUG] Phase 2 grounded brain retry raised "
                    "vault=%s", vault_id,
                )
            return encrypted_reply(_p2_decision.reply_body)

                                                                     
        try:
            from vault_chat_memory import get_last_file_search_results
            from vault_chat_followup import (
                resolve_followup,
                format_followup_disambiguation,
                format_followup_open_one_message,
            )
            last_search = get_last_file_search_results(vault_id)
            followup = resolve_followup(decrypted_message, last_search)
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] followup_resolver failed vault=%s",
                vault_id,
            )
            followup = {"action": "none"}

        if followup.get("action") == "open_one":
            chosen = followup["file"]
                                                               
                                                                   
            try:
                from vault_chat_memory import clear_last_file_search_results
                clear_last_file_search_results(vault_id)
            except Exception:
                pass
            asset = {
                "id":            chosen.get("file_id"),
                "file_name":     chosen.get("file_name") or "file",
                "content_type":  chosen.get("mime_type"),
                "saved_name":    chosen.get("saved_name"),
                "relative_path": chosen.get("relative_path"),
                "asset_type":    "file",
            }
            asset_name = (
                chosen.get("saved_name")
                or chosen.get("file_name")
                or "file"
            )


            try:
                from vault_chat_active_entity import (
                    set_active_entity,
                    ENTITY_FILE,
                    ACTION_SHOW, ACTION_OPEN, ACTION_VIEW,
                    ACTION_RENAME, ACTION_DELETE, ACTION_DOWNLOAD,
                )
                _file_session_id = str(
                    (principal or {}).get("token_id") or "",
                ) or None
                set_active_entity(
                    vault_id,
                    entity_type=ENTITY_FILE,
                    entity_ref={
                        "file_id": str(chosen.get("file_id") or ""),
                        "content_type": str(chosen.get("mime_type") or ""),
                        "relative_path": str(chosen.get("relative_path") or ""),
                    },
                    display_label=asset_name,
                    query=asset_name,
                    allowed_actions=(
                        ACTION_SHOW, ACTION_OPEN, ACTION_VIEW,
                        ACTION_RENAME, ACTION_DELETE, ACTION_DOWNLOAD,
                    ),
                    session_id=_file_session_id,
                )
            except Exception:
                logger.exception(
                    "[CHAT-DEBUG] set_active_entity file failed",
                )

            reply = _build_structured_asset_reply(asset, asset_name)
            return encrypted_reply(reply)
        if followup.get("action") == "disambiguate":
                                                                      
                                                               
            disambig_files = followup["files"]
            confidences = {
                (f.get("confidence") or "").lower() for f in disambig_files
            }
            if confidences == {"strong"}:
                title = "Which strong credential-file match do you mean?"
            else:
                title = "Which file do you mean?"
            reply = _build_file_disambiguation_envelope(
                files=disambig_files,
                title=title,
                message=format_followup_disambiguation(disambig_files),
                context_kind=(
                    (last_search or {}).get("kind") or "credential_files"
                ),
            )
            return encrypted_reply(reply)

                                           
        try:
            from vault_config import vault_brain_chat_enabled
            if not vault_brain_chat_enabled():
                _brain_decision = None
                logger.info(
                    "[CHAT-TRACE] brain_pipeline_disabled vault=%s "
                    "flag=VAULTAI_EXPERIMENTAL_VAULT_BRAIN_CHAT_ENABLED=false",
                    (vault_id or "")[:8] + "...",
                )
            else:
                from vault_brain_chat_pipeline import (
                    run_brain_chat_pipeline as _brain_run,
                )
                from brain_embedder import make_query_embed_fn as _brain_embed
                from vault_analysis import analysis_coverage_for_vault as _brain_cov

                try:
                    _brain_coverage = _brain_cov(vault_id) or {}
                except Exception:
                    _brain_coverage = {}

                _brain_decision = await _brain_run(
                    vault_id=vault_id,
                    message=decrypted_message or "",
                    key=key,
                    embed_fn=_brain_embed(client),
                    coverage=_brain_coverage,
                    pending_file_present=(
                        get_pending_named_file(vault_id) is not None
                    ),
                    has_uploaded_files_in_turn=bool(req.uploaded_file_ids),
                    grounded_answer_client=client,
                )
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] brain_pipeline_raised vault=%s", vault_id,
            )
            _brain_decision = None

                                                                 
        logger.info(
            "[CHAT-TRACE] pipeline=vault_brain vault=%s "
            "handled=%s intent=%s no_evidence=%s "
            "n_evidence=%d retrieval_mode=%s breadth=%s "
            "continuation_available=%s",
            (vault_id or "")[:8] + "…",
            bool(_brain_decision and _brain_decision.handled),
            (_brain_decision.intent if _brain_decision else None),
            bool(_brain_decision and _brain_decision.no_evidence),
            (len(_brain_decision.evidence_rows) if _brain_decision else 0),
            (_brain_decision.retrieval_mode if _brain_decision else None),
            (_brain_decision.breadth if _brain_decision else None),
            bool(_brain_decision and _brain_decision.continuation_available),
        )
        if _brain_decision is not None and _brain_decision.handled:
            envelope = _build_vault_brain_answer_envelope(
                message=_brain_decision.reply_body,
                intent=_brain_decision.intent,
                evidence_rows=[
                    r.to_dict() for r in _brain_decision.evidence_rows
                ],
                no_evidence=_brain_decision.no_evidence,
                coverage_note=_brain_decision.coverage_note,
                retrieval_mode=_brain_decision.retrieval_mode,
                breadth=_brain_decision.breadth,
                brain_coverage=_brain_decision.brain_coverage_dict,
                continuation_available=_brain_decision.continuation_available,
            )
            return encrypted_reply(envelope)

        lower_msg = decrypted_message.lower().strip()

        safe_for_intent = redact_message(decrypted_message)

                                                                
        pending_file_for_intent = get_pending_named_file(vault_id)
        if pending_file_for_intent is not None:
            ct = (pending_file_for_intent.get("content_type") or "").lower()
            if ct.startswith("audio/"):
                pending_asset_kind = "audio recording"
            elif ct.startswith("video/"):
                pending_asset_kind = "video"
            elif ct.startswith("image/"):
                pending_asset_kind = "image"
            else:
                pending_asset_kind = "file"
            memory = dict(memory) if memory else {}
            memory["pending_named_file"] = {
                "file_name": pending_file_for_intent.get("file_name"),
                "asset_kind": pending_asset_kind,
            }

        print("[CHAT-DEBUG] intent_start", flush=True)
        try:
            intent_data = await detect_vault_intent(
                safe_for_intent,
                memory,
                req.vault_name,
            )
        except Exception as e:
            print(
                f"[CHAT-DEBUG] intent_failed error_type={type(e).__name__} repr={e!r}",
                flush=True,
            )
            raise

                                                               
        raw_message_for_log = decrypted_message or ""
        llm_intent_for_log = intent_data.get("intent")

        final_intent_str, override_applied, override_skipped_reason = (
            decide_pending_file_intent_override(
                llm_intent_for_log,
                raw_message_for_log,
                pending_file_present=pending_file_for_intent is not None,
            )
        )
        if pending_file_for_intent is not None:
            intent_data["intent"] = final_intent_str or llm_intent_for_log
                                                                         
                                                                  
            if override_applied and not intent_data.get("asset_name"):
                intent_data["asset_name"] = raw_message_for_log.strip()

                                                                   
        print(
            "[CHAT-DEBUG] pending_file_override "
            f"raw_message_len={len(raw_message_for_log)} "
            f"llm_intent={llm_intent_for_log!r} "
            f"pending_file_present={pending_file_for_intent is not None} "
            f"pending_file_override_applied={override_applied} "
            f"override_skipped_reason={override_skipped_reason!r} "
            f"final_intent={intent_data.get('intent')!r}",
            flush=True,
        )

        intent = intent_data.get("intent")
        print(f"[CHAT-DEBUG] intent_ok intent={intent!r}", flush=True)

                                                                    
        try:
            from vault_brain_intent import is_credential_files_query
            if is_credential_files_query(decrypted_message or ""):
                if intent != "search_files_for_credentials":
                    logger.info(
                        "[CHAT-TRACE] credential_files_override "
                        "vault=%s prev_intent=%r "
                        "forced=search_files_for_credentials "
                        "reason=brain_pattern_match",
                        (vault_id or "")[:8] + "…",
                        intent,
                    )
                    intent = "search_files_for_credentials"
                    intent_data["intent"] = intent
        except Exception:
            logger.exception(
                "[CHAT-DEBUG] credential_files_override raised "
                "vault=%s", vault_id,
            )

                                              
        service = _normalize_service_name(intent_data.get("service"))
        field = intent_data.get("field")
        value = intent_data.get("value")
        asset_name = intent_data.get("asset_name")
        confirmed = intent_data.get("confirmed") == True

        uploaded_files = _get_uploaded_files_for_chat(
            vault_id,
            req.uploaded_file_ids or [],
        )

                                                                 
        if memory.get("pending_edit"):
            pending = memory["pending_edit"]
            pending_service = pending.get("service")
            pending_field = pending.get("field")
            new_value = decrypted_message.strip()

            if pending_service and pending_field and new_value:
                save_secret_tool(
                    vault_id,
                    {
                        "secret_type": "login",
                        "service": pending_service,
                        "fields": {pending_field: new_value},
                    },
                    key,
                )
                memory.pop("pending_edit", None)
                return encrypted_reply(
                    f"Done. I updated your {pending_service.title()} {pending_field}."
                )

                              
        if memory.get("pending_delete_login"):
            pending_service = service if service != "general" else _normalize_service_name(decrypted_message)

            conn = get_db()
            try:
                cursor = conn.cursor(cursor_factory=RealDictCursor)
                cursor.execute(
                    """
                    DELETE FROM vault_items
                    WHERE vault_id = %s
                      AND item_type = 'login'
                      AND LOWER(service) = LOWER(%s)
                    RETURNING service
                    """,
                    (vault_id, pending_service),
                )
                deleted = cursor.fetchone()
                conn.commit()
                memory.pop("pending_delete_login", None)

                if not deleted:
                    return encrypted_reply(f"I could not find a saved login for {pending_service.title()}.")

                                                                        
                try:
                    from vault_intelligence_updater import on_credential_changed
                    on_credential_changed(vault_id)
                except Exception:
                    pass

                return encrypted_reply(f"Deleted your {deleted['service'].title()} login.")
            finally:
                conn.close()

                                  
        if intent == "list_logins":
            print(
                "[CHAT-DEBUG] response_branch=list_logins "
                f"pending_file_present={pending_file_for_intent is not None}",
                flush=True,
            )
            return encrypted_reply(list_logins_tool(vault_id))

        if intent == "retrieve_login":
            if service == "general":
                return encrypted_reply("Which login do you want me to show?")
            result = retrieve_secret_tool(
                vault_id,
                {"service": service},
                key,
            )
            return encrypted_reply(result)

        if intent == "delete_login":
            if service == "general":
                memory["pending_delete_login"] = True
                return encrypted_reply("Okay. Which login do you want me to delete?")

            conn = get_db()
            try:
                cursor = conn.cursor(cursor_factory=RealDictCursor)
                cursor.execute(
                    """
                    DELETE FROM vault_items
                    WHERE vault_id = %s
                      AND item_type = 'login'
                      AND LOWER(service) = LOWER(%s)
                    RETURNING service
                    """,
                    (vault_id, service),
                )
                deleted = cursor.fetchone()
                conn.commit()

                if not deleted:
                    return encrypted_reply(f"I could not find a saved login for {service.title()}.")

                                                                        
                try:
                    from vault_intelligence_updater import on_credential_changed
                    on_credential_changed(vault_id)
                except Exception:
                    pass

                return encrypted_reply(f"Deleted your {deleted['service'].title()} login.")
            finally:
                conn.close()

        if intent == "edit_login":
            if service == "general" or not field:
                return encrypted_reply("Which login and which field do you want me to update?")

            memory["pending_edit"] = {
                "service": service,
                "field": field,
            }
            return encrypted_reply(f"Okay. Send me the new {field} for your {service.title()} login.")

        if intent == "generate_password":
            if service == "general":
                return encrypted_reply("Which login do you want me to generate a password for?")

            new_password = generate_strong_password()

            save_secret_tool(
                vault_id,
                {
                    "secret_type": "login",
                    "service": service,
                    "fields": {"password": new_password},
                },
                key,
                generated=True,                                      
            )

            return encrypted_reply(
                f"I generated and saved a strong password for your {service.title()} login.\n\n"
                f"Password: {new_password}"
            )

        if intent == "generate_login":
            print(
                f"[CHAT-DEBUG] generate_login_start service={service!r} "
                f"parts_wanted={intent_data.get('parts_wanted')!r}",
                flush=True,
            )
            if service == "general":
                print("[CHAT-DEBUG] generate_login_no_service", flush=True)
                return encrypted_reply(
                    "Which service should I create the login for?"
                )

            parts = intent_data.get("parts_wanted") or "full"
            parts = parts.lower() if isinstance(parts, str) else "full"
            if parts not in ("password_only", "username_password", "full"):
                parts = "full"

            existing = _peek_existing_login_fields(
                vault_id, service, key,
            )
            has_username = bool((existing.get("username") or "").strip())
            has_email = bool((existing.get("email") or "").strip())

            wants_username = parts in ("username_password", "full")
                                                                      
                                                                     
            _msg_l = (decrypted_message or "").lower()
            _save_now_phrasing = bool(re.search(
                r"\b("
                r"save\s+it\s+now|"
                r"save\s+now|"
                r"generate\s+and\s+save|"
                r"create\s+and\s+save|"
                r"and\s+save\s+it|"
                r"save\s+this\s+one"
                r")\b",
                _msg_l,
            ))
                                                                     
                                                                       
            _wants_options = bool(re.search(
                r"\b("
                r"options?|"
                r"alternatives?|"
                r"choices?|"
                r"variants?|"
                r"suggestions?|"
                r"different\s+usernames?|"
                r"a\s+few\s+usernames?|"
                r"some\s+usernames?|"
                r"three\s+usernames?"
                r")\b",
                _msg_l,
            ))
            _options_count = 3 if _wants_options else 1

            new_password = generate_strong_password()
            resolved_policy: Optional[UsernamePolicy] = None
            username_options: list[str] = []

            if wants_username and not has_username and not has_email:
                try:
                    resolved_policy = await resolve_username_policy(
                        service,
                        cache=_get_username_policy_cache(),
                        openai_client=client,
                    )
                except Exception as exc:
                    print(
                        f"[CHAT-DEBUG] username_policy_failed service={service!r} "
                        f"error_type={type(exc).__name__} repr={exc!r}",
                        flush=True,
                    )
                    resolved_policy = conservative_username_policy(service)

                if not resolved_policy.email_required:
                    from username_policy import generate_username_options
                                                                     
                                                                 
                    _identity_hints: list[str] = []
                    try:
                        _vault_name = memory.get("vault_name") or ""
                        if not _vault_name:
                            _vault_name = _lookup_vault_name_safe(
                                vault_id,
                            )
                        if _vault_name:
                            _identity_hints.append(_vault_name)
                    except Exception:
                        pass
                    username_options = generate_username_options(
                        resolved_policy,
                        n=_options_count,
                        forbidden_substrings={service},
                        user_identity_hints=_identity_hints,
                    )

                                                                   
            if _save_now_phrasing:
                new_fields = {"password": new_password}
                if username_options:
                    new_fields["username"] = username_options[0]
                save_secret_tool(
                    vault_id,
                    {
                        "secret_type": "login",
                        "service": service,
                        "fields": new_fields,
                    },
                    key,
                    generated=True,
                )
                memory["last_generated_login"] = {
                    "service": service,
                    "generated_fields": list(new_fields.keys()),
                    "ts": int(time.time()),
                    "username_policy": (
                        {
                            "allowed_chars": resolved_policy.allowed_chars,
                            "min_length": resolved_policy.min_length,
                            "max_length": resolved_policy.max_length,
                            "disallow_symbols": resolved_policy.disallow_symbols,
                            "disallow_underscore": resolved_policy.disallow_underscore,
                            "email_required": resolved_policy.email_required,
                            "source_url": resolved_policy.source_url,
                            "confidence": resolved_policy.confidence,
                        }
                        if resolved_policy is not None
                        else None
                    ),
                }
                save_lines = [
                    f"I generated and saved your {service.title()} login.",
                ]
                if new_fields.get("username"):
                    save_lines.append(f"Username: {new_fields['username']}")
                save_lines.append(f"Password: {new_password}")
                return encrypted_reply("\n".join(save_lines))

                                                                      
            try:
                from vault_active_context import (
                    set_active_context,
                    CONTEXT_CREDENTIAL_DRAFT,
                )
                set_active_context(vault_id, CONTEXT_CREDENTIAL_DRAFT)
            except Exception:
                pass
            memory["pending_login_draft"] = {
                "service": service,
                "username_options": list(username_options),
                "password": new_password,
                "policy_email_required": bool(
                    resolved_policy.email_required
                    if resolved_policy is not None else False
                ),
                "has_existing_username": bool(has_username),
                "has_existing_email": bool(has_email),
                "existing_username": str(existing.get("username") or ""),
                "existing_email": str(existing.get("email") or ""),
                "ts": int(time.time()),
            }

            reply_lines: list[str] = []
            if username_options and len(username_options) >= 2 and _wants_options:
                                                               
                reply_lines.append("Here are three username options:")
                reply_lines.append("")
                for i, opt in enumerate(username_options, start=1):
                    reply_lines.append(f"{i}. {opt}")
                reply_lines.append("")
                reply_lines.append(
                    "I also generated a strong password. Pick one "
                    "username (reply with 1, 2, or 3, or say "
                    "\"save the first one\") and I'll save it."
                )
                reply_lines.append("")
                reply_lines.append(f"Password: {new_password}")
            elif username_options:
                                                                   
                _picked = username_options[0]
                reply_lines.append(
                    f"I created a username and strong password "
                    f"for {service.title()}."
                )
                reply_lines.append("")
                reply_lines.append(f"Username: {_picked}")
                reply_lines.append(f"Password: {new_password}")
                reply_lines.append("")
                reply_lines.append(
                    "Say \"save it\" and I'll store it in your vault."
                )
            elif wants_username and has_username:
                reply_lines.append(
                    f"I'll keep your existing {service.title()} "
                    f"username ({existing['username']})."
                )
                reply_lines.append(f"Password: {new_password}")
                reply_lines.append(
                    "Reply \"save it\" and I'll update the password."
                )
            elif wants_username and has_email:
                reply_lines.append(
                    f"I'll keep your existing {service.title()} "
                    f"email ({existing['email']})."
                )
                reply_lines.append(f"Password: {new_password}")
                reply_lines.append(
                    "Reply \"save it\" and I'll update the password."
                )
            elif (
                wants_username
                and resolved_policy is not None
                and resolved_policy.email_required
            ):
                reply_lines.append(
                    f"{service.title()} uses email addresses as "
                    "usernames, so I'll need the email you want to "
                    "use. Send it and I'll save it alongside this "
                    "password."
                )
                reply_lines.append(f"Password: {new_password}")
            else:
                                       
                reply_lines.append(
                    f"I generated a strong password for your "
                    f"{service.title()} login. Reply \"save it\" "
                    "and I'll store it."
                )
                reply_lines.append(f"Password: {new_password}")

            print(
                f"[CHAT-DEBUG] generate_login_drafted service={service!r} "
                f"options_count={len(username_options)} "
                f"wants_options={_wants_options} "
                f"save_now={_save_now_phrasing} "
                f"policy_confidence="
                f"{resolved_policy.confidence if resolved_policy else 'n/a'}",
                flush=True,
            )
            return encrypted_reply("\n".join(reply_lines))

                                                                        
        if intent == "generated_login_repair":
            print(
                f"[CHAT-DEBUG] generated_login_repair_start "
                f"intent_service={intent_data.get('service')!r} "
                f"requested_fields={intent_data.get('fields_to_regenerate')!r} "
                f"last_gen={memory.get('last_generated_login')!r}",
                flush=True,
            )
            raw_fields = intent_data.get("fields_to_regenerate") or []
            if not isinstance(raw_fields, list):
                raw_fields = []
            requested = []
            for f in raw_fields:
                if not isinstance(f, str):
                    continue
                normalized = f.strip().lower()
                if normalized in ("username", "email", "password"):
                    if normalized not in requested:
                        requested.append(normalized)

            last_gen = memory.get("last_generated_login") or {}
            last_service = last_gen.get("service")

                                                                      
            if service != "general":
                target_service = service
            elif last_service:
                target_service = _normalize_service_name(last_service)
            else:
                return encrypted_reply(
                    "Which login should I update? Tell me the service and "
                    "which field to regenerate (username, password, or both)."
                )

                                                                          
            if not requested:
                requested = list(last_gen.get("generated_fields") or [])
                requested = [
                    f for f in requested
                    if f in ("username", "email", "password")
                ]
            if not requested:
                return encrypted_reply(
                    f"Which field on your {target_service.title()} login "
                    "should I regenerate — username, password, or both?"
                )

            new_fields = {}
            regenerated_labels = []
                                                                      
                                                                      
            tightened_policy: Optional[UsernamePolicy] = None
            for f in requested:
                if f == "password":
                    new_fields["password"] = generate_strong_password()
                    regenerated_labels.append(
                        f"Password: {new_fields['password']}"
                    )
                elif f == "username":
                                                                       
                                                                       
                    raw_policy = (last_gen or {}).get("username_policy")
                    if isinstance(raw_policy, dict):
                        base_policy = UsernamePolicy(
                            service=target_service,
                            allowed_chars=raw_policy.get(
                                "allowed_chars", "letters_numbers"
                            ),
                            min_length=int(raw_policy.get("min_length", 8)),
                            max_length=int(raw_policy.get("max_length", 16)),
                            disallow_symbols=bool(
                                raw_policy.get("disallow_symbols", True)
                            ),
                            disallow_underscore=bool(
                                raw_policy.get("disallow_underscore", True)
                            ),
                            email_required=bool(
                                raw_policy.get("email_required", False)
                            ),
                            source_url=raw_policy.get("source_url"),
                            confidence=raw_policy.get("confidence", "unknown"),
                        )
                    else:
                        try:
                            base_policy = await resolve_username_policy(
                                target_service,
                                cache=_get_username_policy_cache(),
                                openai_client=client,
                            )
                        except Exception:
                            base_policy = conservative_username_policy(
                                target_service
                            )

                    if base_policy.email_required:
                                                                       
                                                                    
                        return encrypted_reply(
                            f"{target_service.title()} uses email addresses "
                            "as usernames. Send me the email you want to "
                            "use and I'll save it for this login."
                        )

                    tightened_policy = tighten_username_policy(base_policy)
                    new_fields["username"] = generate_username_for_policy(
                        tightened_policy
                    )
                    regenerated_labels.append(
                        f"Username: {new_fields['username']}"
                    )
                elif f == "email":
                                                                         
                                                                      
                    return encrypted_reply(
                        f"I can regenerate the username or password for "
                        f"your {target_service.title()} login, but I "
                        "can't make up an email address for you — that "
                        "needs to be one you actually control. Send me "
                        "the email you want to use and I'll save it."
                    )

            save_secret_tool(
                vault_id,
                {
                    "secret_type": "login",
                    "service": target_service,
                    "fields": new_fields,
                },
                key,
                generated=True,
            )

                                                                    
            if tightened_policy is not None:
                try:
                    _get_username_policy_cache().put(tightened_policy)
                except Exception:
                    pass

                                                                       
            prev_fields = set(last_gen.get("generated_fields") or [])
            prev_fields.update(new_fields.keys())
                                                                     
                                                                       
            carried_policy = None
            if tightened_policy is not None:
                carried_policy = {
                    "allowed_chars": tightened_policy.allowed_chars,
                    "min_length": tightened_policy.min_length,
                    "max_length": tightened_policy.max_length,
                    "disallow_symbols": tightened_policy.disallow_symbols,
                    "disallow_underscore": tightened_policy.disallow_underscore,
                    "email_required": tightened_policy.email_required,
                    "source_url": tightened_policy.source_url,
                    "confidence": tightened_policy.confidence,
                }
            elif isinstance((last_gen or {}).get("username_policy"), dict):
                carried_policy = last_gen["username_policy"]
            memory["last_generated_login"] = {
                "service": target_service,
                "generated_fields": list(prev_fields),
                "ts": int(time.time()),
                "username_policy": carried_policy,
            }

            kept = []
            if "password" not in new_fields and "password" in prev_fields:
                kept.append("password")
            if "username" not in new_fields and "username" in prev_fields:
                kept.append("username")

            lines = [
                f"I generated a new "
                + " and ".join(new_fields.keys())
                + f" for your {target_service.title()} login"
                + (f" and kept the {', '.join(kept)} unchanged." if kept
                   else "."),
            ]
            lines.extend(regenerated_labels)
            if tightened_policy is not None and "username" in new_fields:
                                                                       
                                                                       
                lines.append(
                    "I tightened the username rules (letters and numbers "
                    "only, shorter length) so this one should be accepted."
                )
            print(
                f"[CHAT-DEBUG] generated_login_repair_ok service={target_service!r} "
                f"regenerated_keys={list(new_fields.keys())} "
                f"tightened={'yes' if tightened_policy is not None else 'no'}",
                flush=True,
            )
            return encrypted_reply("\n".join(lines))

        if intent == "save_login":
            direct_payload = _maybe_extract_login_payload(
                decrypted_message,
                fallback_service=intent_data.get("service"),
            )
            if not direct_payload:
                                                                      
                                                                  
                print(
                    "[CHAT-DEBUG] save_login_no_credentials_detected",
                    flush=True,
                )
                return encrypted_reply(
                    "I couldn't find any credentials in that message, so "
                    "nothing was saved. Send the service name plus a "
                    "username, email, or password — for example: "
                    "\"save my Capital One login: username=foo "
                    "password=Bar123\"."
                )
            try:
                result = save_secret_tool(
                    vault_id, direct_payload, key,
                )
            except SaveSecretStorageLimitError as exc:
                                                                     
                                                                    
                print(
                    f"[CHAT-DEBUG] save_login_storage_limit "
                    f"used={exc.used_bytes} projected={exc.projected_bytes} "
                    f"limit={exc.limit_bytes}",
                    flush=True,
                )
                used_mb = exc.used_bytes / (1024 * 1024)
                limit_mb = exc.limit_bytes / (1024 * 1024)
                raise HTTPException(
                    status_code=413,
                    detail={
                        "code": "vault_storage_limit_exceeded",
                        "message": (
                            f"Vault storage limit reached. "
                            f"Used {used_mb:.0f} MB of {limit_mb:.0f} MB. "
                            f"Delete data before adding more."
                        ),
                        "used_bytes": exc.used_bytes,
                        "projected_bytes": exc.projected_bytes,
                        "limit_bytes": exc.limit_bytes,
                    },
                )
            except SaveSecretInvalidPayloadError as exc:
                                                                         
                                                                    
                print(
                    f"[CHAT-DEBUG] save_login_invalid_payload "
                    f"reason={exc.reason!r}",
                    flush=True,
                )
                raise HTTPException(
                    status_code=400,
                    detail={
                        "code": "save_login_invalid_payload",
                        "message": (
                            "I couldn't save that login because the "
                            "extracted payload was incomplete. Try "
                            "again with the service name and at least "
                            "one of username, email, or password."
                        ),
                        "reason": exc.reason,
                    },
                )
                                                                      
                                                                    
            return encrypted_reply(result)


        pending_file = (
            pending_file_for_intent
            if pending_file_for_intent is not None
            else get_pending_named_file(vault_id)
        )

        if intent == "name_file" and pending_file:


            raw_candidate = asset_name or decrypted_message
            stripped_candidate = _strip_naming_command(raw_candidate)
            clean_name = _normalize_asset_name(stripped_candidate)
            print(
                "[CHAT-DEBUG] response_branch=name_file "
                f"raw={raw_candidate!r} "
                f"stripped={stripped_candidate!r} clean={clean_name!r} "
                f"file_id={pending_file.get('id')}",
                flush=True,
            )
            if clean_name and clean_name != "general":
                result = save_named_uploaded_asset(
                    vault_id=vault_id,
                    file_id=pending_file["id"],
                    saved_name=clean_name,
                    key=key,
                )


                ct = (pending_file.get("content_type") or "").lower()
                if ct.startswith("audio/"):
                    asset_noun = "recording"
                elif ct.startswith("video/"):
                    asset_noun = "video"
                elif ct.startswith("image/"):
                    asset_noun = "image"
                else:
                    asset_noun = "file"
                return encrypted_reply(
                    f"Saved this {asset_noun} as "
                    f"{_title_case_asset(result['saved_name'])}."
                )
        elif pending_file is not None and intent != "name_file":
            print(
                "[CHAT-DEBUG] pending_file_MISSED intent="
                f"{intent!r} file_id="
                f"{pending_file.get('id')!r} raw_message_len="
                f"{len(decrypted_message or '')}",
                flush=True,
            )

                                                                     
        if intent in ("summarize_vault", "list_files"):
            try:
                from vault_inventory import (
                    summarize_vault_contents,
                    format_vault_summary_reply,
                )
                from vault_analysis import (
                    analysis_coverage_for_vault,
                    summarise_coverage_for_chat,
                    chat_coverage_note,
                )
                rows = list_uploaded_files(vault_id)
                summary = summarize_vault_contents(rows)
                raw_coverage = analysis_coverage_for_vault(vault_id)
                chat_coverage = summarise_coverage_for_chat(raw_coverage)
                base_message = format_vault_summary_reply(summary)
                note = chat_coverage_note(raw_coverage)
                full_message = (
                    f"{base_message}\n\n{note}" if note else base_message
                )
                reply = _build_vault_inventory_envelope(
                    summary=summary,
                    message=full_message,
                    analysis_coverage=chat_coverage,
                )
                return encrypted_reply(reply)
            except Exception:
                logger.exception("summarize_vault failed")
                return encrypted_reply(
                    "I couldn't read your vault inventory right now. "
                    "Try again in a moment."
                )

        if intent == "list_folders":
            try:
                from vault_inventory import (
                    list_top_folders,
                    format_folder_list_reply,
                )
                rows = list_uploaded_files(vault_id)
                folders = list_top_folders(rows)
                return encrypted_reply(format_folder_list_reply(folders))
            except Exception:
                logger.exception("list_folders failed")
                return encrypted_reply(
                    "I couldn't read your folder list right now."
                )

        if intent == "list_recent_uploads":
            try:
                from vault_inventory import (
                    list_recent_files,
                    DEFAULT_RECENT_FILES_LIMIT,
                )
                rows = list_uploaded_files(vault_id)
                recent = list_recent_files(
                    rows, limit=DEFAULT_RECENT_FILES_LIMIT,
                )
                if not recent:
                    return encrypted_reply(
                        "I don't see any recent uploads."
                    )
                title = f"Your {len(recent)} most recent files"
                reply = _build_vault_file_list_envelope(
                    recent,
                    title,
                    message=title + ".",
                )
                return encrypted_reply(reply)
            except Exception:
                logger.exception("list_recent_uploads failed")
                return encrypted_reply(
                    "I couldn't read your recent uploads right now."
                )

        if intent == "search_files_for_credentials":
                                              
             
            logger.info(
                "[CHAT-TRACE] credential_files_route "
                "vault=%s route=strict_verifier_then_brain",
                (vault_id or "")[:8] + "…",
            )
            try:
                try:
                    from vault_config import vault_brain_chat_enabled
                    if not vault_brain_chat_enabled():
                        _cred_brain_decision = None
                        logger.info(
                            "[CHAT-TRACE] brain_pipeline_disabled vault=%s "
                            "flag=VAULTAI_EXPERIMENTAL_VAULT_BRAIN_CHAT_ENABLED=false",
                            (vault_id or "")[:8] + "...",
                        )
                    else:
                        from vault_brain_chat_pipeline import (
                            run_brain_chat_pipeline as _cred_brain_run,
                        )
                        from brain_embedder import (
                            make_query_embed_fn as _cred_brain_embed,
                        )
                        from vault_analysis import (
                            analysis_coverage_for_vault as _cred_brain_cov,
                        )
                        try:
                            _cred_brain_coverage = (
                                _cred_brain_cov(vault_id) or {}
                            )
                        except Exception:
                            _cred_brain_coverage = {}
                        _cred_brain_decision = await _cred_brain_run(
                            vault_id=vault_id,
                            message=decrypted_message or "",
                            key=key,
                            embed_fn=_cred_brain_embed(client),
                            coverage=_cred_brain_coverage,
                            pending_file_present=(
                                get_pending_named_file(vault_id) is not None
                            ),
                            has_uploaded_files_in_turn=bool(req.uploaded_file_ids),
                            grounded_answer_client=client,
                        )
                    if _cred_brain_decision is not None and _cred_brain_decision.handled:
                        envelope = _build_vault_brain_answer_envelope(
                            message=_cred_brain_decision.reply_body,
                            intent=_cred_brain_decision.intent,
                            evidence_rows=[
                                r.to_dict()
                                for r in _cred_brain_decision.evidence_rows
                            ],
                            no_evidence=_cred_brain_decision.no_evidence,
                            coverage_note=_cred_brain_decision.coverage_note,
                            retrieval_mode=_cred_brain_decision.retrieval_mode,
                            breadth=_cred_brain_decision.breadth,
                            brain_coverage=(
                                _cred_brain_decision.brain_coverage_dict
                            ),
                            continuation_available=(
                                _cred_brain_decision.continuation_available
                            ),
                        )
                        return encrypted_reply(envelope)
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] credential brain handoff failed "
                        "vault=%s", vault_id,
                    )
                                                                  
                                                                 
                from vault_inventory import (
                    verified_credential_files_report,
                    format_credential_files_reply,
                )
                rows = _list_uploaded_files_for_credential_search(
                    vault_id, key,
                )
                                                            
                                                                   
                try:
                    from vault_analysis import analysis_coverage_for_vault
                    from vault_deep_answer import build_coverage_report
                    raw_cov = analysis_coverage_for_vault(vault_id) or {}
                    coverage = build_coverage_report(raw_cov)
                except Exception:
                    coverage = None
                                                                    
                                                                       
                report = verified_credential_files_report(
                    rows, coverage=coverage,
                )
                matches = report["matches"]
                not_scanned = int(report.get("not_scanned_count") or 0)
                scanned = int(report.get("scanned_count") or 0)
                is_partial = bool(report.get("is_partial"))
                                                                   
                                                                    
                try:
                    from vault_ai_provider import active_provider_name
                    _trace_provider = active_provider_name()
                except Exception:
                    _trace_provider = "unknown"
                logger.info(
                    "[CHAT-TRACE] route_name=search_files_for_credentials "
                    "vault=%s provider_used=%s intent=%s "
                    "evidence_bundle_present=%s strict_verifier_used=%s "
                    "coverage_scanned=%d coverage_not_scanned=%d "
                    "is_partial=%s result_file_count=%d "
                    "fallback_used=%s",
                    (vault_id or "")[:8] + "…",
                    _trace_provider,
                    "search_files_for_credentials",
                    True,
                    True,
                    scanned,
                    not_scanned,
                    is_partial,
                    len(matches or []),
                    False,
                )

                                                                 
                composer_used = False
                composer_message = ""
                try:
                    from vault_credential_files_composer import (
                        compose_credential_files_answer,
                    )
                    composer_message = await compose_credential_files_answer(
                        user_question=decrypted_message or "",
                        vault_id=vault_id,
                        verified_matches=matches,
                        scanned_count=scanned,
                        not_scanned_count=not_scanned,
                        is_partial=is_partial,
                        coverage_extra=coverage,
                        authorized_unlocked=True,
                    )
                    composer_used = bool(composer_message)
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] credential_files composer failed "
                        "vault=%s — falling back to deterministic copy",
                        vault_id,
                    )

                if composer_used:
                    reply_message = composer_message
                else:
                    reply_message = format_credential_files_reply(
                        matches,
                        not_scanned_count=not_scanned,
                        scanned_count=scanned,
                        is_partial=is_partial,
                    )
                logger.info(
                    "[CHAT-TRACE] credential_files_composer "
                    "vault=%s composer_used=%s message_len=%d "
                    "verified_file_count=%d",
                    (vault_id or "")[:8] + "…",
                    composer_used,
                    len(reply_message or ""),
                    len(matches or []),
                )
                reply = _build_credential_files_envelope(
                    matches=matches,
                    message=reply_message,
                    scanned_count=scanned,
                    not_scanned_count=not_scanned,
                    is_partial=is_partial,
                )
                                                                      
                                                                      
                try:
                    from vault_chat_memory import set_last_file_search_results
                    set_last_file_search_results(
                        vault_id,
                        kind="credential_files",
                        results=matches,
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] set_last_file_search_results "
                        "failed vault=%s", vault_id,
                    )
                                                                     
                                                                      
                try:
                    from vault_result_context import (
                        make_credential_search_result,
                        set_last_assistant_result,
                    )
                    file_ids = [
                        str(m.get("file_id") or m.get("id") or "")
                        for m in matches
                    ]
                    set_last_assistant_result(
                        vault_id,
                        make_credential_search_result(
                            query=decrypted_message or "",
                            file_ids_returned=[fid for fid in file_ids if fid],
                            file_ids_excluded=[],
                            total_matches_known=len(file_ids),
                            coverage=dict(coverage or {}) if coverage else {},
                            is_partial=bool(is_partial),
                        ),
                    )
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] set_last_assistant_result (Phase 2) "
                        "failed vault=%s", vault_id,
                    )
                return encrypted_reply(reply)
            except Exception:
                logger.exception("search_files_for_credentials failed")
                return encrypted_reply(
                    "I couldn't scan your files for credentials right now."
                )

        if intent == "extract_logins_from_file":
            try:
                reply = _handle_extract_logins_from_file(
                    vault_id=vault_id,
                    key=key,
                    asset_name=asset_name,
                )
                return encrypted_reply(reply)
            except Exception:
                logger.exception("extract_logins_from_file failed")
                return encrypted_reply(
                    "I couldn't extract logins from that file right now."
                )

        if intent == "search_files_about":
                                                                 
                                                                  
            try:
                from vault_understanding_search import (
                    search_vault_understanding,
                    format_search_reply,
                )
                                                                 
                                                               
                query = (asset_name or decrypted_message or "").strip()
                report = search_vault_understanding(
                    vault_id, query, key=key, limit=15,
                )
                results = report["results"]
                pending_n = int(report.get("pending_count") or 0)
                stale_n = int(report.get("stale_count") or 0)
                pending_emb_n = int(
                    report.get("pending_embedding_count") or 0
                )
                semantic_available = bool(
                    report.get("semantic_available")
                )
                message = format_search_reply(
                    query=query,
                    results=results,
                    pending_count=pending_n,
                    stale_count=stale_n,
                    pending_embedding_count=pending_emb_n,
                    semantic_available=semantic_available,
                )
                reply = _build_file_search_envelope(
                    query=query,
                    results=results,
                    message=message,
                    pending_count=pending_n,
                    stale_count=stale_n,
                    pending_embedding_count=pending_emb_n,
                    semantic_available=semantic_available,
                )
                return encrypted_reply(reply)
            except Exception:
                logger.exception("search_files_about failed")
                return encrypted_reply(
                    "I couldn't search your files right now."
                )

        if intent == "related_files":
                                                                 
                                                              
            try:
                anchor_row = _resolve_file_for_analysis(
                    vault_id,
                    asset_name=asset_name,
                    decrypted_message=decrypted_message,
                    key=None,                                   
                )
                if anchor_row is None:
                    return encrypted_reply(
                        "I couldn't find a file matching that name "
                        "in your vault. Try \"list my files\" to "
                        "see what's available."
                    )

                if not str(anchor_row.get("id") or ""):
                    return encrypted_reply(
                        "I couldn't find the anchor file for that "
                        "request."
                    )

                envelope = _compose_related_files_envelope(
                    vault_id=vault_id, anchor_row=anchor_row,
                )
                return encrypted_reply(json.dumps(envelope))
            except Exception:
                logger.exception("related_files failed")
                return encrypted_reply(
                    "I couldn't load the related-files graph "
                    "right now."
                )

        if intent == "vault_clusters":
                                                            
                                                              
            try:
                envelope = _compose_vault_clusters_envelope(
                    vault_id,
                )
                return encrypted_reply(json.dumps(envelope))
            except Exception:
                logger.exception("vault_clusters failed")
                return encrypted_reply(
                    "I couldn't load your vault clusters right "
                    "now."
                )

        if intent == "analyze_file":
                                                                  
                                                                  
            try:
                from vault_file_analysis import (
                    build_safe_file_analysis,
                    user_explicitly_requested_values,
                )
                from vault_document_purpose import classify_document_purpose
                from vault_inventory import _compute_credential_density_metrics

                row = _resolve_file_for_analysis(
                    vault_id,
                    asset_name=asset_name,
                    decrypted_message=decrypted_message,
                    key=key,
                )
                if row is None:
                    return encrypted_reply(
                        "I couldn't find a file matching that name "
                        "in your vault. Try \"list my files\" to "
                        "see what's available."
                    )

                file_id = str(row.get("id") or "")
                pretty_name = (
                    row.get("saved_name")
                    or row.get("file_name")
                    or "this file"
                )
                show_secret_values = user_explicitly_requested_values(
                    decrypted_message,
                )

                                                               
                understanding_row = None
                if file_id:
                    try:
                        from vault_understanding import (
                            get_understanding_for_file,
                            UNDERSTANDING_STATUS_READY,
                        )
                        understanding_row = get_understanding_for_file(
                            vault_id, file_id,
                        )
                    except Exception:
                        understanding_row = None

                if understanding_row and \
                        understanding_row.get("status") == "ready":
                    purpose_decision = {
                        "purpose":       understanding_row.get("document_purpose"),
                        "purpose_label": understanding_row.get("purpose_label") or "",
                        "confidence":    float(
                            understanding_row.get("purpose_confidence") or 0.0
                        ),
                        "evidence":      [],
                    }
                    metrics = understanding_row.get("credential_signals_jsonb") or {}
                else:
                    metrics = _compute_credential_density_metrics(
                        row.get("extracted_text"),
                    )
                    purpose_decision = classify_document_purpose(
                        row.get("extracted_text"), metrics=metrics,
                    )

                extracted_text = row.get("extracted_text")

                                                                   
                composer_reply = ""
                try:
                    from vault_file_read_composer import (
                        compose_file_read_answer,
                    )
                    composer_reply = await compose_file_read_answer(
                        user_question=decrypted_message or "",
                        vault_id=vault_id,
                        file_id=file_id,
                        file_name=pretty_name,
                        extracted_text=extracted_text,
                        saved_name=row.get("saved_name"),
                        relative_path=row.get("relative_path"),
                        mime_type=row.get("content_type"),
                        asset_type=row.get("asset_type"),
                        purpose=purpose_decision.get("purpose"),
                        purpose_label=purpose_decision.get("purpose_label"),
                        metrics=metrics,
                        authorized_unlocked=True,
                    )
                    composer_reply = (composer_reply or "").strip()
                except Exception:
                    logger.exception(
                        "[CHAT-DEBUG] file-read composer failed "
                        "vault=%s file=%s — falling back to "
                        "deterministic analyzer",
                        vault_id, file_id,
                    )
                    composer_reply = ""

                if composer_reply:
                    logger.info(
                        "[CHAT-TRACE] analyze_file composer "
                        "vault=%s file=%s composer_used=True "
                        "reply_len=%d",
                        (vault_id or "")[:8] + "…",
                        (file_id or "")[:8] + "…",
                        len(composer_reply),
                    )
                    return encrypted_reply(composer_reply)

                reply = build_safe_file_analysis(
                    extracted_text,
                    file_name=pretty_name,
                    metrics=metrics,
                    purpose_decision=purpose_decision,
                    show_secret_values=show_secret_values,
                )
                logger.info(
                    "[CHAT-TRACE] analyze_file composer "
                    "vault=%s file=%s composer_used=False "
                    "fallback=deterministic reply_len=%d",
                    (vault_id or "")[:8] + "…",
                    (file_id or "")[:8] + "…",
                    len(reply or ""),
                )
                return encrypted_reply(reply)
            except Exception:
                logger.exception("analyze_file failed")
                return encrypted_reply(
                    "I couldn't analyze that file right now."
                )

        if intent == "travel_readiness":
            try:
                from vault_inventory import (
                    travel_readiness_check,
                    format_travel_readiness_reply,
                )
                rows = list_uploaded_files(vault_id)
                metadata_rows = _list_document_metadata_for_vault(vault_id)
                report = travel_readiness_check(
                    rows=rows,
                    metadata_rows=metadata_rows,
                )
                reply = _build_travel_readiness_envelope(
                    report=report,
                    message=format_travel_readiness_reply(report),
                )
                return encrypted_reply(reply)
            except Exception:
                logger.exception("travel_readiness failed")
                return encrypted_reply(
                    "I couldn't run your travel readiness check right now."
                )

        if intent == "retrieve_file":
                                                                    
                                                                  
            folder_aware_reply = _try_folder_aware_file_retrieval(
                vault_id=vault_id,
                decrypted_message=decrypted_message,
                asset_name_from_llm=asset_name,
            )
            if folder_aware_reply is not None:
                return encrypted_reply(folder_aware_reply)

            clean_asset_name = _normalize_asset_name(asset_name or decrypted_message)
            asset = retrieve_saved_asset(vault_id, clean_asset_name)

            if asset:
                return encrypted_reply(_build_structured_asset_reply(asset, clean_asset_name))

                                        
            try:
                hybrid_hits = await hybrid_retrieve_files(
                    vault_id,
                    (asset_name or decrypted_message or "").strip(),
                    limit=5,
                )
            except Exception as e:
                logger.warning("retrieve_file hybrid fallback raised: %s", e)
                hybrid_hits = []
            if hybrid_hits:
                top = hybrid_hits[0]
                print(
                    "[CHAT-DEBUG] retrieve_file_hybrid_top "
                    f"file_id={top.get('id')} score={top.get('score')!r} "
                    f"reasons={top.get('reasons')!r}",
                    flush=True,
                )
                full = _lookup_file_by_id(vault_id, top["id"])
                if full:
                    display = (
                        full.get("saved_name")
                        or full.get("file_name")
                        or clean_asset_name
                        or "file"
                    )
                    return encrypted_reply(
                        _build_structured_asset_reply(full, display)
                    )

                                                                 
        if intent == "read_file_text":
            try:
                from vault_read_file import handle_read_file_text_intent
                                                             
                                                                    
                file_query = (
                    (intent_data.get("file_name")
                     or intent_data.get("asset_name")
                     or asset_name
                     or decrypted_message
                     or "").strip()
                )
                reply = handle_read_file_text_intent(
                    vault_id=vault_id,
                    query_name=file_query,
                    key=key,
                )
                return encrypted_reply(reply)
            except Exception:
                logger.exception(
                    "read_file_text intent handler crashed vault=%s",
                    (vault_id or "")[:8],
                )
                return encrypted_reply(
                    "I couldn't read that file right now. Please "
                    "try again or check the file in your vault."
                )

                                                                           
        if intent == "list_by_tag":
            tag_reply = _handle_list_by_tag(
                vault_id, (intent_data.get("tag") or "").strip(),
            )
            if tag_reply:
                return encrypted_reply(tag_reply)
                                

        if intent == "understand_document":
            doc_reply = _handle_understand_document(
                vault_id,
                intent_data.get("doc_type"),
                intent_data.get("doc_field"),
            )
            if doc_reply:
                return encrypted_reply(doc_reply)

                                                         
        if intent == "document_details":
            ent_reply = _handle_document_details(
                vault_id,
                intent_data.get("entity_doc_type"),
                intent_data.get("entity_type"),
                intent_data.get("entity_key"),
            )
            if ent_reply:
                return encrypted_reply(ent_reply)

                                                         
        if intent == "expiry_status":
            stat_reply = _handle_expiry_status(
                vault_id,
                intent_data.get("expiry_filter"),
            )
            if stat_reply:
                return encrypted_reply(stat_reply)

                                                   
        if intent == "expiry_alerts":
            exp_reply = _handle_expiry_alerts(vault_id)
            if exp_reply:
                return encrypted_reply(exp_reply)

                                                               
        if intent == "remember_fact":
            rem_reply = _handle_remember_fact(
                vault_id,
                intent_data.get("memory_type"),
                intent_data.get("memory_key"),
                intent_data.get("memory_value"),
                intent_data.get("memory_event_date"),
            )
            if rem_reply:
                return encrypted_reply(rem_reply)

        if intent == "recall_memory":
            rec_reply = _handle_recall_memory(
                vault_id,
                intent_data.get("memory_type"),
                intent_data.get("memory_query"),
                intent_data.get("memory_anchor"),
                intent_data.get("memory_direction"),
            )
            if rec_reply:
                return encrypted_reply(rec_reply)

                                          
        if intent == "related_items":
            anchor_text = (intent_data.get("anchor_text") or "").strip()
            try:
                from related_files import (
                    handle_related_items_envelope_payload,
                )
                envelope_data = handle_related_items_envelope_payload(
                    vault_id, anchor_text,
                )
            except Exception:
                envelope_data = None
            if envelope_data:
                payload = envelope_data["payload"]
                reply = _build_related_files_envelope(
                    anchor=payload["anchor"],
                    results=payload["results"],
                    message=envelope_data["message"],
                )
                return encrypted_reply(reply)

            try:
                from relationship_builder import handle_related_items
                rel_reply = handle_related_items(
                    vault_id, anchor_text,
                )
            except Exception:
                rel_reply = None
            if rel_reply:
                return encrypted_reply(rel_reply)

                                                 
        if intent == "search_memory":
            print("[CHAT-DEBUG] semantic_start", flush=True)
            try:
                search_reply = await _handle_search_memory_intent(
                    vault_id, decrypted_message, key,
                )
            except Exception as e:
                print(
                    f"[CHAT-DEBUG] semantic_failed error_type={type(e).__name__} repr={e!r}",
                    flush=True,
                )
                raise
            print(
                f"[CHAT-DEBUG] semantic_ok matched={bool(search_reply)}",
                flush=True,
            )
            if search_reply:
                return encrypted_reply(search_reply)
                                                                 

        return _route_to_ai_planner_stream()

    except HTTPException:
                                                                      
                                                                  
        raise
    except Exception as _post_decrypt_exc:
                                                                 
                                                               
        tb_str = traceback.format_exc()
        print(
            f"[CHAT-DEBUG] post_decrypt_caught error_type={type(_post_decrypt_exc).__name__} "
            f"repr={_post_decrypt_exc!r}",
            flush=True,
        )
        print("[CHAT-DEBUG] post_decrypt_traceback BEGIN", flush=True)
        print(tb_str, flush=True)
        print("[CHAT-DEBUG] post_decrypt_traceback END", flush=True)
        logger.error(
            "Chat endpoint failed (post-decrypt)",
            exc_info=_post_decrypt_exc,
        )
        raise HTTPException(
            status_code=500,
            detail="Chat processing error",
        )


if __name__ == "__main__":
                                                                    
                                                                    
    dev_reload = (
        os.getenv("DEV_RELOAD", "false").strip().lower() == "true"
    )
    logger.info(
        "Starting VaultAI server on port 8000 (reload=%s)...", dev_reload,
    )
    uvicorn.run(
        "main:app", host="0.0.0.0", port=8000,
        reload=dev_reload, log_level="info",
    )
